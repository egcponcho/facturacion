"""Plantillas de Excel para las cargas masivas y lectura de los archivos.

Cada plantilla tiene una hoja "Data" con los encabezados (los obligatorios
llevan *), listas desplegables donde hay valores fijos y filas de ejemplo; una
hoja "Instructions" con lo que significa cada columna y una hoja "Values" con
los valores permitidos. Al leer, los encabezados se reconocen sin importar
mayúsculas, acentos, espacios ni el asterisco.
"""
import csv
import io
import unicodedata
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .common import ErrorNegocio

ACENTO = "5B3FD0"
FONDO = "F1EDFE"
REQ = "FDECEA"


def norm(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode().lower().replace("*", "")
    return "_".join("".join(c if c.isalnum() else " " for c in t).split())


def texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return (v.date() if isinstance(v, datetime) else v).isoformat()
    return str(v).strip()


def _hoja_datos(ws, valores, col_val: int, columnas: list[dict], ejemplos: list[list] | None) -> int:
    for i, c in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i, value=c["nombre"] + (" *" if c.get("req") else ""))
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=ACENTO)
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = c.get("ancho") or max(12, min(40, len(c["nombre"]) + 6))
        if c.get("opciones"):
            letra = get_column_letter(col_val)
            valores.cell(row=1, column=col_val, value=c["nombre"]).font = Font(bold=True)
            for j, o in enumerate(c["opciones"], start=2):
                valores.cell(row=j, column=col_val, value=o)
            valores.column_dimensions[letra].width = max(14, min(45, max(len(str(o)) for o in c["opciones"]) + 2))
            dv = DataValidation(type="list", formula1=f"=Values!${letra}$2:${letra}${len(c['opciones']) + 1}",
                                allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(i)}2:{get_column_letter(i)}5000")
            col_val += 1
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"
    for r, fila in enumerate(ejemplos or [], start=2):
        for i, v in enumerate(fila, start=1):
            ws.cell(row=r, column=i, value=v)
    return col_val


def plantilla(titulo: str, columnas: list[dict], ejemplos: list[list] | None = None,
              instrucciones: list[str] | None = None) -> bytes:
    """columnas: [{"nombre", "ayuda", "req", "opciones": [..] | None, "ancho"}]"""
    return plantilla_hojas(titulo, [("Data", columnas, ejemplos)], instrucciones)


def plantilla_hojas(titulo: str, hojas: list[tuple[str, list[dict], list[list] | None]],
                    instrucciones: list[str] | None = None) -> bytes:
    """Varias hojas de datos (p. ej. Generics y Sizes) con una sola hoja de
    valores permitidos y de instrucciones."""
    wb = Workbook()
    primera = True
    hojas_ws = []
    for nombre, _, _ in hojas:
        ws = wb.active if primera else wb.create_sheet(nombre)
        ws.title = nombre
        primera = False
        hojas_ws.append(ws)
    valores = wb.create_sheet("Values")
    guia = wb.create_sheet("Instructions")
    col_val = 1
    for ws, (_, columnas, ejemplos) in zip(hojas_ws, hojas):
        col_val = _hoja_datos(ws, valores, col_val, columnas, ejemplos)
    guia.column_dimensions["A"].width = 30
    guia.column_dimensions["B"].width = 100
    guia.cell(row=1, column=1, value=titulo).font = Font(bold=True, size=14)
    fila = 2
    for t in instrucciones or []:
        guia.cell(row=fila, column=1, value=t).alignment = Alignment(wrap_text=True)
        guia.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
        fila += 1
    for nombre, columnas, _ in hojas:
        fila += 1
        guia.cell(row=fila, column=1, value=f"Sheet {nombre}" if len(hojas) > 1 else "Column").font = Font(bold=True)
        guia.cell(row=fila, column=2, value="What goes in it").font = Font(bold=True)
        fila += 1
        for c in columnas:
            a = guia.cell(row=fila, column=1, value=c["nombre"] + (" *" if c.get("req") else ""))
            if c.get("req"):
                a.fill = PatternFill("solid", fgColor=REQ)
            ayuda = c.get("ayuda") or ""
            if c.get("opciones"):
                ayuda += (" " if ayuda else "") + "Choose from the list (sheet Values)."
            guia.cell(row=fila, column=2, value=ayuda).alignment = Alignment(wrap_text=True)
            fila += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def hojas(nombre: str, contenido: bytes) -> list[str]:
    if not (nombre or "").lower().endswith((".xlsx", ".xlsm")):
        return []
    try:
        return load_workbook(io.BytesIO(contenido), read_only=True).sheetnames
    except Exception:
        return []


def leer(nombre: str, contenido: bytes, alias: dict[str, str], hoja: str | None = None,
         vacio_ok: bool = False) -> list[dict]:
    """Filas del archivo como {campo: texto, "_fila": n}. `alias` lleva el
    encabezado normalizado al nombre del campo."""
    nombre = (nombre or "").lower()
    try:
        if nombre.endswith((".xlsx", ".xlsm")):
            wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
            ws = wb[hoja] if hoja and hoja in wb.sheetnames else wb["Data"] if "Data" in wb.sheetnames else wb.active
            filas = [[texto(v) for v in f] for f in ws.iter_rows(values_only=True)]
        elif nombre.endswith((".csv", ".txt")):
            t = contenido.decode("utf-8-sig", errors="replace")
            delim = ";" if t[:2000].count(";") > t[:2000].count(",") else ","
            filas = [[v.strip() for v in f] for f in csv.reader(io.StringIO(t), delimiter=delim)]
        else:
            raise ErrorNegocio("Upload an Excel (.xlsx) or CSV file.", 422, "formato")
    except ErrorNegocio:
        raise
    except Exception:
        raise ErrorNegocio("The file could not be read. Use the template.", 422, "formato")
    filas = [f for f in filas if any(v for v in f)]
    if len(filas) < 2 and vacio_ok:
        return []
    if len(filas) < 2:
        raise ErrorNegocio("The file has no rows with data.", 422, "archivo_vacio")
    enc = [alias.get(norm(h), norm(h)) for h in filas[0]]
    return [{**{enc[i]: (f[i] if i < len(f) else "") for i in range(len(enc)) if enc[i]}, "_fila": n}
            for n, f in enumerate(filas[1:], start=2)]


def si_no(v) -> bool | None:
    t = norm(v)
    if t in ("1", "yes", "y", "si", "true", "x", "verdadero"):
        return True
    if t in ("0", "no", "n", "false", "falso"):
        return False
    return None


def vista(contenido: bytes) -> dict:
    """Vista previa de una plantilla ya armada: sus hojas de datos con las
    columnas (obligatorias y ayuda) y las filas de ejemplo, y las instrucciones."""
    wb = load_workbook(io.BytesIO(contenido), read_only=False)
    ayudas, instrucciones = {}, []
    if "Instructions" in wb.sheetnames:
        en_columnas = False
        for i, (a, b) in enumerate(wb["Instructions"].iter_rows(min_col=1, max_col=2, values_only=True)):
            if i == 0 or a is None:
                continue
            if b == "What goes in it":
                en_columnas = True
            elif en_columnas:
                ayudas[str(a).removesuffix(" *")] = b or ""
            else:
                instrucciones.append(str(a))
    hojas_out = []
    for ws in wb.worksheets:
        if ws.title in ("Values", "Instructions"):
            continue
        filas = [list(r) for r in ws.iter_rows(values_only=True)]
        if not filas:
            continue
        cab = [str(x) for x in filas[0] if x is not None]
        cols = [{"nombre": c.removesuffix(" *"), "req": c.endswith(" *"), "ayuda": ayudas.get(c.removesuffix(" *"), "")} for c in cab]
        ejemplos = [["" if v is None else str(v) for v in r[:len(cab)]] for r in filas[1:6] if any(v is not None for v in r)]
        hojas_out.append({"nombre": ws.title, "columnas": cols, "filas": ejemplos})
    return {"hojas": hojas_out, "instrucciones": instrucciones}
