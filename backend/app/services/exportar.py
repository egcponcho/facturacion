"""Factura comercial, packing list y reportes en Excel, con la misma estructura
que el PDF (documentos.py): cabecera, partes, condiciones, detalle, totales,
total en letras y declaración. Listos para imprimir en una página de ancho."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .documentos import _fecha, _por_unidad_txt
from . import visibilidad
from .preferencias import fecha_hora_txt

ACENTO = "5B3FD1"
TENUE = "5B6475"
_FONDO = PatternFill("solid", fgColor="F3F1FB")
_FONDO_2 = PatternFill("solid", fgColor="F7F8FA")
_LINEA = Side(style="thin", color="D9DCE3")
_FUERTE = Side(style="medium", color=ACENTO)
_ARRIBA = Alignment(vertical="top", wrap_text=True)
_MONEDA = "#,##0.00"
_ENTERO = "#,##0"
_DECIMAL = "#,##0.000"


class _Hoja:
    """Escribe de arriba abajo sobre una hoja de `columnas` columnas."""

    def __init__(self, ws, anchos: list[float]):
        self.ws = ws
        self.n = len(anchos)
        self.fila = 1
        for i, a in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = a
        ws.sheet_view.showGridLines = False

    def celda(self, fila, col, valor=None, negrita=False, color=None, tam=None, formato=None, fondo=None,
              alinear=None):
        c = self.ws.cell(row=fila, column=col, value=valor)
        c.font = Font(bold=negrita, color=color, size=tam or 9)
        if formato == _ENTERO and isinstance(valor, float) and valor != int(valor):
            formato = _DECIMAL  # una cantidad medida (kg, litros, metros) lleva sus decimales
        c.alignment = alinear or _ARRIBA
        if formato:
            c.number_format = formato
        if fondo:
            c.fill = fondo
        return c

    def unir(self, fila, c1, c2, filas=1):
        if c2 > c1 or filas > 1:
            self.ws.merge_cells(start_row=fila, start_column=c1, end_row=fila + filas - 1, end_column=c2)

    def marco(self, f1, f2, c1, c2, fondo=None, borde=_LINEA):
        for r in range(f1, f2 + 1):
            for c in range(c1, c2 + 1):
                x = self.ws.cell(row=r, column=c)
                x.border = Border(left=borde if c == c1 else None, right=borde if c == c2 else None,
                                  top=borde if r == f1 else None, bottom=borde if r == f2 else None)
                if fondo:
                    x.fill = fondo

    def espacio(self, n=1):
        self.fila += n

    def cabecera(self, d: dict, titulo: str, subtitulo: str, datos: list[tuple[str, str]]):
        """Exportador a la izquierda; título, número y fecha a la derecha."""
        exp = d["exportador"]
        corte = max(2, self.n - 3)
        self.celda(1, 1, exp["nombre"], negrita=True, tam=15)
        self.unir(1, 1, corte)
        self.celda(2, 1, " · ".join(x for x in (exp.get("direccion"), exp.get("pais")) if x), color=TENUE, tam=8)
        self.unir(2, 1, corte)
        self.celda(3, 1, " · ".join(x for x in (f"Tax ID: {exp['id_fiscal']}" if exp.get("id_fiscal") else None,
                                                  exp.get("correos"), exp.get("telefono")) if x), color=TENUE, tam=8)
        self.unir(3, 1, corte)
        self.celda(1, corte + 1, titulo, negrita=True, color=ACENTO, tam=13)
        self.unir(1, corte + 1, self.n)
        self.celda(2, corte + 1, subtitulo, color=TENUE, tam=8)
        self.unir(2, corte + 1, self.n)
        fila = 3
        for k, v in datos:
            self.celda(fila, corte + 1, k, color=TENUE, tam=8)
            self.celda(fila, corte + 2, v, negrita=True)
            self.unir(fila, corte + 2, self.n)
            fila += 1
        self.marco(1, fila - 1, corte + 1, self.n, _FONDO, Side(style="thin", color=ACENTO))
        self.fila = fila + 1

    def partes(self, d: dict):
        """Exportador, importador y consignatario en tres bloques."""
        exp = d["exportador"]
        bloques = [
            ("EXPORTER / SELLER", [f"{exp['codigo']} · {exp['nombre']}", f"Tax ID: {exp['id_fiscal'] or '—'}",
                                       " · ".join(x for x in (exp.get("direccion"), exp.get("pais")) if x),
                                       " · ".join(x for x in (exp.get("contacto"), exp.get("correos"),
                                                               exp.get("telefono")) if x)]),
            ("IMPORTER / BILL TO", _lineas_parte(d["importador"])),
            ("CONSIGNEE / NOTIFY PARTY", _lineas_parte(d["consignatario"])),
        ]
        tercio = self.n // 3
        rangos = [(1, tercio), (tercio + 1, 2 * tercio), (2 * tercio + 1, self.n)]
        f = self.fila
        alto = max(len(b[1]) for b in bloques)
        for (titulo, lineas), (c1, c2) in zip(bloques, rangos):
            self.celda(f, c1, titulo, negrita=True, color=ACENTO, tam=7)
            self.unir(f, c1, c2)
            for i in range(alto):
                v = lineas[i] if i < len(lineas) else None
                self.celda(f + 1 + i, c1, v, negrita=i == 0, color=None if i == 0 else TENUE, tam=9 if i == 0 else 8)
                self.unir(f + 1 + i, c1, c2)
            self.marco(f, f + alto, c1, c2)
        for i in range(alto):
            largo = max(len(b[1][i]) if i < len(b[1]) else 0 for b in bloques)
            if largo > 38:
                self.ws.row_dimensions[f + 1 + i].height = 24
        self.fila = f + alto + 2

    def rejilla(self, pares: list[tuple[str, str]], columnas: int = 4):
        """Pares etiqueta/valor en `columnas` bloques por fila."""
        ancho = self.n // columnas
        for i in range(0, len(pares), columnas):
            f = self.fila
            for j, (k, v) in enumerate(pares[i:i + columnas]):
                c1 = 1 + j * ancho
                c2 = self.n if j == columnas - 1 else c1 + ancho - 1
                self.celda(f, c1, k.upper(), negrita=True, color=ACENTO, tam=7)
                self.unir(f, c1, c2)
                self.celda(f + 1, c1, v if v not in (None, "") else "—")
                self.unir(f + 1, c1, c2)
                self.marco(f, f + 1, c1, c2, _FONDO_2)
            self.fila += 2
        self.fila += 1

    def tabla(self, encabezados: list[tuple[str, str | None]], filas: list[list], pie: list | None = None,
              texto: set[int] = frozenset()):
        """Detalle con encabezado fijo, filas cebra y fila de totales.
        `encabezados`: (título, formato numérico o None). `texto`: columnas
        (desde 0) que van como texto para no perder ceros a la izquierda."""
        f = self.fila
        for i, (t, fmt) in enumerate(encabezados, start=1):
            c = self.celda(f, i, t, negrita=True, color=TENUE, tam=8, fondo=_FONDO,
                           alinear=Alignment(wrap_text=True, vertical="center",
                                             horizontal="right" if fmt else "left"))
            c.border = Border(bottom=_FUERTE)
        self.ws.row_dimensions[f].height = 26
        inicio = f + 1
        for n, valores in enumerate(filas):
            r = inicio + n
            for i, (v, (_, fmt)) in enumerate(zip(valores, encabezados), start=1):
                c = self.celda(r, i, v, formato=fmt if fmt else ("@" if (i - 1) in texto else None),
                               fondo=_FONDO_2 if n % 2 else None)
                c.border = Border(bottom=_LINEA)
        r = inicio + len(filas)
        if pie:
            for i, (v, (_, fmt)) in enumerate(zip(pie, encabezados), start=1):
                c = self.celda(r, i, v, negrita=True, formato=fmt, fondo=_FONDO)
                c.border = Border(top=Side(style="medium", color="1F2430"))
            r += 1
        self.fila = r + 1
        return f

    def parrafo(self, texto: str, negrita=False, color=None, tam=None, alto=None, hasta=None):
        self.celda(self.fila, 1, texto, negrita=negrita, color=color, tam=tam)
        self.unir(self.fila, 1, hasta or self.n)
        if alto:
            self.ws.row_dimensions[self.fila].height = alto
        self.fila += 1

    def firma(self, texto: str):
        self.espacio()
        mitad = max(2, int(self.n * 0.6))
        self.celda(self.fila, 1, texto, color=TENUE, tam=8)
        self.unir(self.fila, 1, mitad, filas=2)
        self.ws.row_dimensions[self.fila].height = 24
        for col in range(mitad + 2, self.n + 1):
            self.ws.cell(row=self.fila + 1, column=col).border = Border(bottom=Side(style="thin", color="1F2430"))
        self.celda(self.fila + 2, mitad + 2, "Name, title and signature of the exporter", color=TENUE, tam=8)
        self.unir(self.fila + 2, mitad + 2, self.n)
        self.fila += 4

    def imprimir(self, titulo_filas: int | None, horizontal: bool, pie: str, borrador: bool):
        ws = self.ws
        ws.page_setup.orientation = "landscape" if horizontal else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.6
        if titulo_filas:
            ws.print_title_rows = f"{titulo_filas}:{titulo_filas}"
        ws.oddFooter.left.text = pie + ("  ·  DRAFT, NOT OFFICIAL" if borrador else "")
        ws.oddFooter.left.size = 7
        ws.oddFooter.right.text = "Page &P of &N"
        ws.oddFooter.right.size = 7


def _lineas_parte(p: dict | None) -> list[str]:
    if not p:
        return ["—"]
    lineas = [f"{p.get('codigo') or ''} · {p.get('razon_social') or p.get('nombre') or ''}"]
    if p.get("id_fiscal"):
        lineas.append(f"Tax ID (NIT / RUC): {p['id_fiscal']}")
    lineas.append(" · ".join(x for x in (p.get("direccion"), p.get("pais")) if x) or "—")
    if p.get("puerto"):
        lineas.append(f"Port of arrival: {p.get('puerto_nombre') or p['puerto']}")
    c = (p.get("contactos") or [None])[0]
    if c:
        lineas.append(" · ".join(x for x in (c.get("nombre"), c.get("correos"), c.get("telefono")) if x))
    elif p.get("correos"):
        lineas.append(p["correos"])
    return lineas


def _guardar(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hoja_titulo(texto: str) -> str:
    for ch in "[]:*?/\\":
        texto = texto.replace(ch, "-")
    return texto[:31] or "Sheet"


def exportar_factura(d: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _hoja_titulo(f"Invoice {d['numero']}")
    h = _Hoja(ws, [13, 7, 14, 15, 16, 10, 7, 30, 9, 11, 10, 7, 12, 14])
    tr = d["transporte"] or {}
    h.cabecera(d, "COMMERCIAL INVOICE", "Customs invoice · original", [
        ("No.", d["numero"]), ("Date", _fecha(d["fecha"])),
        ("Status", "Official" if d["oficial"] else "Draft")])
    h.partes(d)
    h.rejilla([
        ("Incoterm", d["incoterm"]), ("Currency", d["moneda"]), ("Payment terms", d["condiciones"]),
        ("Mode of transport", tr.get("modo")),
        ("Country of origin", d["pais_origen"]), ("Country of shipment", d["pais_procedencia"]),
        ("Country of destination", d["pais_destino"]), ("Carrier", tr.get("transportista")),
        ("Port of loading", tr.get("puerto_origen") or d["puerto_embarque"]),
        ("Port of discharge", tr.get("puerto_destino") or d["consignatario"].get("puerto_nombre")),
        ("B/L / AWB / waybill", tr.get("documento")), ("Packing lists", ", ".join(d["pls"]) or "—"),
    ])
    t = d["totales"]
    filas = [[l["oc"], l["posicion"], l["sku"], l["upc"], l["marca"], l["estilo"], l["talla"],
              f"{l['descripcion'] or ''} · {l['color'] or ''}" + (f" · prepack {l['prepack']}" if l["prepack"] else ""),
              l["origen"], l["partida"], l["cantidad"], l["unidad"], l["precio"], l["total"]] for l in d["lineas"]]
    fila_cab = h.tabla(
        [("PO", None), ("Line", None), ("Item code", None), ("UPC", None), ("Brand", None), ("Style", None),
         ("Size", None), ("Customs description · color", None), ("Origin", None), ("HS code (SAC)", None),
         ("Quantity", _ENTERO), ("UoM", None), ("Unit price", "#,##0.0000"), (f"Amount {d['moneda']}", _MONEDA)],
        filas, pie=["Total", None, f"{len(filas)} lines", None, None, None, None, None, None, None,
                    sum(l["cantidad"] for l in d["lineas"]), None, None, t["importe"]],
        texto={0, 1, 2, 3, 9})
    h.rejilla([
        ("Total quantity", _por_unidad_txt(t["por_unidad"])),
        ("Packages", f"{t['bultos']:,} cartons" + (f" on {t['pallets']} pallets" if t["pallets"] else "")),
        ("Net weight", f"{t['peso_neto']:,.2f} kg"), ("Gross weight", f"{t['peso_bruto']:,.2f} kg"),
        ("Volume", f"{t['cbm']:,.3f} m³"), ("Containers / AWB", ", ".join(tr.get("unidades", [])) or "—"),
        ("Total value " + (d["incoterm"] or ""), f"{d['moneda']} {t['importe']:,.2f}"),
        ("Purchase orders", f"{len(d['ocs'])} (see lines)"),
    ])
    h.parrafo(f"SAY: {d['total_letras']}", negrita=True)
    if d.get("observaciones"):
        h.parrafo(f"Remarks: {d['observaciones']}", color=TENUE, tam=8)
    h.firma("We declare under oath that the information in this invoice is true and correct, that the value is the "
            "price actually paid or payable for the goods and that the declared origin is correct.")
    h.imprimir(fila_cab, False, f"Commercial invoice {d['numero']} · {d['exportador']['nombre']}", not d["oficial"])
    ws.freeze_panes = None
    return _guardar(wb)


def exportar_pl(d: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _hoja_titulo(f"{d['numero']} {d['numero_pl']}")
    h = _Hoja(ws, [8, 6, 12, 6, 14, 15, 14, 10, 6, 8, 10, 8, 6, 8, 8, 8, 9, 9, 10, 10, 8, 7, 11, 18])
    tr = d["transporte"] or {}
    tp = d["totales_pl"]
    h.cabecera(d, "PACKING LIST", "Detailed carton list", [
        ("No.", f"{d['numero']} · {d['numero_pl']}"), ("Date", _fecha(d["fecha"])), ("Invoice", d["numero"]),
        ("Status", "Official" if d["oficial"] else "Draft")])
    h.partes(d)
    h.rejilla([
        ("Mode of transport", tr.get("modo")),
        ("Carrier", tr.get("transportista")), ("B/L / AWB / waybill", tr.get("documento")),
        ("Container / AWB", ", ".join(tr.get("unidades", [])) or "—"),
        ("Port of loading", tr.get("puerto_origen") or d["puerto_embarque"]),
        ("Port of discharge", tr.get("puerto_destino")), ("Country of origin", d["pais_origen"]),
        ("Country of destination", d["pais_destino"]),
        ("Final destination", d["destinos"][0] if len(d["destinos"]) == 1 else "Per carton (see lines)"),
        ("ETD / ETA", f"{_fecha(tr.get('etd'))} / {_fecha(tr.get('eta'))}"),
    ], columnas=5)
    filas = []
    for g in d["grupos"]:
        for i, it in enumerate(g["items"]):
            p = i == 0
            filas.append([
                g["rango"] if p else None, g["num_cajas"] if p else None, it["oc"], it["posicion"], it["sku"],
                it["upc"], it["marca"], it["descripcion"], it["estilo"], it["color"], it["talla"], it["inners"], it["inner_pack"],
                it["por_caja"], it["total"], it["unidad"], g["largo"] if p else None, g["ancho"] if p else None,
                g["alto"] if p else None, g["neto_caja"] if p else None, g["bruto_caja"] if p else None,
                g["neto_total"] if p else None, g["bruto_total"] if p else None, g["cbm"] if p else None,
                (g["pallet"] if g["pallet"] else None) if p else None, g["etiqueta"] if p else None,
                (", ".join(g["ocs"]) + (f" → {g['centro_destino']}" if g["centro_destino"] else "")) if p else None,
            ])
    fila_cab = h.tabla(
        [("Cartons", None), ("Qty", _ENTERO), ("PO", None), ("Line", None), ("Item code", None), ("UPC", None),
         ("Brand", None), ("Customs description", None), ("Style", None), ("Color", None), ("Size", None),
         ("Inner packs per carton", _ENTERO), ("Per inner pack", _ENTERO), ("Total per carton", _ENTERO), ("Total", _ENTERO), ("UoM", None), ("Length cm", "0.0"), ("Width cm", "0.0"), ("Height cm", "0.0"),
         ("Net/ctn kg", "0.00"), ("Gross/ctn kg", "0.00"), ("Net total kg", _MONEDA), ("Gross total kg", _MONEDA),
         ("m³", "0.000"), ("Pallet", None), ("Label", None), ("POs in carton → destination", None)],
        filas, pie=["Total", d["total_cajas"], None, None, None, None, None, None, None, None, None, None, None, None,
                    sum(it["total"] for g in d["grupos"] for it in g["items"]), None, None, None, None, None, None,
                    tp["peso_neto"], tp["peso_bruto"], tp["cbm"], None, None, None],
        texto={0, 2, 3, 4, 5})
    if d["pallets"]:
        h.parrafo("PALLETS", negrita=True, color=ACENTO, tam=8)
        h.tabla([("Pallet", None), ("Cartons", _ENTERO), ("Dimensions cm", None), ("Tare kg", "0.0"), ("Gross kg", "0.00"),
                 ("m³", "0.000")],
                [[p["numero"], p["cajas"], p["medidas"], p["tara"], p["bruto"], p["cbm"]] for p in d["pallets"]])
    if d["sin_caja"]:
        h.parrafo("NOT YET PACKED", negrita=True, color=ACENTO, tam=8)
        h.tabla([("PO", None), ("Item code", None), ("Style", None), ("Size", None), ("Quantity", _ENTERO),
                 ("UM", None)],
                [[x["oc"], x["sku"], x["estilo"], x["talla"], x["cantidad"], x["unidad"]] for x in d["sin_caja"]],
                texto={0, 1})
    h.rejilla([
        ("Total packages", f"{d['total_cajas']:,} cartons" + (f" on {tp['pallets']:,} pallets" if tp["pallets"] else "")),
        ("Quantity", tp["por_unidad_txt"]), ("Net weight", f"{tp['peso_neto']:,.2f} kg"),
        ("Gross weight", f"{tp['peso_bruto']:,.2f} kg"), ("Volume", f"{tp['cbm']:,.3f} m³"),
    ], columnas=5)
    h.parrafo(f"TOTAL PACKAGES: {d['total_bultos_letras']}", negrita=True)
    consig = d["consignatario"]
    h.espacio()
    mitad = h.n // 2
    inicio = h.fila
    h.parrafo("SHIPPING MARKS", negrita=True, color=ACENTO, tam=7, hasta=mitad)
    for i, linea in enumerate((consig.get("nombre"), consig.get("direccion"),
                               f"Invoice {d['numero']} · PO per carton label",
                               f"Carton no. __ of {d['total_cajas']} · Made in {d['pais_origen']}")):
        h.parrafo(linea, negrita=i == 0, hasta=mitad)
    h.marco(inicio, h.fila - 1, 1, mitad)
    h.firma("We declare that the contents, numbering, dimensions and weights of the packages correspond to the goods "
            "shipped. Every unit or pair carries its individual label; inner packs carry an inner pack label with "
            "the product and the quantity inside.")
    h.imprimir(fila_cab, True, f"Packing list {d['numero']} {d['numero_pl']} · {d['exportador']['nombre']}",
               not d["oficial"])
    return _guardar(wb)


def exportar_reporte(titulo: str, subtitulo: str, filtros: str, indicadores: list[tuple[str, str]],
                     columnas: list[tuple[str, float, bool]], filas: list[list], hojas: list[dict] | None = None) -> bytes:
    """Reporte tabular: título, filtros aplicados, indicadores y detalle con
    filtros de Excel. `hojas` agrega hojas de detalle adicionales con
    {"titulo", "columnas", "filas"}. Sin las columnas que el rol no ve."""
    indicadores, columnas, filas, hojas = visibilidad.reporte(indicadores, columnas, filas, hojas)
    wb = Workbook()
    ws = wb.active
    ws.title = _hoja_titulo(titulo)
    _hoja_reporte(ws, titulo, subtitulo, filtros, indicadores, columnas, filas)
    for x in hojas or []:
        _hoja_reporte(wb.create_sheet(_hoja_titulo(x["titulo"])), x["titulo"], subtitulo, filtros, [],
                      x["columnas"], x["filas"])
    return _guardar(wb)


def _hoja_reporte(ws, titulo, subtitulo, filtros, indicadores, columnas, filas):
    h = _Hoja(ws, [max(8, min(40, w * 9)) for _, w, _ in columnas])
    h.celda(1, 1, titulo, negrita=True, tam=15)
    h.unir(1, 1, h.n)
    h.celda(2, 1, subtitulo, color=TENUE, tam=8)
    h.unir(2, 1, h.n)
    h.celda(3, 1, f"Generated {fecha_hora_txt(datetime.now())} · {filtros or 'No filters'}", color=TENUE, tam=8)
    h.unir(3, 1, h.n)
    for c in range(1, h.n + 1):
        ws.cell(row=3, column=c).border = Border(bottom=_FUERTE)
    h.fila = 5
    if indicadores:
        h.rejilla(indicadores, columnas=min(len(indicadores), max(1, h.n), 5))
    fmts = [(t, (_ENTERO if der else None)) for t, _, der in columnas]
    for i, (t, _, der) in enumerate(columnas):
        if der and any(isinstance(f[i], float) and not float(f[i]).is_integer() for f in filas):
            fmts[i] = (t, _MONEDA)
    fila_cab = h.tabla(fmts, filas, texto={i for i, (_, _, der) in enumerate(columnas) if not der})
    if filas:
        ws.auto_filter.ref = f"A{fila_cab}:{get_column_letter(h.n)}{fila_cab + len(filas)}"
    ws.freeze_panes = f"A{fila_cab + 1}"
    h.imprimir(fila_cab, True, titulo, False)


def exportar_ficha(d: dict, x: dict) -> bytes:
    """Ficha técnica en Excel: datos y clasificación en la primera hoja; la
    composición, los códigos por país y las tallas en las siguientes."""
    titulo = f"Technical sheet {d['estilo']} · {d['color']}"
    sub = " · ".join(str(v) for v in (d.get("nombre"), d.get("proveedor"), d.get("marca_nombre") or d.get("marca")) if v)
    estado = f"Version {d.get('version_ficha') or 1} · {d.get('estado_txt') or d.get('estado') or ''}"
    filas = [["Classification", k, str(v)] for k, v in x["clasificacion"]]
    if x["descripcion"]:
        filas.append(["Classification", "Customs description", x["descripcion"]])
    if x.get("descripcion_comercial"):
        filas.append(["Classification", "Commercial description", x["descripcion_comercial"]])
    filas += [["Product data", k, str(v)] for k, v in x["datos"]]
    filas += [["Reasoning", "", r] for r in x["razones"]]
    hojas = [
        {"titulo": "Composition", "columnas": [("Part", 1.6, False), ("Materials", 5, False)], "filas": x["composicion"]},
        {"titulo": "National codes", "columnas": [("Country", 0.8, False), ("Code", 1.6, False), ("Duty (DAI)", 0.9, False),
                                                  ("Status", 0.9, False), ("Source", 0.9, False)], "filas": x["partidas"]},
        {"titulo": "Sizes", "columnas": [("SKU", 1.4, False), ("UPC", 1.4, False), ("Size", 0.7, False), ("Unit", 0.6, False),
                                         ("Description", 3, False)], "filas": x["tallas"]},
    ]
    return exportar_reporte(titulo, sub, estado, [], [("Section", 1.4, False), ("Field", 1.8, False), ("Value", 6, False)],
                            filas, hojas)
