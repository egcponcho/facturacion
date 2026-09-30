import csv
import io
import re
import unicodedata
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from ..config import settings
from ..models import (
    Alerta,
    Almacen,
    Articulo,
    Centro,
    Factura,
    FacturaLinea,
    ImportacionOC,
    OrdenCompra,
    Pais,
    PosicionOC,
    Proveedor,
    Puerto,
    Sociedad,
    Usuario,
    ahora,
)
from .cantidades import facturado_por_posicion, facturas_por_posicion
from .productos import clasificacion_txt, pais_de_centro, partida_para, producto_de
from .common import ErrorNegocio, asegurar_proveedor, exigir, proveedor_filtro, registrar

# Dos liberaciones de dos equipos distintos:
# - Comercial: P (pendiente) o C (liberada; si viene vacío también es C).
# - Logística: 304 no liberada, 300 liberada, 301 liberada con cambios posteriores.
# Sin liberación comercial no hay liberación logística. Solo se factura con C y 300/301.
COMERCIAL_TXT = {"C": "Released by commercial", "P": "Pending commercial"}
LIBERACION_TXT = {
    "300": "Released by logistics",
    "301": "Released with later changes",
    "304": "Not released by logistics",
}


def liberacion_logistica(comercial: str, explicita: str | None, actual: str | None, hubo_cambios: bool) -> str:
    """Código logístico final. Sin comercial (P) siempre es 304. Si el archivo
    trae el código, manda el archivo; si no, se conserva el que tenía (304 si
    es nueva) y una OC ya liberada (300) que cambia pasa a 301."""
    if comercial != "C":
        return "304"
    if explicita in LIBERACION_TXT:
        return explicita
    if actual == "300" and hubo_cambios:
        return "301"
    return actual or "304"


def esta_liberada(comercial: str, logistica: str) -> bool:
    return comercial == "C" and logistica in ("300", "301")


# ---- Vista general de OCs ---------------------------------------------------
ORDEN_OC = {"numero", "fecha", "fecha_xf", "fecha_tienda", "importe", "avance", "proveedor"}


def listar_ordenes(
    db: Session,
    user: Usuario,
    proveedor_id: int | None = None,
    q: str | None = None,
    centro: str | None = None,
    solo_disponible: bool = True,
    page: int = 1,
    size: int = 25,
    sociedad: str | None = None,
    marca: str | None = None,
    liberacion: str | None = None,
    destino: str | None = None,
    puerto: str | None = None,
    orden: str | None = None,
    almacen: str | None = None,
    comercial: str | None = None,
) -> dict:
    exigir(user, "oc.ver")
    prov = proveedor_filtro(user, proveedor_id)

    tot = (
        select(
            PosicionOC.oc_id.label("oc_id"),
            func.sum(PosicionOC.cantidad).label("cantidad"),
            func.sum(PosicionOC.cantidad * PosicionOC.precio).label("importe"),
            func.count(PosicionOC.id).label("n"),
        )
        .group_by(PosicionOC.oc_id)
        .subquery()
    )
    fac = (
        select(
            PosicionOC.oc_id.label("oc_id"),
            func.sum(FacturaLinea.cantidad).label("facturado"),
            func.sum(FacturaLinea.cantidad * PosicionOC.precio).label("importe_facturado"),
        )
        .join(FacturaLinea, FacturaLinea.posicion_oc_id == PosicionOC.id)
        .join(Factura, Factura.id == FacturaLinea.factura_id)
        .where(Factura.estado != "CANCELADA")
        .group_by(PosicionOC.oc_id)
        .subquery()
    )
    facturado = func.coalesce(fac.c.facturado, 0)
    importe_fact = func.coalesce(fac.c.importe_facturado, 0)
    consulta = (
        select(OrdenCompra, Proveedor.nombre, tot.c.importe, tot.c.n, importe_fact)
        .join(Proveedor, Proveedor.id == OrdenCompra.proveedor_id)
        .join(tot, tot.c.oc_id == OrdenCompra.id)
        .outerjoin(fac, fac.c.oc_id == OrdenCompra.id)
    )
    if prov:
        consulta = consulta.where(OrdenCompra.proveedor_id == prov)
    if q:
        patron = f"%{q.strip()}%"
        sub = select(PosicionOC.oc_id).where(
            or_(
                PosicionOC.estilo.ilike(patron),
                PosicionOC.codigo_sap.ilike(patron),
                PosicionOC.upc.ilike(patron),
                PosicionOC.color.ilike(patron),
            )
        )
        consulta = consulta.where(or_(OrdenCompra.numero.ilike(patron), OrdenCompra.id.in_(sub)))
    if centro:
        consulta = consulta.where(OrdenCompra.centro == centro)
    if sociedad:
        consulta = consulta.where(OrdenCompra.sociedad == sociedad)
    if destino:
        consulta = consulta.where(OrdenCompra.centro_destino == destino)
    if puerto:
        consulta = consulta.where(OrdenCompra.puerto_despacho == puerto)
    if liberacion:
        consulta = consulta.where(OrdenCompra.liberacion_logistica == liberacion)
    if comercial:
        consulta = consulta.where(OrdenCompra.liberacion_comercial == comercial)
    if marca:
        consulta = consulta.where(OrdenCompra.id.in_(select(PosicionOC.oc_id).where(PosicionOC.marca == marca)))
    if almacen:
        consulta = consulta.where(OrdenCompra.id.in_(select(PosicionOC.oc_id).where(PosicionOC.almacen == almacen)))
    if solo_disponible:
        consulta = consulta.where(tot.c.cantidad - facturado > 0)

    col, _, direccion = (orden or "").partition(":")
    expr = {
        "numero": OrdenCompra.numero, "fecha": OrdenCompra.fecha, "fecha_xf": OrdenCompra.fecha_xf,
        "fecha_tienda": OrdenCompra.fecha_tienda, "importe": tot.c.importe, "proveedor": Proveedor.nombre,
        "avance": importe_fact * 1.0 / func.nullif(tot.c.importe, 0),
    }.get(col)
    if expr is not None:
        consulta = consulta.order_by(expr.desc().nullslast() if direccion == "desc" else expr.asc().nullslast(),
                                     OrdenCompra.numero)
    else:
        consulta = consulta.order_by(OrdenCompra.fecha.desc().nullslast(), OrdenCompra.numero)

    total = db.scalar(select(func.count()).select_from(consulta.subquery())) or 0
    filas = db.execute(consulta.offset((page - 1) * size).limit(size)).all()

    # Cantidades por unidad de medida: nunca se suman pares con unidades
    ids = [oc.id for oc, *_ in filas]
    por_unidad: dict[int, dict] = {i: {} for i in ids}
    marcas: dict[int, set] = {i: set() for i in ids}
    almacenes: dict[int, set] = {i: set() for i in ids}
    if ids:
        for oc_id, alm in db.execute(
            select(PosicionOC.oc_id, PosicionOC.almacen).where(PosicionOC.oc_id.in_(ids)).distinct()
        ).all():
            if alm:
                almacenes[oc_id].add(alm)
        for oc_id, unidad, marca_oc, cant in db.execute(
            select(PosicionOC.oc_id, PosicionOC.unidad, PosicionOC.marca, func.sum(PosicionOC.cantidad))
            .where(PosicionOC.oc_id.in_(ids)).group_by(PosicionOC.oc_id, PosicionOC.unidad, PosicionOC.marca)
        ).all():
            d = por_unidad[oc_id].setdefault(unidad, {"cantidad": 0, "facturado": 0, "disponible": 0})
            d["cantidad"] += int(cant or 0)
            d["disponible"] += int(cant or 0)
            if marca_oc:
                marcas[oc_id].add(marca_oc)
        for oc_id, unidad, cant in db.execute(
            select(PosicionOC.oc_id, PosicionOC.unidad, func.sum(FacturaLinea.cantidad))
            .join(FacturaLinea, FacturaLinea.posicion_oc_id == PosicionOC.id)
            .join(Factura, Factura.id == FacturaLinea.factura_id)
            .where(Factura.estado != "CANCELADA", PosicionOC.oc_id.in_(ids))
            .group_by(PosicionOC.oc_id, PosicionOC.unidad)
        ).all():
            d = por_unidad[oc_id][unidad]
            d["facturado"] = int(cant or 0)
            d["disponible"] = max(d["cantidad"] - d["facturado"], 0)

    hoy = date.today()
    centros = {c.codigo: c for c in db.scalars(select(Centro))}
    items = []
    for oc, prov_nombre, importe, n, importe_f in filas:
        destino = centros.get(oc.centro_destino)
        importe = float(importe or 0)
        items.append(
            {
                **_cabecera_oc(oc),
                "proveedor": prov_nombre,
                "posiciones": n,
                "marcas": sorted(marcas[oc.id]),
                # Una OC puede repartir sus posiciones entre varios almacenes
                "almacenes": sorted(almacenes[oc.id]),
                "por_unidad": por_unidad[oc.id],
                "importe": round(importe, 2),
                "importe_facturado": round(float(importe_f or 0), 2),
                # Avance por valor: es comparable aunque la OC mezcle pares y unidades
                "avance": round(float(importe_f or 0) * 100 / importe, 1) if importe else 0,
                "dias_tienda": (oc.fecha_tienda - hoy).days if oc.fecha_tienda else None,
                # El centro de destino dice a qué país llega al final
                "destino_nombre": destino.nombre if destino else None,
                "pais_destino": destino.pais if destino else None,
            }
        )
    return {"items": items, "total": total, "page": page, "size": size}


def filtros_ordenes(db: Session, user: Usuario, proveedor_id: int | None = None) -> dict:
    """Valores presentes en las OCs visibles, para armar filtros dinámicos."""
    exigir(user, "oc.ver")
    prov = proveedor_filtro(user, proveedor_id)
    base = select(OrdenCompra)
    if prov:
        base = base.where(OrdenCompra.proveedor_id == prov)
    sub = base.subquery()

    def distintos(col):
        return sorted({v for (v,) in db.execute(select(col).select_from(sub).distinct()) if v})

    def de_posiciones(col):
        return sorted({v for (v,) in db.execute(
            select(col).where(PosicionOC.oc_id.in_(select(sub.c.id))).distinct()) if v})

    marcas = de_posiciones(PosicionOC.marca)
    return {
        "sociedades": distintos(sub.c.sociedad),
        "centros": distintos(sub.c.centro),
        "almacenes": de_posiciones(PosicionOC.almacen),
        "destinos": [{"codigo": d.codigo, "nombre": f"{d.nombre} ({d.pais})"} for d in db.scalars(
            select(Centro).where(Centro.codigo.in_(distintos(sub.c.centro_destino))).order_by(Centro.codigo))],
        "puertos": [{"codigo": p.codigo, "nombre": p.nombre} for p in db.scalars(
            select(Puerto).where(Puerto.codigo.in_(distintos(sub.c.puerto_despacho))))],
        "marcas": marcas,
        "liberaciones": [{"codigo": k, "nombre": v} for k, v in LIBERACION_TXT.items()],
        "comerciales": [{"codigo": k, "nombre": v} for k, v in COMERCIAL_TXT.items()],
    }


def _cabecera_oc(oc: OrdenCompra) -> dict:
    return {
        "id": oc.id,
        "numero": oc.numero,
        "proveedor_id": oc.proveedor_id,
        "sociedad": oc.sociedad,
        "centro": oc.centro,
        "centro_destino": oc.centro_destino,
        "moneda": oc.moneda,
        "incoterm": oc.incoterm,
        "fecha": oc.fecha,
        "puerto_despacho": oc.puerto_despacho,
        "pais_origen": oc.pais_origen,
        "pais_procedencia": oc.pais_procedencia,
        "fecha_xf_original": oc.fecha_xf_original,
        "fecha_xf": oc.fecha_xf,
        "fecha_tienda": oc.fecha_tienda,
        "liberacion_comercial": oc.liberacion_comercial,
        "liberacion_logistica": oc.liberacion_logistica,
        "liberacion_txt": LIBERACION_TXT.get(oc.liberacion_logistica),
        "comercial_txt": COMERCIAL_TXT.get(oc.liberacion_comercial),
        "liberada": oc.liberada,
    }


def posiciones_oc(db: Session, user: Usuario, oc_id: int) -> dict:
    exigir(user, "oc.ver")
    oc = db.get(OrdenCompra, oc_id)
    if not oc:
        raise ErrorNegocio("The purchase order does not exist.", 404, "no_encontrado")
    asegurar_proveedor(user, oc.proveedor_id)
    ids = [p.id for p in oc.posiciones]
    facturado = facturado_por_posicion(db, ids)
    facturas = facturas_por_posicion(db, ids)
    # La partida sale de la clasificación del producto, con el código del país destino
    pais_destino = pais_de_centro(db, oc.centro_destino)
    posiciones = []
    for p in oc.posiciones:
        prod = producto_de(p.articulo)
        fact = facturado.get(p.id, 0)
        fs = facturas.get(p.id, [])
        estado, motivo, restringida = estado_posicion(oc, p, fact, fs)
        posiciones.append(
            {
                "id": p.id,
                "posicion": p.posicion,
                "almacen": p.almacen,
                "codigo_sap": p.codigo_sap,
                "upc": p.upc,
                "estilo": p.estilo,
                "color": p.color,
                "talla": p.talla,
                "descripcion": p.descripcion,
                "marca": p.marca,
                "grupo": p.grupo,
                "categoria": p.categoria,
                "tipo_empaque": p.tipo_empaque,
                "casepack": p.casepack,
                "inner_pack": p.inner_pack,
                "prepack": p.prepack,
                "unidades_por_caja": p.unidades_por_caja,
                "cantidad": p.cantidad,
                "unidad": p.unidad,
                "precio": p.precio,
                "total": round(p.cantidad * p.precio, 2),
                "fecha_entrega": p.fecha_entrega,
                "pais_origen": p.pais_origen,
                "partida_arancelaria": partida_para(prod, pais_destino),
                "clasificacion": clasificacion_txt(prod),
                "facturado": fact,
                "disponible": max(p.cantidad - fact, 0),
                "facturas": fs,
                "estado": estado,
                "motivo": motivo,
                "solo_factura_id": restringida,
            }
        )
    prov = db.get(Proveedor, oc.proveedor_id)
    destino = db.scalar(select(Centro).where(Centro.codigo == oc.centro_destino))
    return {"oc": {**_cabecera_oc(oc), "proveedor": prov.nombre,
                   "centro_destino_nombre": destino.nombre if destino else None,
                   "pais_destino": destino.pais if destino else None},
            "posiciones": posiciones}


def empaque_posicion(db: Session, user: Usuario, oc_id: int, posicion_id: int, casepack: int | None,
                     inner_pack: int | None) -> dict:
    """Casepack e inner pack de una posición (condición de la compra). Se
    cambian mientras no haya nada facturado de esa posición."""
    exigir(user, "oc.empaque")
    p = db.get(PosicionOC, posicion_id)
    if not p or p.oc_id != oc_id:
        raise ErrorNegocio("The purchase order line does not exist.", 404, "no_encontrado")
    errores = validar_empaque(p.tipo_empaque, casepack, inner_pack, p.cantidad)
    if errores:
        raise ErrorNegocio("Check the packing of the line.", 422, "validacion",
                           [{"campo": "inner_pack", "mensaje": e} for e in errores])
    if facturado_por_posicion(db, [p.id]).get(p.id, 0):
        raise ErrorNegocio("This line is already invoiced: its packing can no longer change.", 409, "facturada")
    antes = {"casepack": p.casepack, "inner_pack": p.inner_pack}
    p.casepack, p.inner_pack = casepack, inner_pack
    registrar(db, user, "oc", p.oc_id, "empaque",
              {"posicion": p.posicion, **{k: [antes[k], v] for k, v in (("casepack", casepack),
                                                                           ("inner_pack", inner_pack))
                                          if antes[k] != v}})
    return {"id": p.id, "casepack": p.casepack, "inner_pack": p.inner_pack}


def estado_posicion(oc, p, facturado: int, facturas: list[dict]):
    """Devuelve (estado, motivo, factura a la que queda restringido el saldo)."""
    disponible = p.cantidad - facturado
    if oc.liberacion_comercial != "C":
        return "NO_DISPONIBLE", "No commercial release (P): logistics cannot release yet.", None
    if not oc.liberada:
        return "NO_DISPONIBLE", "No logistics release (304).", None
    if p.bloqueada:
        return "NO_DISPONIBLE", p.motivo_bloqueo or "Line blocked.", None
    if disponible <= 0:
        return "FACTURADA", None, None
    if facturado > 0:
        if not settings.POSICION_EN_VARIAS_FACTURAS and facturas:
            f = facturas[0]
            return "PARCIAL", f"The balance can only be added to {f['nombre']}.", f["id"]
        return "PARCIAL", None, None
    return "DISPONIBLE", None, None


# ---- Importación ------------------------------------------------------------
# El primer alias de cada campo es el nombre de la columna en la plantilla (inglés).
ALIAS = {
    "proveedor": ["supplier", "vendor", "supplier_code", "proveedor", "codigo_proveedor", "proveedor_codigo"],
    "oc": ["po", "po_number", "purchase_order", "oc", "orden", "orden_compra", "numero_oc"],
    "posicion": ["po_line", "line", "item", "posicion", "pos", "linea"],
    "sociedad": ["company", "company_code", "sociedad", "compania"],
    "centro": ["plant", "bonded_warehouse", "centro", "bodega", "bodega_fiscal"],
    "almacen": ["storage_location", "sloc", "almacen"],
    "centro_destino": ["destination", "destination_center", "destination_country", "centro_destino", "pais_destino",
                       "destino", "codigo_destino"],
    "moneda": ["currency", "moneda"],
    "incoterm": ["incoterm"],
    "fecha_oc": ["po_date", "fecha_oc", "fecha"],
    "puerto_despacho": ["port_of_loading", "pol", "port", "puerto_despacho", "puerto", "puerto_embarque"],
    "pais_origen": ["country_of_origin", "coo", "origin", "pais_origen", "origen"],
    "pais_procedencia": ["country_of_shipment", "provenance", "pais_procedencia", "procedencia"],
    "fecha_xf_original": ["xf_date_original", "original_xf_date", "fecha_xf_original", "xf_original"],
    "fecha_xf": ["xf_date", "xf", "updated_xf_date", "fecha_xf", "fecha_xf_actualizada", "xf_actualizada"],
    "fecha_tienda": ["in_store_date", "store_date", "fecha_tienda", "fecha_requerida_tienda", "fecha_requerida"],
    "liberacion_comercial": ["commercial_release", "liberacion_comercial", "lib_comercial", "liberada",
                             "estado_liberacion"],
    "liberacion_logistica": ["logistics_release", "liberacion_logistica", "lib_logistica"],
    "codigo_sap": ["sku", "sap_code", "material", "item_code", "codigo_sap", "sap", "codigo"],
    "cantidad": ["quantity", "qty", "cantidad"],
    "unidad": ["uom", "unit", "um", "unidad"],
    "precio": ["unit_price", "price", "precio", "precio_unitario"],
    "casepack": ["casepack", "case_pack"],
    "inner_pack": ["inner_pack", "inner", "innerpack", "inner_casepack", "pack"],
    "fecha_entrega": ["delivery_date", "fecha_entrega", "entrega"],
}
REQUERIDOS = ["proveedor", "oc", "posicion", "codigo_sap", "cantidad", "precio", "moneda", "sociedad", "centro",
              "centro_destino"]
CAMPOS_CABECERA = ["sociedad", "centro", "centro_destino", "moneda", "incoterm", "fecha", "puerto_despacho",
                   "pais_origen", "pais_procedencia", "fecha_xf_original", "fecha_xf", "fecha_tienda",
                   "liberacion_comercial"]
CAMPOS_POSICION = [
    "almacen", "articulo_id", "codigo_sap", "upc", "estilo", "color", "talla", "descripcion", "marca", "grupo", "categoria",
    "tipo_empaque", "casepack", "inner_pack", "prepack", "unidades_por_caja", "cantidad", "unidad", "precio", "fecha_entrega",
    "pais_origen",
]
UNIDADES = {"PAR": "PAR", "PR": "PAR", "PARES": "PAR", "PRS": "PAR",
            "UN": "UN", "UND": "UN", "UNIDAD": "UN", "UNIDADES": "UN", "EA": "UN", "PC": "UN", "PZA": "UN",
            "CJ": "CJ", "CAJA": "CJ", "CAJAS": "CJ", "CS": "CJ", "CTN": "CJ"}


def _norm(texto: str) -> str:
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode().lower()
    return "_".join("".join(c if c.isalnum() else " " for c in t).split())


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _leer_archivo(nombre: str, contenido: bytes) -> list[dict]:
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb.active
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            return []
        encabezados = [_texto(h) for h in filas[0]]
        datos = [[_texto(v) for v in fila] for fila in filas[1:]]
    else:
        texto = contenido.decode("utf-8-sig", errors="replace")
        muestra = texto[:2000]
        delimitador = ";" if muestra.count(";") > muestra.count(",") else ","
        lector = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
        if not lector:
            return []
        encabezados = lector[0]
        datos = [[v.strip() for v in fila] for fila in lector[1:]]

    mapa = {}
    for i, h in enumerate(encabezados):
        n = _norm(h)
        for campo, alias in ALIAS.items():
            if n in alias and campo not in mapa:
                mapa[campo] = i
    faltan = [c for c in REQUERIDOS if c not in mapa]
    if faltan:
        raise ErrorNegocio(
            "The file is missing required columns: " + ", ".join(ALIAS[c][0] for c in faltan) + ".",
            422,
            "columnas_faltantes",
            {"faltan": faltan, "encontradas": encabezados},
        )
    filas_dict = []
    for n_fila, fila in enumerate(datos, start=2):
        if not any(fila):
            continue
        registro = {campo: (fila[i] if i < len(fila) else "") for campo, i in mapa.items()}
        registro["_fila"] = n_fila
        filas_dict.append(registro)
    return filas_dict


def _fecha(valor: str) -> date | None:
    if not valor:
        return None
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y%m%d"):
        try:
            return datetime.strptime(valor[:10], formato).date()
        except ValueError:
            continue
    raise ValueError(f"invalid date: {valor}")


def _comercial(valor: str) -> str:
    v = (valor or "").strip().upper()
    if v in ("P", "PENDIENTE", "NO", "N", "0", "FALSE", "BLOQUEADA"):
        return "P"
    return "C"  # C, vacío o "sí": liberación comercial completa


class Maestros:
    """Catálogos precargados para validar un archivo completo sin consultas por fila."""

    def __init__(self, db: Session):
        self.proveedores = {p.codigo: p for p in db.scalars(select(Proveedor))}
        self.sociedades = {s.codigo: s for s in db.scalars(select(Sociedad))}
        self.centros = {c.codigo: c for c in db.scalars(select(Centro))}
        self.almacenes = {a.codigo: a for a in db.scalars(select(Almacen))}
        self.puertos = {p.codigo: p for p in db.scalars(select(Puerto))}
        self.paises = {p.codigo for p in db.scalars(select(Pais))}
        self.articulos: dict[str, Articulo] = {}
        self._db = db

    def articulo(self, sku: str) -> Articulo | None:
        if sku not in self.articulos:
            self.articulos[sku] = self._db.scalar(select(Articulo).where(Articulo.sku == sku))
        return self.articulos[sku]


PAIS_TXT = {"pais_origen": "country of origin", "pais_procedencia": "country of shipment"}


def _normalizar(registro: dict, m: Maestros | None = None) -> tuple[dict, list[str]]:
    """Convierte los textos de una fila en valores tipados y la valida contra
    los datos maestros. Los códigos se mantienen como texto."""
    errores = []
    r = {k: (v or "").strip() for k, v in registro.items() if k != "_fila"}
    for campo in REQUERIDOS:
        if not r.get(campo):
            errores.append(f"Falta {campo}.")
    d = {
        "proveedor": r.get("proveedor", "").upper(),
        "oc": r.get("oc", ""),
        "posicion": r.get("posicion", "").lstrip("0") or r.get("posicion", ""),
        "sociedad": r.get("sociedad", "").upper(),
        "centro": (r.get("centro") or "").upper() or None,
        "almacen": (r.get("almacen") or "").upper() or None,
        "centro_destino": r.get("centro_destino") or None,
        "moneda": (r.get("moneda") or "").upper(),
        "incoterm": (r.get("incoterm") or "").upper() or None,
        "puerto_despacho": (r.get("puerto_despacho") or "").upper() or None,
        "pais_origen": (r.get("pais_origen") or "").upper() or None,
        "pais_procedencia": (r.get("pais_procedencia") or r.get("pais_origen") or "").upper() or None,
        "liberacion_comercial": _comercial(r.get("liberacion_comercial", "")),
        "liberacion_logistica_archivo": (r.get("liberacion_logistica") or "").strip() or None,
        "codigo_sap": r.get("codigo_sap", ""),
    }
    log = d["liberacion_logistica_archivo"]
    if log and log not in LIBERACION_TXT:
        errores.append(f"Invalid logistics release {log} (use 304, 300 or 301).")
    elif log in ("300", "301") and d["liberacion_comercial"] != "C":
        errores.append(f"Logistics cannot release ({log}) without commercial release: the PO is in P.")
    if d["oc"] and not re.fullmatch(r"44\d{8}", d["oc"]):
        errores.append(f"PO {d['oc']} does not have the 44 + 8 digits format (for example 4400003856).")
    if d["posicion"] and (not d["posicion"].isdigit() or int(d["posicion"]) % 10):
        errores.append(f"Line {d['posicion']} is not a multiple of 10.")
    try:
        cant = float(r.get("cantidad", "").replace(",", "")) if r.get("cantidad") else None
        if cant is not None and (cant < 0 or not cant.is_integer()):
            raise ValueError
        d["cantidad"] = int(cant) if cant is not None else None
    except ValueError:
        errores.append(f"Invalid quantity: {r.get('cantidad')}.")
        d["cantidad"] = None
    try:
        d["precio"] = float(r.get("precio", "").replace(",", "")) if r.get("precio") else None
        if d["precio"] is not None and d["precio"] < 0:
            raise ValueError
    except ValueError:
        errores.append(f"Invalid price: {r.get('precio')}.")
        d["precio"] = None
    # Empaque de la compra: viene en la posición de la OC, no en el artículo
    for campo, texto in (("casepack", "Casepack"), ("inner_pack", "Inner pack")):
        try:
            d[campo] = int(float(r[campo])) if r.get(campo) else None
            if d[campo] is not None and d[campo] < 1:
                raise ValueError
        except ValueError:
            errores.append(f"Invalid {texto.lower()}: {r.get(campo)}.")
            d[campo] = None
    unidad = UNIDADES.get((r.get("unidad") or "").upper())
    if r.get("unidad") and not unidad:
        errores.append(f"Unknown unit: {r.get('unidad')} (use PAR, UN or CJ).")
    d["unidad"] = unidad
    for campo, destino in (("fecha_oc", "fecha"), ("fecha_entrega", "fecha_entrega"),
                           ("fecha_xf_original", "fecha_xf_original"), ("fecha_xf", "fecha_xf"),
                           ("fecha_tienda", "fecha_tienda")):
        try:
            d[destino] = _fecha(r.get(campo, ""))
        except ValueError as e:
            errores.append(str(e).capitalize() + ".")
            d[destino] = None
    d["fecha_xf_original"] = d["fecha_xf_original"] or d["fecha_xf"]
    d["fecha_xf"] = d["fecha_xf"] or d["fecha_xf_original"]

    if m is None:
        return d, errores
    # Datos maestros
    soc = m.sociedades.get(d["sociedad"])
    if d["sociedad"] and not soc:
        errores.append(f"La sociedad {d['sociedad']} no existe.")
    cen = m.centros.get(d["centro"]) if d["centro"] else None
    if d["centro"] and not cen:
        errores.append(f"Plant {d['centro']} does not exist.")
    elif cen and soc and cen.sociedad_id != soc.id:
        errores.append(f"Plant {d['centro']} does not belong to company {d['sociedad']}.")
    if d["almacen"]:
        alm = m.almacenes.get(d["almacen"])
        if not alm:
            errores.append(f"Storage location {d['almacen']} does not exist.")
        elif soc and alm.sociedad_id != soc.id:
            errores.append(f"Storage location {d['almacen']} does not belong to company {d['sociedad']}.")
    prov = m.proveedores.get(d["proveedor"])
    if prov and soc and prov.sociedades and soc.id not in {x.id for x in prov.sociedades}:
        errores.append(f"Supplier {prov.nombre} does not work with company {d['sociedad']}.")
    if d["centro_destino"] and d["centro_destino"] not in m.centros:
        errores.append(f"Destination center {d['centro_destino']} is not registered in master data.")
    if d["puerto_despacho"] and d["puerto_despacho"] not in m.puertos:
        errores.append(f"Port {d['puerto_despacho']} is not registered.")
    for campo in ("pais_origen", "pais_procedencia"):
        if d[campo] and d[campo] not in m.paises:
            errores.append(f"Country {d[campo]} ({PAIS_TXT[campo]}) is not registered.")
    art = m.articulo(d["codigo_sap"]) if d["codigo_sap"] else None
    if d["codigo_sap"] and not art:
        errores.append(f"SKU {d['codigo_sap']} does not exist in the item master. Load it first.")
    elif art:
        prov = m.proveedores.get(d["proveedor"])
        if prov and art.proveedor_id and art.proveedor_id != prov.id:
            errores.append(f"SKU {art.sku} belongs to another supplier, not {prov.nombre}.")
        if not art.activo:
            errores.append(f"SKU {art.sku} is inactive in the item master.")
        if d["unidad"] and d["unidad"] != art.unidad:
            errores.append(f"Unit {d['unidad']} does not match the item master ({art.unidad}).")
        d.update(
            articulo_id=art.id, upc=art.upc, estilo=art.estilo, color=art.color, talla=art.talla,
            descripcion=art.descripcion, marca=art.marca.codigo, grupo=art.grupo.codigo,
            categoria=art.grupo.categoria, tipo_empaque=art.tipo, unidad=art.unidad,
            prepack=art.prepack.codigo if art.prepack else None,
            unidades_por_caja=art.prepack.total if art.prepack else None,
        )
        errores += validar_empaque(art.tipo, d["casepack"], d["inner_pack"], d["cantidad"])
        d["pais_origen_pos"] = (producto_de(art).pais_origen if producto_de(art) else None) or d["pais_origen"]
    return d, errores


def validar_empaque(tipo: str, casepack: int | None, inner_pack: int | None, cantidad: int | None) -> list[str]:
    """Reglas del empaque de una posición:
    - un prepack ya es una caja definida (su curva): no lleva casepack ni inner pack;
    - con casepack e inner pack, el casepack es múltiplo del inner pack;
    - con inner pack, la cantidad de la posición es múltiplo del inner pack
      (todos los inner packs llevan la misma cantidad)."""
    errores = []
    if tipo == "PREPACK":
        if casepack or inner_pack:
            errores.append("A prepack is already a defined carton (its size run): it takes no casepack or inner pack.")
        return errores
    if casepack and inner_pack and casepack % inner_pack:
        errores.append(f"The casepack ({casepack}) must be a multiple of the inner pack ({inner_pack}).")
    if inner_pack and cantidad and cantidad % inner_pack:
        errores.append(f"The quantity ({cantidad}) must be a multiple of the inner pack ({inner_pack}): "
                       "all inner packs carry the same quantity.")
    return errores


def _valores_posicion(d: dict) -> dict:
    v = {c: d.get(c) for c in CAMPOS_POSICION}
    v["pais_origen"] = d.get("pais_origen_pos") or d.get("pais_origen")
    return v


def _clasificar(db: Session, filas: list[dict]) -> list[dict]:
    m = Maestros(db)
    vistos: set = set()
    cabeceras: dict = {}
    resultado = []
    normalizadas = [(r, *_normalizar(r, m)) for r in filas]

    claves_oc = {(d["proveedor"], d["oc"]) for _, d, _ in normalizadas}
    ocs_existentes = {}
    for codigo, numero in claves_oc:
        prov = m.proveedores.get(codigo)
        if prov:
            oc = db.scalar(select(OrdenCompra).where(OrdenCompra.proveedor_id == prov.id, OrdenCompra.numero == numero))
            if oc:
                ocs_existentes[(codigo, numero)] = oc
    ids_pos = [p.id for oc in ocs_existentes.values() for p in oc.posiciones]
    facturado = facturado_por_posicion(db, ids_pos)

    for registro, d, errores in normalizadas:
        clave = f"{d['proveedor']} / PO {d['oc']} / line {d['posicion']}"
        salida = {"fila": registro.get("_fila"), "clave": clave, "datos": registro,
                  "mensajes": list(errores), "cambios": {}}
        if d["proveedor"] and d["proveedor"] not in m.proveedores:
            salida["mensajes"].append(f"Supplier {d['proveedor']} does not exist.")
        k = (d["proveedor"], d["oc"], d["posicion"])
        if k in vistos:
            salida["mensajes"].append("Line repeated within the file.")
        vistos.add(k)
        cab = {c: d.get(c) for c in CAMPOS_CABECERA}
        previa = cabeceras.setdefault((d["proveedor"], d["oc"]), cab)
        if previa != cab:
            salida["mensajes"].append("The header data does not match other rows of the same PO.")
        if salida["mensajes"]:
            salida["estado"] = "error"
            resultado.append(salida)
            continue

        oc = ocs_existentes.get((d["proveedor"], d["oc"]))
        pos = next((p for p in oc.posiciones if p.posicion == d["posicion"]), None) if oc else None
        if not pos:
            salida["estado"] = "nuevo"
            resultado.append(salida)
            continue

        cambios = {}
        for c in CAMPOS_CABECERA:
            anterior = getattr(oc, c)
            if anterior != d.get(c):
                cambios[c] = {"antes": anterior, "despues": d.get(c)}
        for c, nuevo in _valores_posicion(d).items():
            if c == "articulo_id":
                continue
            anterior = getattr(pos, c)
            igual = abs(anterior - nuevo) < 1e-9 if isinstance(anterior, float) and nuevo is not None else anterior == nuevo
            if not igual:
                cambios[c] = {"antes": anterior, "despues": nuevo}
        salida["cambios"] = cambios
        fact = facturado.get(pos.id, 0)
        if not cambios:
            salida["estado"] = "sin_cambio"
        elif fact and d["cantidad"] is not None and d["cantidad"] < fact:
            salida["estado"] = "conflicto"
            salida["mensajes"].append(f"The new quantity ({d['cantidad']}) is less than what is already invoiced ({fact}).")
        elif fact and any(c in cambios for c in ("sociedad", "moneda", "centro", "unidad", "codigo_sap")):
            salida["estado"] = "conflicto"
            salida["mensajes"].append(
                "Key data (company, currency, plant, SKU or unit) changes on an already invoiced line.")
        elif fact and any(c in cambios for c in ("casepack", "inner_pack")):
            salida["estado"] = "conflicto"
            salida["mensajes"].append("The casepack or inner pack changes on a line that is already invoiced.")
        else:
            salida["estado"] = "cambio"
        resultado.append(salida)
    return resultado


def _resumen(clasificadas: list[dict]) -> dict:
    res = {"nuevo": 0, "cambio": 0, "sin_cambio": 0, "conflicto": 0, "error": 0}
    for c in clasificadas:
        res[c["estado"]] += 1
    res["total"] = len(clasificadas)
    return res


def importar_previa(db: Session, user: Usuario, nombre: str, contenido: bytes) -> dict:
    exigir(user, "oc.importar")
    filas = _leer_archivo(nombre, contenido)
    if not filas:
        raise ErrorNegocio("The file has no rows with data.", 422, "archivo_vacio")
    clasificadas = _clasificar(db, filas)
    imp = ImportacionOC(usuario_id=user.id, nombre_archivo=nombre, filas=filas)
    db.add(imp)
    db.flush()
    return {
        "importacion_id": imp.id,
        "archivo": nombre,
        "resumen": _resumen(clasificadas),
        "filas": [{k: v for k, v in c.items() if k != "datos"} for c in clasificadas],
    }


def importar_aplicar(db: Session, user: Usuario, importacion_id: int) -> dict:
    exigir(user, "oc.importar")
    imp = db.get(ImportacionOC, importacion_id)
    if not imp:
        raise ErrorNegocio("The import does not exist.", 404, "no_encontrado")
    if imp.estado == "APLICADA":
        return {"ya_aplicada": True, **(imp.resultado or {})}
    # Se vuelve a clasificar: los datos pudieron cambiar desde la vista previa
    clasificadas = _clasificar(db, imp.filas)
    m = Maestros(db)
    aplicadas = 0
    ocs_tocadas: dict[int, dict] = {}
    for c in clasificadas:
        if c["estado"] == "conflicto":
            d, _ = _normalizar(c["datos"])
            prov = m.proveedores.get(d["proveedor"])
            db.add(Alerta(proveedor_id=prov.id if prov else None, tipo="conflicto_oc",
                          mensaje=f"{c['clave']}: " + " ".join(c["mensajes"]),
                          referencia={"importacion_id": imp.id, "fila": c["fila"]}))
        if c["estado"] not in ("nuevo", "cambio"):
            continue
        d, _ = _normalizar(c["datos"], m)
        prov = m.proveedores[d["proveedor"]]
        oc = db.scalar(select(OrdenCompra).where(OrdenCompra.proveedor_id == prov.id, OrdenCompra.numero == d["oc"]))
        nueva = oc is None
        if nueva:
            oc = OrdenCompra(proveedor_id=prov.id, numero=d["oc"])
            db.add(oc)
        info = ocs_tocadas.setdefault(id(oc), {
            "oc": oc, "cambios": False, "explicita": d["liberacion_logistica_archivo"],
            "actual": None if nueva else oc.liberacion_logistica})
        info["cambios"] = info["cambios"] or (c["estado"] == "cambio")
        for campo in CAMPOS_CABECERA:
            setattr(oc, campo, d.get(campo))
        oc.actualizado_en = ahora()
        db.flush()
        pos = db.scalar(select(PosicionOC).where(PosicionOC.oc_id == oc.id, PosicionOC.posicion == d["posicion"]))
        if not pos:
            pos = PosicionOC(oc_id=oc.id, posicion=d["posicion"])
            db.add(pos)
        for campo, valor in _valores_posicion(d).items():
            setattr(pos, campo, valor)
        aplicadas += 1
    for info in ocs_tocadas.values():
        oc = info["oc"]
        oc.liberacion_logistica = liberacion_logistica(oc.liberacion_comercial, info["explicita"], info["actual"],
                                                       info["cambios"])
        oc.liberada = esta_liberada(oc.liberacion_comercial, oc.liberacion_logistica)
    resultado = {"resumen": _resumen(clasificadas), "aplicadas": aplicadas}
    imp.estado = "APLICADA"
    imp.resultado = resultado
    registrar(db, user, "importacion_oc", imp.id, "aplicar", resultado)
    return resultado
