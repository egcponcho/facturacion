import csv
import io
import re
import unicodedata
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from app.core import listas
from app.core.empresa import regla
from app.core.errores import ErrorNegocio
from app.modelos import (
    Alerta,
    Almacen,
    Articulo,
    Centro,
    Factura,
    FacturaLinea,
    ImportacionOC,
    OrdenCompra,
    Pais,
    PosicionOC,
    Proveedor,
    Puerto,
    Sociedad,
    Usuario,
    ahora,
)
from app.modelos import cant as cant_norm
from app.modulos.acceso.permisos import asegurar_proveedor, exigir, proveedor_filtro
from app.modulos.acceso.preferencias import leer_fecha
from app.modulos.compras import liberaciones
from app.modulos.comun.historial import registrar
from app.modulos.comun.texto import filtro_texto, terminos
from app.modulos.facturacion.cantidades import facturado_por_posicion, facturas_por_posicion
from app.modulos.maestros.unidades import error_cantidad
from app.modulos.maestros.unidades import normalizar as unidad_de
from app.modulos.productos.productos import clasificacion_txt, partida_para, producto_de

# Dos liberaciones de dos equipos distintos (comercial y logística). Sus
# códigos, nombres y efecto los define cada empresa en Datos maestros →
# Estados de liberación (modulos/compras/liberaciones.py).


# ---- Vista general de OCs ---------------------------------------------------
ORDEN_OC = {"numero", "fecha", "fecha_xf", "fecha_tienda", "importe", "avance", "proveedor"}


def listar_ordenes(
    db: Session,
    user: Usuario,
    proveedor_id: int | None = None,
    q: str | None = None,
    centro: str | None = None,
    solo_disponible: bool = True,
    page: int = 1,
    size: int = 25,
    sociedad: str | None = None,
    marca: str | None = None,
    liberacion: str | None = None,
    destino: str | None = None,
    puerto: str | None = None,
    orden: str | None = None,
    almacen: str | None = None,
    comercial: str | None = None,
    liberada: bool | None = None,
) -> dict:
    exigir(user, "oc.ver")
    lib = liberaciones.de(db)
    prov = proveedor_filtro(user, proveedor_id)

    tot = (
        select(
            PosicionOC.oc_id.label("oc_id"),
            func.sum(PosicionOC.cantidad).label("cantidad"),
            func.sum(PosicionOC.cantidad * PosicionOC.precio).label("importe"),
            func.count(PosicionOC.id).label("n"),
        )
        .group_by(PosicionOC.oc_id)
        .subquery()
    )
    fac = (
        select(
            PosicionOC.oc_id.label("oc_id"),
            func.sum(FacturaLinea.cantidad).label("facturado"),
            func.sum(FacturaLinea.cantidad * PosicionOC.precio).label("importe_facturado"),
        )
        .join(FacturaLinea, FacturaLinea.posicion_oc_id == PosicionOC.id)
        .join(Factura, Factura.id == FacturaLinea.factura_id)
        .where(Factura.estado != "CANCELADA")
        .group_by(PosicionOC.oc_id)
        .subquery()
    )
    facturado = func.coalesce(fac.c.facturado, 0)
    importe_fact = func.coalesce(fac.c.importe_facturado, 0)
    consulta = (
        select(OrdenCompra, Proveedor.nombre, tot.c.importe, tot.c.n, importe_fact)
        .join(Proveedor, Proveedor.id == OrdenCompra.proveedor_id)
        .join(tot, tot.c.oc_id == OrdenCompra.id)
        .outerjoin(fac, fac.c.oc_id == OrdenCompra.id)
    )
    if prov:
        consulta = consulta.where(OrdenCompra.proveedor_id == prov)
    if q:
        consulta = consulta.where(filtro_texto(q, lambda p: [
            OrdenCompra.numero.ilike(p),
            OrdenCompra.id.in_(select(PosicionOC.oc_id).where(or_(
                PosicionOC.estilo.ilike(p), PosicionOC.codigo_sap.ilike(p), PosicionOC.upc.ilike(p), PosicionOC.color.ilike(p)))),
        ]))
    if centro:
        consulta = consulta.where(OrdenCompra.centro == centro)
    if sociedad:
        consulta = consulta.where(OrdenCompra.sociedad == sociedad)
    if destino:
        consulta = consulta.where(OrdenCompra.centro_destino == destino)
    if puerto:
        consulta = consulta.where(OrdenCompra.puerto_despacho == puerto)
    if liberacion:
        consulta = consulta.where(OrdenCompra.liberacion_logistica == liberacion)
    if comercial:
        consulta = consulta.where(OrdenCompra.liberacion_comercial == comercial)
    if liberada is not None:
        consulta = consulta.where(OrdenCompra.liberada.is_(liberada))
    if marca:
        consulta = consulta.where(OrdenCompra.id.in_(select(PosicionOC.oc_id).where(PosicionOC.marca == marca)))
    if almacen:
        consulta = consulta.where(OrdenCompra.id.in_(select(PosicionOC.oc_id).where(PosicionOC.almacen == almacen)))
    if solo_disponible:
        consulta = consulta.where(tot.c.cantidad - facturado > 0)

    col, _, direccion = (orden or "").partition(":")
    expr = {
        "numero": OrdenCompra.numero, "fecha": OrdenCompra.fecha, "fecha_xf": OrdenCompra.fecha_xf,
        "fecha_tienda": OrdenCompra.fecha_tienda, "importe": tot.c.importe, "proveedor": Proveedor.nombre,
        "avance": importe_fact * 1.0 / func.nullif(tot.c.importe, 0),
    }.get(col)
    if expr is not None:
        consulta = consulta.order_by(expr.desc().nullslast() if direccion == "desc" else expr.asc().nullslast(),
                                     OrdenCompra.numero)
    else:
        consulta = consulta.order_by(OrdenCompra.fecha.desc().nullslast(), OrdenCompra.numero)

    total = db.scalar(select(func.count()).select_from(consulta.subquery())) or 0
    filas = db.execute(consulta.offset((page - 1) * size).limit(size)).all()

    # Cantidades por unidad de medida: nunca se suman pares con unidades
    ids = [oc.id for oc, *_ in filas]
    por_unidad: dict[int, dict] = {i: {} for i in ids}
    marcas: dict[int, set] = {i: set() for i in ids}
    almacenes: dict[int, set] = {i: set() for i in ids}
    if ids:
        for oc_id, alm in db.execute(
            select(PosicionOC.oc_id, PosicionOC.almacen).where(PosicionOC.oc_id.in_(ids)).distinct()
        ).all():
            if alm:
                almacenes[oc_id].add(alm)
        for oc_id, unidad, marca_oc, cant in db.execute(
            select(PosicionOC.oc_id, PosicionOC.unidad, PosicionOC.marca, func.sum(PosicionOC.cantidad))
            .where(PosicionOC.oc_id.in_(ids)).group_by(PosicionOC.oc_id, PosicionOC.unidad, PosicionOC.marca)
        ).all():
            d = por_unidad[oc_id].setdefault(unidad, {"cantidad": 0, "facturado": 0, "disponible": 0})
            d["cantidad"] = cant_norm(d["cantidad"] + (cant or 0))
            d["disponible"] = cant_norm(d["disponible"] + (cant or 0))
            if marca_oc:
                marcas[oc_id].add(marca_oc)
        for oc_id, unidad, cant in db.execute(
            select(PosicionOC.oc_id, PosicionOC.unidad, func.sum(FacturaLinea.cantidad))
            .join(FacturaLinea, FacturaLinea.posicion_oc_id == PosicionOC.id)
            .join(Factura, Factura.id == FacturaLinea.factura_id)
            .where(Factura.estado != "CANCELADA", PosicionOC.oc_id.in_(ids))
            .group_by(PosicionOC.oc_id, PosicionOC.unidad)
        ).all():
            d = por_unidad[oc_id][unidad]
            d["facturado"] = cant_norm(cant or 0)
            d["disponible"] = max(d["cantidad"] - d["facturado"], 0)

    hoy = date.today()
    centros = {c.codigo: c for c in db.scalars(select(Centro))}
    from app.modulos.transporte.leadtimes import tiendas_estimadas
    est_tienda = tiendas_estimadas(db, [oc for oc, *_ in filas])
    items = []
    for oc, prov_nombre, importe, n, importe_f in filas:
        destino = centros.get(oc.centro_destino)
        importe = float(importe or 0)
        items.append(
            {
                **_cabecera_oc(oc, lib),
                "proveedor": prov_nombre,
                "posiciones": n,
                "marcas": sorted(marcas[oc.id]),
                # Una OC puede repartir sus posiciones entre varios almacenes
                "almacenes": sorted(almacenes[oc.id]),
                "por_unidad": por_unidad[oc.id],
                "importe": round(importe, 2),
                "importe_facturado": round(float(importe_f or 0), 2),
                # Avance por valor: es comparable aunque la OC mezcle pares y unidades
                "avance": round(float(importe_f or 0) * 100 / importe, 1) if importe else 0,
                "dias_tienda": (oc.fecha_tienda - hoy).days if oc.fecha_tienda else None,
                # Con los lead times de su origen: cuándo estaría en tienda
                **est_tienda.get(oc.id, {}),
                # El centro de destino dice a qué país llega al final
                "destino_nombre": destino.nombre if destino else None,
                "pais_destino": destino.pais if destino else None,
            }
        )
    return {"items": items, "total": total, "page": page, "size": size}


def filtros_ordenes(db: Session, user: Usuario, proveedor_id: int | None = None) -> dict:
    """Valores presentes en las OCs visibles, para armar filtros dinámicos."""
    exigir(user, "oc.ver")
    prov = proveedor_filtro(user, proveedor_id)
    base = select(OrdenCompra)
    if prov:
        base = base.where(OrdenCompra.proveedor_id == prov)
    sub = base.subquery()

    def distintos(col):
        return sorted({v for (v,) in db.execute(select(col).select_from(sub).distinct()) if v})

    def de_posiciones(col):
        return sorted({v for (v,) in db.execute(
            select(col).where(PosicionOC.oc_id.in_(select(sub.c.id))).distinct()) if v})

    marcas = de_posiciones(PosicionOC.marca)
    lib = liberaciones.de(db)
    return {
        "sociedades": distintos(sub.c.sociedad),
        "centros": distintos(sub.c.centro),
        "almacenes": de_posiciones(PosicionOC.almacen),
        "destinos": [{"codigo": d.codigo, "nombre": f"{d.nombre} ({d.pais})"} for d in db.scalars(
            select(Centro).where(Centro.codigo.in_(distintos(sub.c.centro_destino))).order_by(Centro.codigo))],
        "puertos": [{"codigo": p.codigo, "nombre": p.nombre} for p in db.scalars(
            select(Puerto).where(Puerto.codigo.in_(distintos(sub.c.puerto_despacho))))],
        "marcas": marcas,
        "liberaciones": [{"codigo": x.codigo, "nombre": x.nombre, "libera": x.libera} for x in lib.logistica.lista],
        "comerciales": [{"codigo": x.codigo, "nombre": x.nombre, "libera": x.libera} for x in lib.comercial.lista],
    }


def _cabecera_oc(oc: OrdenCompra, lib: liberaciones.Liberaciones) -> dict:
    return {
        "id": oc.id,
        "numero": oc.numero,
        "proveedor_id": oc.proveedor_id,
        "sociedad": oc.sociedad,
        "centro": oc.centro,
        "centro_destino": oc.centro_destino,
        "moneda": oc.moneda,
        "incoterm": oc.incoterm,
        "fecha": oc.fecha,
        "puerto_despacho": oc.puerto_despacho,
        "pais_origen": oc.pais_origen,
        "pais_procedencia": oc.pais_procedencia,
        "fecha_xf_original": oc.fecha_xf_original,
        "fecha_xf": oc.fecha_xf,
        "fecha_tienda": oc.fecha_tienda,
        "liberacion_comercial": oc.liberacion_comercial,
        "liberacion_logistica": oc.liberacion_logistica,
        "fecha_lib_comercial": oc.fecha_lib_comercial,
        "fecha_lib_logistica": oc.fecha_lib_logistica,
        "liberacion_txt": lib.logistica.nombre(oc.liberacion_logistica),
        "comercial_txt": lib.comercial.nombre(oc.liberacion_comercial),
        # Liberación dada, pero la OC cambió después (para avisarlo en pantalla)
        "con_cambios": bool((x := lib.logistica.get(oc.liberacion_logistica)) and x.con_cambios),
        "liberada": oc.liberada,
    }


def posiciones_oc(db: Session, user: Usuario, oc_id: int) -> dict:
    exigir(user, "oc.ver")
    oc = db.get(OrdenCompra, oc_id)
    if not oc:
        raise ErrorNegocio("The purchase order does not exist.", 404, "no_encontrado")
    asegurar_proveedor(user, oc.proveedor_id)
    ids = [p.id for p in oc.posiciones]
    facturado = facturado_por_posicion(db, ids)
    facturas = facturas_por_posicion(db, ids)
    # La partida sale de la clasificación del producto, con el código del país destino
    posiciones = []
    for p in oc.posiciones:
        prod = producto_de(p.articulo)
        fact = facturado.get(p.id, 0)
        fs = facturas.get(p.id, [])
        estado, motivo, restringida = estado_posicion(oc, p, fact, fs)
        posiciones.append(
            {
                "id": p.id,
                "posicion": p.posicion,
                "almacen": p.almacen,
                "codigo_sap": p.codigo_sap,
                "upc": p.upc,
                "estilo": p.estilo,
                "color": p.color,
                "talla": p.talla,
                "descripcion": p.descripcion,
                "marca": p.marca,
                "grupo": p.grupo,
                "categoria": p.categoria,
                "tipo_empaque": p.tipo_empaque,
                "casepack": p.casepack,
                "inner_pack": p.inner_pack,
                "prepack": p.prepack,
                "unidades_por_caja": p.unidades_por_caja,
                "cantidad": p.cantidad,
                "unidad": p.unidad,
                "precio": p.precio,
                "importe": round(p.cantidad * p.precio, 2) if p.precio is not None else None,
                "fecha_entrega": p.fecha_entrega,
                "pais_origen": p.pais_origen,
                "partida_arancelaria": partida_para(prod),  # 6 dígitos: el destino es solo proyectado
                "clasificacion": clasificacion_txt(prod),
                "facturado": fact,
                "disponible": max(p.cantidad - fact, 0),
                "facturas": fs,
                "estado": estado,
                "motivo": motivo,
                "solo_factura_id": restringida,
            }
        )
    prov = db.get(Proveedor, oc.proveedor_id)
    destino = db.scalar(select(Centro).where(Centro.codigo == oc.centro_destino))
    from app.modulos.transporte.leadtimes import tiendas_estimadas
    return {"oc": {**_cabecera_oc(oc, liberaciones.de(db)), "proveedor": prov.nombre, **tiendas_estimadas(db, [oc]).get(oc.id, {}),
                   "centro_destino_nombre": destino.nombre if destino else None,
                   "pais_destino": destino.pais if destino else None},
            "posiciones": posiciones}


def empaque_posicion(db: Session, user: Usuario, oc_id: int, posicion_id: int, casepack: int | None,
                     inner_pack: int | None) -> dict:
    """Casepack e inner pack de una posición (condición de la compra). Se
    cambian mientras no haya nada facturado de esa posición."""
    exigir(user, "oc.empaque")
    p = db.get(PosicionOC, posicion_id)
    if not p or p.oc_id != oc_id:
        raise ErrorNegocio("The purchase order line does not exist.", 404, "no_encontrado")
    errores = validar_empaque(p.tipo_empaque, casepack, inner_pack, p.cantidad)
    if errores:
        raise ErrorNegocio("Check the packing of the line.", 422, "validacion",
                           [{"campo": "inner_pack", "mensaje": e} for e in errores])
    if facturado_por_posicion(db, [p.id]).get(p.id, 0):
        raise ErrorNegocio("This line is already invoiced: its packing can no longer change.", 409, "facturada")
    antes = {"casepack": p.casepack, "inner_pack": p.inner_pack}
    p.casepack, p.inner_pack = casepack, inner_pack
    registrar(db, user, "oc", p.oc_id, "empaque",
              {"posicion": p.posicion, **{k: [antes[k], v] for k, v in (("casepack", casepack),
                                                                           ("inner_pack", inner_pack))
                                          if antes[k] != v}})
    return {"id": p.id, "casepack": p.casepack, "inner_pack": p.inner_pack}


def estado_posicion(oc, p, facturado: int, facturas: list[dict]):
    """Devuelve (estado, motivo, factura a la que queda restringido el saldo)."""
    disponible = p.cantidad - facturado
    if not oc.liberada:
        return "NO_DISPONIBLE", "The PO is not released yet: it cannot be invoiced.", None
    if p.bloqueada:
        return "NO_DISPONIBLE", p.motivo_bloqueo or "Line blocked.", None
    if disponible <= 0:
        return "FACTURADA", None, None
    if facturado > 0:
        if not regla("POSICION_EN_VARIAS_FACTURAS") and facturas:
            f = facturas[0]
            return "PARCIAL", f"The balance can only be added to {f['nombre']}.", f["id"]
        return "PARCIAL", None, None
    return "DISPONIBLE", None, None


# ---- Importación ------------------------------------------------------------
# El primer alias de cada campo es el nombre de la columna en la plantilla (inglés).
ALIAS = {
    "proveedor": ["supplier", "vendor", "supplier_code", "proveedor", "codigo_proveedor", "proveedor_codigo"],
    "oc": ["po", "po_number", "purchase_order", "oc", "orden", "orden_compra", "numero_oc"],
    "posicion": ["po_line", "line", "item", "posicion", "pos", "linea"],
    "sociedad": ["company", "company_code", "sociedad", "compania"],
    "centro": ["plant", "bonded_warehouse", "centro", "bodega", "bodega_fiscal"],
    "almacen": ["storage_location", "sloc", "almacen"],
    "centro_destino": ["destination", "destination_center", "destination_country", "centro_destino", "pais_destino",
                       "destino", "codigo_destino"],
    "moneda": ["currency", "moneda"],
    "incoterm": ["incoterm"],
    "fecha_oc": ["po_date", "fecha_oc", "fecha"],
    "puerto_despacho": ["port_of_loading", "pol", "port", "puerto_despacho", "puerto", "puerto_embarque"],
    "pais_origen": ["country_of_origin", "coo", "origin", "pais_origen", "origen"],
    "pais_procedencia": ["country_of_shipment", "provenance", "pais_procedencia", "procedencia"],
    "fecha_xf_original": ["xf_date_original", "original_xf_date", "fecha_xf_original", "xf_original"],
    "fecha_xf": ["xf_date", "xf", "updated_xf_date", "fecha_xf", "fecha_xf_actualizada", "xf_actualizada"],
    "fecha_tienda": ["in_store_date", "store_date", "fecha_tienda", "fecha_requerida_tienda", "fecha_requerida"],
    "liberacion_comercial": ["commercial_release", "liberacion_comercial", "lib_comercial", "liberada",
                             "estado_liberacion"],
    "liberacion_logistica": ["logistics_release", "liberacion_logistica", "lib_logistica"],
    "fecha_lib_comercial": ["commercial_release_date", "fecha_liberacion_comercial", "fecha_lib_comercial"],
    "fecha_lib_logistica": ["logistics_release_date", "fecha_liberacion_logistica", "fecha_lib_logistica"],
    "codigo_sap": ["sku", "sap_code", "material", "item_code", "codigo_sap", "sap", "codigo"],
    "cantidad": ["quantity", "qty", "cantidad"],
    "unidad": ["uom", "unit", "um", "unidad"],
    "precio": ["unit_price", "price", "precio", "precio_unitario"],
    "casepack": ["casepack", "case_pack"],
    "inner_pack": ["inner_pack", "inner", "innerpack", "inner_casepack", "pack"],
    "fecha_entrega": ["delivery_date", "fecha_entrega", "entrega"],
}
# Lo mínimo para registrar una OC; el resto es opcional y lo que falte para
# facturar (empresa, moneda, precio) se pide al facturar
REQUERIDOS = ["proveedor", "oc", "posicion", "codigo_sap", "cantidad"]
NOMBRE_CAMPO = {"proveedor": "supplier", "oc": "PO number", "posicion": "PO line", "codigo_sap": "item code",
                "cantidad": "quantity"}
CAMPOS_CABECERA = ["sociedad", "centro", "centro_destino", "moneda", "incoterm", "fecha", "puerto_despacho",
                   "pais_origen", "pais_procedencia", "fecha_xf_original", "fecha_xf", "fecha_tienda",
                   "liberacion_comercial"]
CAMPOS_POSICION = [
    "almacen", "articulo_id", "codigo_sap", "upc", "estilo", "color", "talla", "descripcion", "marca", "grupo", "categoria",
    "tipo_empaque", "casepack", "inner_pack", "prepack", "unidades_por_caja", "cantidad", "unidad", "precio", "fecha_entrega",
    "pais_origen",
]


def _norm(texto: str) -> str:
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode().lower()
    return "_".join("".join(c if c.isalnum() else " " for c in t).split())


def _texto(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v).strip()


def _leer_archivo(nombre: str, contenido: bytes) -> list[dict]:
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        ws = wb.active
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            return []
        encabezados = [_texto(h) for h in filas[0]]
        datos = [[_texto(v) for v in fila] for fila in filas[1:]]
    else:
        texto = contenido.decode("utf-8-sig", errors="replace")
        muestra = texto[:2000]
        delimitador = ";" if muestra.count(";") > muestra.count(",") else ","
        lector = list(csv.reader(io.StringIO(texto), delimiter=delimitador))
        if not lector:
            return []
        encabezados = lector[0]
        datos = [[v.strip() for v in fila] for fila in lector[1:]]

    mapa = {}
    for i, h in enumerate(encabezados):
        n = _norm(h)
        for campo, alias in ALIAS.items():
            if n in alias and campo not in mapa:
                mapa[campo] = i
    faltan = [c for c in REQUERIDOS if c not in mapa]
    if faltan:
        raise ErrorNegocio(
            "The file is missing required columns: " + ", ".join(ALIAS[c][0] for c in faltan) + ".",
            422,
            "columnas_faltantes",
            {"faltan": faltan, "encontradas": encabezados},
        )
    filas_dict = []
    for n_fila, fila in enumerate(datos, start=2):
        if not any(fila):
            continue
        registro = {campo: (fila[i] if i < len(fila) else "") for campo, i in mapa.items()}
        registro["_fila"] = n_fila
        filas_dict.append(registro)
    return filas_dict


def _fecha(valor: str) -> date | None:
    # ISO o el formato de fecha del usuario (con su orden día/mes para no confundirlos)
    return leer_fecha(valor)


class Maestros:
    """Catálogos precargados para validar un archivo completo sin consultas por fila."""

    def __init__(self, db: Session):
        self.lib = liberaciones.de(db)
        self.proveedores = {p.codigo: p for p in db.scalars(select(Proveedor))}
        self.sociedades = {s.codigo: s for s in db.scalars(select(Sociedad))}
        self.centros = {c.codigo: c for c in db.scalars(select(Centro))}
        self.almacenes = {a.codigo: a for a in db.scalars(select(Almacen))}
        self.puertos = {p.codigo: p for p in db.scalars(select(Puerto))}
        self.paises = {p.codigo for p in db.scalars(select(Pais))}
        self.articulos: dict[str, Articulo] = {}
        self._db = db
        # Códigos escritos de otra forma (minúsculas, espacios, por nombre) -> código real
        from app.modulos.comun.normalizar import Referencias

        self.refs = {"proveedor": Referencias(db, Proveedor), "sociedad": Referencias(db, Sociedad),
                     "centro": Referencias(db, Centro), "almacen": Referencias(db, Almacen),
                     "puerto": Referencias(db, Puerto), "pais": Referencias(db, Pais)}

    def canon(self, tabla: str, valor):
        obj = self.refs[tabla].buscar(valor)
        return obj.codigo if obj else valor

    def articulo(self, sku: str) -> Articulo | None:
        if sku not in self.articulos:
            self.articulos[sku] = self._db.scalar(select(Articulo).where(Articulo.sku == sku))
        return self.articulos[sku]


PAIS_TXT = {"pais_origen": "country of origin", "pais_procedencia": "country of shipment"}


def _normalizar(registro: dict, m: Maestros) -> tuple[dict, list[str]]:
    """Convierte los textos de una fila en valores tipados y la valida contra
    los datos maestros. Los códigos se mantienen como texto."""
    errores = []
    r = {k: (v or "").strip() for k, v in registro.items() if k != "_fila"}
    for campo in REQUERIDOS:
        if not r.get(campo):
            errores.append(f"Missing: {NOMBRE_CAMPO[campo]}.")
    d = {
        "proveedor": r.get("proveedor", "").upper(),
        "oc": r.get("oc", ""),
        "posicion": r.get("posicion", "").lstrip("0") or r.get("posicion", ""),
        "sociedad": r.get("sociedad", "").upper(),
        "centro": (r.get("centro") or "").upper() or None,
        "almacen": (r.get("almacen") or "").upper() or None,
        "centro_destino": r.get("centro_destino") or None,
        "moneda": (r.get("moneda") or "").upper(),
        "incoterm": (r.get("incoterm") or "").upper() or None,
        "puerto_despacho": (r.get("puerto_despacho") or "").upper() or None,
        "pais_origen": (r.get("pais_origen") or "").upper() or None,
        "pais_procedencia": (r.get("pais_procedencia") or r.get("pais_origen") or "").upper() or None,
        "liberacion_comercial": (r.get("liberacion_comercial") or "").strip() or None,
        "liberacion_logistica_archivo": (r.get("liberacion_logistica") or "").strip() or None,
        "codigo_sap": r.get("codigo_sap", ""),
    }
    com = m.lib.comercial.leer(d["liberacion_comercial"])
    if com is False:
        errores.append(f"Invalid commercial release {d['liberacion_comercial']} (use {m.lib.comercial.ayuda()}).")
        com = None
    d["liberacion_comercial"] = com or m.lib.comercial.predeterminado()
    log = m.lib.logistica.leer(d["liberacion_logistica_archivo"])
    d["liberacion_logistica_archivo"] = log
    if log is False:
        errores.append(f"Invalid logistics release {r.get('liberacion_logistica')} (use {m.lib.logistica.ayuda()}).")
        log = d["liberacion_logistica_archivo"] = None
    elif log and m.lib.logistica.libera(log) and not m.lib.comercial.libera(d["liberacion_comercial"]):
        errores.append(f"Logistics cannot release ({log}) without commercial release.")
    if d["oc"] and not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._\-/]{0,39}", d["oc"]):
        errores.append(f"PO {d['oc']}: letters and numbers (also . - _ /), up to 40 characters.")
    if d["posicion"] and not re.fullmatch(r"[A-Za-z0-9]{1,10}", d["posicion"]):
        errores.append(f"Line {d['posicion']}: letters and numbers, up to 10 characters.")
    try:
        cant = float(r.get("cantidad", "").replace(",", "")) if r.get("cantidad") else None
        if cant is not None and cant < 0:
            raise ValueError
        d["cantidad"] = cant_norm(cant)
    except ValueError:
        errores.append(f"Invalid quantity: {r.get('cantidad')}.")
        d["cantidad"] = None
    try:
        d["precio"] = float(r.get("precio", "").replace(",", "")) if r.get("precio") else None
        if d["precio"] is not None and d["precio"] < 0:
            raise ValueError
    except ValueError:
        errores.append(f"Invalid price: {r.get('precio')}.")
        d["precio"] = None
    # Empaque de la compra: viene en la posición de la OC, no en el artículo
    for campo, texto in (("casepack", "Casepack"), ("inner_pack", "Inner pack")):
        try:
            d[campo] = int(float(r[campo])) if r.get(campo) else None
            if d[campo] is not None and d[campo] < 1:
                raise ValueError
        except ValueError:
            errores.append(f"Invalid {texto.lower()}: {r.get(campo)}.")
            d[campo] = None
    unidad = unidad_de(r.get("unidad"))
    if r.get("unidad") and not unidad:
        errores.append(f"Unknown unit: {r.get('unidad')} (use PAR, UN or CJ).")
    d["unidad"] = unidad
    for campo, destino in (("fecha_oc", "fecha"), ("fecha_entrega", "fecha_entrega"),
                           ("fecha_xf_original", "fecha_xf_original"), ("fecha_xf", "fecha_xf"),
                           ("fecha_tienda", "fecha_tienda"), ("fecha_lib_comercial", "fecha_lib_comercial"),
                           ("fecha_lib_logistica", "fecha_lib_logistica")):
        try:
            d[destino] = _fecha(r.get(campo, ""))
        except ValueError as e:
            errores.append(str(e).capitalize() + ".")
            d[destino] = None
    if d["precio"] is not None and not d["moneda"]:
        errores.append("Enter the currency of the price.")
    if d["cantidad"] == 0:
        errores.append("The quantity must be greater than zero.")
    d["moneda"] = d["moneda"] or None
    d["sociedad"] = d["sociedad"] or None
    d["fecha_xf_original"] = d["fecha_xf_original"] or d["fecha_xf"]
    d["fecha_xf"] = d["fecha_xf"] or d["fecha_xf_original"]

    if m is None:
        return d, errores
    for campo, tabla in (("proveedor", "proveedor"), ("sociedad", "sociedad"), ("centro", "centro"), ("almacen", "almacen"),
                         ("centro_destino", "centro"), ("puerto_despacho", "puerto"), ("pais_origen", "pais"),
                         ("pais_procedencia", "pais")):
        if d.get(campo):
            d[campo] = m.canon(tabla, d[campo])
    # Datos maestros. Todo queda encadenado: el proveedor trabaja con la
    # sociedad; centro, almacén y centro destino son de esa sociedad; el
    # artículo y su marca son del proveedor. Sin sociedad, se toma la del
    # centro o almacén indicado.
    if not d["sociedad"]:
        for campo, tabla in (("centro", m.centros), ("almacen", m.almacenes), ("centro_destino", m.centros)):
            x = tabla.get(d[campo]) if d[campo] else None
            if x and x.sociedad:
                d["sociedad"] = x.sociedad.codigo
                break
    soc = m.sociedades.get(d["sociedad"])
    if d["sociedad"] and not soc:
        errores.append(f"Company {d['sociedad']} does not exist.")
    cen = m.centros.get(d["centro"]) if d["centro"] else None
    if d["centro"] and not cen:
        errores.append(f"Plant {d['centro']} does not exist.")
    elif cen and soc and cen.sociedad_id != soc.id:
        errores.append(f"Plant {d['centro']} does not belong to company {d['sociedad']}.")
    if d["almacen"]:
        alm = m.almacenes.get(d["almacen"])
        if not alm:
            errores.append(f"Storage location {d['almacen']} does not exist.")
        elif soc and alm.sociedad_id != soc.id:
            errores.append(f"Storage location {d['almacen']} does not belong to company {d['sociedad']}.")
    prov = m.proveedores.get(d["proveedor"])
    if prov and not prov.activo:
        errores.append(f"Supplier {prov.nombre} is inactive.")
    if prov and soc and soc.id not in {x.id for x in prov.sociedades}:
        errores.append(f"Supplier {prov.nombre} does not work with company {d['sociedad']}. "
                       "Assign the company to the supplier in master data first.")
    destino = m.centros.get(d["centro_destino"]) if d["centro_destino"] else None
    if d["centro_destino"] and not destino:
        errores.append(f"Destination center {d['centro_destino']} is not registered in master data.")
    elif destino and soc and destino.sociedad_id != soc.id:
        errores.append(f"Destination center {d['centro_destino']} does not belong to company {d['sociedad']}.")
    if d["puerto_despacho"] and d["puerto_despacho"] not in m.puertos:
        errores.append(f"Port {d['puerto_despacho']} is not registered.")
    for campo in ("pais_origen", "pais_procedencia"):
        if d[campo] and d[campo] not in m.paises:
            errores.append(f"Country {d[campo]} ({PAIS_TXT[campo]}) is not registered.")
    art = m.articulo(d["codigo_sap"]) if d["codigo_sap"] else None
    if d["codigo_sap"] and not art:
        errores.append(f"SKU {d['codigo_sap']} does not exist in the item master. Load it first.")
    elif art:
        prov = m.proveedores.get(d["proveedor"])
        if prov and art.proveedor_id and art.proveedor_id != prov.id:
            errores.append(f"SKU {art.sku} belongs to another supplier, not {prov.nombre}.")
        elif prov and art.marca_id not in {x.id for x in prov.marcas}:
            errores.append(f"Brand {art.marca.codigo} of SKU {art.sku} is not a brand of {prov.nombre}.")
        if not art.activo:
            errores.append(f"SKU {art.sku} is inactive in the item master.")
        if d["unidad"] and d["unidad"] != art.unidad:
            errores.append(f"Unit {d['unidad']} does not match the item master ({art.unidad}).")
        d.update(
            articulo_id=art.id, upc=art.upc, estilo=art.estilo, color=art.color, talla=art.talla,
            descripcion=art.descripcion, marca=art.marca.codigo, grupo=art.grupo.codigo,
            categoria=art.grupo.categoria, tipo_empaque=art.tipo, unidad=art.unidad,
            prepack=art.prepack.codigo if art.prepack else None,
            unidades_por_caja=art.prepack.total if art.prepack else None,
        )
        errores += validar_empaque(art.tipo, d["casepack"], d["inner_pack"], d["cantidad"])
        if d["cantidad"] and (msg := error_cantidad(d["cantidad"], art.unidad)):
            errores.append(msg)
        d["pais_origen_pos"] = (producto_de(art).pais_origen if producto_de(art) else None) or d["pais_origen"]
    return d, errores


def validar_empaque(tipo: str, casepack: int | None, inner_pack: int | None, cantidad: int | None) -> list[str]:
    """Reglas del empaque de una posición:
    - un prepack ya es una caja definida (su curva): no lleva casepack ni inner pack;
    - con casepack e inner pack, el casepack es múltiplo del inner pack;
    - con inner pack, la cantidad de la posición es múltiplo del inner pack
      (todos los inner packs llevan la misma cantidad)."""
    errores = []
    if tipo == "PREPACK":
        if casepack or inner_pack:
            errores.append("A prepack is already a defined carton (its size run): it takes no casepack or inner pack.")
        return errores
    if casepack and inner_pack and casepack % inner_pack:
        errores.append(f"The casepack ({casepack}) must be a multiple of the inner pack ({inner_pack}).")
    if inner_pack and cantidad and cantidad % inner_pack:
        errores.append(f"The quantity ({cantidad}) must be a multiple of the inner pack ({inner_pack}): "
                       "all inner packs carry the same quantity.")
    return errores


def _valores_posicion(d: dict) -> dict:
    v = {c: d.get(c) for c in CAMPOS_POSICION}
    v["pais_origen"] = d.get("pais_origen_pos") or d.get("pais_origen")
    return v


def _clasificar(db: Session, filas: list[dict]) -> list[dict]:
    m = Maestros(db)
    vistos: set = set()
    cabeceras: dict = {}
    resultado = []
    normalizadas = [(r, *_normalizar(r, m)) for r in filas]

    claves_oc = {(d["proveedor"], d["oc"]) for _, d, _ in normalizadas}
    ocs_existentes = {}
    for codigo, numero in claves_oc:
        prov = m.proveedores.get(codigo)
        if prov:
            oc = db.scalar(select(OrdenCompra).where(OrdenCompra.proveedor_id == prov.id, OrdenCompra.numero == numero))
            if oc:
                ocs_existentes[(codigo, numero)] = oc
    ids_pos = [p.id for oc in ocs_existentes.values() for p in oc.posiciones]
    facturado = facturado_por_posicion(db, ids_pos)

    for registro, d, errores in normalizadas:
        clave = f"{d['proveedor']} / PO {d['oc']} / line {d['posicion']}"
        salida = {"fila": registro.get("_fila"), "clave": clave, "datos": registro,
                  "mensajes": list(errores), "cambios": {}}
        if d["proveedor"] and d["proveedor"] not in m.proveedores:
            salida["mensajes"].append(f"Supplier {d['proveedor']} does not exist.")
        k = (d["proveedor"], d["oc"], d["posicion"])
        if k in vistos:
            salida["mensajes"].append("Line repeated within the file.")
        vistos.add(k)
        cab = {c: d.get(c) for c in CAMPOS_CABECERA}
        previa = cabeceras.setdefault((d["proveedor"], d["oc"]), cab)
        if previa != cab:
            salida["mensajes"].append("The header data does not match other rows of the same PO.")
        if salida["mensajes"]:
            salida["estado"] = "error"
            resultado.append(salida)
            continue

        oc = ocs_existentes.get((d["proveedor"], d["oc"]))
        pos = next((p for p in oc.posiciones if p.posicion == d["posicion"]), None) if oc else None
        if not pos:
            salida["estado"] = "nuevo"
            resultado.append(salida)
            continue

        cambios = {}
        for c in CAMPOS_CABECERA:
            anterior = getattr(oc, c)
            if anterior != d.get(c):
                cambios[c] = {"antes": anterior, "despues": d.get(c)}
        for c, nuevo in _valores_posicion(d).items():
            if c == "articulo_id":
                continue
            anterior = getattr(pos, c)
            igual = abs(anterior - nuevo) < 1e-9 if isinstance(anterior, float) and nuevo is not None else anterior == nuevo
            if not igual:
                cambios[c] = {"antes": anterior, "despues": nuevo}
        salida["cambios"] = cambios
        fact = facturado.get(pos.id, 0)
        if not cambios:
            salida["estado"] = "sin_cambio"
        elif fact and d["cantidad"] is not None and d["cantidad"] < fact:
            salida["estado"] = "conflicto"
            salida["mensajes"].append(f"The new quantity ({d['cantidad']}) is less than what is already invoiced ({fact}).")
        elif fact and any(c in cambios for c in ("sociedad", "moneda", "centro", "unidad", "codigo_sap")):
            salida["estado"] = "conflicto"
            salida["mensajes"].append(
                "Key data (company, currency, plant, SKU or unit) changes on an already invoiced line.")
        elif fact and any(c in cambios for c in ("casepack", "inner_pack")):
            salida["estado"] = "conflicto"
            salida["mensajes"].append("The casepack or inner pack changes on a line that is already invoiced.")
        else:
            salida["estado"] = "cambio"
        resultado.append(salida)
    return resultado


def _resumen(clasificadas: list[dict]) -> dict:
    res = {"nuevo": 0, "cambio": 0, "sin_cambio": 0, "conflicto": 0, "error": 0}
    for c in clasificadas:
        res[c["estado"]] += 1
    res["total"] = len(clasificadas)
    return res


def importar_previa(db: Session, user: Usuario, nombre: str, contenido: bytes) -> dict:
    exigir(user, "oc.importar")
    filas = _leer_archivo(nombre, contenido)
    if not filas:
        raise ErrorNegocio("The file has no rows with data.", 422, "archivo_vacio")
    clasificadas = _clasificar(db, filas)
    imp = ImportacionOC(usuario_id=user.id, nombre_archivo=nombre, filas=filas)
    db.add(imp)
    db.flush()
    return {
        "importacion_id": imp.id,
        "archivo": nombre,
        "resumen": _resumen(clasificadas),
        "filas": [{k: v for k, v in c.items() if k != "datos"} for c in clasificadas],
    }


def importar_aplicar(db: Session, user: Usuario, importacion_id: int) -> dict:
    exigir(user, "oc.importar")
    imp = db.get(ImportacionOC, importacion_id)
    if not imp:
        raise ErrorNegocio("The import does not exist.", 404, "no_encontrado")
    if imp.estado == "APLICADA":
        return {"ya_aplicada": True, **(imp.resultado or {})}
    # Se vuelve a clasificar: los datos pudieron cambiar desde la vista previa
    clasificadas = _clasificar(db, imp.filas)
    m = Maestros(db)
    aplicadas = 0
    ocs_tocadas: dict[int, dict] = {}
    for c in clasificadas:
        if c["estado"] == "conflicto":
            d, _ = _normalizar(c["datos"], m)
            prov = m.proveedores.get(d["proveedor"])
            db.add(Alerta(proveedor_id=prov.id if prov else None, tipo="conflicto_oc",
                          mensaje=f"{c['clave']}: " + " ".join(c["mensajes"]),
                          referencia={"importacion_id": imp.id, "fila": c["fila"]}))
        if c["estado"] not in ("nuevo", "cambio"):
            continue
        d, _ = _normalizar(c["datos"], m)
        prov = m.proveedores[d["proveedor"]]
        oc = db.scalar(select(OrdenCompra).where(OrdenCompra.proveedor_id == prov.id, OrdenCompra.numero == d["oc"]))
        nueva = oc is None
        if nueva:
            # Sin liberación logística hasta decidirla al final (abajo)
            oc = OrdenCompra(proveedor_id=prov.id, numero=d["oc"], liberacion_logistica=m.lib.logistica.sin_liberar() or "")
            db.add(oc)
        info = ocs_tocadas.setdefault(id(oc), {
            "oc": oc, "cambios": False, "explicita": d["liberacion_logistica_archivo"],
            "actual": None if nueva else oc.liberacion_logistica})
        info["cambios"] = info["cambios"] or (c["estado"] == "cambio")
        for campo in CAMPOS_CABECERA:
            setattr(oc, campo, d.get(campo))
        # Fechas de liberación: las del archivo si vienen; si no, el día en que se liberó
        for campo in ("fecha_lib_comercial", "fecha_lib_logistica"):
            if d.get(campo):
                setattr(oc, campo, d[campo])
        oc.actualizado_en = ahora()
        db.flush()
        pos = db.scalar(select(PosicionOC).where(PosicionOC.oc_id == oc.id, PosicionOC.posicion == d["posicion"]))
        if not pos:
            pos = PosicionOC(oc_id=oc.id, posicion=d["posicion"])
            db.add(pos)
        for campo, valor in _valores_posicion(d).items():
            setattr(pos, campo, valor)
        aplicadas += 1
    for info in ocs_tocadas.values():
        oc = info["oc"]
        oc.liberacion_logistica = m.lib.logistica_final(oc.liberacion_comercial, info["explicita"], info["actual"],
                                                        info["cambios"])
        oc.liberada = m.lib.liberada(oc.liberacion_comercial, oc.liberacion_logistica)
        m.lib.fechar(oc)
    resultado = {"resumen": _resumen(clasificadas), "aplicadas": aplicadas}
    imp.estado = "APLICADA"
    imp.resultado = resultado
    registrar(db, user, "importacion_oc", imp.id, "aplicar", resultado)
    return resultado


# ---- Alta desde el formulario ------------------------------------------------------
# El formulario arma las mismas filas que un archivo de carga: una por línea con
# los datos de la cabecera; se validan y se guardan con el mismo proceso.
CAMPOS_FORM_CAB = ("proveedor", "oc", "sociedad", "centro", "centro_destino", "moneda", "incoterm", "fecha_oc",
                   "puerto_despacho", "pais_origen", "pais_procedencia", "fecha_xf", "fecha_tienda",
                   "liberacion_comercial", "liberacion_logistica")
CAMPOS_FORM_LINEA = ("posicion", "codigo_sap", "cantidad", "precio", "unidad", "casepack", "inner_pack", "almacen",
                     "fecha_entrega")


def crear_oc(db: Session, user: Usuario, datos: dict) -> dict:
    """Orden de compra creada en la plataforma: misma estructura y mismas
    validaciones que la carga masiva."""
    exigir(user, "oc.importar")
    cab = {k: _texto(v) for k, v in (datos.get("cabecera") or {}).items() if k in CAMPOS_FORM_CAB}
    if user.proveedor_id:  # un proveedor solo crea OCs propias
        prov = db.get(Proveedor, user.proveedor_id)
        cab["proveedor"] = prov.codigo if prov else ""
    lineas = [x for x in (datos.get("lineas") or []) if any(_texto(v) for v in x.values())]
    if not lineas:
        raise ErrorNegocio("Add at least one line.", 422, "validacion", [{"campo": "lineas", "mensaje": "Add at least one line."}])
    filas = []
    for i, linea in enumerate(lineas, start=1):
        fila = {**cab, **{k: _texto(v) for k, v in linea.items() if k in CAMPOS_FORM_LINEA}, "_fila": i}
        fila["posicion"] = fila.get("posicion") or str(i * 10)
        filas.append(fila)
    m = Maestros(db)
    prov = m.proveedores.get((cab.get("proveedor") or "").upper())
    if prov and cab.get("oc") and db.scalar(select(OrdenCompra.id).where(
            OrdenCompra.proveedor_id == prov.id, OrdenCompra.numero == cab["oc"])):
        raise ErrorNegocio(f"PO {cab['oc']} already exists for this supplier.", 409, "duplicado",
                           [{"campo": "oc", "mensaje": f"PO {cab['oc']} already exists for this supplier."}])
    clasificadas = _clasificar(db, filas)
    malas = [c for c in clasificadas if c["estado"] in ("error", "conflicto")]
    if malas:
        # Lo que falla en todas las líneas es de la cabecera; lo demás, de su línea
        comunes = set.intersection(*(set(c["mensajes"]) for c in clasificadas)) if len(malas) == len(clasificadas) else set()
        errores = [{"campo": "cabecera", "mensaje": x} for x in sorted(comunes)]
        for c in malas:
            errores += [{"campo": f"lineas.{c['fila'] - 1}", "mensaje": f"Line {c['fila']}: {x}"}
                        for x in c["mensajes"] if x not in comunes]
        raise ErrorNegocio("Check the purchase order.", 422, "validacion", errores)
    imp = ImportacionOC(usuario_id=user.id, nombre_archivo="(form)", filas=filas)
    db.add(imp)
    db.flush()
    resultado = importar_aplicar(db, user, imp.id)
    oc = db.scalar(select(OrdenCompra).where(OrdenCompra.proveedor_id == m.proveedores[cab["proveedor"].upper()].id,
                                             OrdenCompra.numero == cab["oc"]))
    return {"oc_id": oc.id, "numero": oc.numero, "lineas": resultado["aplicadas"]}


def moneda_base() -> str | None:
    """Moneda de la empresa (Configuración → Empresa) o la primera de la lista."""
    from app.core.empresa import configuracion_actual

    return (configuracion_actual().get("preferencias") or {}).get("moneda") or next(iter(listas.codigos("moneda")), None)


def plantilla_oc(db: Session, user: Usuario) -> bytes:
    """Plantilla de carga de OC en Excel, con las fechas de ejemplo en el
    formato del usuario (el mismo con el que se leerán)."""
    from datetime import timedelta

    from app.modulos.acceso.preferencias import actual, fecha_txt
    from app.modulos.documentos.plantillas import plantilla

    hoy = date.today()
    fechas = {"fecha_oc": hoy, "fecha_xf_original": hoy + timedelta(days=45), "fecha_xf": hoy + timedelta(days=50),
              "fecha_tienda": hoy + timedelta(days=100), "fecha_lib_comercial": None, "fecha_lib_logistica": None,
              "fecha_entrega": None}
    ayudas = {"proveedor": "Supplier code.", "oc": "PO number (letters and numbers, up to 40).", "posicion": "Line number.",
              "codigo_sap": "Item code from the item master.", "cantidad": "Quantity in the item unit.",
              "sociedad": "Company (bill to). Optional; needed to invoice.", "moneda": "Needed if there is a price.",
              "precio": "Optional; needed to invoice.", "casepack": "Solids: exact quantity per carton (optional).",
              "inner_pack": "Usually defined later in the packing list.",
              "liberacion_comercial": "",
              "liberacion_logistica": ""}
    # Códigos de liberación de la empresa (Datos maestros → Estados de liberación)
    lib = liberaciones.de(db)
    ayudas["liberacion_comercial"] = f"{lib.comercial.ayuda()} (or the words)."
    ayudas["liberacion_logistica"] = f"{lib.logistica.ayuda()} (or the words)."
    formato = actual()["formato_fecha"]
    columnas = []
    for campo, alias in ALIAS.items():
        ayuda = ayudas.get(campo, "")
        if campo.startswith("fecha"):
            ayuda = (ayuda + " " if ayuda else "") + f"Date as {formato} (your profile setting) or YYYY-MM-DD."
        columnas.append({"nombre": alias[0], "req": campo in REQUERIDOS, "ayuda": ayuda})
    base = {"proveedor": "SUPPLIER", "oc": "PO-0001", "sociedad": "",
            "moneda": moneda_base() or "", "incoterm": next(iter(listas.codigos("incoterm")), ""),
            "liberacion_comercial": lib.comercial.predeterminado() or "",
            "liberacion_logistica": next((x.codigo for x in lib.logistica.lista if x.libera and not x.con_cambios), ""), "unidad": "", "precio": 10.5, "casepack": 12}
    ejemplos = []
    for linea, cant in ((10, 48), (20, 36)):
        fila = {**base, "posicion": linea, "codigo_sap": "ITEM-CODE", "cantidad": cant,
                **{k: fecha_txt(v) for k, v in fechas.items() if v}}
        ejemplos.append([fila.get(c, "") for c in ALIAS])
    return plantilla("Purchase orders", columnas, ejemplos, [
        "One row per PO line; the header data (supplier, PO, company…) is repeated on each line.",
        "Columns with * are required. Company, currency and price can be completed later; they are needed to invoice.",
        f"Dates: {formato} — the format chosen in your profile — or YYYY-MM-DD. Excel date cells are also read.",
    ])


def opciones_formulario(db: Session, user: Usuario) -> dict:
    """Listas para el formulario de OC (códigos, como en el archivo de carga)."""
    exigir(user, "oc.importar")
    provs = db.scalars(select(Proveedor).where(Proveedor.activo.is_(True),
                                               *([Proveedor.id == user.proveedor_id] if user.proveedor_id else []))
                       .order_by(Proveedor.nombre)).all()
    op = lambda xs, f=lambda x: x.nombre: [{"valor": x.codigo, "texto": f"{x.codigo} · {f(x)}"} for x in xs]  # noqa: E731
    # Las monedas de la lista y las que ya traen las OCs (de antes de la lista)
    usadas = {m for (m,) in db.execute(select(OrdenCompra.moneda).distinct()) if m}
    monedas = listas.codigos("moneda") + sorted(usadas - set(listas.codigos("moneda")))
    return {
        # Cada proveedor lleva sus sociedades: el formulario solo ofrece esas
        "proveedores": [{**o, "sociedades": [x.codigo for x in p.sociedades]} for o, p in zip(op(provs), provs)],
        "sociedades": op(db.scalars(select(Sociedad).order_by(Sociedad.codigo))),
        "centros": [{"valor": c.codigo, "texto": f"{c.codigo} · {c.nombre}", "sociedad": c.sociedad.codigo if c.sociedad else None}
                    for c in db.scalars(select(Centro).order_by(Centro.codigo))],
        "almacenes": [{"valor": a.codigo, "texto": f"{a.codigo} · {a.nombre}", "sociedad": a.sociedad.codigo if a.sociedad else None}
                      for a in db.scalars(select(Almacen).order_by(Almacen.codigo))],
        "puertos": op(db.scalars(select(Puerto).order_by(Puerto.codigo))),
        "paises": op(db.scalars(select(Pais).order_by(Pais.nombre))),
        "monedas": [{"valor": m, "texto": f"{m} · {listas.nombre('moneda', m)}" if listas.valor("moneda", m) else m}
                    for m in monedas],
        "incoterms": [{"valor": x["codigo"], "texto": f"{x['codigo']} · {x['nombre']}"} for x in listas.valores("incoterm")],
        "moneda_base": moneda_base(),
    }


def articulos_formulario(db: Session, user: Usuario, proveedor: str, q: str = "") -> list[dict]:
    """Artículos activos del proveedor para las líneas de la OC."""
    exigir(user, "oc.importar")
    prov = db.scalar(select(Proveedor).where(Proveedor.codigo == (proveedor or "").upper()))
    if not prov or (user.proveedor_id and prov.id != user.proveedor_id):
        return []
    consulta = select(Articulo).where(Articulo.proveedor_id == prov.id, Articulo.activo.is_(True))
    if terminos(q):
        consulta = consulta.where(filtro_texto(q, lambda t: [Articulo.sku.ilike(t), Articulo.estilo.ilike(t), Articulo.color.ilike(t),
                                                             Articulo.sku_proveedor.ilike(t), Articulo.upc.ilike(t)]))
    return [{"valor": a.sku, "texto": f"{a.sku} · {a.estilo} {a.color or ''} {a.talla or ''}".strip(), "unidad": a.unidad,
             "tipo": a.tipo} for a in db.scalars(consulta.order_by(Articulo.estilo, Articulo.color, Articulo.sku).limit(300))]
