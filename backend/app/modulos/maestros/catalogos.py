"""Mantenimiento de datos maestros desde un solo lugar.

Cada catálogo se describe una vez (campos, tipos, obligatorios, filtros) y
el mismo código sirve para listar, filtrar, ordenar, crear, editar y
eliminar. La interfaz usa esa misma descripción para dibujar formularios y
tablas, así que agregar un campo aquí lo agrega en pantalla.
"""
import csv
import io
import json

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.core.errores import ErrorNegocio
from app.modelos import (
    AcuerdoComercial,
    Almacen,
    Articulo,
    CategoriaArticulo,
    Centro,
    Contacto,
    EscalaTalla,
    EstadoLiberacion,
    GrupoArticulo,
    Marca,
    OrdenCompra,
    Pais,
    PasoLeadTime,
    PosicionOC,
    Prepack,
    PrepackComponente,
    Proveedor,
    Puerto,
    RegionLeadTime,
    ReglaLeadTime,
    Sociedad,
    TipoEmpaque,
    TipoUnidad,
    Transportista,
    Usuario,
)
from app.modulos.acceso.permisos import exigir
from app.modulos.comun.historial import registrar
from app.modulos.comun.normalizar import Referencias
from app.modulos.comun.normalizar import nombre as nombre_fmt
from app.modulos.comun.normalizar import texto as texto_fmt
from app.modulos.comun.texto import filtro_texto
from app.modulos.maestros.unidades import opciones as _opciones_unidad
from app.modulos.productos.productos import (
    asegurar_producto,
    codigo_valido,
    fmt_codigo,
    generico_de,
    producto_de,
    producto_por_generico,
)
from app.modulos.transporte import reglas_lt as rlt

UNIDADES = _opciones_unidad()


def c(nombre, etiqueta, tipo="texto", obligatorio=False, **extra):
    return {"nombre": nombre, "etiqueta": etiqueta, "tipo": tipo, "obligatorio": obligatorio, **extra}


# Dependencias entre campos (configurables, no propias de una empresa):
#   depende={"campo": <campo del mismo formulario>, "clave": <dato de la opción>}
# La opción solo vale si su dato `clave` coincide con el valor del campo (o lo
# contiene, si es una lista). Ej.: la marca de un artículo depende del
# proveedor (la marca trae la lista de proveedores que la manejan). El
# formulario filtra las opciones y el servidor lo valida al guardar.
# tipo: texto | entero | numero | bool | opcion (opciones) | ref (catalogo: guarda el id)
#       | codigo (catalogo: guarda el código, p. ej. país ISO) | correos
#       | multi (catalogo: varios registros, p. ej. las marcas de un proveedor)
MODOS = [["MARITIMO", "Ocean"], ["AEREO", "Air"], ["TERRESTRE", "Road"]]
CATALOGOS = {
    "sociedades": {
        "modelo": Sociedad, "titulo": "Companies", "singular": "company",
        "ayuda": "Companies that are invoiced (the PO company). Each one has its assigned plants.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=10, mayus=True),
            c("nombre", "Name", obligatorio=True, formato="nombre"),
            c("razon_social", "Legal name"),
            c("id_fiscal", "Tax ID"),
            c("pais", "Country", "codigo", catalogo="paises", filtro=True),
            c("moneda", "Currency", max=3, mayus=True, ayuda="Default currency of its POs and invoices."),
            c("direccion", "Address"),
            c("correos", "Billing emails", "correos", ayuda="One or more, separated by commas."),
            c("activa", "Active", "bool", filtro=True),
        ],
        "extras": [{"nombre": "centros_txt", "etiqueta": "Plants", "catalogo": "centros", "filtro": "sociedad_id"},
                   {"nombre": "contactos", "etiqueta": "Contacts", "catalogo": "contactos", "filtro": "sociedad_id"}],
        "buscar": ["codigo", "nombre", "razon_social", "id_fiscal"],
    },
    "centros": {
        "modelo": Centro, "titulo": "Plants", "singular": "plant",
        "ayuda": "Plants assigned to each company. A plant is the shipment notify party and, as the PO destination center "
                 "(e.g. 2220), it tells which country the goods go to. Its port is the arrival port.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=10, mayus=True),
            c("sociedad_id", "Company", "ref", obligatorio=True, catalogo="sociedades", filtro=True),
            c("nombre", "Name", obligatorio=True, formato="nombre"),
            c("pais", "Country", "codigo", obligatorio=True, catalogo="paises", filtro=True),
            c("puerto", "Arrival port", "codigo", catalogo="puertos", filtro=True, depende={"campo": "pais", "clave": "pais"}),
            c("tipo", "Type", "opcion", obligatorio=True,
              opciones=[["BODEGA_FISCAL", "Bonded warehouse"], ["ZONA_FRANCA", "Free trade zone"], ["LOCAL", "Local warehouse"],
                        ["TIENDA", "Store / DC"]]),
            c("puertos", "Other arrival ports", "multi", catalogo="puertos", depende={"campo": "pais", "clave": "pais"},
              ayuda="Besides the main one. The shipment suggests the main port and lets you switch between these."),
            c("direccion", "Address"),
            c("correos", "Emails (notify)", "correos", ayuda="One or more, separated by commas."),
            c("activo", "Active", "bool", filtro=True),
        ],
        "extras": [{"nombre": "contactos", "etiqueta": "Contacts", "catalogo": "contactos", "filtro": "centro_id"}],
        "buscar": ["codigo", "nombre"],
    },
    "almacenes": {
        "modelo": Almacen, "titulo": "Storage locations", "singular": "storage location",
        "ayuda": "Inventory split in the system (virtual, retail, wholesale); it is not a physical place. "
                 "Each PO line states its own.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=10, mayus=True),
            c("sociedad_id", "Company", "ref", obligatorio=True, catalogo="sociedades", filtro=True),
            c("nombre", "Name", obligatorio=True, formato="nombre"),
            c("tipo", "Type", "opcion", obligatorio=True, filtro=True,
              opciones=[["VIRTUAL", "Virtual"], ["DETALLE", "Retail"], ["MAYOREO", "Wholesale"]]),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre"],
    },
    "contactos": {
        "modelo": Contacto, "titulo": "Contacts", "singular": "contact",
        "ayuda": "Contact people of a company (billing) or of a plant (notify party). "
                 "They appear on the invoice and the packing list.",
        "campos": [
            c("nombre", "Name", obligatorio=True, formato="nombre"),
            c("cargo", "Job title", formato="nombre"),
            c("rol", "Role", "opcion", obligatorio=True, filtro=True,
              opciones=[["FACTURACION", "Billing"], ["NOTIFY", "Notify party"], ["LOGISTICA", "Logistics"]]),
            c("sociedad_id", "Company", "ref", catalogo="sociedades", filtro=True),
            c("centro_id", "Plant", "ref", catalogo="centros", filtro=True, depende={"campo": "sociedad_id", "clave": "sociedad_id"}),
            c("correos", "Emails", "correos", ayuda="One or more, separated by commas."),
            c("telefono", "Phone"),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["nombre", "cargo", "correos"],
    },
    "paises": {
        "modelo": Pais, "titulo": "Countries", "singular": "country",
        "ayuda": "Countries of origin and shipment (2-letter ISO).",
        "campos": [
            c("codigo", "ISO code", obligatorio=True, max=2, mayus=True, patron=r"^[A-Z]{2}$",
              mensaje_patron="Use the 2-letter ISO code."),
            c("nombre", "Name", obligatorio=True, formato="nombre"),
            c("region", "Lead time region", "codigo", catalogo="regiones", filtro=True,
              ayuda="Origin region (e.g. ASIA). Lead time plans can apply to the whole region."),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre"],
    },
    "acuerdos": {
        "modelo": AcuerdoComercial, "titulo": "Trade agreements", "singular": "agreement",
        "ayuda": "Trade agreements in force: goods originating in the origin countries enter the destination countries "
                 "with a preferential duty when the proof of origin is presented. The technical sheet shows them "
                 "for the product's country of origin.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=20, mayus=True),
            c("nombre", "Name", obligatorio=True),
            c("origenes", "Origin countries (ISO)", obligatorio=True, mayus=True, patron=r"^[A-Z]{2}(\s*,\s*[A-Z]{2})*$",
              mensaje_patron="ISO codes separated by commas, e.g. US, DO.", ayuda="ISO codes separated by commas."),
            c("destinos", "Destination countries (ISO)", obligatorio=True, mayus=True, patron=r"^[A-Z]{2}(\s*,\s*[A-Z]{2})*$",
              mensaje_patron="ISO codes separated by commas, e.g. GT, SV.", ayuda="Destination countries where it applies."),
            c("prueba", "Proof of origin"),
            c("nota", "Note"),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre", "origenes", "destinos"],
    },
    "regiones": {
        "modelo": RegionLeadTime, "titulo": "Regions", "singular": "region",
        "ayuda": "Groups of origin countries (Asia, Central America…) to filter, compare and apply one lead time "
                 "plan to the whole region.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=10, mayus=True),
            c("nombre", "Name", obligatorio=True, formato="nombre"),
            c("predeterminada", "Default for other origins", "bool"),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre"],
    },
    "pasos_lt": {
        "modelo": PasoLeadTime, "titulo": "Lead time steps", "singular": "lead time step",
        "ayuda": "Catalog of the steps a lead time can use (booking, release, XF, ETD, ETA, customs…). Each "
                 "configuration chooses which steps apply, their reference and days. Linking a step to a "
                 "measured date lets the system compare the plan with what really happened.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=20, mayus=True),
            c("nombre", "Name", obligatorio=True),
            c("hito", "Measured date", "opcion", filtro=True, opciones=[[h, n] for h, n in rlt.HITOS],
              ayuda="Optional. The date the system records for this step (to compare plan and reality)."),
            c("descripcion", "Description"),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre", "descripcion"],
    },
    "leadtimes": {
        "modelo": ReglaLeadTime, "titulo": "Lead time rules", "singular": "lead time rule",
        "ayuda": "Lead times by geographic level with inheritance: Port > Country > Region > Global. Each rule only "
                 "defines what it changes (steps it adds, overrides or removes, and the order); the rest is "
                 "inherited from the level above. See the result in Lead time › Effective lead time.",
        "campos": [
            c("nivel", "Level", "opcion", obligatorio=True, filtro=True, opciones=[[n, t] for n, t in rlt.NIVELES]),
            c("region", "Region", "codigo", catalogo="regiones", filtro=True, mostrar_si={"campo": "nivel", "valores": ["REGION"]}),
            c("pais", "Country", "codigo", catalogo="paises", filtro=True, mostrar_si={"campo": "nivel", "valores": ["PAIS"]}),
            c("puerto", "Port", "codigo", catalogo="puertos", mostrar_si={"campo": "nivel", "valores": ["PUERTO"]}),
            c("nombre", "Name", obligatorio=True),
            c("pasos", "Steps", "regla_lt", obligatorio=True,
              ayuda="Inherited steps appear in grey: override only what changes at this level."),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["nombre", "region", "pais", "puerto"],
    },
    "puertos": {
        "modelo": Puerto, "titulo": "Ports", "singular": "port",
        "ayuda": "Sea ports and airports (UN/LOCODE).",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=10, mayus=True),
            c("nombre", "Name", obligatorio=True, formato="nombre"),
            c("pais", "Country", "codigo", obligatorio=True, catalogo="paises", filtro=True),
            c("tipo", "Type", "opcion", obligatorio=True, filtro=True,
              opciones=[["MARITIMO", "Ocean"], ["AEREO", "Air"], ["TERRESTRE", "Road"]]),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre"],
    },
    "tipos_empaque": {
        "modelo": TipoEmpaque, "titulo": "Packaging types", "singular": "packaging type",
        "ayuda": "Configurable logistics units (inner pack, carton, pallet, bag, drum…) and which ones each can "
                 "contain. A packing list is a tree of packaging inside packaging with the products inside: the "
                 "net weight comes from each item's unit weight, and each level adds its tare.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=20, mayus=True),
            c("nombre", "Name", obligatorio=True),
            c("nivel", "Level", "entero", obligatorio=True, minimo=1, ayuda="1 = the innermost."),
            c("prefijo", "Label prefix", max=6, mayus=True, ayuda="Numbering of each unit, e.g. C → C1, C2…"),
            c("cuenta_como", "Counts as", "opcion", obligatorio=True, filtro=True, opciones=[
                ["BULTO", "Package (carton)"], ["INTERIOR", "Inner unit (inside a package)"],
                ["SOPORTE", "Support that carries packages (pallet)"]]),
            c("contiene", "Can contain", "multi", catalogo="tipos_empaque",
              ayuda="Packaging types that can go inside this one."),
            c("contiene_productos", "Holds products directly", "bool"),
            c("largo", "Length (cm)", "numero", minimo=0, ayuda="Outer dimensions by default; each unit can change them."),
            c("ancho", "Width (cm)", "numero", minimo=0),
            c("alto", "Height (cm)", "numero", minimo=0),
            c("tara", "Tare (kg)", "numero", minimo=0, ayuda="Weight of the empty packaging."),
            c("peso_max", "Max gross weight (kg)", "numero", minimo=0),
            c("max_unidades", "Max product units", "entero", minimo=1),
            c("max_contenido", "Max inner packaging", "entero", minimo=1),
            c("mezcla_productos", "Can mix products", "bool"),
            c("mezcla_tallas", "Can mix sizes", "bool"),
            c("mezcla_oc", "Can mix POs", "bool"),
            c("uom", "UoM", max=10, mayus=True),
            c("identificador", "Identifier", "opcion", opciones=[
                ["SERIAL", "Serial number"], ["CODIGO_BARRAS", "Barcode"], ["SSCC", "SSCC (GS1)"]]),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre", "prefijo"],
    },
    "marcas": {
        "modelo": Marca, "titulo": "Brands", "singular": "brand",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=10, mayus=True),
            c("nombre", "Name", obligatorio=True),
            c("activa", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre"],
    },
    "escalas": {
        "modelo": EscalaTalla, "titulo": "Size scales", "singular": "size scale",
        "ayuda": "Reusable size structures for any product (footwear, apparel, accessories…). A generic starts from a "
                 "scale: its sizes come in order and each size code is generated with the scale's rule or typed in the list.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=20, mayus=True),
            c("nombre", "Name", obligatorio=True),
            c("categoria", "Category", "codigo", catalogo="categorias", filtro=True),
            c("regla", "Size code rule", "opcion", obligatorio=True, filtro=True, opciones=[
                ["MULTIPLICAR", "Size × factor (7.5 × 10 → 075)"], ["CONSECUTIVO", "Consecutive (001, 002…)"],
                ["TALLA", "Same as the size (S, M, XL)"]]),
            c("factor", "Factor", "entero", minimo=1, ayuda="For Size × factor, e.g. 10."),
            c("longitud", "Code digits", "entero", minimo=1, ayuda="Zeros on the left, e.g. 3 → 070."),
            c("tallas", "Sizes", obligatorio=True,
              ayuda="In order, separated by commas. Ranges allowed (6-10, 6.5-9.5). A fixed code with = (7=070)."),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre", "tallas"],
    },
    "categorias": {
        "modelo": CategoriaArticulo, "titulo": "Item categories", "singular": "category",
        "ayuda": "Your company's product lines (footwear, apparel, accessories…). Item groups and size scales use them.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=10, mayus=True),
            c("nombre", "Name", obligatorio=True),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre"],
    },
    "liberaciones": {
        "modelo": EstadoLiberacion, "titulo": "Release statuses", "singular": "release status",
        "ayuda": "The codes your ERP sends for the commercial and logistics release of a PO (e.g. SAP: C/P and "
                 "300/301/304). A PO can be invoiced only when both releases are given.",
        "campos": [
            c("tipo", "Release", "opcion", obligatorio=True, filtro=True,
              opciones=[["COMERCIAL", "Commercial"], ["LOGISTICA", "Logistics"]]),
            c("codigo", "Code", obligatorio=True, max=10, mayus=True, ayuda="As it comes from your ERP."),
            c("nombre", "Name", obligatorio=True),
            c("libera", "Released", "bool", ayuda="With this status the release is given."),
            c("con_cambios", "Released with later changes", "bool",
              ayuda="Logistics: status a released PO moves to when it changes afterwards."),
            c("predeterminado", "Default", "bool", ayuda="Used when the file does not bring this release."),
            c("alias", "Other words", max=300, ayuda="Other words the importer accepts, separated by commas."),
            c("orden", "Order", "entero", minimo=0),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre", "alias"],
    },
    "grupos": {
        "modelo": GrupoArticulo, "titulo": "Item groups", "singular": "group",
        "ayuda": "Each item belongs to a single group. The category sets the packing rule; it does not define the "
                 "product type, which comes from the technical sheet.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=15, mayus=True),
            c("nombre", "Name", obligatorio=True, formato="nombre"),
            c("categoria", "Category", "codigo", catalogo="categorias", obligatorio=True, filtro=True),
            c("dias_extra", "Extra days after arrival", "numero", minimo=0,
              ayuda="Handling this product type needs after the port (inspection, labeling, permits). "
                    "It is added to the in-store estimate; empty = none."),
            c("activo", "Active", "bool", filtro=True),
        ],
        "extras": [{"nombre": "articulos", "etiqueta": "Items", "catalogo": "articulos", "filtro": "grupo_id"}],
        "buscar": ["codigo", "nombre"],
    },
    "proveedores": {
        "modelo": Proveedor, "titulo": "Suppliers", "singular": "supplier",
        "ayuda": "Exporters. Each one handles its own brands and items and works with the assigned companies.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=30, mayus=True),
            c("nombre", "Name", obligatorio=True),
            c("razon_social", "Legal name"),
            c("id_fiscal", "Tax ID"),
            c("pais", "Country", "codigo", catalogo="paises", filtro=True),
            c("direccion", "Address"),
            c("contacto", "Contact", formato="nombre"),
            c("correos", "Emails", "correos", ayuda="One or more, separated by commas."),
            c("telefono", "Phone"),
            c("marcas", "Brands handled", "multi", catalogo="marcas", filtro=True),
            c("sociedades", "Companies it works with", "multi", catalogo="sociedades", filtro=True),
            c("activo", "Active", "bool", filtro=True),
        ],
        "extras": [{"nombre": "articulos", "etiqueta": "Items", "catalogo": "articulos", "filtro": "proveedor_id"}],
        "buscar": ["codigo", "nombre", "razon_social"],
    },
    "transportistas": {
        "modelo": Transportista, "titulo": "Carriers", "singular": "carrier",
        "ayuda": "Registered shipping lines, airlines and road carriers, with the companies they work for.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=20, mayus=True),
            c("nombre", "Name", obligatorio=True),
            c("tipo", "Type", "opcion", obligatorio=True, filtro=True, opciones=MODOS + [["MULTIMODAL", "Multimodal"]]),
            c("codigo_internacional", "SCAC / IATA", max=10, mayus=True,
              ayuda="Shipping line SCAC or airline IATA prefix."),
            c("id_fiscal", "Tax ID"),
            c("pais", "Country", "codigo", catalogo="paises"),
            c("contacto", "Contact", formato="nombre"),
            c("correos", "Emails", "correos"),
            c("telefono", "Phone"),
            c("sociedades", "Companies it works with", "multi", catalogo="sociedades", ayuda="Empty = it works with every company.",
              filtro=True),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre", "codigo_internacional"],
    },
    "tipos_unidad": {
        "modelo": TipoUnidad, "titulo": "Unit types", "singular": "unit type",
        "ayuda": "Load units by transport mode with their capacity. A shipment only offers those of its mode; "
                 "the service (FCL, LCL…) belongs to each unit, so a shipment can be mixed.",
        "campos": [
            c("codigo", "Code", obligatorio=True, max=10, mayus=True),
            c("nombre", "Name", obligatorio=True),
            c("modo", "Transport mode", "opcion", obligatorio=True, filtro=True, opciones=MODOS),
            c("modalidad", "Service", "opcion", obligatorio=True, filtro=True,
              opciones=[["FCL", "FCL · full container"], ["LCL", "LCL · consolidated cargo"],
                        ["AEREO", "Air cargo"], ["FTL", "FTL · full truck"], ["LTL", "LTL · partial load"]]),
            c("capacidad_cbm", "Max volume (m³)", "numero", minimo=0),
            c("capacidad_kg", "Max weight (kg)", "numero", minimo=0),
            c("requiere_sello", "Requires seal", "bool"),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "nombre"],
    },
    "articulos": {
        "modelo": Articulo, "titulo": "Items", "singular": "item",
        "ayuda": "Master data of each item: the item code (your company's own format) and the supplier SKU are different. "
                 "The description, the product type and the HS code come from the product's technical sheet (Products), "
                 "not from this form. Prepacks are created in the Prepacks tab and take the classification of their solids.",
        "campos": [
            c("sku", "Item code", obligatorio=True, max=40, mayus=True, patron=r"^[A-Za-z0-9][A-Za-z0-9._\-/]{0,39}$",
              mensaje_patron="Letters and numbers (also . - _ /), up to 40 characters."),
            c("generico", "Generic (style-color)", max=40, mayus=True,
              ayuda="Groups the sizes and prepacks of one style and color: they share the technical sheet. Empty = style-color."),
            c("sku_proveedor", "Supplier SKU", max=60, mayus=True,
              ayuda="The supplier's own code, e.g. VN0A4BV4W00-7."),
            c("estilo", "Style", obligatorio=True, mayus=True),
            c("color", "Color", ayuda="Empty if the item has no color."),
            c("talla", "Size / prepack ID", mayus=True,
              ayuda="For a prepack it is its prepack ID, e.g. AB12."),
            c("marca_id", "Brand", "ref", obligatorio=True, catalogo="marcas", filtro=True,
              depende={"campo": "proveedor_id", "clave": "proveedores"}),
            c("grupo_id", "Item group", "ref", obligatorio=True, catalogo="grupos", filtro=True,
              ayuda="Internal grouping and packing rule; the product type is set by the technical sheet."),
            c("proveedor_id", "Supplier", "ref", obligatorio=True, catalogo="proveedores", filtro=True,
              ayuda="Each supplier handles its items; the brand must be one of theirs."),
            c("tipo", "Type", "opcion", obligatorio=True, filtro=True,
              opciones=[["SOLIDO", "Solid"], ["PREPACK", "Prepack"]]),
            c("unidad", "Unit", "opcion", obligatorio=True, opciones=UNIDADES, filtro=True),
            c("peso_unitario", "Unit weight (kg)", "numero", minimo=0,
              ayuda="Net weight of one unit (pair, piece or, for a prepack, the whole size run). The packaging "
                    "weight is not included: each packaging level adds its own tare. Empty in a prepack = sum of its solids."),
            c("upc", "UPC"),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["sku", "sku_proveedor", "estilo", "color", "upc", "descripcion"],
    },
    "prepacks": {
        "modelo": Prepack, "titulo": "Prepacks (assortments)", "singular": "prepack",
        "ayuda": "A prepack is an item with its own product code. Its assortment (breakdown) spreads "
                 "sizes and quantities per master carton with solids of the same style and color, and its prepack ID "
                 "(usually 2 letters and 2 digits, e.g. AB12) is its size. The breakdown can be viewed but not changed.",
        "campos": [
            c("codigo", "Prepack ID", obligatorio=True, max=10, mayus=True, patron=r"^[A-Z0-9]{2,10}$",
              mensaje_patron="Letters and numbers only, e.g. AB12."),
            c("estilo", "Style", obligatorio=True, mayus=True, filtro=True),
            c("color", "Color", obligatorio=True),
            c("descripcion", "Description"),
            c("activo", "Active", "bool", filtro=True),
        ],
        "buscar": ["codigo", "estilo", "color", "descripcion"],
    },
}
ORDEN_CATALOGOS = ["articulos", "prepacks", "escalas", "tipos_empaque", "marcas", "grupos", "proveedores", "sociedades", "centros",
                   "contactos", "almacenes", "transportistas", "tipos_unidad", "paises", "puertos", "regiones", "pasos_lt", "leadtimes", "acuerdos"]
CORREO = r"^[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+$"


def _cat(tipo: str) -> dict:
    if tipo not in CATALOGOS:
        raise ErrorNegocio("The catalog does not exist.", 404, "no_encontrado")
    return CATALOGOS[tipo]


def _mostrar(obj) -> str:
    """Texto corto para mostrar un registro referido."""
    if isinstance(obj, Prepack):
        return f"{obj.codigo} · {obj.estilo} {obj.color or ''}".strip()
    if isinstance(obj, Contacto):
        return obj.nombre
    for a, b in (("codigo", "nombre"), ("sku", "estilo"), ("codigo", "descripcion")):
        if hasattr(obj, a):
            extra = getattr(obj, b, None)
            return f"{getattr(obj, a)} · {extra}" if extra else str(getattr(obj, a))
    return str(obj.id)


def meta(db: Session, user: Usuario) -> list[dict]:
    exigir(user, "catalogos.ver")
    return [{"tipo": t, "titulo": CATALOGOS[t]["titulo"], "singular": CATALOGOS[t]["singular"],
             "ayuda": CATALOGOS[t].get("ayuda"), "campos": CATALOGOS[t]["campos"],
             "extras": CATALOGOS[t].get("extras", []),
             "total": db.scalar(select(func.count()).select_from(CATALOGOS[t]["modelo"])) or 0}
            for t in ORDEN_CATALOGOS]


def _fila(cat: dict, obj, refs: dict) -> dict:
    fila = {"id": obj.id}
    for campo in cat["campos"]:
        v = getattr(obj, campo["nombre"])
        if campo["tipo"] == "multi":
            fila[campo["nombre"]] = [x.id for x in v]
            fila[campo["nombre"] + "_txt"] = ", ".join(x.codigo for x in v) or None
            continue
        if campo["tipo"] == "regla_lt":
            fila[campo["nombre"]] = rlt.cargar(v)
            n = len(fila[campo["nombre"]]["pasos"])
            fila[campo["nombre"] + "_txt"] = f"{n} change" if n == 1 else f"{n} changes"
            continue
        fila[campo["nombre"]] = v
        if campo["tipo"] == "ref" and v:
            fila[campo["nombre"] + "_txt"] = refs.get((campo["catalogo"], v))
    if isinstance(obj, Articulo):
        # La ficha técnica y la partida viven en el producto (estilo-color)
        prod = producto_de(obj)
        fila["producto_id"] = prod.id if prod else None
        # Descripción que se arma con la ficha técnica (comercial + talla)
        base = prod.descripcion_comercial or prod.nombre if prod else None
        fila["descripcion"] = f"{base}, {'prepack' if obj.tipo == 'PREPACK' else 'size'} {obj.talla}" if base and obj.talla \
            else (base or obj.descripcion)
        fila["partida_txt"] = fmt_codigo(prod.codigo) if prod and prod.codigo else None
        fila["clasificacion"] = prod.estado if prod else None
    if isinstance(obj, Prepack):
        fila["total"] = obj.total
        fila["componentes"] = len(obj.componentes)
        fila["sku"] = refs.get(("_sku_prepack", obj.id))
    # Relaciones hijas: centros de la sociedad, artículos del grupo, contactos
    if isinstance(obj, Sociedad):
        fila["centros_txt"] = ", ".join(c.codigo for c in obj.centros) or None
    for ex in cat.get("extras", []):
        if ex["nombre"] != "centros_txt":
            fila[ex["nombre"]] = refs.get(("_" + ex["nombre"], obj.id), 0)
    return fila


def _refs(db: Session, cat: dict, objs: list) -> dict:
    refs = {}
    ids_obj = [o.id for o in objs]
    if cat["modelo"] is Prepack and ids_obj:
        for pid, sku in db.execute(select(Articulo.prepack_id, Articulo.sku).where(Articulo.prepack_id.in_(ids_obj))):
            refs[("_sku_prepack", pid)] = sku
    for ex in cat.get("extras", []):
        if ex["nombre"] == "centros_txt" or not ids_obj:
            continue
        hijo = CATALOGOS[ex["catalogo"]]["modelo"]
        col = getattr(hijo, ex["filtro"])
        for padre, n in db.execute(select(col, func.count()).where(col.in_(ids_obj)).group_by(col)):
            refs[("_" + ex["nombre"], padre)] = n
    for campo in cat["campos"]:
        if campo["tipo"] != "ref":
            continue
        ids = {getattr(o, campo["nombre"]) for o in objs} - {None}
        if not ids:
            continue
        modelo = CATALOGOS[campo["catalogo"]]["modelo"]
        for r in db.scalars(select(modelo).where(modelo.id.in_(ids))).all():
            refs[(campo["catalogo"], r.id)] = _mostrar(r)
    return refs


def listar(db: Session, user: Usuario, tipo: str, q: str | None, filtros: dict, orden: str | None,
           page: int, size: int) -> dict:
    exigir(user, "catalogos.ver")
    cat = _cat(tipo)
    modelo = cat["modelo"]
    consulta = select(modelo)
    if q:
        consulta = consulta.where(filtro_texto(q, lambda p: [getattr(modelo, b).ilike(p) for b in cat["buscar"]]))
    nombres = {x["nombre"]: x for x in cat["campos"]}
    for k, v in filtros.items():
        if k not in nombres or v in (None, ""):
            continue
        campo = nombres[k]
        if campo["tipo"] == "multi":
            consulta = consulta.where(getattr(modelo, k).any(id=int(v)))
            continue
        if campo["tipo"] == "bool":
            v = str(v).lower() in ("1", "true", "si", "sí")
        elif campo["tipo"] in ("ref", "entero"):
            v = int(v)
        consulta = consulta.where(getattr(modelo, k) == v)
    col, _, direccion = (orden or "").partition(":")
    if col in nombres and nombres[col]["tipo"] != "multi":
        expr = getattr(modelo, col)
        consulta = consulta.order_by(expr.desc() if direccion == "desc" else expr.asc(), modelo.id)
    else:
        consulta = consulta.order_by(modelo.id)
    total = db.scalar(select(func.count()).select_from(consulta.subquery())) or 0
    objs = list(db.scalars(consulta.offset((page - 1) * size).limit(size)).all())
    refs = _refs(db, cat, objs)
    return {"items": [_fila(cat, o, refs) for o in objs], "total": total, "page": page, "size": size}


def opciones(db: Session, user: Usuario, tipo: str) -> list[dict]:
    """Lista corta (id, código, texto) para selects y filtros."""
    exigir(user, "catalogos.ver")
    cat = _cat(tipo)
    modelo = cat["modelo"]
    objs = db.scalars(select(modelo).order_by(modelo.id)).all()
    clave = "sku" if tipo == "articulos" else "codigo"
    # Datos para filtrar campos dependientes: sus referencias y, al revés, qué
    # registros de otros catálogos lo incluyen (p. ej. los proveedores de una marca)
    propios = [x for x in cat["campos"] if x["tipo"] in ("ref", "codigo", "multi", "opcion")]
    inversas = [(t, x["nombre"]) for t, c2 in CATALOGOS.items() for x in c2["campos"]
                if x["tipo"] == "multi" and x.get("catalogo") == tipo]
    inv: dict[str, dict[int, list[int]]] = {}
    for t, campo in inversas:
        m = inv.setdefault(t, {})
        for otro in db.scalars(select(CATALOGOS[t]["modelo"])).all():
            for rel in getattr(otro, campo):
                m.setdefault(rel.id, []).append(otro.id)
    # Nombres de los registros referidos (país de un puerto, región de un país…):
    # van en `sub` para que la búsqueda los encuentre ("vietnam cat" → Cat Lai)
    nombres_ref: dict[str, dict] = {}
    for x in propios:
        if x["tipo"] in ("ref", "codigo") and x.get("catalogo") and x["catalogo"] != tipo:
            m = CATALOGOS[x["catalogo"]]["modelo"]
            k = "id" if x["tipo"] == "ref" else ("sku" if x["catalogo"] == "articulos" else "codigo")
            nombres_ref[x["nombre"]] = {getattr(r, k): (getattr(r, "nombre", None) or _mostrar(r))
                                        for r in db.scalars(select(m)).all()}
    res = []
    for o in objs:
        fila = {"id": o.id, "codigo": getattr(o, clave), "texto": _mostrar(o)}
        sub = [nombres_ref[n].get(getattr(o, n)) for n in nombres_ref if getattr(o, n, None) is not None]
        if any(sub):
            fila["sub"] = " · ".join(x for x in sub if x)
        if hasattr(o, "activo") or hasattr(o, "activa"):
            fila["activo"] = bool(getattr(o, "activo", getattr(o, "activa", True)))
        for x in propios:
            v = getattr(o, x["nombre"])
            fila[x["nombre"]] = [y.id for y in v] if x["tipo"] == "multi" else v
        for t, m in inv.items():
            fila[t] = m.get(o.id, [])
        res.append(fila)
    return res


def _valor_dato(db: Session, catalogo: str, valor, clave: str):
    """Dato `clave` de la opción elegida (id o código) de un catálogo."""
    modelo = CATALOGOS[catalogo]["modelo"]
    obj = db.get(modelo, valor) if isinstance(valor, int) else db.scalar(select(modelo).where(modelo.codigo == valor))
    if not obj:
        return None
    if hasattr(obj, clave):
        v = getattr(obj, clave)
        return [y.id for y in v] if isinstance(v, list) else v
    # Relación inversa: qué registros de `clave` (otro catálogo) incluyen esta opción
    for x in CATALOGOS.get(clave, {}).get("campos", []):
        if x["tipo"] == "multi" and x.get("catalogo") == catalogo:
            m = CATALOGOS[clave]["modelo"]
            return [o.id for o in db.scalars(select(m)).all() if obj in getattr(o, x["nombre"])]
    return None


def _dependencias(db: Session, cat: dict, final: dict) -> list[dict]:
    """Valida los campos dependientes declarados en el catálogo."""
    errores = []
    for x in cat["campos"]:
        dep = x.get("depende")
        v, base = final.get(x["nombre"]), final.get((dep or {}).get("campo"))
        if not dep or v in (None, "", []) or base in (None, ""):
            continue
        for elegido in (v if isinstance(v, list) else [v]):
            ident = elegido.id if hasattr(elegido, "id") else elegido
            dato = _valor_dato(db, x["catalogo"], ident, dep["clave"])
            ok = (base in dato) if isinstance(dato, list) else (dato == base)
            if not ok:
                nombre_base = next((y["etiqueta"] for y in cat["campos"] if y["nombre"] == dep["campo"]), dep["campo"])
                errores.append({"campo": x["nombre"], "mensaje":
                                f"{x['etiqueta']}: the option chosen does not belong to the selected {nombre_base.lower()}."})
                break
    return errores


def _limpiar(db: Session, cat: dict, datos: dict, parcial: bool, actual=None) -> dict:
    import re

    errores = []
    limpio = {}
    for campo in cat["campos"]:
        n = campo["nombre"]
        if n not in datos:
            if not parcial and campo["obligatorio"]:
                errores.append({"campo": n, "mensaje": f"{campo['etiqueta']} is required."})
            continue
        v = datos[n]
        if campo["tipo"] == "regla_lt":
            regla, errs = rlt.validar(db, v)
            if errs:
                errores.extend({"campo": n, "mensaje": e} for e in errs)
            else:
                limpio[n] = json.dumps(regla, ensure_ascii=False)
            continue
        if isinstance(v, str):
            v = nombre_fmt(v) if campo.get("formato") == "nombre" else texto_fmt(v)
            if campo.get("mayus"):
                v = v.upper()
        if v in ("", None) or (campo["tipo"] == "multi" and v == [] and not campo["obligatorio"]):
            if campo["tipo"] == "multi":
                if campo["obligatorio"]:
                    errores.append({"campo": n, "mensaje": f"{campo['etiqueta']}: choose at least one."})
                else:
                    limpio[n] = []
                continue
            v = None
        if v is None:
            if campo["obligatorio"]:
                errores.append({"campo": n, "mensaje": f"{campo['etiqueta']} is required."})
            limpio[n] = False if campo["tipo"] == "bool" else None
            continue
        t = campo["tipo"]
        try:
            if t == "entero":
                v = int(v)
                if campo.get("minimo") is not None and v < campo["minimo"]:
                    raise ValueError
            elif t == "numero":
                v = float(v)
            elif t == "bool":
                v = v if isinstance(v, bool) else str(v).lower() in ("1", "true", "si", "sí")
            elif t == "ref":
                v = int(v)
                if not db.get(CATALOGOS[campo["catalogo"]]["modelo"], v):
                    raise ValueError
            elif t == "codigo":
                # El código o el nombre, escrito de cualquier forma, se enlaza con el registro
                modelo = CATALOGOS[campo["catalogo"]]["modelo"]
                obj = db.scalar(select(modelo).where(modelo.codigo == str(v).upper())) or \
                    Referencias(db, modelo).buscar(v)
                if not obj:
                    errores.append({"campo": n, "mensaje": f"{campo['etiqueta']}: {v} is not in the catalog."})
                    continue
                v = obj.codigo
            elif t == "opcion" and v not in [o[0] for o in campo["opciones"]]:
                raise ValueError
            elif t == "multi":
                ids = v if isinstance(v, list) else [x for x in str(v).split(",") if x.strip()]
                modelo = CATALOGOS[campo["catalogo"]]["modelo"]
                objs = list(db.scalars(select(modelo).where(modelo.id.in_([int(x) for x in ids])))) if ids else []
                if len(objs) != len(set(int(x) for x in ids)):
                    raise ValueError
                if campo["obligatorio"] and not objs:
                    errores.append({"campo": n, "mensaje": f"{campo['etiqueta']}: choose at least one."})
                    continue
                v = objs
            elif t == "correos":
                lista = [x.strip().lower() for x in re.split(r"[,;\s]+", str(v)) if x.strip()]
                malos = [x for x in lista if not re.match(CORREO, x)]
                if malos:
                    errores.append({"campo": n, "mensaje": f"Invalid email: {', '.join(malos)}."})
                    continue
                v = ", ".join(dict.fromkeys(lista))
            else:
                v = str(v)
        except (TypeError, ValueError):
            errores.append({"campo": n, "mensaje": f"{campo['etiqueta']}: invalid value."})
            continue
        if campo.get("max") and isinstance(v, str) and len(v) > campo["max"]:
            errores.append({"campo": n, "mensaje": f"{campo['etiqueta']}: {campo['max']} characters maximum."})
            continue
        if campo.get("patron") and isinstance(v, str) and not re.match(campo["patron"], v):
            errores.append({"campo": n, "mensaje": f"{campo['etiqueta']}: {campo.get('mensaje_patron', 'invalid format')}"})
            continue
        limpio[n] = v

    final = {**({c["nombre"]: getattr(actual, c["nombre"]) for c in cat["campos"]} if actual else {}), **limpio}
    # Reglas propias de los artículos: el prepack se enlaza por estilo, color
    # y prepack ID (la talla); debe existir su curva
    if cat["modelo"] is Articulo:
        if final.get("tipo") == "PREPACK" and not (actual and actual.tipo == "PREPACK"):
            errores.append({"campo": "tipo", "mensaje":
                            "Prepacks are created in the Prepacks tab, with their product code and breakdown."})
        elif actual and actual.tipo == "PREPACK":
            # La explosión es fija: no cambia lo que la enlaza
            fijos = [c for c in ("tipo", "estilo", "color", "talla", "unidad") if c in limpio and limpio[c] != getattr(actual, c)]
            if fijos:
                errores.append({"campo": fijos[0], "mensaje":
                                "A prepack and its breakdown cannot change (style, color, prepack ID, unit and type). "
                                "Create a new prepack."})
        elif final.get("tipo") == "SOLIDO":
            if actual and db.scalar(select(PrepackComponente.id).where(PrepackComponente.articulo_id == actual.id)):
                fijos = [c for c in ("estilo", "color", "talla", "unidad") if c in limpio and limpio[c] != getattr(actual, c)]
                if fijos:
                    errores.append({"campo": fijos[0], "mensaje":
                                    "This solid is part of a prepack breakdown: its style, "
                                    "color, size and unit cannot change."})
            limpio["prepack_id"] = None
            if final.get("unidad") == "CJ":
                errores.append({"campo": "unidad", "mensaje": "A solid is handled in pairs or units."})
    # El genérico (estilo-color) agrupa sus tallas, sólidos y prepacks: todos
    # comparten estilo, color, marca, grupo y proveedor
    if cat["modelo"] is Articulo and generico_de(final):
        gen = generico_de(final)
        ref = db.scalar(select(Articulo).where(Articulo.generico == gen, Articulo.tipo == "SOLIDO",
                                               *([Articulo.id != actual.id] if actual else [])).limit(1))
        prod = producto_por_generico(db, gen)
        base = ({"estilo": ref.estilo, "color": ref.color, "marca_id": ref.marca_id, "proveedor_id": ref.proveedor_id}
                if ref else {"estilo": prod.estilo, "color": prod.color, "marca_id": prod.marca_id,
                             "proveedor_id": prod.proveedor_id} if prod else None)
        if base:
            distintos = [k for k, v in base.items() if v is not None and final.get(k) not in (None, "") and final.get(k) != v]
            if distintos:
                nombres = {"estilo": "style", "color": "color", "marca_id": "brand", "proveedor_id": "supplier"}
                errores.append({"campo": distintos[0], "mensaje":
                                f"Generic {gen} is {base['estilo']} {base['color'] or ''}: every size must have the same "
                                f"{', '.join(nombres[k] for k in distintos)}. Use another generic."})
    if cat["modelo"] is Articulo and final.get("proveedor_id") and final.get("marca_id"):
        prov = db.get(Proveedor, final["proveedor_id"])
        if prov and final["marca_id"] not in {m.id for m in prov.marcas}:
            errores.append({"campo": "marca_id", "mensaje":
                            f"The brand does not belong to {prov.nombre}; its brands are {', '.join(m.codigo for m in prov.marcas)}."
                            if prov.marcas else f"{prov.nombre} has no brands yet: assign them in Suppliers first."})
    if cat["modelo"] is Articulo and actual and "proveedor_id" in limpio and limpio["proveedor_id"] != actual.proveedor_id \
            and db.scalar(select(PosicionOC.id).where(PosicionOC.articulo_id == actual.id).limit(1)):
        errores.append({"campo": "proveedor_id", "mensaje": "The item is on purchase orders: its supplier cannot change."})
    if cat["modelo"] is Proveedor and actual and "marcas" in limpio:
        quedan = {m.id for m in limpio["marcas"]}
        usadas = {m for (m,) in db.execute(select(Articulo.marca_id).where(Articulo.proveedor_id == actual.id).distinct())}
        if usadas - quedan:
            nombres_m = [m.codigo for m in db.scalars(select(Marca).where(Marca.id.in_(usadas - quedan)))]
            errores.append({"campo": "marcas", "mensaje":
                            f"The supplier has items of {', '.join(nombres_m)}: those brands cannot be removed."})
    if cat["modelo"] is Proveedor and actual and "sociedades" in limpio:
        quedan = {x.codigo for x in limpio["sociedades"]}
        usadas = {c for (c,) in db.execute(select(OrdenCompra.sociedad).where(
            OrdenCompra.proveedor_id == actual.id, OrdenCompra.sociedad.is_not(None)).distinct())}
        if usadas - quedan:
            errores.append({"campo": "sociedades", "mensaje":
                            f"The supplier has purchase orders with {', '.join(sorted(usadas - quedan))}: "
                            "those companies cannot be removed."})
    if cat["modelo"] is ReglaLeadTime:
        errores += _validar_regla_lt(db, final, actual, limpio)
    if cat["modelo"] is EstadoLiberacion:
        otros = [x for x in db.scalars(select(EstadoLiberacion).where(EstadoLiberacion.tipo == final.get("tipo")))
                 if not actual or x.id != actual.id]
        if final.get("predeterminado") and any(x.predeterminado for x in otros):
            errores.append({"campo": "predeterminado", "mensaje": "There is already a default status for this release."})
        if final.get("con_cambios") and (final.get("tipo") != "LOGISTICA" or not final.get("libera")):
            errores.append({"campo": "con_cambios", "mensaje": "Only a released logistics status can be the one with later changes."})
        if final.get("con_cambios") and any(x.con_cambios for x in otros):
            errores.append({"campo": "con_cambios", "mensaje": "There is already a status for released POs with later changes."})
        if actual and "codigo" in limpio and limpio["codigo"] != actual.codigo:
            col = OrdenCompra.liberacion_comercial if actual.tipo == "COMERCIAL" else OrdenCompra.liberacion_logistica
            if db.scalar(select(OrdenCompra.id).where(col == actual.codigo).limit(1)):
                errores.append({"campo": "codigo", "mensaje": "Purchase orders use this code: it cannot change."})
    if cat["modelo"] is EscalaTalla and final.get("tallas"):
        from app.modulos.maestros.tallas import validar as validar_escala

        errores += [{"campo": "tallas", "mensaje": m} for m in validar_escala(final["tallas"])]
        if final.get("regla") == "MULTIPLICAR" and not final.get("factor"):
            errores.append({"campo": "factor", "mensaje": "Enter the factor for the rule Size × factor."})
    if cat["modelo"] is Contacto and not final.get("sociedad_id") and not final.get("centro_id"):
        errores.append({"campo": "sociedad_id", "mensaje": "Enter the contact's company or plant."})
    if cat["modelo"] in (Centro, Almacen) and actual and "sociedad_id" in limpio and limpio["sociedad_id"] != actual.sociedad_id:
        campo_oc = OrdenCompra.centro_destino if cat["modelo"] is Centro else None
        en_uso = db.scalar(select(OrdenCompra.id).where(
            (OrdenCompra.centro == actual.codigo) | (campo_oc == actual.codigo) if campo_oc is not None
            else OrdenCompra.id.in_(select(PosicionOC.oc_id).where(PosicionOC.almacen == actual.codigo))).limit(1))
        if en_uso:
            errores.append({"campo": "sociedad_id", "mensaje": "It is used on purchase orders: its company cannot change."})
    if cat["modelo"] is Prepack and actual:
        fijos = [c for c in ("codigo", "estilo", "color") if c in limpio and limpio[c] != getattr(actual, c)]
        if fijos:
            errores.append({"campo": fijos[0], "mensaje":
                            "The prepack ID, style and color cannot change; only the description and whether it is active."})
    ya = {e.get("campo") for e in errores}
    errores += [e for e in _dependencias(db, cat, final) if e["campo"] not in ya]
    if errores:
        raise ErrorNegocio("Check the data.", 422, "validacion", errores)
    return limpio


def crear(db: Session, user: Usuario, tipo: str, datos: dict) -> dict:
    exigir(user, "catalogos.crear")
    if tipo == "prepacks":
        return crear_prepack(db, user, datos)
    cat = _cat(tipo)
    limpio = _limpiar(db, cat, datos, parcial=False)
    for campo in cat["campos"]:
        if campo["tipo"] == "bool" and campo["nombre"] not in datos:
            limpio[campo["nombre"]] = True
    obj = cat["modelo"](**limpio)
    try:
        with db.begin_nested():
            db.add(obj)
            db.flush()
    except IntegrityError:
        raise ErrorNegocio(f"A {cat['singular']} with that code already exists.", 409, "duplicado") from None
    if isinstance(obj, Articulo):
        asegurar_producto(db, obj)
    registrar(db, user, tipo, obj.id, "crear", {"codigo": _mostrar(obj)})
    return _fila(cat, obj, _refs(db, cat, [obj]))


def actualizar(db: Session, user: Usuario, tipo: str, obj_id: int, datos: dict) -> dict:
    exigir(user, "catalogos.editar")
    cat = _cat(tipo)
    obj = db.get(cat["modelo"], obj_id)
    if not obj:
        raise ErrorNegocio(f"The {cat['singular']} does not exist.", 404, "no_encontrado")
    limpio = _limpiar(db, cat, datos, parcial=True, actual=obj)
    cambios = {}
    for k, v in limpio.items():
        if getattr(obj, k) != v:
            antes = getattr(obj, k)
            if isinstance(v, list):  # relaciones: se guardan los códigos en el historial
                cambios[k] = [[x.codigo for x in antes], [x.codigo for x in v]]
            else:
                cambios[k] = [antes, v]
            setattr(obj, k, v)
    try:
        with db.begin_nested():
            db.flush()
    except IntegrityError:
        raise ErrorNegocio(f"A {cat['singular']} with that code already exists.", 409, "duplicado") from None
    if isinstance(obj, Articulo) and {"estilo", "color", "proveedor_id"} & set(cambios):
        asegurar_producto(db, obj)
    if cambios:
        registrar(db, user, tipo, obj.id, "editar", cambios)
    return _fila(cat, obj, _refs(db, cat, [obj]))


def _validar_regla_lt(db: Session, final: dict, actual, limpio: dict) -> list[dict]:
    """Una regla por ámbito, con su ámbito según el nivel, y la cadena que
    resulta (con lo heredado) debe ser válida."""
    errores = []
    nivel = final.get("nivel")
    campo = rlt.CAMPO_NIVEL.get(nivel)
    for otro in rlt.CAMPO_NIVEL.values():
        if otro != campo:
            final[otro] = limpio[otro] = None
    ambito = final.get(campo) if campo else None
    if campo and not ambito:
        return [{"campo": campo, "mensaje": "Choose the scope of this level."}]
    filtro = [ReglaLeadTime.nivel == nivel] + ([getattr(ReglaLeadTime, campo) == ambito] if campo else [])
    otra = db.scalar(select(ReglaLeadTime.id).where(*filtro, ReglaLeadTime.id != (actual.id if actual else 0)))
    if otra:
        errores.append({"campo": campo or "nivel", "mensaje": "There is already a rule for that scope; edit it."})
    regla = rlt.cargar(final.get("pasos"))
    kw = {campo: ambito} if campo else {}
    res = rlt.efectivo(db, **kw, con_regla=(nivel, ambito, regla))
    errores += [{"campo": "pasos", "mensaje": m} for m in res["errores"]]
    return errores


def eliminar(db: Session, user: Usuario, tipo: str, obj_id: int) -> dict:
    exigir(user, "catalogos.eliminar")
    cat = _cat(tipo)
    obj = db.get(cat["modelo"], obj_id)
    if not obj:
        raise ErrorNegocio(f"The {cat['singular']} does not exist.", 404, "no_encontrado")
    if isinstance(obj, PasoLeadTime):
        usan = [r.nombre for r in db.scalars(select(ReglaLeadTime)) if any(
            e["paso"] == obj.codigo or e.get("ref") == obj.codigo for e in rlt.cargar(r.pasos)["pasos"])]
        if usan:
            raise ErrorNegocio(f"It cannot be deleted: the step is used in {', '.join(usan)}. Deactivate it instead.",
                               409, "en_uso")
    try:
        with db.begin_nested():
            db.delete(obj)
            db.flush()
    except IntegrityError:
        raise ErrorNegocio(f"It cannot be deleted: the {cat['singular']} is in use. Deactivate it instead.",
                           409, "en_uso") from None
    registrar(db, user, tipo, obj_id, "eliminar", None)
    return {"ok": True}


# ---- Prepacks: componentes -------------------------------------------------
def componentes(db: Session, user: Usuario, prepack_id: int, validar: bool = True) -> dict:
    if validar:
        exigir(user, "catalogos.ver")
    p = db.get(Prepack, prepack_id)
    if not p:
        raise ErrorNegocio("The prepack does not exist.", 404, "no_encontrado")
    art = db.scalar(select(Articulo).where(Articulo.prepack_id == p.id))
    return {
        "sku": art.sku if art else None, "unidad": "CJ", "marca": art.marca.codigo if art else None,
        "id": p.id, "codigo": p.codigo, "estilo": p.estilo, "color": p.color, "descripcion": p.descripcion,
        "total": p.total,
        "unidad_componentes": p.componentes[0].articulo.unidad if p.componentes else None,
        "componentes": [{"articulo_id": x.articulo_id, "sku": x.articulo.sku, "estilo": x.articulo.estilo,
                         "color": x.articulo.color, "talla": x.articulo.talla, "unidad": x.articulo.unidad,
                         "cantidad": x.cantidad} for x in p.componentes],
    }


def _validar_curva(db: Session, estilo: str, color: str, items: list[dict]) -> tuple[list, list]:
    """Solo sólidos del mismo estilo y color del prepack, con cantidad, sin
    repetir artículo y sin mezclar pares con unidades."""
    errores = []
    arts = []
    for it in items:
        a = db.get(Articulo, int(it.get("articulo_id") or 0))
        try:
            cant = int(it.get("cantidad") or 0)
        except (TypeError, ValueError):
            cant = 0
        if not a:
            errores.append({"mensaje": "One of the items does not exist."})
            continue
        if a.tipo != "SOLIDO":
            errores.append({"mensaje": f"{a.sku}: a breakdown only carries solid items."})
        if cant < 1:
            errores.append({"mensaje": f"{a.sku}: the quantity must be greater than zero."})
        arts.append((a, cant))
    if not arts:
        errores.append({"mensaje": "The breakdown needs at least one solid item."})
    if len({a.id for a, _ in arts}) != len(arts):
        errores.append({"mensaje": "An item appears twice in the breakdown."})
    otros = [a for a, _ in arts if (a.estilo, a.color) != (estilo, color)]
    if otros:
        errores.append({"mensaje": f"The prepack is {estilo} {color}: "
                                   + ", ".join(f"{a.sku} ({a.estilo} {a.color})" for a in otros)
                                   + " is not the same style and color."})
    if len({a.unidad for a, _ in arts}) > 1:
        errores.append({"mensaje": "Pairs and units cannot be mixed in a breakdown."})
    return arts, errores


def crear_prepack(db: Session, user: Usuario, datos: dict) -> dict:
    """Crea el prepack completo en un paso: su curva (explosión) y el artículo
    prepack con su propio código de producto. La explosión ya no cambia."""
    import re

    exigir(user, "catalogos.crear")
    sku = str(datos.get("sku") or "").strip()
    codigo = str(datos.get("codigo") or "").strip().upper()
    estilo = str(datos.get("estilo") or "").strip().upper()
    color = str(datos.get("color") or "").strip()
    errores = []
    if not codigo_valido(sku):
        errores.append({"campo": "sku", "mensaje": "Item code: letters and numbers (also . - _ /), up to 40 characters."})
    elif db.scalar(select(Articulo.id).where(Articulo.sku == sku)):
        errores.append({"campo": "sku", "mensaje": f"Code {sku} already exists in the item master."})
    if not re.fullmatch(r"[A-Z0-9]{2,10}", codigo):
        errores.append({"campo": "codigo", "mensaje": "Prepack ID: letters and numbers, for example AB12."})
    if not estilo or not color:
        errores.append({"campo": "estilo", "mensaje": "Enter the prepack's style and color."})
    elif db.scalar(select(Prepack.id).where(Prepack.estilo == estilo, Prepack.color == color, Prepack.codigo == codigo)):
        errores.append({"campo": "codigo", "mensaje": f"Prepack {codigo} already exists for {estilo} {color}."})
    arts, err_curva = _validar_curva(db, estilo, color, datos.get("componentes") or [])
    errores += err_curva
    if errores:
        raise ErrorNegocio("The prepack is not valid.", 422, "validacion", errores)
    base = arts[0][0]
    p = Prepack(codigo=codigo, estilo=estilo, color=color, activo=True,
                descripcion=(datos.get("descripcion") or "").strip()
                or f"{estilo} {color} prepack {codigo} ({sum(c for _, c in arts)} {base.unidad.lower()})")
    for a, cant in arts:
        p.componentes.append(PrepackComponente(articulo_id=a.id, cantidad=cant))
    db.add(p)
    db.flush()
    art = Articulo(sku=sku.upper(), generico=base.generico or generico_de(base), upc=(datos.get("upc") or "").strip() or None, estilo=estilo, color=color, talla=codigo,
                   descripcion=p.descripcion, marca_id=base.marca_id, grupo_id=base.grupo_id,
                   proveedor_id=base.proveedor_id, unidad="CJ", tipo="PREPACK", prepack_id=p.id,
                   activo=True)
    db.add(art)
    db.flush()
    asegurar_producto(db, art)
    registrar(db, user, "prepacks", p.id, "crear",
              {"sku": sku, "prepack": codigo, "explosion": {a.talla: c for a, c in arts}})
    return componentes(db, user, p.id)


def explosion(db: Session, user: Usuario, sku: str) -> dict:
    """Explosión de un artículo prepack, de solo lectura (para OC, factura,
    packing list y seguimiento)."""
    exigir(user, "oc.ver")
    a = db.scalar(select(Articulo).where(Articulo.sku == sku))
    if not a or a.tipo != "PREPACK" or not a.prepack:
        raise ErrorNegocio("That code is not a prepack item.", 404, "no_encontrado")
    return componentes(db, user, a.prepack_id, validar=False)


# ---- Carga masiva ----------------------------------------------------------
# Columnas de las plantillas en inglés -> nombre interno (también se aceptan los internos)
COLUMNAS_EN = {"style": "estilo", "size": "talla", "description": "descripcion", "brand": "marca", "group": "grupo",
               "supplier": "proveedor", "type": "tipo", "uom": "unidad", "unit": "unidad",
               "hs_code": "partida_arancelaria", "country_of_origin": "pais_origen", "prepack_sku": "sku_prepack",
               "quantity": "cantidad", "qty": "cantidad"}
CAMPO_EN = {"marca": "Brand", "grupo": "Group", "proveedor": "Supplier"}


def _filas_archivo(nombre: str, contenido: bytes) -> list[dict]:
    from app.modulos.compras.ordenes import _norm, _texto

    if nombre.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        filas = [[_texto(v) for v in f] for f in wb.active.iter_rows(values_only=True)]
    else:
        texto = contenido.decode("utf-8-sig", errors="replace")
        delim = ";" if texto[:2000].count(";") > texto[:2000].count(",") else ","
        filas = [[v.strip() for v in f] for f in csv.reader(io.StringIO(texto), delimiter=delim)]
    if not filas:
        return []
    enc = [COLUMNAS_EN.get(_norm(h), _norm(h)) for h in filas[0]]
    return [{**{enc[i]: (f[i] if i < len(f) else "") for i in range(len(enc))}, "_fila": n}
            for n, f in enumerate(filas[1:], start=2) if any(f)]


def importar_prepacks(db: Session, user: Usuario, nombre: str, contenido: bytes) -> dict:
    """Columnas: sku_prepack (código de producto del prepack), prepack_id,
    descripcion, sku (sólido) y cantidad. Una fila por talla de la explosión.
    Estilo y color salen de los sólidos. Un prepack que ya existe no cambia."""
    exigir(user, "catalogos.crear")
    filas = _filas_archivo(nombre, contenido)
    grupos: dict[str, dict] = {}
    errores = []
    for f in filas:
        sku_pp = (f.get("sku_prepack") or f.get("codigo_prepack") or "").strip()
        codigo = (f.get("prepack_id") or f.get("prepack") or "").strip().upper()
        sku = (f.get("sku") or "").strip()
        if not sku_pp or not codigo or not sku:
            errores.append({"fila": f["_fila"], "mensaje": "Missing prepack_sku, prepack_id or sku."})
            continue
        a = db.scalar(select(Articulo).where(Articulo.sku == sku))
        if not a:
            errores.append({"fila": f["_fila"], "mensaje": f"SKU {sku} does not exist in the item master."})
            continue
        g = grupos.setdefault(sku_pp, {"codigo": codigo, "estilo": a.estilo, "color": a.color,
                                       "descripcion": f.get("descripcion") or None, "items": []})
        g["items"].append({"articulo_id": a.id, "cantidad": f.get("cantidad") or 0})
    creados = sin_cambio = 0
    for sku_pp, g in grupos.items():
        existente = db.scalar(select(Articulo).where(Articulo.sku == sku_pp))
        if existente:
            actual = {x.articulo_id: x.cantidad for x in existente.prepack.componentes} if existente.prepack else {}
            nueva = {it["articulo_id"]: int(it["cantidad"] or 0) for it in g["items"]}
            if actual == nueva:
                sin_cambio += 1
            else:
                errores.append({"fila": sku_pp, "mensaje": "The prepack already exists with another breakdown; "
                                                           "the breakdown cannot change. Use a new code."})
            continue
        try:
            with db.begin_nested():
                crear_prepack(db, user, {"sku": sku_pp, **g, "componentes": g["items"]})
            creados += 1
        except ErrorNegocio as e:
            errores.append({"fila": sku_pp, "mensaje": "; ".join(d["mensaje"] for d in e.detalle or []) or e.mensaje})
    return {"creados": creados, "actualizados": 0, "sin_cambio": sin_cambio, "errores": errores[:200]}
