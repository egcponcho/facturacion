"""Capa oficial del arancel: fuentes, versiones de datos, países (esquema de
código), control de capítulos, dominios de clasificación y atributos.

Se carga desde los paquetes Excel oficiales (hojas Sources, Countries,
Versions, Chapter_Control, Domain_Chapter_Map, Domains, Attributes,
Attribute_Options, Attribute_Scope, Classification_Rules, Rule_Conditions;
paquete nacional: Country_Source_Map, National_Codes, Regulations, Taxes). La carga es
idempotente: actualiza por clave natural (código de fuente, ISO, código de
versión, capítulo, dominio) y nunca borra lo que ya existe.
"""
import io
import re
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import func, insert, select
from sqlalchemy.orm import Session

from app.core.errores import ErrorNegocio
from app.instalacion.datos import MOTOR, OFICIAL
from app.modelos import (
    ControlCapitulo,
    DominioCapitulo,
    DominioClasificacion,
    FuenteOficial,
    PaisArancel,
    Usuario,
    VersionDataset,
    ahora,
)
from app.modulos.acceso.permisos import exigir
from app.modulos.comun.historial import registrar
from app.modulos.comun.texto import filtro_texto
from app.modulos.documentos.plantillas import norm

CARPETA = OFICIAL
# El motor (02) trae los dominios que el paquete oficial (01) relaciona con capítulos
# y el nacional (03) usa países, fuentes y versiones del 01
# Paquetes incluidos y en qué orden: el 02 (configuración del motor) vive aparte, en data/motor
PAQUETES = [MOTOR / "02_carga_motor_dinamico_v3.xlsx", OFICIAL / "01_carga_oficial_catalogos_v3.xlsx",
            OFICIAL / "03_carga_nacional_regulaciones_v3.xlsx"]
HOJAS = ("Sources", "Versions", "Countries", "Chapter_Control", "Domains", "Domain_Chapter_Map", "Material_Classes", "Categories", "Attribute_Scope_Conditions",
         "Attributes", "Attribute_Options", "Attribute_Scope", "Classification_Rules", "Rule_Conditions",
         "Country_Source_Map", "National_Codes", "Regulations", "Taxes")
ESTADOS_VERSION = {"PUBLISHED": "PUBLICADA", "PUBLICADA": "PUBLICADA", "DYNAMIC": "DINAMICA", "DINAMICA": "DINAMICA",
                   "DRAFT": "BORRADOR", "BORRADOR": "BORRADOR", "ARCHIVED": "ARCHIVADA", "ARCHIVADA": "ARCHIVADA"}


# ---- Lectura -----------------------------------------------------------------
def _hojas(contenido: bytes) -> dict[str, list[dict]]:
    """Cada hoja como lista de filas {encabezado normalizado: valor}."""
    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise ErrorNegocio("The file could not be read. Use the Excel (.xlsx) package.", 422, "archivo_invalido") from None
    out = {}
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        if not filas:
            continue
        enc = [norm(c) if c is not None else "" for c in filas[0]]
        datos = []
        for i, f in enumerate(filas[1:], start=2):
            if not any(v not in (None, "") for v in f):
                continue
            d = {enc[j]: f[j] for j in range(min(len(enc), len(f))) if enc[j]}
            d["_fila"] = i
            datos.append(d)
        out[ws.title] = datos
    return out


def _txt(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    t = re.sub(r"\s+", " ", str(v)).strip()
    return t or None


def _si(v) -> bool:
    return str(v or "").strip().lower() in ("yes", "si", "sí", "true", "1", "y", "x")


def _fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _txt(v)
    if not t:
        return None
    try:
        return date.fromisoformat(t[:10])
    except ValueError:
        raise ValueError(f"invalid date {t}") from None


def _capitulo(v) -> str | None:
    t = _txt(v)
    if not t or not re.fullmatch(r"\d{1,2}", t):
        return None
    return t.zfill(2)


# ---- Carga ---------------------------------------------------------------------
def importar(db: Session, contenido: bytes, usuario: Usuario | None = None, nombre: str = "") -> dict:
    """Carga las hojas reconocidas del paquete. Devuelve lo creado/actualizado
    por hoja y los errores fila por fila (las filas con error no se cargan)."""
    hojas = _hojas(contenido)
    res: dict[str, dict] = {}
    errores: list[dict] = []

    def cuenta(hoja, creado):
        r = res.setdefault(hoja, {"creados": 0, "actualizados": 0})
        r["creados" if creado else "actualizados"] += 1

    for hoja in HOJAS:
        if hoja in hojas:
            res[hoja] = {"creados": 0, "actualizados": 0}

    def error(hoja, fila, msg):
        errores.append({"fila": f"{hoja} {fila}", "mensaje": msg})

    # Fuentes
    for f in hojas.get("Sources", []):
        cod = _txt(f.get("source_id"))
        if not cod:
            error("Sources", f["_fila"], "Source ID is required.")
            continue
        x = db.scalar(select(FuenteOficial).where(FuenteOficial.codigo == cod))
        nuevo = x is None
        x = x or FuenteOficial(codigo=cod)
        x.ambito = _txt(f.get("country_region")) or "REGIONAL"
        x.autoridad = _txt(f.get("authority")) or cod
        x.dataset = _txt(f.get("official_dataset")) or cod
        x.uso, x.url = _txt(f.get("use")), _txt(f.get("official_url"))
        x.acceso, x.autenticacion = _txt(f.get("access_mode")), _txt(f.get("authentication"))
        x.nota_version, x.verificacion = _txt(f.get("version_status_note")), _txt(f.get("verification"))
        x.documento = _txt(f.get("official_document")) or x.documento
        try:
            # Fecha de verificación: la columna propia o la que dice el texto de verificación («Verified 2026-10-01»)
            verificado = _fecha(f.get("verified_on")) or _fecha_en_texto(x.verificacion)
        except ValueError as e:
            error("Sources", f["_fila"], str(e))
            continue
        if verificado and verificado > date.today():
            error("Sources", f["_fila"], f"The verification date {verificado.isoformat()} is in the future.")
            continue
        x.verificado_en = verificado or x.verificado_en
        x.verificado_por = _txt(f.get("verified_by")) or x.verificado_por
        db.add(x)
        cuenta("Sources", nuevo)
    db.flush()
    fuentes = {x.codigo: x for x in db.scalars(select(FuenteOficial))}

    # Versiones
    for f in hojas.get("Versions", []):
        cod = _txt(f.get("version_id"))
        src = _txt(f.get("source_id"))
        if not cod:
            error("Versions", f["_fila"], "Version ID is required.")
            continue
        if src and src not in fuentes:
            error("Versions", f["_fila"], f"Source {src} does not exist.")
            continue
        try:
            desde, hasta = _fecha(f.get("valid_from")), _fecha(f.get("valid_to"))
        except ValueError as e:
            error("Versions", f["_fila"], str(e))
            continue
        if desde and hasta and hasta < desde:
            error("Versions", f["_fila"], "Valid to must be on or after valid from.")
            continue
        estado = ESTADOS_VERSION.get((_txt(f.get("status")) or "").upper(), "BORRADOR")
        # Una versión publicada o dinámica respalda datos oficiales: fuente trazable y, si es publicada, vigencia
        if estado in ("PUBLICADA", "DINAMICA"):
            faltas = ([] if src else [f"Version {cod} needs its official source (Source ID)."]) + \
                     (problemas_fuente(fuentes[src]) if src else []) + \
                     ([f"Version {cod} is published without its validity (valid from)."] if estado == "PUBLICADA" and not desde else [])
            if faltas:
                error("Versions", f["_fila"], " ".join(faltas))
                continue
        x = db.scalar(select(VersionDataset).where(VersionDataset.codigo == cod))
        nuevo = x is None
        x = x or VersionDataset(codigo=cod)
        x.dataset = _txt(f.get("dataset")) or cod
        x.etiqueta = _txt(f.get("version_label")) or cod
        x.estado = estado
        x.vigente_desde, x.vigente_hasta = desde, hasta
        x.fuente = fuentes.get(src)
        x.ambito = ambito_version(cod, x.fuente, _txt(f.get("scope")))
        x.nota = _txt(f.get("notes"))
        x.importado_en = ahora()
        db.add(x)
        cuenta("Versions", nuevo)
    db.flush()
    versiones = {x.codigo: x for x in db.scalars(select(VersionDataset))}

    # Países: esquema de código configurable (sin longitud fija)
    for f in hojas.get("Countries", []):
        iso = (_txt(f.get("iso")) or "").upper()
        if not re.fullmatch(r"[A-Z]{2}", iso):
            error("Countries", f["_fila"], "ISO must be 2 letters.")
            continue
        x = db.scalar(select(PaisArancel).where(PaisArancel.iso == iso))
        nuevo = x is None
        if nuevo:
            x = PaisArancel(iso=iso, digitos=10, orden=(db.scalar(select(func.max(PaisArancel.orden))) or 0) + 1)
        x.nombre = _txt(f.get("country")) or x.nombre or iso
        x.contexto = _txt(f.get("integration_context"))
        x.modelo_arancel = _txt(f.get("tariff_model"))
        largo = _txt(f.get("national_code_length")) or ""
        if re.fullmatch(r"\d+(\s*,\s*\d+)*", largo):
            x.longitudes = ",".join(p.strip() for p in largo.split(","))
        src = _txt(f.get("primary_source"))
        if src and src in fuentes:
            x.fuente_id = fuentes[src].id
        # La base legal solo si el paquete la declara (nunca se arma con el nombre de la fuente)
        if _txt(f.get("legal_basis")):
            x.base_legal = _txt(f.get("legal_basis"))[:300]
        x.mcca = "common market" in (x.contexto or "").lower()
        x.nota = _txt(f.get("implementation_note")) or x.nota
        if f.get("active") is not None:
            x.activo = _si(f.get("active"))
        db.add(x)
        cuenta("Countries", nuevo)

    # Control de capítulos
    for f in hojas.get("Chapter_Control", []):
        cap = _capitulo(f.get("chapter"))
        if not cap:
            error("Chapter_Control", f["_fila"], "Chapter must be 2 digits.")
            continue
        ver = _txt(f.get("version"))
        if ver and ver not in versiones:
            error("Chapter_Control", f["_fila"], f"Version {ver} does not exist.")
            continue
        x = db.scalar(select(ControlCapitulo).where(ControlCapitulo.capitulo == cap))
        nuevo = x is None
        x = x or ControlCapitulo(capitulo=cap)
        x.seccion = _txt(f.get("section"))
        x.titulo = _txt(f.get("official_working_title")) or f"Chapter {cap}"
        x.activo, x.clasificacion = _si(f.get("active")), _si(f.get("classification_enabled"))
        x.candidato_auto, x.solo_manual = _si(f.get("auto_candidate")), _si(f.get("manual_only"))
        x.archivado = _si(f.get("archived"))
        x.alcance = _txt(f.get("initial_scope"))
        x.version_id = versiones[ver].id if ver else None
        src = _txt(f.get("source_id"))
        x.fuente_id = fuentes[src].id if src in fuentes else None
        x.nota = _txt(f.get("notes"))
        db.add(x)
        cuenta("Chapter_Control", nuevo)

    # Dominios
    for i, f in enumerate(hojas.get("Domains", [])):
        cod = (_txt(f.get("domain_code")) or "").upper()
        if not cod:
            error("Domains", f["_fila"], "Domain code is required.")
            continue
        x = db.scalar(select(DominioClasificacion).where(DominioClasificacion.codigo == cod))
        nuevo = x is None
        x = x or DominioClasificacion(codigo=cod, orden=(i + 1) * 10)
        x.nombre = _txt(f.get("label")) or cod
        x.descripcion = _txt(f.get("description"))
        x.modo = (_txt(f.get("default_mode")) or "AUTO").upper()
        x.activo = _si(f.get("active")) if f.get("active") is not None else True
        db.add(x)
        cuenta("Domains", nuevo)
    db.flush()
    dominios = {x.codigo: x for x in db.scalars(select(DominioClasificacion))}

    for f in hojas.get("Domain_Chapter_Map", []):
        dom = (_txt(f.get("domain")) or "").upper()
        cap = _capitulo(f.get("chapter"))
        if dom not in dominios or not cap:
            error("Domain_Chapter_Map", f["_fila"], f"Domain {dom or '(empty)'} or chapter does not exist.")
            continue
        d = dominios[dom]
        x = db.scalar(select(DominioCapitulo).where(DominioCapitulo.dominio_id == d.id, DominioCapitulo.capitulo == cap))
        nuevo = x is None
        x = x or DominioCapitulo(dominio=d, capitulo=cap)
        x.relevancia = "SECONDARY" if (_txt(f.get("relevance")) or "").upper().startswith("SEC") else "PRIMARY"
        x.habilitado = _si(f.get("enabled_initially")) if f.get("enabled_initially") is not None else True
        x.proposito = _txt(f.get("purpose"))
        db.add(x)
        cuenta("Domain_Chapter_Map", nuevo)
    db.flush()
    from app.modulos.clasificacion import (
        atributos,
        categorias,
        materiales,
        nacional,
        reglas,  # usan dominios, países, fuentes y versiones ya cargados
    )

    materiales.importar_hojas(db, hojas, cuenta, error)  # las clases de material que nombran las derivaciones
    categorias.importar_hojas(db, hojas, cuenta, error)  # antes de los ámbitos y las reglas que las nombran  # antes de los ámbitos y las reglas que las nombran
    atributos.importar_hojas(db, hojas, cuenta, error)
    categorias.importar_hojas(db, hojas, cuenta, error, fase=2)  # su plantilla aduanera nombra atributos y partes
    reglas.importar_hojas(db, hojas, cuenta, error)
    nacional.importar_hojas(db, hojas, cuenta, error)
    if not res:
        raise ErrorNegocio(f"No known sheet was found ({', '.join(HOJAS)}).", 422, "validacion")
    if usuario:
        registrar(db, usuario, "aranceles", 0, "importar_oficial", {"archivo": nombre, "hojas": res, "errores": len(errores)})
    return {"hojas": res, "errores": errores, "creados": sum(r["creados"] for r in res.values()),
            "actualizados": sum(r["actualizados"] for r in res.values())}


def cargar_paquetes_base(db: Session, paquetes: list | None = None) -> dict:
    """Carga paquetes incluidos (app.cargas decide cuáles: el 02 es del motor, el 01 y el 03 oficiales)."""
    out = {}
    for ruta in paquetes or PAQUETES:
        nombre = ruta.name
        if ruta.exists():
            out[nombre] = r = importar(db, ruta.read_bytes(), nombre=nombre)
            # Un paquete incluido con errores no se carga a medias: es un defecto del paquete
            if r["errores"]:
                raise ValueError(f"Official package {nombre} has errors: " + "; ".join(f"{e['fila']}: {e['mensaje']}" for e in r["errores"][:5]))
    return out


# ---- Datos oficiales incluidos -------------------------------------------------------
DATOS = OFICIAL


def cargar_lineas_regionales(db: Session, version: str = "SAC-2025-V6") -> int:
    """Los países que aplican tal cual las líneas del SAC regional a 10 dígitos
    (nivel_base SAC10, configuración del país) reciben como líneas nacionales
    las líneas oficiales del árbol de esa versión: con su fuente (la de la
    versión regional), su versión, vigencia y DAI. Nada sale de datos de la
    empresa. Las condiciones que el clasificador deduce del texto oficial (p.
    ej. «para hombres») van como reglas del motor (CLASSIFIER), no en el dato."""
    import json

    from app.modelos import IncisoNacional, NodoArancel, VersionDataset

    v = db.scalar(select(VersionDataset).where(VersionDataset.codigo == version))
    if not v:
        return 0
    paises = [p.iso for p in db.scalars(select(PaisArancel).where(PaisArancel.nivel_base == "SAC10"))]
    if not paises:
        return 0
    ya = {(i, c) for i, c in db.execute(select(IncisoNacional.pais, IncisoNacional.codigo).where(IncisoNacional.version_id == v.id))}
    nodos = db.execute(select(NodoArancel.codigo_norm, NodoArancel.descripcion, NodoArancel.dai).where(
        NodoArancel.version_id == v.id, NodoArancel.nivel == "INCISO")).all()
    # Lo que el clasificador lee del texto oficial vive en el motor, no en el archivo oficial
    interpretacion = json.loads((MOTOR / "interpretacion_aci.json").read_text(encoding="utf-8"))["condiciones"]
    fuente = db.get(FuenteOficial, v.fuente_id) if v.fuente_id else None
    nota = f"{fuente.dataset} — {v.etiqueta}" if fuente else v.etiqueta
    filas, con_cond = [], []
    for iso in paises:
        for cod, desc, dai in nodos:
            if (iso, cod) in ya:
                continue
            # DAI de la Parte II: distinto por país, no está en el texto regional (queda vacío, con la nota)
            parte2 = bool(dai) and dai.upper().startswith("PARTE II")
            f = {"pais": iso, "codigo": cod, "sub6": cod[:6], "dai": None if parte2 else (dai or None), "descripcion": (desc or "")[:300],
                 "fuente": "oficial", "fuente_id": v.fuente_id, "version_id": v.id, "vigente_desde": v.vigente_desde,
                 "vigente_hasta": v.vigente_hasta, "codigo_base": cod[:8], "url": fuente.url if fuente else None, "activo": True,
                 "nota": (f"{nota} · DAI in Part II of the ACI (country-specific)" if parte2 else nota)[:300]}
            (con_cond if cod in interpretacion else filas).append(f)
    if filas:
        db.execute(insert(IncisoNacional), filas)
    for f in con_cond:
        x = IncisoNacional(**f, cond=interpretacion[f["codigo"]])
        x.regla.tipo_fuente = "CLASSIFIER"  # lo que el clasificador lee del texto oficial: motor, no ley ni empresa
        db.add(x)
    db.flush()
    return len(filas) + len(con_cond)


def cargar_notas_incluidas(db: Session, version: str = "SAC-2025-V6") -> int:
    """Notas legales oficiales del ACI (RGI, sección, capítulo, subpartida y
    complementarias) con su fuente y versión, y aparte la guía interna del
    clasificador (resúmenes propios de las Notas Explicativas)."""
    import json

    from app.modelos import NotaSAC, VersionDataset

    if db.scalar(select(func.count()).select_from(NotaSAC)):
        return 0
    v = db.scalar(select(VersionDataset).where(VersionDataset.codigo == version))
    leer = lambda nombre: json.loads((DATOS / nombre).read_text(encoding="utf-8"))  # noqa: E731
    n = 0
    for x in leer("sac_notas.json"):
        db.add(NotaSAC(ambito=x["ambito"], codigo=x["codigo"], numero=x["numero"], texto=x["texto"], capitulos=x.get("capitulos") or [],
                       claves=x.get("claves") or [], tipo_fuente="OFFICIAL_LEGAL", version_id=v.id if v else None, fuente_id=v.fuente_id if v else None,
                       vigente_desde=v.vigente_desde if v else None))
        n += 1
    for x in json.loads((MOTOR / "sac_explicativas.json").read_text(encoding="utf-8")):
        db.add(NotaSAC(ambito=x["ambito"], codigo=x["codigo"], numero=x["numero"], texto=x["texto"], capitulos=x.get("capitulos") or [],
                       claves=x.get("claves") or [], tipo_fuente="CLASSIFIER_GUIDANCE"))
        n += 1
    db.flush()
    return n


# ---- Consulta y edición ----------------------------------------------------------
def ambito_version(codigo: str, fuente=None, explicito: str | None = None) -> str:
    """REGIONAL o el ISO del país de una versión: lo que diga el paquete, si
    no el ámbito de su fuente, si no el prefijo del código (CR-ATENA → CR)."""
    for v in (explicito, fuente.ambito if fuente else None):
        v = (v or "").strip().upper()
        if re.fullmatch(r"[A-Z]{2}", v):
            return v
        if v.startswith("REGION"):
            return "REGIONAL"
    m = re.match(r"^([A-Z]{2})-", codigo or "")
    return m.group(1) if m and not codigo.startswith("SAC") else "REGIONAL"


def _fecha_en_texto(t: str | None) -> date | None:
    m = re.search(r"(\d{4}-\d{2}-\d{2})", t or "")
    return date.fromisoformat(m.group(1)) if m else None


# ---- Trazabilidad: qué le falta a una fuente o versión para respaldar datos oficiales -----
VERIFICACION_MAX_DIAS = 365


def problemas_fuente(f: FuenteOficial | None) -> list[str]:
    if not f:
        return ["No official source."]
    out = []
    if not (f.dataset or f.documento):
        out.append(f"Source {f.codigo} does not name its official document or dataset.")
    if not (f.url or f.documento):
        out.append(f"Source {f.codigo} has no official link or document reference.")
    if not f.verificado_en:
        out.append(f"Source {f.codigo} has not been verified against the official publication (no verification date).")
    if not f.activo:
        out.append(f"Source {f.codigo} is inactive.")
    return out


def problemas_version(v: VersionDataset | None) -> list[str]:
    """Una versión respalda datos oficiales solo si su fuente es trazable y,
    si está publicada, tiene vigencia."""
    if not v:
        return ["No official version."]
    out = [] if v.fuente else [f"Version {v.codigo} has no official source."]
    out += problemas_fuente(v.fuente) if v.fuente else []
    if v.estado == "PUBLICADA" and not v.vigente_desde:
        out.append(f"Version {v.codigo} is published without its validity (valid from).")
    return out


def exigir_trazable(v: VersionDataset | None) -> None:
    if p := problemas_version(v):
        raise ErrorNegocio("Official data needs a traceable source: " + " ".join(p), 422, "fuente_no_trazable", [{"mensaje": x} for x in p])


def verificar_fuente(db: Session, user: Usuario, fuente_id: int, datos: dict) -> dict:
    """Registra que alguien comparó la fuente con la publicación oficial: el
    documento o dataset exacto, el enlace y la fecha. Queda en la bitácora."""
    exigir(user, "aranceles.editar")
    f = db.get(FuenteOficial, fuente_id)
    if not f:
        raise ErrorNegocio("The source does not exist.", 404, "no_encontrado")
    fecha = datos.get("verificado_en") or date.today()
    if isinstance(fecha, str):
        fecha = date.fromisoformat(fecha[:10])
    if fecha > date.today():
        raise ErrorNegocio("The verification date cannot be in the future.", 422, "validacion")
    if datos.get("documento"):
        f.documento = datos["documento"].strip()[:300]
    if datos.get("url"):
        f.url = datos["url"].strip()[:400]
    if not (f.url or f.documento):
        raise ErrorNegocio("Give the official document or the link that was checked.", 422, "validacion")
    f.verificado_en, f.verificado_por = fecha, (user.nombre or user.email)[:120]
    f.verificacion = f"Verified {fecha.isoformat()} by {f.verificado_por}"[:120]
    db.flush()
    registrar(db, user, "aranceles", f.id, "fuente_verificada", {"fuente": f.codigo, "fecha": fecha.isoformat(), "documento": f.documento,
                                                                 "url": f.url})
    return _fuente_dict(f)


def _fuente_dict(x: FuenteOficial) -> dict:
    d = {c: getattr(x, c) for c in ("id", "codigo", "ambito", "autoridad", "dataset", "uso", "url", "acceso",
                                     "autenticacion", "nota_version", "verificacion", "documento", "verificado_en", "verificado_por", "activo")}
    d["problemas"] = problemas_fuente(x)
    return d


def _version_dict(x: VersionDataset) -> dict:
    return {"id": x.id, "codigo": x.codigo, "dataset": x.dataset, "etiqueta": x.etiqueta, "estado": x.estado,
            "vigente_desde": x.vigente_desde, "vigente_hasta": x.vigente_hasta, "nota": x.nota,
            "fuente": x.fuente.codigo if x.fuente else None, "importado_en": x.importado_en}


def fuentes_y_versiones(db: Session, user: Usuario) -> dict:
    exigir(user, "aranceles.ver")
    return {"fuentes": [_fuente_dict(x) for x in db.scalars(select(FuenteOficial).order_by(FuenteOficial.ambito, FuenteOficial.codigo))],
            "versiones": [_version_dict(x) for x in db.scalars(select(VersionDataset).order_by(VersionDataset.codigo))]}


def _capitulo_dict(x: ControlCapitulo, dominios: dict) -> dict:
    return {"id": x.id, "capitulo": x.capitulo, "seccion": x.seccion, "titulo": x.titulo, "activo": x.activo,
            "clasificacion": x.clasificacion, "candidato_auto": x.candidato_auto, "solo_manual": x.solo_manual,
            "archivado": x.archivado, "alcance": x.alcance, "nota": x.nota,
            "version": x.version.codigo if x.version else None, "dominios": dominios.get(x.capitulo, [])}


def capitulos(db: Session, user: Usuario, q: str | None = None, estado: str | None = None,
              dominio: str | None = None) -> dict:
    exigir(user, "aranceles.ver")
    dominios: dict[str, list] = {}
    for dc in db.scalars(select(DominioCapitulo)):
        dominios.setdefault(dc.capitulo, []).append({"codigo": dc.dominio.codigo, "nombre": dc.dominio.nombre,
                                                       "relevancia": dc.relevancia, "habilitado": dc.habilitado})
    consulta = select(ControlCapitulo).order_by(ControlCapitulo.capitulo)
    if q:
        consulta = consulta.where(filtro_texto(q, lambda p: [ControlCapitulo.capitulo.ilike(p), ControlCapitulo.titulo.ilike(p),
                                                             ControlCapitulo.seccion.ilike(p), ControlCapitulo.alcance.ilike(p)]))
    filas = [_capitulo_dict(x, dominios) for x in db.scalars(consulta)]
    if estado == "habilitados":
        filas = [f for f in filas if f["activo"] and f["clasificacion"]]
    elif estado == "inactivos":
        filas = [f for f in filas if not f["activo"]]
    elif estado == "manuales":
        filas = [f for f in filas if f["solo_manual"]]
    if dominio:
        filas = [f for f in filas if any(d["codigo"] == dominio for d in f["dominios"])]
    total = db.scalar(select(func.count()).select_from(ControlCapitulo)) or 0
    habilitados = db.scalar(select(func.count()).select_from(ControlCapitulo).where(
        ControlCapitulo.activo.is_(True), ControlCapitulo.clasificacion.is_(True))) or 0
    return {"items": filas, "total": total, "habilitados": habilitados}


CAMPOS_CAPITULO = ("activo", "clasificacion", "candidato_auto", "solo_manual", "archivado")


def actualizar_capitulos(db: Session, user: Usuario, ids: list[int], campos: dict) -> dict:
    """Cambia en bloque los controles de varios capítulos. Mantiene la
    coherencia: inactivo o archivado no clasifica; solo manual no genera
    candidatos automáticos."""
    exigir(user, "aranceles.editar")
    cambios = {k: bool(v) for k, v in campos.items() if k in CAMPOS_CAPITULO and v is not None}
    if not cambios or not ids:
        raise ErrorNegocio("Choose chapters and what to change.", 422, "validacion")
    caps = list(db.scalars(select(ControlCapitulo).where(ControlCapitulo.id.in_(ids))))
    for x in caps:
        for k, v in cambios.items():
            setattr(x, k, v)
        if cambios.get("solo_manual"):
            x.candidato_auto = False
        if cambios.get("candidato_auto"):
            x.solo_manual = False
        if cambios.get("clasificacion") or cambios.get("candidato_auto"):
            x.activo, x.archivado = True, False  # habilitar implica activar
        if not x.activo or x.archivado:
            x.clasificacion = False
            x.candidato_auto = False
    registrar(db, user, "aranceles", 0, "capitulos", {"capitulos": [x.capitulo for x in caps], "cambios": cambios})
    return {"actualizados": len(caps)}


def dominios(db: Session, user: Usuario) -> list[dict]:
    exigir(user, "clasificacion.ver")
    caps = {x.capitulo: x for x in db.scalars(select(ControlCapitulo))}
    return [{"id": d.id, "codigo": d.codigo, "nombre": d.nombre, "descripcion": d.descripcion, "modo": d.modo,
             "activo": d.activo, "estado": d.estado,
             "capitulos": sorted([{"id": c.id, "capitulo": c.capitulo, "relevancia": c.relevancia, "habilitado": c.habilitado,
                                   "titulo": caps[c.capitulo].titulo if c.capitulo in caps else "",
                                   "capitulo_habilitado": bool(caps.get(c.capitulo) and caps[c.capitulo].clasificacion)}
                                  for c in d.capitulos], key=lambda c: (c["relevancia"], c["capitulo"]))}
            for d in db.scalars(select(DominioClasificacion).order_by(DominioClasificacion.orden, DominioClasificacion.codigo))]


def guardar_dominio_capitulo(db: Session, user: Usuario, dominio_id: int, capitulo: str, relevancia: str | None,
                             habilitado: bool | None, quitar: bool = False) -> dict:
    exigir(user, "clasificacion.configurar")
    d = db.get(DominioClasificacion, dominio_id)
    cap = _capitulo(capitulo)
    if not d or not cap or not db.scalar(select(ControlCapitulo.id).where(ControlCapitulo.capitulo == cap)):
        raise ErrorNegocio("The domain or chapter does not exist.", 404, "no_encontrado")
    x = db.scalar(select(DominioCapitulo).where(DominioCapitulo.dominio_id == d.id, DominioCapitulo.capitulo == cap))
    if quitar:
        if x:
            db.delete(x)
        return {"ok": True}
    x = x or DominioCapitulo(dominio=d, capitulo=cap)
    if relevancia:
        x.relevancia = "SECONDARY" if relevancia.upper().startswith("SEC") else "PRIMARY"
    if habilitado is not None:
        x.habilitado = habilitado
    db.add(x)
    db.flush()
    return {"id": x.id}
