"""Factura comercial, packing list y reportes en Excel, con la misma estructura
que el PDF (documentos.py): cabecera, partes, condiciones, detalle, totales,
total en letras y declaración. Listos para imprimir en una página de ancho."""
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.modulos.acceso import visibilidad
from app.modulos.acceso.preferencias import fecha_hora_txt
from app.modulos.documentos import idioma_doc
from app.modulos.documentos.documentos import _acento_hex, _fecha, _por_unidad_txt, declaracion
from app.modulos.documentos.idioma_doc import L


def _acento() -> str:
    """Color de la empresa (o el de fábrica) en formato de Excel."""
    return _acento_hex().lstrip("#").upper()


TENUE = "5B6475"
_FONDO = PatternFill("solid", fgColor="F3F1FB")
_FONDO_2 = PatternFill("solid", fgColor="F7F8FA")
_LINEA = Side(style="thin", color="D9DCE3")
_ARRIBA = Alignment(vertical="top", wrap_text=True)
_MONEDA = "#,##0.00"
_ENTERO = "#,##0"
_DECIMAL = "#,##0.000"


class _Hoja:
    """Escribe de arriba abajo sobre una hoja de `columnas` columnas."""

    def __init__(self, ws, anchos: list[float]):
        self.ws = ws
        self.n = len(anchos)
        self.fila = 1
        for i, a in enumerate(anchos, start=1):
            ws.column_dimensions[get_column_letter(i)].width = a
        ws.sheet_view.showGridLines = False

    def celda(self, fila, col, valor=None, negrita=False, color=None, tam=None, formato=None, fondo=None,
              alinear=None):
        c = self.ws.cell(row=fila, column=col, value=valor)
        c.font = Font(bold=negrita, color=color, size=tam or 9)
        if formato == _ENTERO and isinstance(valor, float) and valor != int(valor):
            formato = _DECIMAL  # una cantidad medida (kg, litros, metros) lleva sus decimales
        c.alignment = alinear or _ARRIBA
        if formato:
            c.number_format = formato
        if fondo:
            c.fill = fondo
        return c

    def unir(self, fila, c1, c2, filas=1):
        if c2 > c1 or filas > 1:
            self.ws.merge_cells(start_row=fila, start_column=c1, end_row=fila + filas - 1, end_column=c2)

    def marco(self, f1, f2, c1, c2, fondo=None, borde=_LINEA):
        for r in range(f1, f2 + 1):
            for c in range(c1, c2 + 1):
                x = self.ws.cell(row=r, column=c)
                x.border = Border(left=borde if c == c1 else None, right=borde if c == c2 else None,
                                  top=borde if r == f1 else None, bottom=borde if r == f2 else None)
                if fondo:
                    x.fill = fondo

    def espacio(self, n=1):
        self.fila += n

    def cabecera(self, d: dict, titulo: str, subtitulo: str, datos: list[tuple[str, str]]):
        """Exportador a la izquierda; título, número y fecha a la derecha."""
        exp = d["exportador"]
        corte = max(2, self.n - 3)
        self.celda(1, 1, exp["nombre"], negrita=True, tam=15)
        self.unir(1, 1, corte)
        self.celda(2, 1, " · ".join(x for x in (exp.get("direccion"), exp.get("pais")) if x), color=TENUE, tam=8)
        self.unir(2, 1, corte)
        self.celda(3, 1, " · ".join(x for x in (L("Tax ID: {0}", exp['id_fiscal']) if exp.get("id_fiscal") else None,
                                                  exp.get("correos"), exp.get("telefono")) if x), color=TENUE, tam=8)
        self.unir(3, 1, corte)
        self.celda(1, corte + 1, titulo, negrita=True, color=_acento(), tam=13)
        self.unir(1, corte + 1, self.n)
        self.celda(2, corte + 1, subtitulo, color=TENUE, tam=8)
        self.unir(2, corte + 1, self.n)
        fila = 3
        for k, v in datos:
            self.celda(fila, corte + 1, k, color=TENUE, tam=8)
            self.celda(fila, corte + 2, v, negrita=True)
            self.unir(fila, corte + 2, self.n)
            fila += 1
        self.marco(1, fila - 1, corte + 1, self.n, _FONDO, Side(style="thin", color=_acento()))
        self.fila = fila + 1

    def partes(self, d: dict):
        """Exportador, importador y consignatario en tres bloques."""
        exp = d["exportador"]
        bloques = [
            (L("EXPORTER / SELLER"), [f"{exp['codigo']} · {exp['nombre']}", L("Tax ID: {0}", exp['id_fiscal'] or '—'),
                                       " · ".join(x for x in (exp.get("direccion"), exp.get("pais")) if x),
                                       " · ".join(x for x in (exp.get("contacto"), exp.get("correos"),
                                                               exp.get("telefono")) if x)]),
            (L("IMPORTER / BILL TO"), _lineas_parte(d["importador"])),
            (L("CONSIGNEE / NOTIFY PARTY"), _lineas_parte(d["consignatario"])),
        ]
        tercio = self.n // 3
        rangos = [(1, tercio), (tercio + 1, 2 * tercio), (2 * tercio + 1, self.n)]
        f = self.fila
        alto = max(len(b[1]) for b in bloques)
        for (titulo, lineas), (c1, c2) in zip(bloques, rangos, strict=False):
            self.celda(f, c1, titulo, negrita=True, color=_acento(), tam=7)
            self.unir(f, c1, c2)
            for i in range(alto):
                v = lineas[i] if i < len(lineas) else None
                self.celda(f + 1 + i, c1, v, negrita=i == 0, color=None if i == 0 else TENUE, tam=9 if i == 0 else 8)
                self.unir(f + 1 + i, c1, c2)
            self.marco(f, f + alto, c1, c2)
        for i in range(alto):
            largo = max(len(b[1][i]) if i < len(b[1]) else 0 for b in bloques)
            if largo > 38:
                self.ws.row_dimensions[f + 1 + i].height = 24
        self.fila = f + alto + 2

    def rejilla(self, pares: list[tuple[str, str]], columnas: int = 4):
        """Pares etiqueta/valor en `columnas` bloques por fila."""
        ancho = self.n // columnas
        for i in range(0, len(pares), columnas):
            f = self.fila
            for j, (k, v) in enumerate(pares[i:i + columnas]):
                c1 = 1 + j * ancho
                c2 = self.n if j == columnas - 1 else c1 + ancho - 1
                self.celda(f, c1, k.upper(), negrita=True, color=_acento(), tam=7)
                self.unir(f, c1, c2)
                self.celda(f + 1, c1, v if v not in (None, "") else "—")
                self.unir(f + 1, c1, c2)
                self.marco(f, f + 1, c1, c2, _FONDO_2)
            self.fila += 2
        self.fila += 1

    def tabla(self, encabezados: list[tuple[str, str | None]], filas: list[list], pie: list | None = None,
              texto: set[int] = frozenset()):
        """Detalle con encabezado fijo, filas cebra y fila de totales.
        `encabezados`: (título, formato numérico o None). `texto`: columnas
        (desde 0) que van como texto para no perder ceros a la izquierda."""
        f = self.fila
        for i, (t, fmt) in enumerate(encabezados, start=1):
            c = self.celda(f, i, t, negrita=True, color=TENUE, tam=8, fondo=_FONDO,
                           alinear=Alignment(wrap_text=True, vertical="center",
                                             horizontal="right" if fmt else "left"))
            c.border = Border(bottom=Side(style="medium", color=_acento()))
        self.ws.row_dimensions[f].height = 26
        inicio = f + 1
        for n, valores in enumerate(filas):
            r = inicio + n
            for i, (v, (_, fmt)) in enumerate(zip(valores, encabezados, strict=False), start=1):
                c = self.celda(r, i, v, formato=fmt if fmt else ("@" if (i - 1) in texto else None),
                               fondo=_FONDO_2 if n % 2 else None)
                c.border = Border(bottom=_LINEA)
        r = inicio + len(filas)
        if pie:
            for i, (v, (_, fmt)) in enumerate(zip(pie, encabezados, strict=False), start=1):
                c = self.celda(r, i, v, negrita=True, formato=fmt, fondo=_FONDO)
                c.border = Border(top=Side(style="medium", color="1F2430"))
            r += 1
        self.fila = r + 1
        return f

    def parrafo(self, texto: str, negrita=False, color=None, tam=None, alto=None, hasta=None):
        self.celda(self.fila, 1, texto, negrita=negrita, color=color, tam=tam)
        self.unir(self.fila, 1, hasta or self.n)
        if alto:
            self.ws.row_dimensions[self.fila].height = alto
        self.fila += 1

    def firma(self, texto: str):
        self.espacio()
        mitad = max(2, int(self.n * 0.6))
        self.celda(self.fila, 1, texto, color=TENUE, tam=8)
        self.unir(self.fila, 1, mitad, filas=2)
        self.ws.row_dimensions[self.fila].height = 24
        for col in range(mitad + 2, self.n + 1):
            self.ws.cell(row=self.fila + 1, column=col).border = Border(bottom=Side(style="thin", color="1F2430"))
        self.celda(self.fila + 2, mitad + 2, L("Name, title and signature of the exporter"), color=TENUE, tam=8)
        self.unir(self.fila + 2, mitad + 2, self.n)
        self.fila += 4

    def imprimir(self, titulo_filas: int | None, horizontal: bool, pie: str, borrador: bool):
        ws = self.ws
        ws.page_setup.orientation = "landscape" if horizontal else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = ws.page_margins.right = 0.4
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.6
        if titulo_filas:
            ws.print_title_rows = f"{titulo_filas}:{titulo_filas}"
        ws.oddFooter.left.text = pie + (L("  ·  DRAFT, NOT OFFICIAL") if borrador else "")
        ws.oddFooter.left.size = 7
        ws.oddFooter.right.text = L("Page &P of &N")
        ws.oddFooter.right.size = 7


def _lineas_parte(p: dict | None) -> list[str]:
    if not p:
        return ["—"]
    lineas = [f"{p.get('codigo') or ''} · {p.get('razon_social') or p.get('nombre') or ''}"]
    if p.get("id_fiscal"):
        lineas.append(L("Tax ID (NIT / RUC): {0}", p['id_fiscal']))
    lineas.append(" · ".join(x for x in (p.get("direccion"), p.get("pais")) if x) or "—")
    if p.get("puerto"):
        lineas.append(L("Port of arrival: {0}", p.get('puerto_nombre') or p['puerto']))
    c = (p.get("contactos") or [None])[0]
    if c:
        lineas.append(" · ".join(x for x in (c.get("nombre"), c.get("correos"), c.get("telefono")) if x))
    elif p.get("correos"):
        lineas.append(p["correos"])
    return lineas


def _guardar(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _hoja_titulo(texto: str) -> str:
    for ch in "[]:*?/\\":
        texto = texto.replace(ch, "-")
    return texto[:31] or L("Sheet")


def exportar_factura(d: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _hoja_titulo(L("Invoice {0}", d['numero']))
    h = _Hoja(ws, [13, 7, 14, 15, 16, 10, 7, 30, 9, 11, 10, 7, 12, 14])
    tr = d["transporte"] or {}
    h.cabecera(d, L("COMMERCIAL INVOICE"), L("Customs invoice · original"), [
        (L("No."), d["numero"]), (L("Date"), _fecha(d["fecha"])),
        (L("Status"), L("Official") if d["oficial"] else L("Draft"))])
    h.partes(d)
    h.rejilla([
        (L("Incoterm"), d["incoterm"]), (L("Currency"), d["moneda"]), (L("Payment terms"), d["condiciones"]),
        (L("Mode of transport"), tr.get("modo")),
        (L("Country of origin"), d["pais_origen"]), (L("Country of shipment"), d["pais_procedencia"]),
        (L("Country of destination"), d["pais_destino"]), (L("Carrier"), tr.get("transportista")),
        (L("Port of loading"), tr.get("puerto_origen") or d["puerto_embarque"]),
        (L("Port of discharge"), tr.get("puerto_destino") or d["consignatario"].get("puerto_nombre")),
        (L("B/L / AWB / waybill"), tr.get("documento")), (L("Packing lists"), ", ".join(d["pls"]) or "—"),
        *d.get("propios", []),
    ])
    t = d["totales"]
    filas = [[l["oc"], l["posicion"], l["sku"], l["upc"], l["marca"], l["estilo"], l["talla"],
              f"{l['descripcion'] or ''} · {l['color'] or ''}" + (L(" · prepack {0}", l['prepack']) if l["prepack"] else ""),
              l["origen"], l["partida"], l["cantidad"], l["unidad"], l["precio"], l["total"]] for l in d["lineas"]]
    fila_cab = h.tabla(
        [(L("PO"), None), (L("Line"), None), (L("Item code"), None), (L("UPC"), None), (L("Brand"), None), (L("Style"), None),
         (L("Size"), None), (L("Customs description · color"), None), (L("Origin"), None), (L("HS code (SAC)"), None),
         (L("Quantity"), _ENTERO), (L("UoM"), None), (L("Unit price"), "#,##0.0000"), (L("Amount {0}", d['moneda']), _MONEDA)],
        filas, pie=[L("Total"), None, L("{0} lines", len(filas)), None, None, None, None, None, None, None,
                    sum(l["cantidad"] for l in d["lineas"]), None, None, t["importe"]],
        texto={0, 1, 2, 3, 9})
    h.rejilla([
        (L("Total quantity"), _por_unidad_txt(t["por_unidad"])),
        (L("Packages"), L("{0} cartons", f"{t['bultos']:,}") + (L(" on {0} pallets", t['pallets']) if t["pallets"] else "")),
        (L("Net weight"), f"{t['peso_neto']:,.2f} kg"), (L("Gross weight"), f"{t['peso_bruto']:,.2f} kg"),
        (L("Volume"), f"{t['cbm']:,.3f} m³"), (L("Containers / AWB"), ", ".join(tr.get("unidades", [])) or "—"),
        (L("Total value ") + (d["incoterm"] or ""), f"{d['moneda']} {t['importe']:,.2f}"),
        (L("Purchase orders"), L("{0} (see lines)", len(d['ocs']))),
    ])
    h.parrafo(L("SAY: {0}", d['total_letras']), negrita=True)
    if d.get("observaciones"):
        h.parrafo(L("Remarks: {0}", d['observaciones']), color=TENUE, tam=8)
    h.firma(declaracion("declaracion_factura", L(
        "We declare under oath that the information in this invoice is true and correct, that the value is the price "
        "actually paid or payable for the goods and that the declared origin is correct.")))
    h.imprimir(fila_cab, False, L("Commercial invoice {0} · {1}", d['numero'], d['exportador']['nombre']), not d["oficial"])
    ws.freeze_panes = None
    return _guardar(wb)


def exportar_pl(d: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = _hoja_titulo(f"{d['numero']} {d['numero_pl']}")
    h = _Hoja(ws, [8, 6, 12, 6, 14, 15, 14, 10, 6, 8, 10, 8, 6, 8, 8, 8, 9, 9, 10, 10, 8, 7, 11, 18])
    tr = d["transporte"] or {}
    tp = d["totales_pl"]
    h.cabecera(d, L("PACKING LIST"), L("Detailed carton list"), [
        (L("No."), f"{d['numero']} · {d['numero_pl']}"), (L("Date"), _fecha(d["fecha"])), (L("Invoice"), d["numero"]),
        (L("Status"), L("Official") if d["oficial"] else L("Draft"))])
    h.partes(d)
    h.rejilla([
        (L("Mode of transport"), tr.get("modo")),
        (L("Carrier"), tr.get("transportista")), (L("B/L / AWB / waybill"), tr.get("documento")),
        (L("Container / AWB"), ", ".join(tr.get("unidades", [])) or "—"),
        (L("Port of loading"), tr.get("puerto_origen") or d["puerto_embarque"]),
        (L("Port of discharge"), tr.get("puerto_destino")), (L("Country of origin"), d["pais_origen"]),
        (L("Country of destination"), d["pais_destino"]),
        (L("Final destination"), d["destinos"][0] if len(d["destinos"]) == 1 else L("Per carton (see lines)")),
        (L("ETD / ETA"), f"{_fecha(tr.get('etd'))} / {_fecha(tr.get('eta'))}"),
    ], columnas=5)
    filas = []
    for g in d["grupos"]:
        for i, it in enumerate(g["items"]):
            p = i == 0
            filas.append([
                g["rango"] if p else None, g["num_cajas"] if p else None, it["oc"], it["posicion"], it["sku"],
                it["upc"], it["marca"], it["descripcion"], it["estilo"], it["color"], it["talla"], it["inners"], it["inner_pack"],
                it["por_caja"], it["total"], it["unidad"], g["largo"] if p else None, g["ancho"] if p else None,
                g["alto"] if p else None, g["neto_caja"] if p else None, g["bruto_caja"] if p else None,
                g["neto_total"] if p else None, g["bruto_total"] if p else None, g["cbm"] if p else None,
                (g["pallet"] if g["pallet"] else None) if p else None, g["etiqueta"] if p else None,
                (", ".join(g["ocs"]) + (f" → {g['centro_destino']}" if g["centro_destino"] else "")) if p else None,
            ])
    fila_cab = h.tabla(
        [(L("Cartons"), None), (L("Qty"), _ENTERO), (L("PO"), None), (L("Line"), None), (L("Item code"), None), (L("UPC"), None),
         (L("Brand"), None), (L("Customs description"), None), (L("Style"), None), (L("Color"), None), (L("Size"), None),
         (L("Inner packs per carton"), _ENTERO), (L("Per inner pack"), _ENTERO), (L("Total per carton"), _ENTERO), (L("Total"), _ENTERO), (L("UoM"), None), (L("Length cm"), "0.0"), (L("Width cm"), "0.0"), (L("Height cm"), "0.0"),
         (L("Net/ctn kg"), "0.00"), (L("Gross/ctn kg"), "0.00"), (L("Net total kg"), _MONEDA), (L("Gross total kg"), _MONEDA),
         ("m³", "0.000"), (L("Pallet"), None), (L("Label"), None), (L("POs in carton → destination"), None)],
        filas, pie=[L("Total"), d["total_cajas"], None, None, None, None, None, None, None, None, None, None, None, None,
                    sum(it["total"] for g in d["grupos"] for it in g["items"]), None, None, None, None, None, None,
                    tp["peso_neto"], tp["peso_bruto"], tp["cbm"], None, None, None],
        texto={0, 2, 3, 4, 5})
    if d["pallets"]:
        h.parrafo(L("PALLETS"), negrita=True, color=_acento(), tam=8)
        h.tabla([(L("Pallet"), None), (L("Cartons"), _ENTERO), (L("Dimensions cm"), None), (L("Tare kg"), "0.0"), (L("Gross kg"), "0.00"),
                 ("m³", "0.000")],
                [[p["numero"], p["cajas"], p["medidas"], p["tara"], p["bruto"], p["cbm"]] for p in d["pallets"]])
    if d["sin_caja"]:
        h.parrafo(L("NOT YET PACKED"), negrita=True, color=_acento(), tam=8)
        h.tabla([(L("PO"), None), (L("Item code"), None), (L("Style"), None), (L("Size"), None), (L("Quantity"), _ENTERO),
                 (L("UM"), None)],
                [[x["oc"], x["sku"], x["estilo"], x["talla"], x["cantidad"], x["unidad"]] for x in d["sin_caja"]],
                texto={0, 1})
    h.rejilla([
        (L("Total packages"), L("{0} cartons", f"{d['total_cajas']:,}") + (L(" on {0} pallets", f"{tp['pallets']:,}") if tp["pallets"] else "")),
        (L("Quantity"), tp["por_unidad_txt"]), (L("Net weight"), f"{tp['peso_neto']:,.2f} kg"),
        (L("Gross weight"), f"{tp['peso_bruto']:,.2f} kg"), (L("Volume"), f"{tp['cbm']:,.3f} m³"),
    ], columnas=5)
    h.parrafo(L("TOTAL PACKAGES: {0}", d['total_bultos_letras']), negrita=True)
    consig = d["consignatario"]
    h.espacio()
    mitad = h.n // 2
    inicio = h.fila
    h.parrafo(L("SHIPPING MARKS"), negrita=True, color=_acento(), tam=7, hasta=mitad)
    for i, linea in enumerate((consig.get("nombre"), consig.get("direccion"),
                               L("Invoice {0} · PO per carton label", d['numero']),
                               L("Carton no. __ of {0} · Made in {1}", d['total_cajas'], d['pais_origen']))):
        h.parrafo(linea, negrita=i == 0, hasta=mitad)
    h.marco(inicio, h.fila - 1, 1, mitad)
    h.firma(declaracion("declaracion_packing", L(
        "We declare that the contents, numbering, dimensions and weights of the packages correspond to the goods "
        "shipped. Every unit or pair carries its individual label; inner packs carry an inner pack label with the "
        "product and the quantity inside.")))
    h.imprimir(fila_cab, True, L("Packing list {0} {1} · {2}", d['numero'], d['numero_pl'], d['exportador']['nombre']),
               not d["oficial"])
    return _guardar(wb)


def exportar_reporte(titulo: str, subtitulo: str, filtros: str, indicadores: list[tuple[str, str]],
                     columnas: list[tuple[str, float, bool]], filas: list[list], hojas: list[dict] | None = None) -> bytes:
    """Reporte tabular: título, filtros aplicados, indicadores y detalle con
    filtros de Excel. `hojas` agrega hojas de detalle adicionales con
    {"titulo", "columnas", "filas"}. Sin las columnas que el rol no ve."""
    indicadores, columnas, filas, hojas = visibilidad.reporte(indicadores, columnas, filas, hojas)
    titulo, subtitulo, indicadores, columnas, filas, hojas = idioma_doc.reporte(titulo, subtitulo, indicadores, columnas, filas, hojas)
    wb = Workbook()
    ws = wb.active
    ws.title = _hoja_titulo(titulo)
    _hoja_reporte(ws, titulo, subtitulo, filtros, indicadores, columnas, filas)
    for x in hojas or []:
        _hoja_reporte(wb.create_sheet(_hoja_titulo(x["titulo"])), x["titulo"], subtitulo, filtros, [],
                      x["columnas"], x["filas"])
    return _guardar(wb)


def _hoja_reporte(ws, titulo, subtitulo, filtros, indicadores, columnas, filas):
    h = _Hoja(ws, [max(8, min(40, w * 9)) for _, w, _ in columnas])
    h.celda(1, 1, titulo, negrita=True, tam=15)
    h.unir(1, 1, h.n)
    h.celda(2, 1, subtitulo, color=TENUE, tam=8)
    h.unir(2, 1, h.n)
    h.celda(3, 1, L("Generated {0} · {1}", fecha_hora_txt(datetime.now()), filtros or L("No filters")), color=TENUE, tam=8)
    h.unir(3, 1, h.n)
    for c in range(1, h.n + 1):
        ws.cell(row=3, column=c).border = Border(bottom=Side(style="medium", color=_acento()))
    h.fila = 5
    if indicadores:
        h.rejilla(indicadores, columnas=min(len(indicadores), max(1, h.n), 5))
    fmts = [(t, (_ENTERO if der else None)) for t, _, der in columnas]
    for i, (t, _, der) in enumerate(columnas):
        if der and any(isinstance(f[i], float) and not float(f[i]).is_integer() for f in filas):
            fmts[i] = (t, _MONEDA)
    fila_cab = h.tabla(fmts, filas, texto={i for i, (_, _, der) in enumerate(columnas) if not der})
    if filas:
        ws.auto_filter.ref = f"A{fila_cab}:{get_column_letter(h.n)}{fila_cab + len(filas)}"
    ws.freeze_panes = f"A{fila_cab + 1}"
    h.imprimir(fila_cab, True, titulo, False)


def exportar_ficha(d: dict, x: dict) -> bytes:
    """Ficha técnica en Excel: datos y clasificación en la primera hoja; la
    composición, los códigos por país y las tallas en las siguientes."""
    titulo = L("Technical sheet {0} · {1}", d['estilo'], d['color'])
    sub = " · ".join(str(v) for v in (d.get("nombre"), d.get("proveedor"), d.get("marca_nombre") or d.get("marca")) if v)
    estado = L("Version {0} · {1}", d.get('version_ficha') or 1, d.get('estado_txt') or d.get('estado') or '')
    filas = [[L("Classification"), k, str(v)] for k, v in x["clasificacion"]]
    if x["descripcion"]:
        filas.append([L("Classification"), L("Customs description"), x["descripcion"]])
    if x.get("descripcion_comercial"):
        filas.append([L("Classification"), L("Commercial description"), x["descripcion_comercial"]])
    filas += [[L("Product data"), k, str(v)] for k, v in x["datos"]]
    filas += [[L("Reasoning"), "", r] for r in x["razones"]]
    hojas = [
        {"titulo": L("Composition"), "columnas": [(L("Part"), 1.6, False), (L("Materials"), 5, False)], "filas": x["composicion"]},
        {"titulo": L("National codes"), "columnas": [(L("Country"), 0.8, False), (L("Code"), 1.6, False), (L("Duty (DAI)"), 0.9, False),
                                                  (L("Status"), 0.9, False), (L("Source"), 0.9, False)], "filas": x["partidas"]},
        {"titulo": L("Sizes"), "columnas": [(L("SKU"), 1.4, False), (L("UPC"), 1.4, False), (L("Size"), 0.7, False), (L("Unit"), 0.6, False),
                                         (L("Description"), 3, False)], "filas": x["tallas"]},
    ]
    return exportar_reporte(titulo, sub, estado, [], [(L("Section"), 1.4, False), (L("Field"), 1.8, False), (L("Value"), 6, False)],
                            filas, hojas)
