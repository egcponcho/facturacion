"""Marca y documentos de la empresa (Configuración → Empresa): color, nombre
del sistema y textos del ingreso; papel y declaraciones de los documentos.
Las cuentas de demostración solo se ofrecen en la demostración."""
import io
import re

from openpyxl import load_workbook

from app.core.config import settings


def _textos(contenido: bytes) -> str:
    wb = load_workbook(io.BytesIO(contenido))
    return " ".join(str(c.value) for ws in wb for fila in ws.iter_rows() for c in fila if c.value)


def test_pantalla_de_ingreso_sin_sesion(client):
    r = client.get("/api/publico/empresa").json()
    assert r["nombre"] and set(r["marca"]) >= {"color", "titulo", "ingreso_titulo", "ingreso_texto", "ingreso_ayuda"}
    assert r["demo"]["password"] and ["admin@demo.com", "Administrator", "Full access"] in r["demo"]["cuentas"]


def test_sin_cuentas_de_demostracion_en_produccion(client, monkeypatch):
    monkeypatch.setattr(settings, "SEED_DEMO", False)
    r = client.get("/api/publico/empresa").json()
    assert r["demo"] is None and "password" not in str(r)


def test_marca_valida_y_se_ve_en_el_menu(admin):
    assert admin.put("/organizacion", {"marca": {"color": "azul"}}).status_code == 422
    assert admin.put("/organizacion", {"marca": {"titulo": "x" * 61}}).status_code == 422
    try:
        r = admin.put("/organizacion", {"marca": {"color": "#0A7A5C", "titulo": "Portal de proveedores", "otra": "x"}})
        assert r.status_code == 200 and r.json()["marca"]["color"] == "#0A7A5C" and "otra" not in r.json()["marca"]
        org = admin.get("/auth/me").json()["organizacion"]
        assert org["marca"]["titulo"] == "Portal de proveedores"
    finally:
        admin.put("/organizacion", {"marca": {"color": "", "titulo": ""}})


def test_papel_y_declaraciones_de_los_documentos(admin, interno):
    assert admin.put("/organizacion", {"documentos": {"papel": "A3"}}).status_code == 422
    assert admin.put("/organizacion", {"documentos": {"logo_en_reportes": "si"}}).status_code == 422
    fid = interno.get("/facturas").json()["items"][0]["id"]
    pdf = lambda: interno.get(f"/facturas/{fid}/exportar", params={"formato": "pdf"}).content  # noqa: E731
    carta = re.search(rb"/MediaBox \[ 0 0 ([\d.]+) ([\d.]+) \]", pdf())
    assert carta and float(carta.group(1)) == 612  # carta (8.5 in) de fábrica
    try:
        assert admin.put("/organizacion", {"documentos": {"papel": "A4", "declaracion_factura": "Declaramos que todo es cierto."}}
                         ).status_code == 200
        a4 = re.search(rb"/MediaBox \[ 0 0 ([\d.]+) ([\d.]+) \]", pdf())
        assert a4 and round(float(a4.group(1))) == 595  # A4 (210 mm)
        assert "Declaramos que todo es cierto." in _textos(interno.get(f"/facturas/{fid}/exportar").content)
    finally:
        admin.put("/organizacion", {"documentos": {"papel": "LETTER", "declaracion_factura": ""}})
    assert "We declare under oath" in _textos(interno.get(f"/facturas/{fid}/exportar").content)


def test_reporte_con_y_sin_logo(admin, interno):
    """Con logo, el reporte lo lleva; un logo que no es una imagen real se rechaza al guardarlo."""
    png = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")
    try:
        for logo in ("data:image/png;base64,no-es-una-imagen", "data:image/png;base64,JVBERi0xLjQ=",
                     "data:image/svg+xml;base64,PHN2Zz48L3N2Zz4="):
            assert admin.put("/organizacion", {"logo": logo}).status_code == 422
        for logo, con_imagen in ((None, False), (png, True)):
            assert admin.put("/organizacion", {"logo": logo}).status_code == 200
            r = interno.get("/seguimiento/ordenes/exportar", params={"formato": "pdf"})
            assert r.status_code == 200 and r.content[:4] == b"%PDF"
            assert (b"/Subtype /Image" in r.content) is con_imagen
        assert admin.put("/organizacion", {"logo": png, "documentos": {"logo_en_reportes": False}}).status_code == 200
        assert b"/Subtype /Image" not in interno.get("/seguimiento/ordenes/exportar", params={"formato": "pdf"}).content
    finally:
        admin.put("/organizacion", {"logo": None, "documentos": {"logo_en_reportes": True}})
