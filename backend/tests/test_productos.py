"""Productos: ficha técnica, clasificación arancelaria y su paso a la OC y la factura."""
import io

from openpyxl import load_workbook

from test_flujo import _oc

PNG = (b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15\xc4\x89"
       b"\x00\x00\x00\rIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18\xd8N\x00\x00\x00\x00IEND\xaeB`\x82")


def _producto(api, estilo, color):
    items = api.get("/productos", params={"q": estilo, "size": 50}).json()["items"]
    return next(p for p in items if p["estilo"] == estilo and p["color"] == color)


def test_lista_contexto_y_separacion(tnf, vans, interno):
    r = interno.get("/productos", params={"size": 100}).json()
    assert r["kpis"]["total"] >= 7 and r["kpis"]["pendientes"] >= 2 and r["kpis"]["aprobados"] >= 5
    # Cada talla (artículo) pertenece a su producto estilo-color
    old = _producto(interno, "VN000EE3", "BLK Black")
    assert old["estado"] == "aprobado" and old["codigo"] == "6404.19" and old["skus"] >= 6
    # El prepack no se clasifica: toma el producto (y la partida) de sus sólidos
    assert old["n_prepacks"] >= 1 and old["rango_tallas"] == "7–12"
    assert old["descripcion_comercial"] == "CALZADO VANS" and old["codigo_generico"] == "30095125"
    assert old["paises_ok"] == old["paises_total"] == 6
    # El proveedor solo ve lo suyo y recibe 404 en lo ajeno
    assert all(p["proveedor"] == "Vans" for p in vans.get("/productos", params={"size": 100}).json()["items"])
    assert vans.get(f"/productos/{_producto(interno, 'NF0A5GLL', 'JK3 TNF Black')['id']}").status_code == 404
    # Filtros de estado
    pend = interno.get("/productos", params={"estado": "pendientes"}).json()["items"]
    assert {p["estado"] for p in pend} <= {"borrador", "sugerida", "observado", "revision"}
    ctx = interno.get("/clasificacion/contexto").json()
    assert len(ctx["destinos"]) == 6 and ctx["pais_base"] == "SV" and ctx["puede_aprobar"]
    # El contexto es solo para mostrar: ni historial ni códigos nacionales (los usa el motor del servidor)
    assert "incisos" not in ctx and "recs" not in ctx and ctx["categorias"] and ctx["capitulos"]
    assert not vans.get("/clasificacion/contexto").json()["puede_aprobar"]


def test_ficha_aprobacion_y_documentos(tnf, vans, interno):
    p = _producto(vans, "VN0A4BV4", "White")
    det = vans.get(f"/productos/{p['id']}").json()
    assert det["estado"] == "sugerida" and det["articulos"] and det["prepacks"][0]["codigo"] == "CD08"

    # La OC muestra la clasificación del producto; sin aprobar no hay partida
    oc = _oc(vans, "4400003902")
    pos = next(x for x in oc["posiciones"] if x["disponible"] > 0)
    assert pos["partida_arancelaria"] is None and pos["clasificacion"]["estado"] == "sugerida"
    # Se factura igual, pero no se puede finalizar sin la partida aprobada
    f = vans.post("/facturas", {"lineas": [{"posicion_id": pos["id"], "cantidad": pos["disponible"]}]}).json()
    f = vans.get(f"/facturas/{f['id']}").json()
    vans.patch(f"/facturas/{f['id']}", {"version": f["version"], "numero": "VN-TEST-CLASIF", "fecha": "2026-09-01",
                                       "incoterm": "FOB"})
    fd = vans.get(f"/facturas/{f['id']}").json()
    r = vans.post(f"/facturas/{f['id']}/finalizar", {"version": fd["version"]})
    assert r.status_code == 422
    assert any(e.get("codigo") == "sin_clasificar" for e in r.json()["detalle"])

    # El proveedor completa la ficha; aprobar es del equipo interno
    det = vans.put(f"/productos/{p['id']}/ficha", {
        "version": det["version"], "tipo": "calzado", "pais_origen": "CN",
        "ficha": {**det["ficha"], "uso": "Casual canvas sneaker"}}).json()
    assert det["estado"] == "sugerida" and det["sugerido"] == "6404.19" and det["partidas"]["PA"]["codigo"] == "640419970000"
    assert vans.post(f"/productos/{p['id']}/aprobar", {"version": det["version"]}).status_code == 403
    # Borrador → enviar a revisión: queda cerrada para el proveedor hasta que se revise o la retire
    assert vans.post("/productos/enviar", {"ids": [p["id"]]}).json()["enviados"] == 1
    det = vans.get(f"/productos/{p['id']}").json()
    assert det["estado"] == "revision" and any(h["accion"] == "enviado" for h in det["historial"])
    r = vans.put(f"/productos/{p['id']}/ficha", {"version": det["version"], "ficha": det["ficha"]})
    assert r.status_code == 409 and r.json()["codigo"] == "en_revision"
    assert vans.post(f"/productos/{p['id']}/retirar").json()["estado"] == "sugerida"
    vans.post("/productos/enviar", {"ids": [p["id"]]})
    revisar = interno.get("/productos", params={"estado": "revision"}).json()
    assert revisar["kpis"]["revision"] >= 1 and any(x["id"] == p["id"] for x in revisar["items"])
    det = interno.get(f"/productos/{p['id']}").json()
    # Un código nacional que no empieza con la subpartida se rechaza
    r = interno.post(f"/productos/{p['id']}/aprobar", {"version": det["version"], "codigo": "640419",
                                                        "partidas": {"SV": {"codigo": "6402991000", "estado": "ok"}}})
    assert r.status_code == 422
    det = interno.post(f"/productos/{p['id']}/aprobar", {
        "version": det["version"], "codigo": "6404.19",
        "partidas": {"SV": {"codigo": "6404199000", "estado": "ok", "fuente": "base"},
                     "PA": {"codigo": "640419970000", "estado": "auto", "fuente": "base"}}}).json()
    assert det["estado"] == "aprobado" and det["codigo"] == "6404.19" and det["revisado_por"]
    assert any(h["accion"] == "aprobado" for h in det["historial"]), det["historial"]

    # Ahora la OC lleva el código del país destino y la factura se completa al finalizar
    assert next(x for x in _oc(vans, "4400003902")["posiciones"] if x["id"] == pos["id"])["partida_arancelaria"] \
        == "6404.19.90.00"
    fd = vans.get(f"/facturas/{f['id']}").json()
    r = vans.post(f"/facturas/{f['id']}/finalizar", {"version": fd["version"]})
    assert r.status_code == 200, r.text
    linea = vans.get(f"/facturas/{f['id']}").json()["lineas"][0]
    assert linea["partida_arancelaria"] == "6404.19.90.00" and linea["pais_origen"]

    # Una ficha aprobada no se edita: se abre una versión nueva que vuelve a borrador
    assert vans.put(f"/productos/{p['id']}/ficha", {"version": det["version"], "ficha": {}}).status_code == 409
    det = vans.post(f"/productos/{p['id']}/versiones", {"version": det["version"], "motivo": "New outsole"}).json()
    assert det["estado"] == "borrador" and det["version_ficha"] == 2 and det["versiones"][0]["codigo"] == "6404.19"
    assert not det["partidas"]
    # La versión anterior se puede ver y descargar tal como quedó
    v1 = vans.get(f"/productos/{p['id']}/versiones/1").json()
    assert v1["version"] == 1 and v1["clasificacion"][0][1] == "6404.19"
    assert any(x[0] == "SV" and x[1] == "6404.19.90.00" for x in v1["partidas"]) and v1["composicion"]
    assert vans.get(f"/productos/{p['id']}/versiones/9").status_code == 404
    x = vans.get(f"/productos/{p['id']}/ficha", params={"formato": "xlsx", "version": 1})
    assert x.status_code == 200 and x.content[:2] == b"PK"
    wb = load_workbook(io.BytesIO(x.content))
    assert {"Composition", "National codes", "Sizes"} <= set(wb.sheetnames)
    assert vans.get(f"/productos/{p['id']}/ficha", params={"formato": "pdf"}).content[:4] == b"%PDF"


def test_devolver_aprobar_lote_y_aprendizaje(vans, interno):
    p = _producto(interno, "NF0A5GLL", "Summit blue")
    det = interno.get(f"/productos/{p['id']}").json()
    assert det["estado"] == "observado" and det["observaciones"]
    # Devolver exige decir qué corregir
    r = interno.post(f"/productos/{p['id']}/observar", {"version": det["version"], "devolver": True, "observaciones": ""})
    assert r.status_code == 422
    # Aprobar una ficha incompleta no se permite
    r = interno.post(f"/productos/{p['id']}/aprobar", {"version": det["version"], "codigo": "620140"})
    assert r.status_code == 422 and r.json()["codigo"] == "ficha_incompleta"
    # Aprobación en lote: solo las sugeridas con ficha completa
    r = interno.post("/productos/aprobar", {"ids": [p["id"]]}).json()
    assert r["aprobados"] == 0 and r["errores"]

    # Enseñar un código nacional: dígitos del país y condiciones
    assert interno.post("/clasificacion/incisos", {"pais": "PA", "codigo": "6404199"}).status_code == 422
    r = interno.post("/clasificacion/incisos", {"pais": "PA", "codigo": "640419990000", "cond": {"edadNac": "bebe"}})
    assert r.status_code == 200
    assert vans.post("/clasificacion/incisos", {"pais": "PA", "codigo": "640419990000"}).status_code == 403
    assert interno.post("/clasificacion/palabras", {"frase": "old skool", "tipo": "calzado",
                                                    "atributos": {"estiloCalz": "tenis"}}).status_code == 200
    assert vans.post("/clasificacion/sinonimos", {"palabra": "cordura", "equivale": "nylon"}).status_code == 200
    # Lo aprendido lo usa el motor: la palabra clave da categoría y estilo
    s = interno.post("/clasificacion/sesion", {"nombre": "Old Skool", "paises": False}).json()
    assert s["categoria"]["codigo"] == "calzado" and s["ficha"].get("estiloCalz") == "tenis"
    s = interno.post("/clasificacion/sesion", {"categoria": "mochila", "ficha": {"comp": {"exterior": "100% cordura"}}, "paises": False}).json()
    assert s["hechos"].get("exterior") == "textil"


def test_fotos(vans, interno):
    p = _producto(vans, "VN000EE3", "BLK Black")
    r = vans.c.post(f"/api/productos/{p['id']}/fotos", headers=vans.h,
                    files={"archivo": ("front.png", io.BytesIO(PNG), "image/png")})
    assert r.status_code == 200, r.text
    foto = r.json()["id"]
    r = vans.get(f"/productos/fotos/{foto}")
    assert r.status_code == 200 and r.content.startswith(b"\x89PNG")
    assert vans.c.post(f"/api/productos/{p['id']}/fotos", headers=vans.h,
                       files={"archivo": ("x.txt", io.BytesIO(b"hola"), "text/plain")}).status_code == 422
    assert _producto(vans, "VN000EE3", "BLK Black")["foto_id"] == foto
    assert vans.delete_(f"/productos/{p['id']}/fotos/{foto}").status_code == 200
    # El especialista con Claude solo existe si hay clave configurada
    assert interno.get("/productos/opciones").json()["especialista"] is False


def test_exportar(vans, interno):
    p = _producto(vans, "VN000EE3", "BLK Black")
    r = vans.get(f"/productos/{p['id']}/pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    for formato in ("xlsx", "pdf"):
        r = interno.get("/productos/exportar", params={"formato": formato, "estado": "pendientes"})
        assert r.status_code == 200 and len(r.content) > 1000
    # El reporte completo lleva la composición y el código guardado por país
    wb = load_workbook(io.BytesIO(interno.get("/productos/exportar", params={"formato": "xlsx", "estado": "aprobados"}).content))
    assert {"Composition", "National codes"} <= set(wb.sheetnames)
    cab = [c.value for fila in wb.worksheets[0].iter_rows(max_row=20) for c in fila]
    assert "Composition" in cab and "Code SV" in cab and "Code PA" in cab
    assert any(f[3].value == "SV" for f in wb["National codes"].iter_rows() if len(f) > 3)
    x = vans.get(f"/productos/{p['id']}/ficha", params={"formato": "xlsx"})
    assert x.status_code == 200 and "Composition" in load_workbook(io.BytesIO(x.content)).sheetnames
