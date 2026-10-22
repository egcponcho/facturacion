"""Los documentos (PDF y Excel) salen en el idioma que elige cada persona
(Perfil → idioma de los documentos) o el que pide la descarga (?idioma=)."""
import io

from openpyxl import load_workbook

from app.modulos.documentos import idioma_doc, letras


def _textos(contenido: bytes) -> str:
    wb = load_workbook(io.BytesIO(contenido))
    return " ".join(str(c.value) for ws in wb for fila in ws.iter_rows() for c in fila if c.value)


def test_factura_en_espanol_e_ingles(interno):
    fid = interno.get("/facturas").json()["items"][0]["id"]
    en = _textos(interno.get(f"/facturas/{fid}/exportar", params={"idioma": "en"}).content)
    es = _textos(interno.get(f"/facturas/{fid}/exportar", params={"idioma": "es"}).content)
    assert "COMMERCIAL INVOICE" in en and "SAY:" in en and "DOLLARS" in en
    assert "FACTURA COMERCIAL" in es and "SON:" in es and "DÓLARES" in es and "COMMERCIAL INVOICE" not in es
    # El PDF también se genera en los dos idiomas
    for idioma in ("es", "en"):
        r = interno.get(f"/facturas/{fid}/exportar", params={"idioma": idioma, "formato": "pdf"})
        assert r.status_code == 200 and r.content[:4] == b"%PDF"


def test_idioma_de_documentos_del_perfil(interno):
    fid = interno.get("/facturas").json()["items"][0]["id"]
    descargar = lambda: interno.c.get(f"/api/facturas/{fid}/exportar", headers=interno.h).content  # noqa: E731  (sin ?idioma=)
    assert interno.patch("/perfil", {"idioma_documentos": "en"}).json()["preferencias"]["idioma_documentos"] == "en"
    assert "COMMERCIAL INVOICE" in _textos(descargar())
    interno.patch("/perfil", {"idioma_documentos": "es"})
    assert "FACTURA COMERCIAL" in _textos(descargar())
    assert interno.patch("/perfil", {"idioma_documentos": "fr"}).status_code == 422
    interno.patch("/perfil", {"idioma_documentos": None})


def test_reporte_en_espanol(interno):
    es = _textos(interno.get("/seguimiento/ordenes/exportar", params={"idioma": "es"}).content)
    assert "Seguimiento de órdenes de compra" in es
    assert "Purchase order tracking" not in es


def test_montos_en_letras():
    idioma_doc.usar("es")
    assert letras.monto_en_letras(1_000_000, "USD") == "UN MILLÓN DE DÓLARES CON 00/100"
    assert letras.monto_en_letras(21.5, "USD") == "VEINTIÚN DÓLARES CON 50/100"
    assert letras.monto_en_letras(1, "XYZ") == "UN XYZ CON 00/100"
    idioma_doc.usar("en")
    assert letras.monto_en_letras(4850, "USD") == "FOUR THOUSAND EIGHT HUNDRED FIFTY US DOLLARS AND 00/100"


def test_nombres_aduaneros_del_pais_de_destino(admin, interno):
    """La nomenclatura y la identificación fiscal que se leen en la factura son
    datos del país (Arancel → país; catálogo de países), no textos fijos."""
    # Los países de la demostración (los que otras pruebas agregan no traen estos nombres)
    paises = [p for p in admin.get("/aranceles/paises").json() if p["iso"] in {"GT", "SV", "HN", "NI", "CR", "PA"}]
    assert {p["nomenclatura"] for p in paises} == {"SAC"} and {p["nombre_derecho"] for p in paises} == {"DAI"}
    fid = interno.get("/facturas").json()["items"][0]["id"]
    en = _textos(interno.get(f"/facturas/{fid}/exportar", params={"idioma": "en"}).content)
    assert "HS code (SAC)" in en
    # Otro nombre en todos los países: la factura lo toma, sin cambiar el código
    for p in paises:
        cuerpo = {k: p[k] for k in ("iso", "nombre", "digitos", "mcca", "impuesto", "nota", "base_legal", "activo",
                                     "nivel_base", "modelo_arancel", "contexto")}
        assert admin.put(f"/aranceles/paises/{p['id']}", {**cuerpo, "nomenclatura": "XYZ", "nombre_derecho": "IMP"}).status_code == 200
    try:
        en = _textos(interno.get(f"/facturas/{fid}/exportar", params={"idioma": "en"}).content)
        assert "HS code (XYZ)" in en and "HS code (SAC)" not in en
    finally:
        for p in paises:
            cuerpo = {k: p[k] for k in ("iso", "nombre", "digitos", "mcca", "impuesto", "nota", "base_legal", "activo",
                                         "nivel_base", "modelo_arancel", "contexto", "nomenclatura", "nombre_derecho")}
            admin.put(f"/aranceles/paises/{p['id']}", cuerpo)
    # La identificación fiscal se nombra como en el país del exportador o del importador
    assert any(x in en for x in ("Tax ID (NIT)", "Tax ID (RUC)", "Tax ID (RTN)", "Tax ID:"))
