"""Un producto de cualquier rubro se crea con su núcleo (código, nombre, unidad)
y lo demás lo pide la empresa o la familia (productos/perfil.py):

- sin proveedor, marca ni grupo, el producto y sus variantes se crean y se listan;
- un dato que la empresa vuelve obligatorio (Datos maestros → Artículos) se exige;
- un dato que la familia vuelve obligatorio se exige al crear sus variantes, también en la carga masiva;
- la plantilla de carga ya no pide datos de un rubro (estilo, color, talla)."""
import io

from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.db import SessionLocal
from app.modelos import Articulo, Producto


def _borrar(*gens):
    with SessionLocal() as db:
        for gen in gens:
            for a in db.scalars(select(Articulo).where(Articulo.generico == gen)):
                db.delete(a)
            db.flush()
            p = db.scalar(select(Producto).where(Producto.codigo_generico == gen))
            if p is not None:
                db.delete(p)
        db.commit()


def _atributo(api, codigo):
    r = api.post("/catalogos/atributos_variante", {"codigo": codigo, "nombre": f"Attr {codigo}", "orden": 60})
    assert r.status_code == 200, r.text
    return r.json()


def test_producto_solo_con_su_nucleo(interno):
    atr = _atributo(interno, "PRESQ")
    try:
        r = interno.post("/catalogos/genericos", {"generico": "HARINA-01", "unidad": "UN",
                                                   "combinar": {"PRESQ": ["1 kg", "25 kg"]}})
        assert r.status_code == 200, r.text
        d = r.json()
        # Sin nombre, el código; sin proveedor, marca ni grupo
        assert d["nombre"] == "HARINA-01" and d["proveedor_id"] is None and d["marca_id"] is None and d["grupo_id"] is None
        assert len(d["tallas"]) == 2
        with SessionLocal() as db:
            arts = list(db.scalars(select(Articulo).where(Articulo.generico == "HARINA-01")))
            assert arts and all(a.marca_id is None and a.grupo_id is None and a.proveedor_id is None for a in arts)
        # Se lista y se abre como cualquier otro
        lista = interno.get("/productos", params={"q": "HARINA-01"}).json()
        assert [x["codigo_generico"] for x in lista["items"]] == ["HARINA-01"]
        detalle = interno.get(f"/productos/{d['producto_id']}")
        assert detalle.status_code == 200, detalle.text
        # Se edita sin inventar datos
        r = interno.put("/catalogos/genericos/HARINA-01", {"nombre": "Wheat flour", "unidad": "UN"})
        assert r.status_code == 200, r.text
        assert r.json()["nombre"] == "Wheat flour"
    finally:
        _borrar("HARINA-01")
        interno.delete_(f"/catalogos/atributos_variante/{atr['id']}")


def test_la_empresa_vuelve_obligatoria_la_marca(admin, interno):
    antes = interno.get("/catalogos").json()
    gob = next(c for c in antes if c["tipo"] == "articulos")["gobierno"]
    assert "marca_id" in {x["nombre"] for x in gob["obligables"]}
    r = admin.put("/catalogos/articulos/gobierno", {"responsables": [], "obligatorios": ["marca_id"]})
    assert r.status_code == 200, r.text
    try:
        r = interno.post("/catalogos/genericos", {"generico": "SINMARCA-1", "unidad": "UN"})
        assert r.status_code == 422 and any(e["campo"] == "marca_id" for e in r.json()["detalle"]), r.text
    finally:
        admin.put("/catalogos/articulos/gobierno", {"responsables": gob["responsables"] and [x["id"] for x in gob["responsables"]],
                                                    "obligatorios": gob["obligatorios"]})
        _borrar("SINMARCA-1")


def test_la_familia_exige_el_codigo_de_barras(admin, interno):
    atr = _atributo(interno, "VOLTQ")
    fam = admin.post("/grupos-producto", {"clase": "FAMILIA", "codigo": "ELECTRO-Q", "nombre": "Appliances",
                                          "ejes_variante": ["VOLTQ"],
                                          "campos": [{"campo": "articulo.upc", "obligatorio": True}]})
    assert fam.status_code == 200, fam.text
    fam = fam.json()
    try:
        assert fam["campos"] == [{"campo": "articulo.upc", "obligatorio": True}]
        perfil = interno.get("/productos/perfil", params={"grupo_id": fam["id"]}).json()
        assert perfil["nucleo"] == ["generico", "nombre", "unidad"] and perfil["ejes"] == ["VOLTQ"]
        assert [x["clave"] for x in perfil["familia"]] == ["articulo.upc"] and perfil["familia"][0]["al_guardar"]
        # Sin código de barras: no se crea nada
        r = interno.post("/catalogos/genericos", {"generico": "LICUA-1", "nombre": "Blender", "unidad": "UN",
                                                   "grupo_producto_id": fam["id"], "tallas": [{"valores": {"VOLTQ": "110 V"}}]})
        assert r.status_code == 422 and r.json()["codigo"] == "perfil_incompleto", r.text
        with SessionLocal() as db:
            assert db.scalar(select(Producto.id).where(Producto.codigo_generico == "LICUA-1")) is None
        # Con código de barras: sí
        r = interno.post("/catalogos/genericos", {"generico": "LICUA-1", "nombre": "Blender", "unidad": "UN",
                                                   "grupo_producto_id": fam["id"],
                                                   "tallas": [{"valores": {"VOLTQ": "110 V"}, "upc": "7501234567890"}]})
        assert r.status_code == 200, r.text
        # Una variante nueva sin código de barras tampoco
        r = interno.post("/catalogos/genericos/LICUA-1/tallas", {"tallas": [{"valores": {"VOLTQ": "220 V"}}]})
        assert r.status_code == 422 and "LICUA-1" in r.text
        # Un campo que no existe no entra en la familia
        r = admin.patch(f"/grupos-producto/{fam['id']}", {"campos": [{"campo": "articulo.no_existe", "obligatorio": True}]})
        assert r.status_code == 422
    finally:
        _borrar("LICUA-1")
        admin.delete_(f"/grupos-producto/{fam['id']}")
        interno.delete_(f"/catalogos/atributos_variante/{atr['id']}")


def _libro(productos: list[list], variantes: list[list]) -> bytes:
    wb = Workbook()
    hp = wb.active
    hp.title = "Products"
    hp.append(["Product code", "Name", "Unit", "Supplier"])
    for f in productos:
        hp.append(f)
    hv = wb.create_sheet("Variants")
    hv.append(["Product code", "Attr CARGQ", "Barcode"])
    for f in variantes:
        hv.append(f)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_carga_masiva_de_productos_sin_datos_de_rubro(admin, interno):
    atr = _atributo(interno, "CARGQ")
    fam = admin.post("/grupos-producto", {"clase": "FAMILIA", "codigo": "CARGA-Q", "nombre": "Bulk family",
                                          "ejes_variante": ["CARGQ"],
                                          "campos": [{"campo": "articulo.upc", "obligatorio": True}]}).json()
    try:
        # La plantilla pide el núcleo y deja lo demás opcional
        plantilla = interno.get("/catalogos/articulos/plantilla")
        assert plantilla.status_code == 200
        wb = load_workbook(io.BytesIO(plantilla.content))
        assert {"Products", "Variants"} <= set(wb.sheetnames)
        cabecera = [c.value for c in wb["Products"][1] if c.value]
        assert cabecera[:3] == ["Product code *", "Name", "Unit *"] or cabecera[:3] == ["Product code", "Name", "Unit"], cabecera
        assert not any(str(c).startswith(("Style", "Size")) for c in cabecera)
        archivo = _libro([["CARGA-1", "Cement", "UN", ""]], [["CARGA-1", "25 kg", "7790001112223"], ["CARGA-1", "50 kg", ""]])
        r = interno.post("/catalogos/articulos/importar", files={"archivo": ("productos.xlsx", archivo)})
        assert r.status_code == 200, r.text
        d = r.json()
        # Sin familia nada pide el código de barras: entran las dos variantes
        assert d["genericos_creados"] == 1 and d["creados"] == 2, d
        # Con una familia que lo pide, la variante sin código de barras no entra
        with SessionLocal() as db:
            p = db.scalar(select(Producto).where(Producto.codigo_generico == "CARGA-1"))
            assert p.nombre == "Cement" and p.proveedor_id is None
            p.grupo_producto_id = fam["id"]
            db.commit()
        archivo = _libro([], [["CARGA-1", "75 kg", ""]])
        r = interno.post("/catalogos/articulos/importar", files={"archivo": ("productos.xlsx", archivo)})
        assert r.status_code == 200, r.text
        assert r.json()["creados"] == 0 and any("Barcode" in e["mensaje"] or "UPC" in e["mensaje"] for e in r.json()["errores"]), r.json()
    finally:
        _borrar("CARGA-1")
        admin.delete_(f"/grupos-producto/{fam['id']}")
        interno.delete_(f"/catalogos/atributos_variante/{atr['id']}")


def test_sin_comercio_exterior_nada_espera_una_partida(admin, interno):
    """Con el complemento de comercio exterior apagado (módulo «clasificacion»), ni la factura
    pide país de origen y partida ni un producto sin clasificar detiene una factura."""
    from app.core.empresa import usar_configuracion
    from app.modulos.acceso.permisos import comercio_exterior, exige_datos_aduana
    from app.modulos.productos.productos import ids_bloquean_facturas

    assert interno.get("/auth/me").json()["config"]["requerir_datos_aduana"] is True
    try:
        assert admin.put("/organizacion", {"modulos": {"clasificacion": False}}).status_code == 200
        yo = interno.get("/auth/me").json()
        assert yo["config"]["modulos"]["clasificacion"] is False and yo["config"]["requerir_datos_aduana"] is False
        usar_configuracion({"modulos": {"clasificacion": False}})
        assert not comercio_exterior() and not exige_datos_aduana()
        with SessionLocal() as db:
            assert ids_bloquean_facturas(db) == set()
    finally:
        usar_configuracion(None)
        admin.put("/organizacion", {"modulos": {"clasificacion": True}})
    assert interno.get("/auth/me").json()["config"]["requerir_datos_aduana"] is True
