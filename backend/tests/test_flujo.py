"""Recorre el flujo completo OC -> factura -> PL -> cajas -> unidad de carga -> salida."""
import os
from datetime import date

HOY = date.today().isoformat()


def _oc(api, numero):
    oc = next(o for o in api.get("/ordenes", params={"solo_disponible": False}).json()["items"] if o["numero"] == numero)
    return api.get(f"/ordenes/{oc['id']}/posiciones").json()


def _pos(det, talla, estilo=None):
    return next(p for p in det["posiciones"] if p["talla"] == talla and (estilo is None or p["estilo"] == estilo))


estado = {}


def test_alcance_por_proveedor(tnf, vans):
    ocs_tnf = {o["numero"] for o in tnf.get("/ordenes", params={"solo_disponible": False}).json()["items"]}
    ocs_vans = {o["numero"] for o in vans.get("/ordenes", params={"solo_disponible": False}).json()["items"]}
    assert "4400003845" in ocs_tnf and "4400003901" not in ocs_tnf
    assert ocs_tnf.isdisjoint(ocs_vans)
    oc_tnf = next(o for o in tnf.get("/ordenes").json()["items"] if o["numero"] == "4400003845")
    assert vans.get(f"/ordenes/{oc_tnf['id']}/posiciones").status_code == 404


def test_factura_parcial_y_reglas(tnf):
    det = _oc(tnf, "4400003845")
    s, m, xl = _pos(det, "S"), _pos(det, "M"), _pos(det, "XL")
    zap10 = _pos(det, "10", "NF0A7W4G")
    # Factura con cantidad parcial de S y completa de XL (27) y calzado talla 10
    r = tnf.post("/facturas", {"lineas": [
        {"posicion_id": s["id"], "cantidad": 25},
        {"posicion_id": xl["id"], "cantidad": 27},
        {"posicion_id": zap10["id"], "cantidad": 50},
    ]}, clave="crear-1")
    assert r.status_code == 200, r.text
    fid = r.json()["id"]
    estado["fid"] = fid
    # Idempotencia: la misma clave no crea otra factura
    r2 = tnf.post("/facturas", {"lineas": [{"posicion_id": m["id"], "cantidad": 1}]}, clave="crear-1")
    assert r2.json()["id"] == fid
    assert _pos(_oc(tnf, "4400003845"), "S")["disponible"] == 15

    # Exceso sobre lo disponible
    f = tnf.get(f"/facturas/{fid}").json()
    r = tnf.post(f"/facturas/{fid}/lineas", {"version": f["version"], "lineas": [{"posicion_id": s["id"], "cantidad": 99}]})
    assert r.status_code == 422

    # Regla: el saldo de S no puede ir a otra factura mientras esta siga activa
    r = tnf.post("/facturas", {"lineas": [{"posicion_id": s["id"], "cantidad": 5}]})
    assert r.status_code == 422 and "is already in" in r.json()["detalle"][0]["mensaje"]

    # Pero sí se agrega a la misma factura (suma a la línea existente)
    r = tnf.post(f"/facturas/{fid}/lineas", {"version": f["version"], "lineas": [
        {"posicion_id": s["id"], "cantidad": 15}, {"posicion_id": m["id"], "cantidad": 60}]})
    assert r.status_code == 200, r.text
    assert r.json() == {"agregadas": 1, "aumentadas": 1, "advertencias": []}

    # Compatibilidad: no se mezclan centros (PA10 con PA20)
    pa20 = _oc(tnf, "4400003850")["posiciones"][0]
    f = tnf.get(f"/facturas/{fid}").json()
    r = tnf.post(f"/facturas/{fid}/lineas", {"version": f["version"], "lineas": [{"posicion_id": pa20["id"], "cantidad": 1}]})
    assert r.status_code == 422 and any("plants" in d["mensaje"] for d in r.json()["detalle"])


def test_version_conflicto(tnf):
    fid = estado["fid"]
    r = tnf.patch(f"/facturas/{fid}", {"version": 1, "observaciones": "x"})
    assert r.status_code == 409 and r.json()["codigo"] == "conflicto_version"


def test_precio_requiere_motivo(tnf):
    fid = estado["fid"]
    f = tnf.get(f"/facturas/{fid}").json()
    linea = f["lineas"][0]
    r = tnf.patch(f"/facturas/{fid}/lineas", {"version": f["version"], "cambios": [
        {"linea_id": linea["id"], "precio_unitario": linea["precio_oc"] + 1}]})
    assert r.status_code == 422 and r.json()["detalle"][0]["codigo"] == "motivo_precio"
    r = tnf.patch(f"/facturas/{fid}/lineas", {"version": f["version"], "cambios": [
        {"linea_id": linea["id"], "precio_unitario": linea["precio_oc"] + 1, "motivo_precio": "Ajuste acordado"}]})
    assert r.status_code == 200, r.text


def _empacar(api, pl_id, version, filas, sobrante="caja_parcial"):
    """filas: [(pl_linea_id, plantilla_id)]"""
    return api.post(f"/packing-lists/{pl_id}/empaque/aplicar", {
        "version": version, "sobrante": sobrante,
        "filas": [{"pl_linea_id": l, "plantilla_id": t} for l, t in filas]})


def test_packing_list_pendientes_mover(tnf):
    fid = estado["fid"]
    r = tnf.post(f"/facturas/{fid}/packing-lists", {})
    assert r.status_code == 200, r.text
    pl_id = r.json()["id"]
    estado["pl"] = pl_id
    # Sin saldo: segundo PL con pendientes explica dónde está todo
    r = tnf.post(f"/facturas/{fid}/packing-lists", {})
    assert r.status_code == 409 and r.json()["codigo"] == "sin_saldo"

    pl = tnf.get(f"/packing-lists/{pl_id}").json()
    assert pl["plantillas"] and all(t["proveedor_id"] if "proveedor_id" in t else True for t in pl["plantillas"])
    zap = next(l for l in pl["lineas"] if l["unidad"] == "PAR")
    # Mover solo 14 de los 50 pares a un PL nuevo (sin dividir antes)
    r = tnf.post(f"/packing-lists/{pl_id}/mover", {"version": pl["version"],
                 "movimientos": [{"pl_linea_id": zap["id"], "cantidad": 14}], "destino_pl_id": None})
    assert r.status_code == 200, r.text
    estado["pl2"] = r.json()["destino_id"]
    assert r.json()["destino_numero"] == "PL-002"
    pl = tnf.get(f"/packing-lists/{pl_id}").json()
    assert next(l for l in pl["lineas"] if l["unidad"] == "PAR")["cantidad"] == 36


def test_plantilla_con_sobrante(tnf):
    pl_id = estado["pl"]
    pl = tnf.get(f"/packing-lists/{pl_id}").json()
    plantillas = tnf.get("/plantillas").json()
    chaqueta = next(t for t in plantillas if t["nombre"] == "Jacket carton 10 units")
    xl = next(l for l in pl["lineas"] if l["talla"] == "XL")
    assert xl["cantidad"] == 27
    previa = tnf.post(f"/packing-lists/{pl_id}/empaque/previa",
                      {"filas": [{"pl_linea_id": xl["id"], "plantilla_id": chaqueta["id"]}]}).json()
    assert previa["resumen"]["cajas_completas"] == 2 and previa["resumen"]["sobrante_total"] == 7
    # Una plantilla de pares no aplica a una fila en unidades: se omite, no falla
    calzado = next(t for t in plantillas if t["nombre"] == "Footwear carton 12 pairs")
    previa = tnf.post(f"/packing-lists/{pl_id}/empaque/previa",
                      {"filas": [{"pl_linea_id": xl["id"], "plantilla_id": calzado["id"]}]}).json()
    assert previa["resumen"]["filas"] == 0 and previa["filas"][0]["omitida"]
    r = _empacar(tnf, pl_id, pl["version"], [(xl["id"], chaqueta["id"])])
    assert r.status_code == 200, r.text
    pl = tnf.get(f"/packing-lists/{pl_id}").json()
    grupos = [g for g in pl["grupos"] if g["items"][0]["pl_linea_id"] == xl["id"]]
    assert [(g["num_cajas"], g["items"][0]["cantidad_por_caja"], g["es_parcial"]) for g in grupos] == [(2, 10, False), (1, 7, True)]
    parcial = grupos[1]
    # El peso sale del peso unitario del artículo (0.9 kg) y la tara de la caja (1.2 kg): no es estimado
    assert not parcial["peso_estimado"] and parcial["peso_neto_caja"] == 6.3 and parcial["peso_bruto_caja"] == 7.5
    assert next(l for l in pl["lineas"] if l["id"] == xl["id"])["sin_caja"] == 0


def test_reducir_factura_con_pl(tnf):
    fid = estado["fid"]
    f = tnf.get(f"/facturas/{fid}").json()
    m = next(l for l in f["lineas"] if l["talla"] == "M")
    r = tnf.patch(f"/facturas/{fid}/lineas", {"version": f["version"], "cambios": [{"linea_id": m["id"], "cantidad": 50}]})
    assert r.status_code == 409 and r.json()["codigo"] == "requiere_ajuste_pl"
    assert r.json()["detalle"][0]["hay_que_liberar"] == 10
    r = tnf.patch(f"/facturas/{fid}/lineas", {"version": f["version"], "ajuste_pl": "automatico",
                  "cambios": [{"linea_id": m["id"], "cantidad": 50}]})
    assert r.status_code == 200, r.text
    # Lo empacado no se libera automáticamente
    xl = next(l for l in f["lineas"] if l["talla"] == "XL")
    f = tnf.get(f"/facturas/{fid}").json()
    r = tnf.patch(f"/facturas/{fid}/lineas", {"version": f["version"], "ajuste_pl": "automatico",
                  "cambios": [{"linea_id": xl["id"], "cantidad": 20}]})
    assert r.status_code == 409 and r.json()["codigo"] == "ajuste_insuficiente"


def test_empacar_todo_y_finalizar(tnf):
    pl_id = estado["pl"]
    plantillas = {t["nombre"]: t for t in tnf.get("/plantillas").json()}
    pl = tnf.get(f"/packing-lists/{pl_id}").json()
    # Todo en un solo paso: cada fila con la plantilla de su unidad
    por_unidad = {"UN": plantillas["Jacket carton 10 units"]["id"], "PAR": plantillas["Footwear carton 12 pairs"]["id"]}
    filas = [(l["id"], por_unidad[l["unidad"]]) for l in pl["lineas"] if l["sin_caja"]]
    r = _empacar(tnf, pl_id, pl["version"], filas)
    assert r.status_code == 200, r.text
    v = r.json()["version"]
    # Los pesos salen solos de la estructura: neto = artículos, bruto = neto + tara de cada nivel
    pl = tnf.get(f"/packing-lists/{pl_id}").json()
    for g in pl["grupos"]:
        neto = sum(i["cantidad_por_caja"] * i["peso_unitario"] for i in g["contenido"])
        assert g["peso_calculado"] and abs(g["peso_neto_caja"] - neto) < 1e-6
        assert abs(g["peso_bruto_caja"] - neto - g["tara"]) < 1e-6
    t = pl["totales"]
    assert abs(t["peso_bruto"] - sum(g["peso_bruto_total"] for g in pl["grupos"] if not g["padre_id"])) < 1e-3
    r = tnf.post(f"/packing-lists/{pl_id}/finalizar", {"version": v})
    assert r.status_code == 200, r.text

    # PL-002: 14 pares -> 1 caja de 12 + caja parcial de 2. La plantilla se
    # sugiere sola porque el mismo estilo ya se empacó con ella en PL-001.
    pl2 = tnf.get(f"/packing-lists/{estado['pl2']}").json()
    fila = pl2["lineas"][0]
    assert fila["plantilla_sugerida_id"] == plantillas["Footwear carton 12 pairs"]["id"]
    r = _empacar(tnf, pl2["id"], pl2["version"], [(fila["id"], fila["plantilla_sugerida_id"])])
    assert r.status_code == 200 and r.json()["resumen"]["cajas_completas"] == 1, r.text
    pl2 = tnf.get(f"/packing-lists/{pl2['id']}").json()
    assert [g["num_cajas"] for g in pl2["grupos"]] == [1, 1] and pl2["grupos"][1]["es_parcial"]
    r = tnf.post(f"/packing-lists/{pl2['id']}/finalizar", {"version": pl2["version"]})
    assert r.status_code == 200, r.text


def test_finalizar_factura(tnf):
    fid = estado["fid"]
    f = tnf.get(f"/facturas/{fid}").json()
    r = tnf.post(f"/facturas/{fid}/finalizar", {"version": f["version"]})
    assert r.status_code == 422 and any("number" in e["mensaje"] for e in r.json()["detalle"])
    r = tnf.patch(f"/facturas/{fid}", {"version": f["version"], "numero": "INV-2026-001", "fecha": "2026-09-20"})
    assert r.status_code == 200
    r = tnf.post(f"/facturas/{fid}/finalizar", {"version": r.json()["version"]})
    assert r.status_code == 200, r.text
    f = tnf.get(f"/facturas/{fid}").json()
    assert f["estado"] == "FINALIZADA" and f["lista_transporte"]


def test_transporte_y_salida(interno, tnf):
    e = interno.get("/embarques").json()[0]
    det = interno.get(f"/embarques/{e['id']}").json()
    unidad = det["unidades"][0]
    disp = interno.get(f"/unidades/{unidad['id']}/disponibles").json()
    grupo = next(g for g in disp if g["factura_id"] == estado["fid"])
    pl_ids = [p["id"] for p in grupo["packing_lists"]]
    r = interno.post(f"/unidades/{unidad['id']}/asignar", {"pl_ids": pl_ids})
    assert r.status_code == 200, r.text
    assert r.json()["confirmados"] == len(pl_ids) and r.json()["tentativos"] == 0
    # Sin BL, contenedor ni sello no hay salida: son datos obligatorios del transporte
    r = interno.post(f"/embarques/{e['id']}/eventos", {"tipo": "SALIDA", "fecha": f"{HOY}T08:00:00"})
    assert r.status_code == 422 and r.json()["codigo"] == "datos_transporte"
    assert len(r.json()["detalle"]) == 3
    assert interno.patch(f"/embarques/{e['id']}", {"documento_numero": "MAEU 123"}).status_code == 200
    assert interno.patch(f"/unidades/{unidad['id']}", {"numero": "MSKU 1234567", "sello": "S-1"}).status_code == 200
    r = interno.post(f"/embarques/{e['id']}/eventos", {"tipo": "SALIDA", "fecha": f"{HOY}T08:00:00"})
    assert r.status_code == 200 and r.json()["estado"] == "EN_TRANSITO"
    # El proveedor ve el seguimiento desde su factura
    f = tnf.get(f"/facturas/{estado['fid']}").json()
    assert f["packing_lists"][0]["transporte"]["estado"] == "EN_TRANSITO"
    # Y no puede usar el módulo de transporte
    assert tnf.get("/embarques").status_code == 403
    # Después de la salida la carga queda cerrada: no se quita, no se agrega, no se reabre
    r = interno.post(f"/unidades/{unidad['id']}/desasignar", {"pl_ids": pl_ids[:1], "motivo": "x"})
    assert r.status_code == 409 and r.json()["codigo"] == "embarque_cerrado"
    r = interno.post(f"/unidades/{unidad['id']}/asignar", {"pl_ids": pl_ids[:1], "modo": "AUTO", "motivo": "x"})
    assert r.status_code == 409
    r = interno.post(f"/packing-lists/{pl_ids[0]}/reabrir", {"motivo": "x"})
    assert r.status_code == 409
    assert interno.post(f"/embarques/{e['id']}/unidades", {"tipo": "20GP"}).status_code == 409
    # Los eventos siguen el orden: no hay entrega antes del arribo
    r = interno.post(f"/embarques/{e['id']}/eventos", {"tipo": "ENTREGA", "fecha": f"{HOY}T09:00:00"})
    assert r.status_code == 409 and r.json()["codigo"] == "orden_eventos"
    # Todo lo que zarpó quedó marcado como recolectado
    pl = interno.get(f"/packing-lists/{pl_ids[0]}").json()
    assert pl["recolectado_en"] == HOY


def test_eliminar_linea_con_cascada(vans):
    det = _oc(vans, "4400003901")
    p7, p8 = _pos(det, "7"), _pos(det, "8")
    fid = vans.post("/facturas", {"lineas": [{"posicion_id": p7["id"], "cantidad": 36},
                                            {"posicion_id": p8["id"], "cantidad": 48}]}).json()["id"]
    pl_id = vans.post(f"/facturas/{fid}/packing-lists", {}).json()["id"]
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    t = next(t for t in vans.get("/plantillas").json() if t["nombre"] == "Master 12 pairs")
    _empacar(vans, pl_id, pl["version"], [(l["id"], t["id"]) for l in pl["lineas"]], "sin_caja")
    f = vans.get(f"/facturas/{fid}").json()
    linea7 = next(l for l in f["lineas"] if l["talla"] == "7")
    r = vans.post(f"/facturas/{fid}/lineas/eliminar", {"version": f["version"], "linea_ids": [linea7["id"]]})
    assert r.status_code == 409 and r.json()["codigo"] == "requiere_confirmacion"
    r = vans.post(f"/facturas/{fid}/lineas/eliminar", {"version": f["version"], "linea_ids": [linea7["id"]],
                  "confirmar_cascada": True})
    assert r.status_code == 200, r.text
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    assert len(pl["lineas"]) == 1 and pl["totales"]["cajas"] == 4
    assert _pos(_oc(vans, "4400003901"), "7")["disponible"] == 36


def test_mover_cajas(vans):
    fs = vans.get("/facturas").json()["items"]
    fid = fs[0]["id"]
    f = vans.get(f"/facturas/{fid}").json()
    pl_id = f["packing_lists"][0]["id"]
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    g = pl["grupos"][0]
    r = vans.post(f"/packing-lists/{pl_id}/mover-cajas", {"version": pl["version"],
                  "grupos": [{"grupo_id": g["id"], "num_cajas": 1}], "destino_pl_id": None})
    assert r.status_code == 200, r.text
    destino = vans.get(f"/packing-lists/{r.json()['destino_id']}").json()
    origen = vans.get(f"/packing-lists/{pl_id}").json()
    assert destino["totales"]["cajas"] == 1 and destino["lineas"][0]["cantidad"] == 12
    assert origen["totales"]["cajas"] == 3 and origen["lineas"][0]["cantidad"] == 36
    assert origen["lineas"][0]["sin_caja"] == 0


def test_importacion_oc(interno):
    oc = _oc(interno, "4400003901")
    pos10 = oc["posiciones"][0]
    cab = oc["oc"]
    enc = ("proveedor,oc,posicion,sociedad,centro,almacen,centro_destino,moneda,incoterm,sku,cantidad,precio,"
           "puerto,pais_origen,fecha_xf_original,fecha_tienda,liberacion_comercial\n")
    fila = lambda oc_n, pos, sku, cant, lib="C", alm="BF20": (  # noqa: E731
        f"VANS,{oc_n},{pos},8000,8020,{alm},2220,USD,FOB,{sku},{cant},25.5,VNSGN,VN,{cab['fecha_xf_original']},"
        f"{cab['fecha_tienda']},{lib}\n")
    csv = (enc + fila("4400009999", "10", pos10["codigo_sap"], 24, "P")
           + fila("4400003901", "10", pos10["codigo_sap"], 10)
           + "XXX,1,1,8000,PA10,,2220,USD,FOB,1,1,1,,,,,\n")
    r = interno.c.post("/api/ordenes/importar/previa", headers=interno.h,
                       files={"archivo": ("ocs.csv", csv.encode(), "text/csv")})
    assert r.status_code == 200, r.text
    res = r.json()["resumen"]
    assert res["nuevo"] == 1 and res["error"] == 1
    r = interno.post(f"/ordenes/importar/{r.json()['importacion_id']}/aplicar")
    assert r.status_code == 200, r.text
    ocs = interno.get("/ordenes", params={"q": "4400009999", "solo_disponible": False}).json()["items"]
    assert ocs and ocs[0]["por_unidad"]["PAR"]["cantidad"] == 24
    # Liberación comercial pendiente: logística 304 y no se puede facturar
    assert ocs[0]["liberacion_comercial"] == "P" and ocs[0]["liberacion_logistica"] == "304"
    det = interno.get(f"/ordenes/{ocs[0]['id']}/posiciones").json()
    assert det["posiciones"][0]["codigo_sap"] == pos10["codigo_sap"] and det["posiciones"][0]["estado"] == "NO_DISPONIBLE"
    # La OC ya liberada que cambia después queda como 301
    assert _oc(interno, "4400003901")["oc"]["liberacion_logistica"] == "301"
    # Un SKU que no está en el maestro no entra
    r = interno.c.post("/api/ordenes/importar/previa", headers=interno.h, files={"archivo": ("x.csv", (
        enc + fila("4400009998", "10", "NOEXISTE", 1)).encode(), "text/csv")})
    assert r.json()["resumen"]["error"] == 1 and "item master" in r.json()["filas"][0]["mensajes"][0]
    # El SKU debe ser del proveedor de la OC
    r = interno.c.post("/api/ordenes/importar/previa", headers=interno.h, files={"archivo": ("p.csv", (
        enc + fila("4400009995", "10", "30095120001", 5)).encode(), "text/csv")})
    assert r.json()["resumen"]["error"] == 1 and "another supplier" in " ".join(r.json()["filas"][0]["mensajes"])
    # Dos liberaciones: sin comercial (P) no puede haber logística 300/301; nueva sin código logístico = 304
    enc_log = enc.strip() + ",liberacion_logistica\n"
    fila_log = lambda oc_n, lib, log: fila(oc_n, "10", pos10["codigo_sap"], 12, lib).strip() + f",{log}\n"  # noqa: E731
    r = interno.c.post("/api/ordenes/importar/previa", headers=interno.h, files={"archivo": ("l.csv", (
        enc_log + fila_log("4400009990", "P", "300") + fila_log("4400009991", "", "") + fila_log("4400009992", "C", "300")
    ).encode(), "text/csv")})
    res = r.json()
    assert res["resumen"]["error"] == 1 and "without commercial release" in res["filas"][0]["mensajes"][0]
    interno.post(f"/ordenes/importar/{res['importacion_id']}/aplicar")
    ocs = {o["numero"]: o for o in interno.get("/ordenes", params={"q": "44000099", "solo_disponible": False}).json()["items"]}
    assert ocs["4400009991"]["liberacion_comercial"] == "C" and ocs["4400009991"]["liberacion_logistica"] == "304"
    assert not ocs["4400009991"]["liberada"] and ocs["4400009992"]["liberada"]
    # Misma sociedad y centro, cada posición en su almacén; un almacén de otra sociedad no entra
    r = interno.c.post("/api/ordenes/importar/previa", headers=interno.h, files={"archivo": ("a.csv", (
        enc + fila("4400009997", "10", pos10["codigo_sap"], 12, alm="BF19")
        + fila("4400009997", "20", pos10["codigo_sap"], 12, alm="BF20")
        + fila("4400009996", "10", pos10["codigo_sap"], 12, alm="BF01")).encode(), "text/csv")})
    res = r.json()
    assert res["resumen"]["nuevo"] == 2 and res["resumen"]["error"] == 1, res
    assert "does not belong to company" in res["filas"][-1]["mensajes"][0]
    interno.post(f"/ordenes/importar/{res['importacion_id']}/aplicar")
    oc = interno.get("/ordenes", params={"q": "4400009997", "solo_disponible": False}).json()["items"][0]
    assert oc["almacenes"] == ["BF19", "BF20"]
    det = interno.get(f"/ordenes/{oc['id']}/posiciones").json()["posiciones"]
    assert [p["almacen"] for p in det] == ["BF19", "BF20"]
    assert interno.get("/ordenes", params={"almacen": "BF19", "q": "4400009997", "solo_disponible": False}).json()["total"] == 1


def test_exportar(tnf):
    r = tnf.get(f"/packing-lists/{estado['pl']}/exportar")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    r = tnf.get(f"/facturas/{estado['fid']}/exportar")
    assert r.status_code == 200


def test_asignacion_automatica(interno, vans):
    """AUTO confirma lo que está listo y deja tentativo lo demás, en un paso."""
    # El puerto de destino debe ser el del centro de llegada
    r = interno.post("/embarques", {"tipo_transporte": "MARITIMO", "modalidad": "FCL", "centro": "8010",
                                    "puerto_destino": "PAONX"})
    assert r.status_code == 422 and "choose one of those ports" in r.json()["detalle"][0]["mensaje"]
    # Puertos sugeridos del centro: se puede cambiar a otro de sus puertos del mismo modo
    r = interno.post("/embarques", {"tipo_transporte": "MARITIMO", "centro": "8010", "puerto_destino": "SVLUN"})
    assert r.status_code == 200, r.text
    r = interno.post("/embarques", {"tipo_transporte": "MARITIMO", "puerto_origen": "HKG"})
    assert r.status_code == 422 and "air" in r.json()["detalle"][0]["mensaje"]
    # Transportista del modo y de la sociedad del centro
    trans = {t["codigo"]: t["id"] for t in interno.get("/catalogos/transportistas").json()["items"]}
    r = interno.post("/embarques", {"tipo_transporte": "MARITIMO", "centro": "8010", "transportista_id": trans["AVCG"]})
    assert r.status_code == 422 and "air" in r.json()["detalle"][0]["mensaje"]
    r = interno.post("/embarques", {"tipo_transporte": "MARITIMO", "centro": "8010", "transportista_id": trans["CMDU"]})
    assert r.status_code == 422 and "does not work with company 8000" in r.json()["detalle"][0]["mensaje"]
    aereo = interno.post("/embarques", {"tipo_transporte": "AEREO", "centro": "8010", "transportista_id": trans["AVCG"]}).json()
    det = interno.get(f"/embarques/{aereo['id']}").json()
    assert det["puerto_destino"] == "SAL" and det["transportista"] == "Avianca Cargo"
    assert {t["codigo"] for t in det["tipos_unidad"]} == {"AWB"}
    # Solo unidades del modo; la modalidad es de cada unidad (un marítimo puede ser mixto)
    assert interno.post(f"/embarques/{aereo['id']}/unidades", {"tipo": "40HC"}).status_code == 422
    mar = interno.post("/embarques", {"tipo_transporte": "MARITIMO"}).json()
    interno.post(f"/embarques/{mar['id']}/unidades", {"tipo": "40HC"})
    assert interno.get(f"/embarques/{mar['id']}").json()["modalidad"] == "FCL"
    interno.post(f"/embarques/{mar['id']}/unidades", {"tipo": "LCL"})
    assert interno.get(f"/embarques/{mar['id']}").json()["modalidad"] == "MIXTO"
    e = interno.post("/embarques", {"tipo_transporte": "MARITIMO", "modalidad": "FCL", "centro": "8010"}).json()
    assert interno.get(f"/embarques/{e['id']}").json()["puerto_destino"] == "SVAQJ"
    u8010 = interno.post(f"/embarques/{e['id']}/unidades", {"tipo": "20GP"}).json()["id"]
    assert all(g["centro"] == "8010" for g in interno.get(f"/unidades/{u8010}/disponibles").json())
    # Sin centro, lo define la primera carga y no se mezclan centros
    e = interno.post("/embarques", {"tipo_transporte": "MARITIMO", "modalidad": "FCL"}).json()
    nueva = interno.post(f"/embarques/{e['id']}/unidades", {"tipo": "20GP"}).json()["id"]
    disp = interno.get(f"/unidades/{nueva}/disponibles").json()
    centros = {g["centro"] for g in disp}
    if len(centros) > 1:
        r = interno.post(f"/unidades/{nueva}/asignar", {"pl_ids": [p["id"] for g in disp for p in g["packing_lists"]]})
        assert r.status_code == 422 and r.json()["codigo"] == "centro_distinto"
    # Solo se ofrece y se acepta lo finalizado (factura y PL)
    assert all(p["puede_confirmar"] for g in disp for p in g["packing_lists"])
    centro = sorted(centros)[0]
    pls = [p for g in disp if g["centro"] == centro for p in g["packing_lists"]]
    r = interno.post(f"/unidades/{nueva}/asignar", {"pl_ids": [p["id"] for p in pls]})
    assert r.status_code == 200, r.text
    assert r.json() == {"asignados": len(pls), "confirmados": len(pls), "tentativos": 0}
    borrador = next((f for f in interno.get("/seguimiento/documentos", params={"size": 200}).json()["items"]
                     if f["pl_id"] and f["estado_pl"] in ("BORRADOR", "EN_CORRECCION")), None)
    assert borrador
    r = interno.post(f"/unidades/{nueva}/asignar", {"pl_ids": [borrador["pl_id"]]})
    assert r.status_code == 422
    # Al reabrir un PL cargado, sale de la unidad
    r = interno.post(f"/packing-lists/{pls[0]['id']}/reabrir", {"motivo": "Corrección"})
    assert r.status_code == 200 and "Removed from" in r.json()["nota"]
    assert pls[0]["id"] not in {p["id"] for p in interno.get(f"/unidades/{nueva}").json()["asignados"]}
    if pls[1:]:
        r = interno.post(f"/unidades/{nueva}/desasignar", {"pl_ids": [p["id"] for p in pls[1:]], "motivo": "Prueba"})
        assert r.status_code == 200, r.text


def test_dashboard(interno, tnf):
    d = interno.get("/dashboard").json()
    claves = {k["clave"] for k in d["kpis"]}
    assert {"por_facturar", "listas", "sin_contenedor", "en_camino"} <= claves
    assert d["proveedores"] and {p["nombre"] for p in d["proveedores"]} >= {"The North Face", "Vans"}
    assert len(d["facturado_mes"]) == 6 and d["contenedores"] is not None
    # El proveedor solo ve lo suyo y no recibe datos internos
    p = tnf.get("/dashboard").json()
    assert "listas" not in {k["clave"] for k in p["kpis"]}
    assert p["proveedores"] == [] and p["contenedores"] == [] and p["alertas"] == []
    assert any(e["estado"] == "EN_TRANSITO" for e in p["envios"])
    flujo = p["flujo"]
    assert set(flujo) <= {"PAR", "UN"} and flujo["UN"]["embarcado"] > 0


def test_historial(tnf):
    h = tnf.get(f"/facturas/{estado['fid']}/historial").json()
    acciones = {x["accion"] for x in h}
    assert {"crear", "finalizar", "aplicar_plantilla", "asignar_unidad"} <= acciones


def test_reglas_de_empaque(vans):
    """Prepack: una curva por caja. Casepack: cantidad exacta, sin mezclar.
    Nunca se mezclan países de destino en una caja."""
    prepack = _oc(vans, "4400003903")["posiciones"][0]
    solido = _pos(_oc(vans, "4400003901"), "11")
    assert prepack["tipo_empaque"] == "PREPACK" and prepack["unidades_por_caja"] == 12
    fid = vans.post("/facturas", {"lineas": [{"posicion_id": prepack["id"], "cantidad": 10},
                                            {"posicion_id": solido["id"], "cantidad": 24}]}).json()["id"]
    pl_id = vans.post(f"/facturas/{fid}/packing-lists", {}).json()["id"]
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    lp = next(l for l in pl["lineas"] if l["regla"] == "PREPACK")
    ls = next(l for l in pl["lineas"] if l["regla"] == "CASEPACK")
    # Mezclar el prepack con otra fila (y otro destino) en una caja no se permite
    r = vans.post(f"/packing-lists/{pl_id}/cajas", {"version": pl["version"], "num_cajas": 1, "items": [
        {"pl_linea_id": lp["id"], "cantidad_por_caja": 1}, {"pl_linea_id": ls["id"], "cantidad_por_caja": 12}]})
    assert r.status_code == 422 and r.json()["codigo"] == "regla_empaque"
    assert any("destination" in d["mensaje"] for d in r.json()["detalle"])
    # El casepack no se reduce: 10 por caja cuando el casepack es 12
    r = vans.post(f"/packing-lists/{pl_id}/cajas", {"version": pl["version"], "num_cajas": 2, "items": [
        {"pl_linea_id": ls["id"], "cantidad_por_caja": 10}]})
    assert r.status_code == 422
    # Empaque automático sin plantilla: el artículo define la cantidad por caja
    r = _empacar(vans, pl_id, pl["version"], [(lp["id"], None), (ls["id"], None)])
    assert r.status_code == 200, r.text
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    cajas = {g["items"][0]["pl_linea_id"]: (g["num_cajas"], g["items"][0]["cantidad_por_caja"]) for g in pl["grupos"]}
    assert cajas[lp["id"]] == (10, 1) and cajas[ls["id"]] == (2, 12)
    assert all(g["etiqueta"]["tipo"] == "ESTANDAR" for g in pl["grupos"])
    # La OC ya dice a quién se factura (sociedad) y a quién se notifica (centro)
    assert pl["partes"]["facturar_a"]["codigo"] == "8000" and pl["partes"]["facturar_a"]["correos"]
    assert pl["partes"]["notify"]["codigo"] == "8020" and pl["partes"]["notify"]["contactos"]
    # Paletizar: un pallet nuevo necesita medidas; luego se agregan cajas a él
    ids = [g["id"] for g in pl["grupos"]]
    r = vans.post(f"/packing-lists/{pl_id}/pallets", {"version": pl["version"], "grupo_ids": ids[:1]})
    assert r.status_code == 422
    r = vans.post(f"/packing-lists/{pl_id}/pallets", {"version": pl["version"], "grupo_ids": ids[:1],
                                                      "largo": 120, "ancho": 100, "alto": 150, "peso_tara": 20})
    assert r.status_code == 200, r.text
    pallet = r.json()["pallet_id"]
    r = vans.post(f"/packing-lists/{pl_id}/pallets", {"version": r.json()["version"], "grupo_ids": ids[1:],
                                                      "pallet_id": pallet})
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    assert len(pl["pallets"]) == 1 and pl["pallets"][0]["cajas"] == 12 and pl["totales"]["pallets"] == 1
    assert pl["totales"]["cbm"] == 1.8  # el volumen es el del pallet
    # Bruto del pallet = bruto de sus cajas (artículos + tara de cada caja) + tara del pallet
    cajas_bruto = sum(g["peso_bruto_total"] for g in pl["grupos"] if g["padre_id"] == pallet)
    assert abs(pl["pallets"][0]["peso_bruto"] - (cajas_bruto + 20)) < 1e-3
    assert abs(pl["totales"]["peso_bruto"] - pl["pallets"][0]["peso_bruto"]) < 1e-3
    # Sacar todas las cajas elimina el pallet vacío
    r = vans.post(f"/packing-lists/{pl_id}/pallets/quitar", {"version": pl["version"], "pallet_id": pallet})
    assert r.status_code == 200 and vans.get(f"/packing-lists/{pl_id}").json()["pallets"] == []


def test_catalogos(interno, tnf):
    assert tnf.get("/catalogos").status_code == 403
    r = interno.post("/catalogos/marcas", {"codigo": "col", "nombre": "Columbia"})
    assert r.status_code == 200 and r.json()["codigo"] == "COL" and r.json()["activa"] is True
    assert interno.post("/catalogos/marcas", {"codigo": "COL", "nombre": "Otra"}).status_code == 409
    marca = r.json()["id"]
    assert interno.patch(f"/catalogos/marcas/{marca}", {"nombre": "Columbia Sportswear"}).json()["nombre"] == "Columbia Sportswear"
    lista = interno.get("/catalogos/marcas", params={"q": "colum"}).json()
    assert lista["total"] == 1
    assert interno.delete_(f"/catalogos/marcas/{marca}").status_code == 200
    # Una marca en uso no se elimina: se desactiva
    tnf_marca = next(m for m in interno.get("/catalogos/marcas", params={"q": "TNF"}).json()["items"])
    r = interno.delete_(f"/catalogos/marcas/{tnf_marca['id']}")
    assert r.status_code == 409 and r.json()["codigo"] == "en_uso"
    # El número de artículo es numérico (p. ej. 30095120001)
    grupos = interno.get("/catalogos/grupos").json()["items"]
    provs = {p["codigo"]: p for p in interno.get("/catalogos/proveedores").json()["items"]}
    base = {"estilo": "E1", "color": "Rojo", "marca_id": tnf_marca["id"], "grupo_id": grupos[0]["id"],
            "proveedor_id": provs["TNF"]["id"]}
    # El proveedor maneja sus marcas: un artículo de TNF no puede ser de otra marca
    vans_marca = next(m for m in interno.get("/catalogos/marcas").json()["items"] if m["codigo"] == "VANS")
    r = interno.post("/catalogos/articulos", {**base, "marca_id": vans_marca["id"], "sku": "30099990009", "talla": "9",
                                              "tipo": "SOLIDO", "unidad": "PAR"})
    assert r.status_code == 422 and any(d["campo"] == "marca_id" for d in r.json()["detalle"])
    # Y no puede dejar de manejar una marca de la que tiene artículos
    r = interno.patch(f"/catalogos/proveedores/{provs['TNF']['id']}", {"marcas": [vans_marca["id"]]})
    assert r.status_code == 422 and "has items" in r.json()["detalle"][0]["mensaje"]
    r = interno.patch(f"/catalogos/proveedores/{provs['TNF']['id']}", {"marcas": []})
    assert r.status_code == 422 and "has items" in r.json()["detalle"][0]["mensaje"]
    assert provs["TNF"]["marcas_txt"] == "TNF" and "8000" in provs["TNF"]["sociedades_txt"]
    r = interno.post("/catalogos/articulos", {**base, "sku": "X 1!", "talla": "9", "tipo": "SOLIDO", "unidad": "PAR"})
    assert r.status_code == 422 and r.json()["detalle"][0]["campo"] == "sku"
    # Un prepack no se crea como artículo suelto: se crea con su código y su explosión
    r = interno.post("/catalogos/articulos", {**base, "sku": "30099990001", "talla": "AB12", "tipo": "PREPACK",
                                              "unidad": "CJ"})
    assert r.status_code == 422 and any(d["campo"] == "tipo" for d in r.json()["detalle"])
    sol = interno.post("/catalogos/articulos", {**base, "sku": "30099990002", "talla": "9", "tipo": "SOLIDO",
                                                "unidad": "PAR", "casepack": 6}).json()
    otro = interno.get("/catalogos/articulos", params={"q": "VN000EE3"}).json()["items"][0]
    datos_pp = {"sku": "30099990001", "codigo": "ab12", "estilo": "E1", "color": "Rojo"}
    r = interno.post("/catalogos/prepacks", {**datos_pp, "componentes": [
        {"articulo_id": sol["id"], "cantidad": 6}, {"articulo_id": otro["id"], "cantidad": 1}]})
    assert r.status_code == 422 and any("same style and color" in d["mensaje"] for d in r.json()["detalle"])
    r = interno.post("/catalogos/prepacks", {**datos_pp, "componentes": [{"articulo_id": sol["id"], "cantidad": 6}]})
    assert r.status_code == 200, r.text
    pp = r.json()
    assert pp["codigo"] == "AB12" and pp["sku"] == "30099990001" and pp["total"] == 6
    art = interno.get("/catalogos/articulos", params={"q": "30099990001"}).json()["items"][0]
    assert art["tipo"] == "PREPACK" and art["talla"] == "AB12" and art["unidad"] == "CJ"
    # La explosión se ve (también el proveedor) pero no se modifica
    assert tnf.get("/catalogos/explosion/30099990001").json()["componentes"][0]["cantidad"] == 6
    assert interno.patch(f"/catalogos/prepacks/{pp['id']}", {"codigo": "ZZ99"}).status_code == 422
    assert interno.patch(f"/catalogos/articulos/{art['id']}", {"talla": "ZZ99"}).status_code == 422
    assert interno.patch(f"/catalogos/articulos/{sol['id']}", {"color": "Azul"}).status_code == 422
    assert interno.patch(f"/catalogos/prepacks/{pp['id']}", {"descripcion": "Curva roja"}).status_code == 200
    # Contactos: correos válidos y ligados a una sociedad o centro
    r = interno.post("/catalogos/contactos", {"nombre": "X", "rol": "NOTIFY", "correos": "a@b.com, malo"})
    assert r.status_code == 422
    # Filtros de un catálogo por referencia; la sociedad muestra sus centros
    centros = interno.get("/catalogos/centros", params={"orden": "codigo:desc"}).json()["items"]
    assert [c["codigo"] for c in centros][:4] == ["PA20", "PA10", "8020", "8010"]
    assert centros[0]["sociedad_id_txt"].startswith("PA01") and centros[0]["puerto"] == "PABLB"
    soc = interno.get("/catalogos/sociedades", params={"q": "8000"}).json()["items"][0]
    assert soc["centros_txt"] == "2220, 8010, 8020" and soc["contactos"] == 1


def _rol_proveedor(admin):
    return next(r for r in admin.get("/roles").json()["roles"] if r["nombre"] == "Supplier")


def test_seguimiento(tnf, interno, admin):
    # A los proveedores no les interesa: su rol no lo trae, pero se les puede dar
    assert tnf.get("/seguimiento").status_code == 403
    rol = _rol_proveedor(admin)
    assert admin.patch(f"/roles/{rol['id']}", {"permisos": rol["permisos"] + ["seguimiento.ver"]}).status_code == 200
    s = tnf.get("/seguimiento").json()
    assert s["total"] and all(f["proveedor"] == "The North Face" for f in s["items"])
    assert "VANS" not in s["opciones"]["marcas"]
    en_transito = tnf.get("/seguimiento", params={"etapa": "EN_TRANSITO"}).json()
    assert en_transito["total"] and all(f["embarque"] for f in en_transito["items"])
    pendientes = interno.get("/seguimiento", params={"etapa": "PEND_LIBERACION"}).json()
    assert {f["oc"] for f in pendientes["items"]} >= {"4400003851"}
    ordenadas = interno.get("/seguimiento", params={"orden": "cantidad:desc", "size": 5}).json()["items"]
    assert [f["cantidad"] for f in ordenadas] == sorted([f["cantidad"] for f in ordenadas], reverse=True)
    # Por documento de transporte, grupo de artículos y rango de ETA
    bl = interno.get("/seguimiento", params={"documento": "COSU 640018225"}).json()
    assert bl["total"] and all(f["contenedor"] == "TGHU 772104-3" for f in bl["items"])
    assert "CALZ-CAS" in bl["opciones"]["grupos"]
    grupo = interno.get("/seguimiento", params={"grupo": "CHAQ"}).json()
    assert grupo["total"] and all(f["grupo"] == "CHAQ" for f in grupo["items"])
    hoy = date.today()
    eta = interno.get("/seguimiento", params={"eta_desde": hoy.isoformat()}).json()
    assert eta["total"] and all(f["eta"] >= hoy.isoformat() for f in eta["items"])
    # Tablero de embarques: un renglón por embarque, sus unidades y la explosión por OC
    emb = interno.get("/seguimiento/embarques").json()
    assert emb["kpis"]["embarques"] == emb["total"] and emb["total"] >= 2
    cosu = next(e for e in emb["items"] if e["documento"] == "COSU 640018225")
    assert cosu["estado"] == "EN_TRANSITO" and cosu["modo"] == "MARITIMO" and cosu["ocs"] >= 2
    tghu = next(u for u in cosu["detalle_unidades"] if u["contenedor"] == "TGHU 772104-3")
    assert tghu["modalidad"] in ("FCL", "LCL") and tghu["tipo_nombre"]
    assert interno.get("/seguimiento/embarques", params={"marca": "VANS", "documento": "COSU 640018225"}).json()["total"] == 1
    assert all(e["modo"] == "AEREO" for e in interno.get("/seguimiento/embarques", params={"modo": "AEREO"}).json()["items"])
    exp = interno.get(f"/seguimiento/unidades/{tghu['unidad_id']}/explosion").json()
    assert {o["oc"] for o in exp["ocs"]} >= {"4400003703", "4400003752"}
    assert all(l["sku"] and l["cantidad"] for o in exp["ocs"] for l in o["lineas"])
    # Detalle por SKU de una OC (se abre desde el tablero de OCs)
    oc_id = exp["ocs"][0]["oc_id"]
    det = interno.get("/seguimiento", params={"oc_id": oc_id, "size": 200}).json()
    assert det["total"] and all(f["oc_id"] == oc_id for f in det["items"])
    # Tablero de OCs: liberadas o no, avance y estados
    ocs = interno.get("/seguimiento/ordenes", params={"size": 200}).json()
    por_oc = {o["oc"]: o for o in ocs["items"]}
    assert por_oc["4400003851"]["estado"] == "SIN_COMERCIAL" and por_oc["4400003701"]["estado"] == "RECIBIDA"
    assert ocs["kpis"]["sin_liberar"] >= 2
    solo = interno.get("/seguimiento/ordenes", params={"estado": "SIN_COMERCIAL"}).json()["items"]
    assert solo and all(o["liberacion_comercial"] == "P" for o in solo)
    # Seguimiento de facturación y empaque: una fila por PL
    docs = interno.get("/seguimiento/documentos").json()
    assert docs["total"] and {e["clave"] for e in docs["etapas"]} >= {"EMPACANDO", "EN_CAMINO", "RECIBIDO"}
    assert docs["kpis"]["facturas"] >= 1 and docs["kpis"]["cajas"] > 0
    camino = interno.get("/seguimiento/documentos", params={"etapa": "EN_CAMINO"}).json()["items"]
    assert camino and all(f["estado_embarque"] in ("EN_TRANSITO", "ARRIBADO", "ENTREGADO") for f in camino)
    assert all(f["proveedor"] == "The North Face" for f in tnf.get("/seguimiento/documentos").json()["items"])
    admin.patch(f"/roles/{rol['id']}", {"permisos": rol["permisos"]})
    assert tnf.get("/seguimiento/documentos").status_code == 403


def test_documentos_y_reportes(tnf, interno):
    """Factura y packing list en PDF y Excel; reportes de seguimiento con filtros."""
    from io import BytesIO

    from openpyxl import load_workbook

    from app.modulos.documentos.documentos import monto_en_letras

    assert monto_en_letras(4850, "USD") == "FOUR THOUSAND EIGHT HUNDRED FIFTY US DOLLARS AND 00/100"
    assert monto_en_letras(1_021_001.5, "USD") == "ONE MILLION TWENTY-ONE THOUSAND ONE US DOLLARS AND 50/100"
    f = tnf.get("/facturas").json()["items"][0]
    pdf = tnf.get(f"/facturas/{f['id']}/exportar", params={"formato": "pdf"})
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    assert pdf.headers["content-type"] == "application/pdf" and ".pdf" in pdf.headers["content-disposition"]
    xl = tnf.get(f"/facturas/{f['id']}/exportar", params={"formato": "xlsx"})
    ws = load_workbook(BytesIO(xl.content)).active
    textos = {str(c.value) for fila in ws.iter_rows() for c in fila if c.value}
    assert "COMMERCIAL INVOICE" in textos and "EXPORTER / SELLER" in textos and "HS code (SAC)" in textos
    assert any(t.startswith("SAY: ") for t in textos)
    # Las OCs van en el detalle, no en la cabecera
    assert "OC" not in textos and "PO" in textos
    assert tnf.get(f"/facturas/{f['id']}/exportar", params={"formato": "doc"}).status_code == 422
    pl = interno.get("/seguimiento/documentos", params={"etapa": "RECIBIDO"}).json()["items"][0]
    for formato in ("pdf", "xlsx"):
        r = interno.get(f"/packing-lists/{pl['pl_id']}/exportar", params={"formato": formato})
        assert r.status_code == 200 and len(r.content) > 2000
    ws = load_workbook(BytesIO(r.content)).active
    textos = {str(c.value) for fila in ws.iter_rows() for c in fila if c.value}
    assert "PACKING LIST" in textos and any(t.startswith("TOTAL PACKAGES: ") for t in textos)
    assert "Inner packs per carton" in textos and "Per inner pack" in textos and "FINAL DESTINATION" in textos
    for vista in ("ordenes", "embarques", "documentos"):
        for formato in ("pdf", "xlsx"):
            r = interno.get(f"/seguimiento/{vista}/exportar", params={"formato": formato, "marca": "TNF"})
            assert r.status_code == 200, (vista, formato, r.text[:200])
    wb = load_workbook(BytesIO(interno.get("/seguimiento/ordenes/exportar", params={"marca": "VANS"}).content))
    assert wb.sheetnames[1] == "Detail by SKU"
    marcas = {fila[3] for fila in wb["Detail by SKU"].iter_rows(min_row=6, values_only=True) if fila[0]}
    assert marcas == {"VANS"}
    assert interno.get("/seguimiento/otra/exportar").status_code == 404


def test_inner_pack_y_casepack_de_la_oc(tnf, vans, interno):
    """Casepack e inner pack son de la posición de la OC. Todo se mueve en inner
    packs enteros; el casepack es múltiplo del inner pack."""
    det = _oc(tnf, "4400003846")
    oc_id = next(o for o in interno.get("/ordenes", params={"solo_disponible": False}).json()["items"]
                 if o["numero"] == "4400003846")["id"]
    mochila, fleece_s, fleece_l = _pos(det, "OS"), _pos(det, "S"), _pos(det, "L")
    assert (mochila["casepack"], mochila["inner_pack"]) == (20, 5) and fleece_s["inner_pack"] == 5
    assert "casepack" not in interno.get("/catalogos/articulos", params={"size": 1}).json()["items"][0]
    # Editar el empaque de una posición: casepack múltiplo del inner pack
    url = f"/ordenes/{oc_id}/posiciones/{fleece_l['id']}/empaque"
    assert interno.put(url, {"casepack": 12, "inner_pack": 5}).status_code == 422
    assert tnf.put(url, {"casepack": 15, "inner_pack": 5}).status_code == 403
    assert interno.put(url, {"casepack": 15, "inner_pack": 5}).json()["casepack"] == 15
    assert interno.put(url, {"casepack": None, "inner_pack": 5}).status_code == 200
    prepack = _oc(vans, "4400003903")["posiciones"][0]
    oc_pp = next(o for o in interno.get("/ordenes", params={"solo_disponible": False}).json()["items"]
                 if o["numero"] == "4400003903")["id"]
    assert interno.put(f"/ordenes/{oc_pp}/posiciones/{prepack['id']}/empaque", {"casepack": 2}).status_code == 422
    # Facturar: solo inner packs enteros
    r = tnf.post("/facturas", {"lineas": [{"posicion_id": fleece_s["id"], "cantidad": 12}]})
    assert r.status_code == 422 and "inner pack" in r.json()["detalle"][0]["mensaje"]
    fid = tnf.post("/facturas", {"lineas": [{"posicion_id": fleece_s["id"], "cantidad": 30},
                                           {"posicion_id": mochila["id"], "cantidad": 60}]}).json()["id"]
    assert interno.put(f"/ordenes/{oc_id}/posiciones/{mochila['id']}/empaque",
                       {"casepack": 10, "inner_pack": 5}).status_code == 409
    pl_id = tnf.post(f"/facturas/{fid}/packing-lists", {}).json()["id"]
    pl = tnf.get(f"/packing-lists/{pl_id}").json()
    lf = next(l for l in pl["lineas"] if l["talla"] == "S")
    lm = next(l for l in pl["lineas"] if l["talla"] == "OS")
    assert (lf["regla"], lf["inner_pack"]) == ("LIBRE", 5) and lm["regla"] == "CASEPACK"
    # Sin casepack: la caja lleva inner packs enteros
    r = tnf.post(f"/packing-lists/{pl_id}/cajas", {"version": pl["version"], "num_cajas": 1, "items": [
        {"pl_linea_id": lf["id"], "cantidad_por_caja": 12}]})
    assert r.status_code == 422 and "inner pack" in r.json()["detalle"][0]["mensaje"]
    # Una plantilla que no es múltiplo del inner pack se omite en el empaque automático
    t12 = tnf.post("/plantillas", {"nombre": "Caja fleece 12 un", "cantidad_por_caja": 12, "unidad": "UN",
                                   "largo": 50, "ancho": 40, "alto": 30, "peso_neto": 6, "peso_bruto": 7}).json()
    previa = tnf.post(f"/packing-lists/{pl_id}/empaque/previa",
                      {"filas": [{"pl_linea_id": lf["id"], "plantilla_id": t12["id"]}]}).json()
    assert previa["resumen"]["filas"] == 0 and "inner pack" in previa["filas"][0]["omitida"]
    r = tnf.post(f"/packing-lists/{pl_id}/cajas", {"version": pl["version"], "num_cajas": 2, "items": [
        {"pl_linea_id": lf["id"], "cantidad_por_caja": 15}]})
    assert r.status_code == 200, r.text
    r = _empacar(tnf, pl_id, r.json()["version"], [(lm["id"], None)])
    assert r.status_code == 200, r.text
    pl = tnf.get(f"/packing-lists/{pl_id}").json()
    cajas = [g for g in pl["grupos"] if g["cuenta_como"] == "BULTO"]
    inners = {g["contenido"][0]["pl_linea_id"]: (g["num_cajas"], g["contenido"][0]["inner_packs_por_caja"]) for g in cajas}
    assert inners == {lf["id"]: (2, 3), lm["id"]: (3, 4)}
    # Los inner packs son empaques físicos dentro de cada caja, con su tara
    caja = next(g for g in cajas if g["contenido"][0]["pl_linea_id"] == lf["id"])
    inner = next(g for g in pl["grupos"] if g["padre_id"] == caja["id"])
    assert inner["num_cajas"] == 6 and inner["por_padre"] == 3 and inner["items"][0]["cantidad_por_caja"] == 5
    assert abs(caja["peso_bruto_caja"] - (15 * 0.5 + 3 * inner["tara"] + caja["tara"])) < 1e-6
    # Por destino y sugerencia de unidades de carga
    assert [d["centro_destino"] for d in pl["destinos"]] == ["2220"] and pl["destinos"][0]["cajas"] == 5
    # Las cajas toman las medidas del tipo de empaque y el peso sale de los artículos: ya hay volumen y peso
    t = pl["totales"]
    assert t["cbm"] == 5 * 0.096 and pl["sugerencia_unidades"]["kg"] == round(t["peso_bruto"], 2) > 0


def test_sugerencia_de_unidades(interno):
    s = interno.get("/sugerencia-unidades", params={"cbm": 10, "kg": 2000}).json()["modos"]
    assert s["MARITIMO"][0]["recomendada"] and s["MARITIMO"][0]["texto"] == "1 × LCL"
    assert s["TERRESTRE"][0]["texto"] == "1 × LTL"
    s = interno.get("/sugerencia-unidades", params={"cbm": 60, "kg": 9000, "modo": "MARITIMO"}).json()["modos"]
    assert s["MARITIMO"][0]["texto"] == "1 × 40HC" and list(s) == ["MARITIMO"]
    s = interno.get("/sugerencia-unidades", params={"cbm": 150, "kg": 30000}).json()["modos"]
    assert s["MARITIMO"][0]["texto"] == "2 × 40HC + 1 × 20GP"
    aereo = interno.get("/sugerencia-unidades", params={"cbm": 2, "kg": 500}).json()["modos"]["AEREO"][0]
    assert aereo["peso_cobrable"] == 500
    assert interno.get("/sugerencia-unidades", params={"cbm": 1, "modo": "BARCO"}).status_code == 422


def test_plantillas_de_carga_en_ingles(interno):
    """Las plantillas CSV descargables (columnas en inglés) se aceptan tal cual."""
    base = os.path.join(os.path.dirname(__file__), "..", "..", "frontend", "public")

    def subir(ruta, archivo):
        with open(os.path.join(base, archivo), "rb") as f:
            return interno.c.post("/api" + ruta, headers=interno.h, files={"archivo": (archivo, f.read(), "text/csv")})

    # La plantilla de OC se genera con las fechas en el formato del usuario y se lee tal cual
    plantilla = interno.c.get("/api/ordenes/plantilla", headers=interno.h).content
    r = interno.c.post("/api/ordenes/importar/previa", headers=interno.h,
                       files={"archivo": ("plantilla.xlsx", plantilla, "application/octet-stream")})
    assert r.status_code == 200, r.text
    filas = r.json()["filas"]
    assert filas and not any("column" in m or "date" in m for f in filas for m in f["mensajes"])
    r = subir("/catalogos/articulos/importar", "plantilla_articulos.csv")
    assert r.status_code == 200 and not r.json()["errores"] and r.json()["creados"] == 3, r.text
    r = subir("/catalogos/prepacks/importar", "plantilla_prepacks.csv")
    assert r.status_code == 200 and not r.json()["errores"] and r.json()["creados"] == 1, r.text


def test_acceso_seguro(client):
    """Dos pasos por SMS, cookie httpOnly, CSRF, bloqueo por intentos,
    política de contraseñas y cierre de sesiones."""
    from conftest import PASSWORD, Api, iniciar_sesion

    from app.modulos.acceso.limites import reiniciar

    reiniciar()
    # Paso 1: la contraseña correcta solo abre el desafío; no hay sesión todavía
    r = client.post("/api/auth/login", json={"email": "interno@demo.com", "password": PASSWORD})
    d = r.json()
    assert d["dos_pasos"] and d["telefono"].startswith("+503") and "•" in d["telefono"] and "sesion" not in r.cookies
    assert client.get("/api/auth/me").status_code == 401
    # Código incorrecto: se cuenta el intento
    r = client.post("/api/auth/verificar", json={"desafio": d["desafio"], "codigo": "000000"})
    assert r.status_code == 401 and r.json()["codigo"] == "codigo_incorrecto"
    # Reenvío inmediato: hay que esperar
    assert client.post("/api/auth/reenviar", json={"desafio": d["desafio"]}).status_code == 429
    # Paso 2: cookie httpOnly y SameSite=Strict
    r = client.post("/api/auth/verificar", json={"desafio": d["desafio"], "codigo": d["codigo_demo"]})
    assert r.status_code == 200 and r.json()["usuario"]["email"] == "interno@demo.com"
    cookie = r.headers["set-cookie"].lower()
    assert "httponly" in cookie and "samesite=strict" in cookie
    # El código no se reutiliza
    assert client.post("/api/auth/verificar", json={"desafio": d["desafio"],
                                                    "codigo": d["codigo_demo"]}).status_code == 401
    assert client.get("/api/auth/me").status_code == 200
    # Con cookie, cambiar datos exige el encabezado de la aplicación (CSRF)
    assert client.post("/api/alertas/999/resolver").status_code == 403
    assert client.post("/api/alertas/999/resolver", headers={"X-Requested-With": "fetch"}).status_code != 403
    # Encabezados de seguridad
    h = client.get("/api/auth/me").headers
    assert h["x-frame-options"] == "DENY" and "frame-ancestors 'none'" in h["content-security-policy"]
    # Cerrar sesión revoca la sesión
    client.post("/api/auth/logout", headers={"X-Requested-With": "fetch"})
    client.cookies.clear()
    assert client.get("/api/auth/me").status_code == 401

    # Rutas restringidas: el proveedor no entra a administración ni a mantenimiento
    tnf = Api(client, "tnf@demo.com")
    assert tnf.get("/usuarios").status_code == 403 and tnf.get("/catalogos/sociedades").status_code == 403
    assert client.get("/api/facturas").status_code == 401

    # Política de contraseñas y alta con celular
    admin = Api(client, "admin@demo.com")
    r = admin.post("/usuarios", {"email": "nuevo@demo.com", "nombre": "Nuevo", "rol": "interno", "password": "abc12"})
    assert r.status_code == 422 and r.json()["codigo"] == "password_debil"
    r = admin.post("/usuarios", {"email": "nuevo@demo.com", "nombre": "Nuevo", "rol": "interno",
                                 "password": "Seguridad-2026", "telefono": "7000"})
    assert r.status_code == 422 and r.json()["detalle"][0]["campo"] == "telefono"
    r = admin.post("/usuarios", {"email": "nuevo@demo.com", "nombre": "Nuevo", "rol": "interno",
                                 "password": "Seguridad-2026", "telefono": "+503 7000-1234"})
    assert r.status_code == 200, r.text
    nuevo_id = r.json()["id"]
    nuevo = next(u for u in admin.get("/usuarios").json() if u["id"] == nuevo_id)
    assert nuevo["telefono"] == "+50370001234" and nuevo["dos_pasos"]
    # Sin celular registrado no se puede entrar con dos pasos
    admin.patch(f"/usuarios/{nuevo_id}", {"telefono": None})
    r = client.post("/api/auth/login", json={"email": "nuevo@demo.com", "password": "Seguridad-2026"})
    assert r.status_code == 403 and r.json()["codigo"] == "sin_telefono"
    admin.patch(f"/usuarios/{nuevo_id}", {"telefono": "+50370001234"})
    # Cambiar la contraseña propia cierra las otras sesiones
    reiniciar()
    t1 = iniciar_sesion(client, "nuevo@demo.com", "Seguridad-2026")
    t2 = iniciar_sesion(client, "nuevo@demo.com", "Seguridad-2026")
    h1, h2 = {"Authorization": f"Bearer {t1}"}, {"Authorization": f"Bearer {t2}"}
    r = client.post("/api/auth/password", json={"actual": "Seguridad-2026", "nueva": "corta"}, headers=h1)
    assert r.status_code == 422
    r = client.post("/api/auth/password", json={"actual": "Seguridad-2026", "nueva": "Otra#Clave2027"}, headers=h1)
    assert r.status_code == 200
    assert client.get("/api/auth/me", headers=h1).status_code == 200
    assert client.get("/api/auth/me", headers=h2).status_code == 401
    # Bloqueo tras 5 contraseñas incorrectas, aun con la correcta después. Quien
    # no sabe la contraseña ve el mismo error que con un correo inexistente.
    for _ in range(6):
        r = client.post("/api/auth/login", json={"email": "nuevo@demo.com", "password": "mala"})
    assert r.status_code == 401 and r.json()["codigo"] == "credenciales"
    r = client.post("/api/auth/login", json={"email": "nuevo@demo.com", "password": "Otra#Clave2027"})
    assert r.status_code == 423 and r.json()["codigo"] == "bloqueado"
    assert next(u for u in admin.get("/usuarios").json() if u["id"] == nuevo_id)["bloqueado"]
    # El administrador desbloquea al restablecer la contraseña
    admin.patch(f"/usuarios/{nuevo_id}", {"password": "Restablecida-2028"})
    reiniciar()
    assert iniciar_sesion(client, "nuevo@demo.com", "Restablecida-2028")
    # Mismo mensaje para correo inexistente y contraseña incorrecta
    a = client.post("/api/auth/login", json={"email": "nadie@demo.com", "password": "x"}).json()["mensaje"]
    b = client.post("/api/auth/login", json={"email": "admin@demo.com", "password": "x"}).json()["mensaje"]
    assert a == b
    reiniciar()


def test_fecha_estimada_en_tienda(interno):
    """Con los lead times de su origen cada OC tiene su fecha estimada en tienda."""
    from datetime import date, timedelta
    ocs = interno.get("/ordenes", params={"size": 50}).json()["items"]
    con = [o for o in ocs if o.get("tienda_estimada")]
    assert con and all(o["tienda_estimada"] >= (o.get("arribo_estimado") or "") for o in con)
    o = next(o for o in con if o["fecha_tienda"])
    dif = (date.fromisoformat(o["tienda_estimada"]) - date.fromisoformat(o["fecha_tienda"])).days
    assert o["dias_vs_tienda"] == dif
    det = interno.get(f"/ordenes/{o['id']}/posiciones").json()["oc"]
    assert det["tienda_estimada"] == o["tienda_estimada"]
    # Sin embarque: XF (u hoy) + tránsito + puerto, ingreso y reexportación de su región
    lt = interno.get("/seguimiento/leadtimes", params={"size": 100}).json()
    x = next(i for i in lt["items"] if i["oc_id"] == o["id"])
    assert x["tienda_estimada"] == o["tienda_estimada"]
    fila = next(f for f in interno.get("/seguimiento", params={"size": 200}).json()["items"] if f["tienda_estimada"])
    assert fila["tienda_estimada"] > (date.today() - timedelta(days=400)).isoformat()
    por_oc = interno.get("/seguimiento/ordenes", params={"size": 100}).json()["items"]
    assert any(p["tienda_estimada"] for p in por_oc)


def test_crear_oc_desde_formulario(interno, vans):
    """La OC se puede crear en la plataforma con la misma estructura y
    validaciones que la carga masiva; precio, moneda y empresa son opcionales
    al crearla, pero se exigen al facturar."""
    sku = _oc(interno, "4400003901")["posiciones"][0]["codigo_sap"]
    cab = {"proveedor": "VANS", "oc": "PO-FORM-1", "centro_destino": "2220", "liberacion_comercial": "C",
           "liberacion_logistica": "Released"}
    # Lo mínimo: proveedor, número, artículo y cantidad
    r = interno.post("/ordenes", {"cabecera": cab, "lineas": [{"codigo_sap": sku}]})
    assert r.status_code == 422 and any("quantity" in d["mensaje"] for d in r.json()["detalle"])
    r = interno.post("/ordenes", {"cabecera": cab, "lineas": [{"codigo_sap": "NOEXISTE", "cantidad": 5}]})
    assert r.status_code == 422 and "item master" in r.json()["detalle"][0]["mensaje"]
    r = interno.post("/ordenes", {"cabecera": cab, "lineas": [{"codigo_sap": sku, "cantidad": 12, "casepack": 6}]})
    assert r.status_code == 200, r.text
    oc = r.json()
    assert oc["numero"] == "PO-FORM-1" and oc["lineas"] == 1
    det = interno.get(f"/ordenes/{oc['oc_id']}/posiciones").json()
    assert det["posiciones"][0]["posicion"] == "10" and det["posiciones"][0]["casepack"] == 6
    assert det["posiciones"][0]["importe"] is None  # sin precio: el valor queda pendiente
    # El mismo número no se repite para el proveedor
    assert interno.post("/ordenes", {"cabecera": cab, "lineas": [{"codigo_sap": sku, "cantidad": 1}]}).status_code == 409
    # Sin precio, moneda ni empresa no se puede facturar: se dice qué falta
    r = vans.post("/facturas", {"lineas": [{"posicion_id": det["posiciones"][0]["id"], "cantidad": 6}]})
    assert r.status_code == 422 and "missing currency, price" in r.json()["detalle"][0]["mensaje"]
    # La sociedad se toma del centro destino: el proveedor solo trabaja con las suyas
    assert det["oc"]["sociedad"] == "8000"
    r = interno.post("/ordenes", {"cabecera": {**cab, "oc": "PO-FORM-3", "centro_destino": "3200"},
                                  "lineas": [{"codigo_sap": sku, "cantidad": 2}]})
    assert r.status_code == 422 and "does not work with company GT01" in str(r.json()["detalle"])
    r = interno.post("/ordenes", {"cabecera": {**cab, "oc": "PO-FORM-3", "sociedad": "8000", "centro_destino": "PA20"},
                                  "lineas": [{"codigo_sap": sku, "cantidad": 2}]})
    assert r.status_code == 422 and "does not belong to company 8000" in str(r.json()["detalle"])
    r = interno.post("/ordenes", {"cabecera": {**cab, "oc": "PO-FORM-3", "sociedad": "8000", "centro": "PA10"},
                                  "lineas": [{"codigo_sap": sku, "cantidad": 2}]})
    assert r.status_code == 422 and "does not belong to company 8000" in str(r.json()["detalle"])
    # Un artículo de otro proveedor no entra en la OC
    otro = _oc(interno, "4400003850")["posiciones"][0]["codigo_sap"]
    r = interno.post("/ordenes", {"cabecera": {**cab, "oc": "PO-FORM-3"}, "lineas": [{"codigo_sap": otro, "cantidad": 2}]})
    assert r.status_code == 422 and "another supplier" in str(r.json()["detalle"])
    # Las opciones del formulario dicen con qué sociedades trabaja cada proveedor
    op = interno.get("/ordenes/formulario").json()
    assert next(p for p in op["proveedores"] if p["valor"] == "VANS")["sociedades"] == ["8000", "PA01"]
    # Precio sin moneda no es válido
    r = interno.post("/ordenes", {"cabecera": {**cab, "oc": "PO-FORM-2"}, "lineas": [{"codigo_sap": sku, "cantidad": 2, "precio": 10}]})
    assert r.status_code == 422 and "currency" in r.json()["detalle"][0]["mensaje"]


def test_inner_pack_en_el_packing_list(interno, vans):
    """La OC trae el casepack (sólidos) o la curva (prepacks); el inner pack se
    define al armar el packing list, mientras la línea no tenga cajas."""
    sku = _oc(interno, "4400003901")["posiciones"][0]["codigo_sap"]
    cab = {"proveedor": "VANS", "oc": "PO-INNER-1", "sociedad": "8000", "moneda": "USD", "centro_destino": "2220",
           "liberacion_comercial": "C", "liberacion_logistica": "300"}
    r = interno.post("/ordenes", {"cabecera": cab, "lineas": [{"codigo_sap": sku, "cantidad": 24, "precio": 10, "casepack": 12}]})
    assert r.status_code == 200, r.text
    pos = interno.get(f"/ordenes/{r.json()['oc_id']}/posiciones").json()["posiciones"][0]
    assert pos["inner_pack"] is None
    fid = vans.post("/facturas", {"lineas": [{"posicion_id": pos["id"], "cantidad": 24}]}).json()["id"]
    # La descripción de la factura es la aduanera, sin repetir la marca (que tiene su columna)
    lf = vans.get(f"/facturas/{fid}").json()["lineas"][0]
    assert lf["marca"] == "VANS" and "VANS" not in (lf["descripcion_comercial"] or "").upper()
    pl_id = vans.post(f"/facturas/{fid}/packing-lists", {}).json()["id"]
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    linea = pl["lineas"][0]
    assert linea["inner_editable"] and linea["inner_pack"] is None
    url = f"/packing-lists/{pl_id}/lineas/{linea['id']}/inner"
    r = vans.put(url, {"version": pl["version"], "inner_pack": 5})
    assert r.status_code == 422 and "multiple of the inner pack" in r.json()["mensaje"]
    r = vans.put(url, {"version": pl["version"], "inner_pack": 4})
    assert r.status_code == 200, r.text
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    assert pl["lineas"][0]["inner_pack"] == 4
    # Con cajas ya no se cambia
    r = _empacar(vans, pl_id, pl["version"], [(pl["lineas"][0]["id"], None)])
    assert r.status_code == 200, r.text
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    assert not pl["lineas"][0]["inner_editable"]
    grupo = next(g for g in pl["grupos"] if g["cuenta_como"] == "BULTO")
    assert grupo["contenido"][0]["inner_packs_por_caja"] == 3
    assert vans.put(url, {"version": pl["version"], "inner_pack": 6}).status_code == 409


def test_estado_del_embarque_frente_a_tienda(interno):
    """El embarque se califica con la fecha estimada en tienda: en tiempo, en
    riesgo o atrasado. El tipo de producto puede sumar días después del puerto."""
    from datetime import date
    embs = interno.get("/embarques").json()
    assert all("estado_tiempo" in e for e in embs)
    con = [e for e in embs if e["estado_tiempo"]]
    assert con and all(e["estado_tiempo"] in ("A_TIEMPO", "JUSTO", "ATRASO") for e in con)
    det = interno.get(f"/embarques/{con[0]['id']}").json()
    assert det["estado_tiempo"] == con[0]["estado_tiempo"]
    # Días extra del tipo de producto: la estimación en tienda se corre
    ocs = interno.get("/ordenes", params={"size": 50}).json()["items"]
    o = next(o for o in ocs if o.get("tienda_estimada"))
    grupo = interno.get(f"/ordenes/{o['id']}/posiciones").json()["posiciones"][0]["grupo"]
    g = next(x for x in interno.get("/catalogos/grupos").json()["items"] if x["codigo"] == grupo)
    assert interno.patch(f"/catalogos/grupos/{g['id']}", {"dias_extra": 10}).status_code == 200
    o2 = next(x for x in interno.get("/ordenes", params={"size": 50}).json()["items"] if x["id"] == o["id"])
    assert (date.fromisoformat(o2["tienda_estimada"]) - date.fromisoformat(o["tienda_estimada"])).days == 10
    interno.patch(f"/catalogos/grupos/{g['id']}", {"dias_extra": None})


def test_consistencia_de_maestros(interno):
    """Los maestros quedan encadenados: no se quita a un proveedor una sociedad
    con la que tiene OCs, y el contacto de un centro es de la sociedad de ese centro."""
    socs = {s["codigo"]: s["id"] for s in interno.get("/catalogos/sociedades", params={"size": 100}).json()["items"]}
    vans = next(p for p in interno.get("/catalogos/proveedores").json()["items"] if p["codigo"] == "VANS")
    r = interno.patch(f"/catalogos/proveedores/{vans['id']}", {"sociedades": [socs["PA01"]]})
    assert r.status_code == 422 and "8000" in r.json()["detalle"][0]["mensaje"]
    cen = {c["codigo"]: c["id"] for c in interno.get("/catalogos/centros", params={"size": 100}).json()["items"]}
    r = interno.post("/catalogos/contactos", {"nombre": "Test", "rol": "NOTIFY", "sociedad_id": socs["8000"],
                                              "centro_id": cen["PA10"]})
    assert r.status_code == 422 and "does not belong" in r.json()["detalle"][0]["mensaje"]


def test_numero_propio_del_pl(interno, vans):
    sku = _oc(interno, "4400003901")["posiciones"][0]["codigo_sap"]
    cab = {"proveedor": "VANS", "oc": "PO-PLNUM-1", "sociedad": "8000", "moneda": "USD", "centro_destino": "2220",
           "liberacion_comercial": "C", "liberacion_logistica": "300"}
    oc = interno.post("/ordenes", {"cabecera": cab, "lineas": [{"codigo_sap": sku, "cantidad": 24, "precio": 10}]}).json()
    pos = interno.get(f"/ordenes/{oc['oc_id']}/posiciones").json()["posiciones"][0]
    fid = vans.post("/facturas", {"lineas": [{"posicion_id": pos["id"], "cantidad": 24}]}).json()["id"]
    pl_id = vans.post(f"/facturas/{fid}/packing-lists", {}).json()["id"]
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    r = vans.put(f"/packing-lists/{pl_id}/numero", {"version": pl["version"], "numero": "VN-PL 2026/88"})
    assert r.status_code == 200 and r.json()["numero"] == "VN-PL 2026/88"
    assert vans.get(f"/packing-lists/{pl_id}").json()["numero"] == "VN-PL 2026/88"
    r = vans.put(f"/packing-lists/{pl_id}/numero", {"numero": "<script>"})
    assert r.status_code == 422
    otro = vans.post(f"/facturas/{fid}/packing-lists", {})
    if otro.status_code == 200:
        r = vans.put(f"/packing-lists/{otro.json()['id']}/numero", {"numero": "vn-pl 2026/88"})
        assert r.status_code == 409


def test_estructura_fisica_por_carga(interno, vans):
    """El PL se arma desde un archivo con el identificador de cada nivel:
    unidades iguales se agrupan y los pesos salen de la estructura."""
    import io

    from openpyxl import Workbook, load_workbook

    sku = _oc(interno, "4400003901")["posiciones"][0]["codigo_sap"]
    cab = {"proveedor": "VANS", "oc": "PO-ESTR-1", "sociedad": "8000", "moneda": "USD", "centro_destino": "2220",
           "liberacion_comercial": "C", "liberacion_logistica": "300"}
    r = interno.post("/ordenes", {"cabecera": cab, "lineas": [{"codigo_sap": sku, "cantidad": 24, "precio": 10}]})
    pos = interno.get(f"/ordenes/{r.json()['oc_id']}/posiciones").json()["posiciones"][0]
    fid = vans.post("/facturas", {"lineas": [{"posicion_id": pos["id"], "cantidad": 24}]}).json()["id"]
    pl_id = vans.post(f"/facturas/{fid}/packing-lists", {}).json()["id"]
    plantilla = vans.c.get(f"/api/packing-lists/{pl_id}/estructura/plantilla", headers=vans.h)
    enc = [c.value for c in load_workbook(io.BytesIO(plantilla.content))["Data"][1]]
    assert enc[:3] == ["Pallet", "Master carton", "Inner pack"] and "Item code *" in enc

    def subir(filas):
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(["Pallet", "Master carton", "Inner pack", "Item code", "PO", "Quantity"])
        for f in filas:
            ws.append(f)
        b = io.BytesIO()
        wb.save(b)
        return vans.c.post(f"/api/packing-lists/{pl_id}/estructura/importar", headers=vans.h,
                           files={"archivo": ("e.xlsx", io.BytesIO(b.getvalue()), "application/octet-stream")})

    # Más de lo que tiene el PL o una relación que los tipos no permiten: no cambia nada
    r = subir([["P001", "C001", "PK001", sku, "", 30]])
    assert r.status_code == 422 and "more than" in r.text
    r = subir([["P001", "", "PK001", sku, "", 6]])
    assert r.status_code == 422 and "Not allowed" in r.text
    filas = [["P001", f"C00{c}", f"PK00{(c - 1) * 2 + k}", sku, "", 6] for c in (1, 2) for k in (1, 2)]
    r = subir(filas)
    assert r.status_code == 200, r.text
    pl = vans.get(f"/packing-lists/{pl_id}").json()
    por_tipo = {g["tipo"]: g for g in pl["grupos"]}
    pallet, caja, inner = por_tipo["Pallet"], por_tipo["Master carton"], por_tipo["Inner pack"]
    assert (pallet["num_cajas"], caja["num_cajas"], inner["num_cajas"]) == (1, 2, 4)
    assert caja["padre_id"] == pallet["id"] and inner["padre_id"] == caja["id"] and inner["items"][0]["cantidad_por_caja"] == 6
    peso = inner["items"][0]["peso_unitario"]
    assert abs(caja["peso_bruto_caja"] - (12 * peso + 2 * inner["tara"] + caja["tara"])) < 1e-6
    assert abs(pl["totales"]["peso_bruto"] - (2 * caja["peso_bruto_caja"] + pallet["tara"])) < 1e-3
    assert pl["lineas"][0]["sin_caja"] == 0 and pl["totales"]["cajas"] == 2 and pl["totales"]["pallets"] == 1
