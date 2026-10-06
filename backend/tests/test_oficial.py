"""Capa oficial del arancel (paquetes Excel 01 y 02) y migraciones."""
import io
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

RAIZ = Path(__file__).resolve().parent.parent


def test_migraciones_al_dia_con_los_modelos():
    """Las migraciones de Alembic crean exactamente el esquema de los modelos:
    un cambio de modelo sin su migración hace fallar esta prueba."""
    with tempfile.TemporaryDirectory() as tmp:
        env = {"DATABASE_URL": f"sqlite:///{tmp}/m.db", "PATH": "/usr/bin:/bin"}
        alembic = [sys.executable, "-m", "alembic"]
        for args in (["upgrade", "head"], ["check"], ["downgrade", "base"], ["upgrade", "head"]):
            r = subprocess.run(alembic + args, cwd=RAIZ, env=env, capture_output=True, text=True)
            assert r.returncode == 0, r.stdout + r.stderr


def test_paquete_oficial_cargado(interno):
    caps = interno.get("/aranceles/oficial/capitulos").json()
    assert caps["total"] == 99 and caps["habilitados"] > 60
    por = {c["capitulo"]: c for c in caps["items"]}
    assert por["64"]["clasificacion"] and por["01"]["activo"] is False
    assert {d["codigo"] for d in por["28"]["dominios"]} == {"CHEMICALS", "RAW_MATERIALS"}
    assert por["57"]["solo_manual"] and not por["57"]["candidato_auto"]
    fv = interno.get("/aranceles/oficial/fuentes").json()
    assert {"SRC-SIECA-ACI", "SRC-CR-ATENA", "SRC-PA-ANA"} <= {f["codigo"] for f in fv["fuentes"]}
    sac = next(v for v in fv["versiones"] if v["codigo"] == "SAC-2025-V6")
    assert sac["estado"] == "PUBLICADA" and sac["fuente"] == "SRC-SIECA-ACI" and sac["vigente_desde"] == "2025-08-01"
    doms = {d["codigo"]: d for d in interno.get("/aranceles/oficial/dominios").json()}
    assert set(doms) == {"CHEMICALS", "RAW_MATERIALS", "FOOTWEAR", "APPAREL", "ACCESSORIES_MERCH"}
    quim = {c["capitulo"]: c["relevancia"] for c in doms["CHEMICALS"]["capitulos"]}
    assert quim["29"] == "PRIMARY" and quim["39"] == "SECONDARY"
    # Búsqueda inteligente en capítulos
    assert [c["capitulo"] for c in interno.get("/aranceles/oficial/capitulos", params={"q": "quimicos organicos"}).json()["items"]] == ["28", "29"]
    assert [c["capitulo"] for c in interno.get("/aranceles/oficial/capitulos", params={"q": "calzado polainas"}).json()["items"]] == ["64"]


def test_control_de_capitulos_en_bloque(interno):
    caps = {c["capitulo"]: c for c in interno.get("/aranceles/oficial/capitulos").json()["items"]}
    ids = [caps["93"]["id"], caps["97"]["id"]]
    r = interno.patch("/aranceles/oficial/capitulos", {"ids": ids, "clasificacion": True, "candidato_auto": True})
    assert r.status_code == 200, r.text
    caps = {c["capitulo"]: c for c in interno.get("/aranceles/oficial/capitulos").json()["items"]}
    assert caps["97"]["activo"] and caps["97"]["clasificacion"]  # habilitar activa el capítulo
    r = interno.patch("/aranceles/oficial/capitulos", {"ids": ids, "archivado": True})
    caps = {c["capitulo"]: c for c in interno.get("/aranceles/oficial/capitulos").json()["items"]}
    assert caps["97"]["archivado"] and not caps["97"]["clasificacion"] and not caps["97"]["candidato_auto"]


def test_reimportar_es_idempotente_y_valida(interno):
    paquete = (RAIZ / "app/data/oficial/01_carga_oficial_catalogos_v3.xlsx").read_bytes()
    r = interno.c.post("/api/aranceles/oficial/importar", headers=interno.h,
                       files={"archivo": ("01.xlsx", io.BytesIO(paquete), "application/octet-stream")})
    assert r.status_code == 200 and r.json()["creados"] == 0 and not r.json()["errores"], r.text
    wb = Workbook()
    ws = wb.active
    ws.title = "Versions"
    ws.append(["Version ID", "Dataset", "Version label", "Status", "Valid from", "Valid to", "Source ID"])
    ws.append(["GT-2026", "Guatemala national overlay", "2026 snapshot", "Draft", "2026-01-01", "2025-01-01", "SRC-GT-SAT"])
    ws.append(["X-1", "Nothing", "x", "Draft", "", "", "SRC-NO-EXISTE"])
    b = io.BytesIO()
    wb.save(b)
    r = interno.c.post("/api/aranceles/oficial/importar", headers=interno.h,
                       files={"archivo": ("v.xlsx", io.BytesIO(b.getvalue()), "application/octet-stream")})
    errores = [e["mensaje"] for e in r.json()["errores"]]
    assert any("Valid to" in m for m in errores) and any("SRC-NO-EXISTE" in m for m in errores)


def test_arbol_arancelario_completo(interno):
    r = interno.get("/aranceles/arbol/resumen").json()
    assert r["version"] == "SAC-2025-V6" and len(r["checksum"]) == 64
    assert r["niveles"]["CAPITULO"] == 99 and r["niveles"]["PARTIDA"] > 1000 and r["niveles"]["INCISO"] > 8000
    caps = interno.get("/aranceles/arbol").json()["items"]
    c64 = next(c for c in caps if c["codigo"] == "64")
    assert c64["capitulo_habilitado"] and c64["hijos"] == 6
    # Capítulos fuera del motor actual también tienen sus incisos (p. ej. químicos, cap. 29)
    c29 = next(c for c in caps if c["codigo"] == "29")
    partidas = interno.get("/aranceles/arbol", params={"padre_id": c29["id"]}).json()["items"]
    assert partidas[0]["codigo"].startswith("29") and partidas[0]["nivel"] == "PARTIDA"
    # Búsqueda por código (con puntos o varios) y por palabras en cualquier orden
    r = interno.get("/aranceles/arbol", params={"q": "6404.19"}).json()
    assert r["items"][0]["codigo"] == "6404.19" and all(i["codigo_norm"].startswith("640419") for i in r["items"])
    r = interno.get("/aranceles/arbol", params={"q": "6404.11 4202.92"}).json()
    assert {i["codigo_norm"][:6] for i in r["items"]} == {"640411", "420292"}
    r = interno.get("/aranceles/arbol", params={"q": "caucho suela calzado"}).json()
    assert r["total"] > 0 and any(i["codigo_norm"].startswith("64") for i in r["items"])
    n = interno.get(f"/aranceles/arbol/{r['items'][0]['id']}").json()
    assert n["ruta"][0]["nivel"] == "CAPITULO" and n["version"]["fuente"] == "SRC-SIECA-ACI"
    sub = next(i for i in interno.get("/aranceles/arbol", params={"q": "6404.19"}).json()["items"] if i["codigo"] == "6404.19")
    d = interno.get(f"/aranceles/arbol/{sub['id']}").json()
    assert d["hijos_lista"] and d["hijos_lista"][0]["dai"] is not None
    assert any(p["iso"] == "SV" and p["codigos"] for p in d["paises"]) and d["notas"]


def test_atributos_en_base_de_datos(interno):
    """Atributos oficiales (paquete 02) y de la ficha del motor, con opciones y ámbitos."""
    r = interno.get("/aranceles/atributos").json()
    assert r["por_origen"]["OFICIAL"] == 38 and r["por_origen"]["MOTOR"] == 79
    por = {a["codigo"]: a for a in r["items"]}
    assert por["cas_number"]["dominio"] == "CHEMICALS" and por["material_composition"]["tipo_dato"] == "composition"
    assert por["estiloCalz"]["dominio"] == "FOOTWEAR" and por["tejido"]["dominio"] == "APPAREL"
    d = interno.get(f"/aranceles/atributos/{por['physical_state']['id']}").json()
    assert [o["codigo"] for o in d["opciones"]] == ["SOLID", "LIQUID", "GAS", "POWDER", "PASTE"]
    assert d["ambitos"][0]["tipo_ambito"] == "DOMAIN" and d["ambitos"][0]["codigo_ambito"] == "CHEMICALS"
    # Ámbitos del motor: categorías donde aplica y respuestas que lo activan
    t = interno.get(f"/aranceles/atributos/{por['tejido']['id']}").json()
    cint = next(x for x in t["ambitos"] if x["codigo_ambito"] == "cinturon")
    assert {"campo": "materialCinturon", "operador": "EQUAL", "valor": "textil"}.items() <= cint["condicion"][0].items()
    assert next(x for x in t["ambitos"] if x["codigo_ambito"] == "camiseta")["condicion"] is None
    # Búsqueda inteligente
    assert {a["codigo"] for a in interno.get("/aranceles/atributos", params={"q": "name chemic"}).json()["items"]} == {"chemical_name"}
    # La ficha recibe el catálogo en el contexto
    ctx = interno.get("/clasificacion/contexto").json()["atributos"]
    assert ctx["motor"]["tejido"]["opciones"]["punto"]["activo"] and len(ctx["genericos"]) == 38


def test_editar_atributos_opciones_y_ambitos(interno):
    por = {a["codigo"]: a for a in interno.get("/aranceles/atributos").json()["items"]}
    aid = por["physical_state"]["id"]
    d = interno.post(f"/aranceles/atributos/{aid}/opciones", {"codigo": "GRANULE", "etiqueta": "Granules", "alias": "pellets; granules"}).json()
    assert d["opciones"][-1]["codigo"] == "GRANULE"
    gas = next(o for o in d["opciones"] if o["codigo"] == "GAS")
    d = interno.patch(f"/aranceles/atributos/{aid}/opciones/{gas['id']}", {"activo": False}).json()
    assert not next(o for o in d["opciones"] if o["codigo"] == "GAS")["activo"]
    d = interno.post(f"/aranceles/atributos/{aid}/ambitos", {"tipo_ambito": "CHAPTER", "codigo_ambito": "28", "modo": "REQUIRE", "prioridad": 950}).json()
    amb = next(x for x in d["ambitos"] if x["tipo_ambito"] == "CHAPTER")
    assert amb["modo"] == "REQUIRE"
    r = interno.post(f"/aranceles/atributos/{aid}/ambitos", {"tipo_ambito": "CHAPTER", "codigo_ambito": "28"})
    assert r.status_code == 422
    d = interno.patch(f"/aranceles/atributos/{aid}/ambitos/{amb['id']}", {"quitar": True}).json()
    assert not any(x["tipo_ambito"] == "CHAPTER" for x in d["ambitos"])
    # Etiqueta editada y atributo apagado llegan a la ficha
    interno.patch(f"/aranceles/atributos/{por['polo']['id']}", {"etiqueta": "Polo collar", "activo": False})
    ctx = interno.get("/clasificacion/contexto").json()["atributos"]["motor"]["polo"]
    assert ctx["etiqueta"] == "Polo collar" and ctx["activo"] is False
    interno.patch(f"/aranceles/atributos/{por['polo']['id']}", {"etiqueta": "Has a collar and a buttoned placket at the neck (polo style)", "activo": True})
    # Nuevo atributo del usuario
    d = interno.post("/aranceles/atributos", {"codigo": "flash_point", "etiqueta": "Flash point", "tipo_dato": "number", "unidad": "°C", "dominio": "CHEMICALS"}).json()
    assert d["origen"] == "USUARIO" and d["unidad"] == "°C"
    assert interno.post("/aranceles/atributos", {"codigo": "flash_point", "etiqueta": "x"}).status_code == 422
    # Recargar el paquete no duplica
    paquete = (RAIZ / "app/data/oficial/02_carga_motor_dinamico_v3.xlsx").read_bytes()
    r = interno.c.post("/api/aranceles/oficial/importar", headers=interno.h,
                       files={"archivo": ("02.xlsx", io.BytesIO(paquete), "application/octet-stream")}).json()
    assert r["hojas"]["Attributes"]["creados"] == 0 and r["hojas"]["Attribute_Options"]["creados"] == 0 and not r["errores"]


def test_reglas_del_sistema_y_seleccion_nacional(interno):
    r = interno.get("/aranceles/reglas", params={"tipo": "HARD_CONSTRAINT"}).json()
    sis = {x["codigo"]: x for x in r["items"]}
    assert {"R-SYS-001", "R-SYS-006", "R-SYS-009"} <= set(sis)
    assert [(c["campo"], c["valor"]) for c in sis["R-SYS-001"]["condiciones"]] == [("chapter.active", True), ("chapter.classification_enabled", True)]
    assert r["por_tipo"]["NATIONAL_SELECT"] > 50 and r["por_tipo"]["REVIEW_GATE"] == 1
    # Las condiciones de los códigos nacionales son reglas NATIONAL_SELECT
    nac = interno.get("/aranceles/reglas", params={"tipo": "NATIONAL_SELECT", "pais": "SV", "q": "genero"}).json()["items"]
    assert nac and all(x["pais"] == "SV" and x["inciso"]["codigo"].startswith(x["codigo_ambito"]) for x in nac)
    assert any(c["campo"] == "genero" for c in nac[0]["condiciones"]) and nac[0]["cond_txt"]
    # Búsqueda por código nacional
    cod = nac[0]["inciso"]["codigo"]
    assert any(x["id"] == nac[0]["id"] for x in interno.get("/aranceles/reglas", params={"q": cod[:8]}).json()["items"])


def test_editar_regla_nacional_llega_al_motor(interno):
    # Un código con condiciones creado desde Aranceles queda como regla
    r = interno.post("/aranceles/codigos", {"pais": "SV", "codigo": "6404.19.90.99", "descripcion": "Prueba regla",
                                            "cond": {"genero": "F", "cifMax": 15}, "prio": 3})
    assert r.status_code == 200, r.text
    iid = r.json()["id"]
    regla = next(x for x in interno.get("/aranceles/reglas", params={"q": "6404199099"}).json()["items"] if x["inciso"]["id"] == iid)
    assert regla["prioridad"] == 3 and {(c["campo"], c["operador"], c["valor"]) for c in regla["condiciones"]} == {
        ("genero", "EQUAL", "F"), ("valorCIF", "LTE", 15)}
    # Editar las condiciones de la regla cambia lo que recibe la ficha
    r = interno.patch(f"/aranceles/reglas/{regla['id']}", {"condiciones": [
        {"campo": "genero", "operador": "IN", "valor": ["F", "U"]}, {"campo": "valorCIF", "operador": "GT", "valor": 10}]})
    assert r.status_code == 200, r.text
    ctx = next(x for x in interno.get("/clasificacion/contexto").json()["incisos"] if x["id"] == iid)
    assert ctx["cond"] == {"genero": ["F", "U"], "cifMin": 10} and ctx["prio"] == 3
    # Validaciones: en selección nacional solo los operadores del motor y un solo grupo
    assert interno.patch(f"/aranceles/reglas/{regla['id']}", {"condiciones": [{"campo": "genero", "operador": "NOT_EQUAL", "valor": "M"}]}).status_code == 422
    assert interno.patch(f"/aranceles/reglas/{regla['id']}", {"condiciones": [{"campo": "genero", "operador": "IN", "valor": "M"}]}).status_code == 422
    assert interno.patch(f"/aranceles/reglas/{regla['id']}", {"condiciones": [
        {"grupo": 1, "campo": "genero", "valor": "M"}, {"grupo": 2, "campo": "genero", "valor": "F"}]}).status_code == 422
    # Apagar la regla apaga su código
    interno.patch(f"/aranceles/reglas/{regla['id']}", {"activo": False})
    assert not any(x["id"] == iid for x in interno.get("/clasificacion/contexto").json()["incisos"])
    # Quitar condiciones y prioridad desde el código deja el código sin regla
    interno.put(f"/aranceles/codigos/{iid}", {"pais": "SV", "codigo": "6404.19.90.99", "cond": {}, "prio": 0, "activo": True})
    assert not [x for x in interno.get("/aranceles/reglas", params={"q": "6404199099"}).json()["items"] if x.get("inciso", {}).get("id") == iid]


def _libro(hojas: dict) -> io.BytesIO:
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(nombre)
        for f in filas:
            ws.append(f)
    b = io.BytesIO()
    wb.save(b)
    return io.BytesIO(b.getvalue())


def _cargar(api, hojas):
    r = api.c.post("/api/aranceles/oficial/importar", headers=api.h,
                   files={"archivo": ("p.xlsx", _libro(hojas), "application/octet-stream")})
    assert r.status_code == 200, r.text
    return r.json()


def test_paquete_nacional_codigos_regulaciones_impuestos(interno):
    # El paquete 03 incluido trae el mapa de fuentes por país
    paises = {p["iso"]: p for p in interno.get("/aranceles/paises").json()}
    assert paises["CR"]["modelo_arancel"]
    enc_cod = ["National code ID", "Country", "Version", "Base HS6", "Base SAC code", "Precision/additional code", "Full/display code",
               "Parent node ID", "Official description", "DAI %", "Active", "Status", "Valid from", "Valid to", "Source ID", "Source URL", "Internal note"]
    enc_reg = ["Regulation ID", "Country", "Version", "Scope type", "Code/pattern", "Regulation type", "Requirement name", "Authority",
               "Permit/license code", "Mandatory", "Condition JSON", "Legal basis", "Active", "Valid from", "Valid to", "Source ID", "Source URL", "Notes"]
    enc_imp = ["Tax rule ID", "Country", "Version", "Code/pattern", "Tax type", "Rate %", "Basis", "Threshold from", "Threshold to",
               "Formula / rule", "Active", "Valid from", "Valid to", "Source ID", "Source URL", "Legal basis / notes"]
    r = _cargar(interno, {
        "Versions": [["Version ID", "Dataset", "Version label", "Status", "Valid from", "Valid to", "Source ID"],
                     ["CR-2026", "Costa Rica national tariff", "2026", "Published", "2026-01-01", None, "SRC-CR-ATENA"]],
        "National_Codes": [enc_cod,
                           ["CR-1", "CR", "CR-2026", "330499", "33049900", None, "3304.99.00.00.10", None, "Cremas de belleza", "14", "Yes", "Published", "2026-01-01", None, "SRC-CR-ATENA", None, None],
                           ["CR-2", "CR", "CR-2026", None, "33049900", None, "3304.99.00.00.10", None, "Duplicado", "14", "Yes", None, None, None, "SRC-CR-ATENA", None, None],
                           ["CR-3", "CR", "CR-2026", None, "33049900", None, "3305.10.00.00.00", None, "Otra base", "x", "Yes", None, None, None, None, None, None],
                           ["CR-4", "CR", "NO-EXISTE", None, None, None, "3304.99.00.00.20", None, "Sin versión", None, None, None, None, None, None, None, None],
                           ["CR-5", "CR", "CR-2026", None, None, None, "3304990", None, "Corto", None, None, None, None, None, None, None, None],
                           ["CR-6", "CR", "CR-2026", None, None, None, "9999.99.00.00", None, "Sin padre", None, None, None, None, None, None, None, None]],
        "Regulations": [enc_reg,
                        ["REG-CR-1", "CR", None, "HEADING", "3304*", "SANITARY", "Registro sanitario de cosméticos", "Ministerio de Salud", "RS",
                         "Yes", '{"uso": "cosmetico"}', "Reglamento de cosméticos", "Yes", None, None, "SRC-CR-ATENA", None, None],
                        ["REG-CR-2", "CR", None, None, "", "PERMIT", "Sin patrón", None, None, None, None, None, None, None, None, None, None, None],
                        ["REG-CR-3", "CR", None, None, "33", "PERMIT", "JSON malo", None, None, None, "{malo", None, None, None, None, None, None, None]],
        "Taxes": [enc_imp,
                  ["TAX-CR-SEL", "CR", None, "3304", "SELECTIVO", "10", "CIF + DAI", None, None, None, "Yes", None, None, None, None, "Ley de impuesto selectivo de consumo"],
                  ["TAX-CR-X", "CR", None, "3304", "IVA", "13", None, None, None, None, None, None, None, None, None, None]],
    })
    assert r["hojas"]["National_Codes"]["creados"] == 1 and r["hojas"]["Regulations"]["creados"] == 1 and r["hojas"]["Taxes"]["creados"] == 1
    msgs = " | ".join(e["mensaje"] for e in r["errores"])
    for esperado in ("Duplicate country + version + code", "does not start with its base code", "Version NO-EXISTE does not exist",
                     "8 to 14 digits", "does not exist in the tariff tree", "needs a code or pattern", "not valid JSON", "official source or its legal basis"):
        assert esperado in msgs, esperado
    # El código oficial queda con versión, fuente, vigencia y código base
    cod = interno.get("/aranceles/codigos", params={"pais": "CR", "q": "330499000010"}).json()["items"][0]
    assert cod["dai"] == "14" and cod["descripcion"] == "Cremas de belleza"
    # Requisitos de ese código en Costa Rica: IVA general + selectivo + registro sanitario
    req = interno.get("/aranceles/requisitos", params={"pais": "CR", "codigo": "3304.99.00.00.10"}).json()
    tipos = {i["tipo"]: i for i in req["impuestos"]}
    assert tipos["IVA"]["tasa"] == 13 and tipos["SELECTIVO"]["tasa"] == 10
    assert [x["codigo"] for x in req["regulaciones"]] == ["REG-CR-1"] and req["regulaciones"][0]["condicion"] == {"uso": "cosmetico"}
    # Otro capítulo no lleva el registro sanitario
    assert not interno.get("/aranceles/requisitos", params={"pais": "CR", "codigo": "6404.19"}).json()["regulaciones"]
    # El árbol muestra impuestos y regulaciones por país
    sub = next(i for i in interno.get("/aranceles/arbol", params={"q": "3304.99"}).json()["items"] if i["codigo"] == "3304.99")
    cr = next(p for p in interno.get(f"/aranceles/arbol/{sub['id']}").json()["paises"] if p["iso"] == "CR")
    assert cr["regulaciones"] and {i["tipo"] for i in cr["impuestos"]} >= {"IVA", "SELECTIVO"}


def test_regulaciones_e_impuestos_crud(interno):
    r = interno.post("/aranceles/regulaciones", {"pais": "GT", "patron": "6404", "tipo": "labeling", "nombre": "Etiquetado de calzado",
                                                 "autoridad": "DIACO"})
    assert r.status_code == 200, r.text
    reg = r.json()
    assert reg["tipo"] == "LABELING" and reg["patron"] == "6404" and reg["codigo"].startswith("REG-GT-")
    assert interno.post("/aranceles/regulaciones", {"pais": "GT", "patron": "6", "tipo": "PERMIT", "nombre": "x"}).status_code == 422
    # Nunca se borra: se desactiva
    reg = interno.patch(f"/aranceles/regulaciones/{reg['id']}", {"activo": False}).json()
    assert not interno.get("/aranceles/requisitos", params={"pais": "GT", "codigo": "6404199000"}).json()["regulaciones"]
    assert any(x["id"] == reg["id"] for x in interno.get("/aranceles/regulaciones", params={"pais": "GT", "q": "etiquetado"}).json()["items"])
    imp = interno.post("/aranceles/impuestos", {"pais": "GT", "patron": "*", "tipo": "OTRO", "tasa": 1, "base_legal": "Prueba"}).json()
    assert imp["patron"] == "*" and imp["fuente"]
    # Impuesto general sembrado (demo): IVA de Guatemala
    assert any(x["tipo"] == "IVA" and x["tasa"] == 12 for x in interno.get("/aranceles/impuestos", params={"pais": "GT"}).json()["items"])


def test_longitud_de_codigo_configurable(interno):
    # Sin esquema configurado: 8 a 14 dígitos (no un número fijo)
    r = interno.post("/aranceles/codigos", {"pais": "SV", "codigo": "6404.19.90.00.01"})
    assert r.status_code == 200, r.text
    # Con esquema: solo las longitudes declaradas
    _cargar(interno, {"Countries": [["ISO", "Country", "National code length"], ["SV", "El Salvador", "10"]]})
    r = interno.post("/aranceles/codigos", {"pais": "SV", "codigo": "6404.19.90.00.02"})
    assert r.status_code == 422 and "10 digits" in r.json()["mensaje"]
    _cargar(interno, {"Countries": [["ISO", "Country", "National code length"], ["SV", "El Salvador", "CONFIGURABLE"]]})


def test_ruta_generica_quimicos_y_materias_primas(interno):
    r = interno.post("/clasificacion/generico", {"texto": "ácido acético glacial", "dominio": "CHEMICALS",
                                                 "respuestas": {"substance_or_mixture": "SUBSTANCE", "physical_state": "LIQUID"}})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["candidatos"][0]["codigo"] == "291521", [c["codigo"] for c in d["candidatos"]]
    assert d["revision"] and d["candidatos"][0]["incisos"] and d["candidatos"][0]["dominio"]
    # Pregunta lo obligatorio y lo del dominio; marca lo ya respondido
    pregs = {p["codigo"]: p for p in d["preguntas"]}
    assert pregs["product_name"]["modo"] == "REQUIRE" and pregs["physical_state"]["respondida"]
    assert "chemical_name" in pregs and "upper_material" not in pregs
    # Materias primas: el dominio ordena, no obliga (otro capítulo habilitado también puede salir)
    r = interno.post("/clasificacion/generico", {"texto": "hilados de algodón crudo", "dominio": "RAW_MATERIALS"}).json()
    assert "52" in [c["capitulo"] for c in r["candidatos"][:3]]
    # Capítulos no habilitados nunca son candidatos automáticos
    caps = {c["capitulo"]: c for c in interno.get("/aranceles/oficial/capitulos").json()["items"]}
    assert all(caps[c["capitulo"]]["clasificacion"] for c in r["candidatos"])
    assert interno.get("/clasificacion/contexto").json()["dominios_genericos"][0]["codigo"] == "CHEMICALS"


def test_carga_por_etapas_previa_diferencias_publicar(interno):
    def subir(hojas):
        r = interno.c.post("/api/aranceles/oficial/previa", headers=interno.h,
                           files={"archivo": ("p.xlsx", _libro(hojas), "application/octet-stream")})
        assert r.status_code == 200, r.text
        return r.json()

    dom = [["Domain code", "Label", "Description", "Active", "Default mode"]]
    # Sin cambios: la previa no encuentra diferencias
    lote = subir({"Domains": dom + [["CHEMICALS", "Chemicals", "Chemical substances, mixtures and chemical preparations; dynamic questions based on remaining candidates.", "Yes", "AUTO"]]})
    assert lote["estado"] == "PREVIA" and not [f for f in lote["filas"] if f["accion"] != "NUEVO"]
    # Un cambio y un nuevo: se ven antes/después y no se aplican hasta publicar
    lote = subir({"Domains": dom + [["CHEMICALS", "Química", None, "Yes", "AUTO"], ["PLASTICS", "Plastics", "Plastic articles", "Yes", "MANUAL"]],
})
    cambio = next(f for f in lote["filas"] if f["clave"] == "CHEMICALS")
    assert cambio["accion"] == "CAMBIO" and cambio["antes"]["nombre"] == "Chemicals" and cambio["despues"]["nombre"] == "Química"
    assert any(f["accion"] == "NUEVO" and f["clave"] == "PLASTICS" for f in lote["filas"])
    assert lote["resumen"]["tablas"]["Domains"]["NUEVO"] == 1
    doms = {d["codigo"]: d for d in interno.get("/aranceles/oficial/dominios").json()}
    assert doms["CHEMICALS"]["nombre"] == "Chemicals" and "PLASTICS" not in doms
    # Errores de validación también se ven en la previa
    malo = subir({"Domain_Chapter_Map": [["Domain", "Chapter", "Relevance"], ["NOPE", "39", "PRIMARY"]]})
    assert malo["errores"] and not malo["filas"]
    assert interno.post(f"/aranceles/oficial/lotes/{malo['id']}/descartar").json()["estado"] == "DESCARTADA"
    assert interno.post(f"/aranceles/oficial/lotes/{malo['id']}/publicar").status_code == 422
    # Publicar aplica la previa
    r = interno.post(f"/aranceles/oficial/lotes/{lote['id']}/publicar")
    assert r.status_code == 200 and r.json()["estado"] == "PUBLICADA", r.text
    doms = {d["codigo"]: d for d in interno.get("/aranceles/oficial/dominios").json()}
    assert doms["CHEMICALS"]["nombre"] == "Química" and "PLASTICS" in doms
    # Una versión publicada es inmutable: cambiarla no se publica
    inm = subir({"Versions": [["Version ID", "Dataset", "Version label", "Status", "Valid from", "Valid to", "Source ID"],
                              ["SAC-2025-V6", "SAC", "Cambiada", "Published", "2025-08-01", None, "SRC-SIECA-ACI"]]})
    assert inm["resumen"]["bloqueos"] == 1 and inm["filas"][0]["advertencia"]
    r = interno.post(f"/aranceles/oficial/lotes/{inm['id']}/publicar")
    assert r.status_code == 422 and r.json()["codigo"] == "version_publicada", r.text
    # Cerrar su vigencia sí se permite
    v = next(x for x in interno.get("/aranceles/oficial/fuentes").json()["versiones"] if x["codigo"] == "SAC-2025-V6")
    cierre = subir({"Versions": [["Version ID", "Dataset", "Version label", "Status", "Valid from", "Valid to", "Source ID", "Notes"],
                                 ["SAC-2025-V6", v["dataset"], v["etiqueta"], "Published", "2025-08-01", "2030-12-31", "SRC-SIECA-ACI", v["nota"]]]})
    assert cierre["resumen"]["bloqueos"] == 0, cierre["filas"]
    # Lo vigente cambió desde una previa vieja: hay que volver a revisar
    viejo = subir({"Domains": dom + [["CHEMICALS", "Chemicals", None, "Yes", "AUTO"]]})
    interno.c.post("/api/aranceles/oficial/importar", headers=interno.h,
                   files={"archivo": ("x.xlsx", _libro({"Domains": dom + [["CHEMICALS", "Chemicals", None, "Yes", "AUTO"]]}), "application/octet-stream")})
    assert interno.post(f"/aranceles/oficial/lotes/{viejo['id']}/publicar").status_code == 409


def test_aprobar_respeta_control_de_capitulos(interno):
    """R-SYS-001: un código de un capítulo no habilitado no se aprueba."""
    caps = {c["capitulo"]: c for c in interno.get("/aranceles/oficial/capitulos").json()["items"]}
    assert not caps["01"]["clasificacion"]
    ctx = interno.get("/clasificacion/contexto").json()
    assert {c["capitulo"]: c["habilitado"] for c in ctx["capitulos"]}["64"] is True
    p = next(x for x in interno.get("/productos", params={"size": 50}).json()["items"] if x["estado"] not in ("aprobado", "corregido"))
    det = interno.get(f"/productos/{p['id']}").json()
    r = interno.post(f"/productos/{p['id']}/aprobar", {"version": det["version"], "codigo": "0101210000", "forzar": True})
    assert r.status_code == 422 and r.json()["codigo"] == "capitulo_no_habilitado", r.text


def test_producto_guarda_hs6_y_cada_pais_su_linea_con_evidencia(interno):
    """El producto guarda el HS6; la línea SAC va aparte y solo si existe en el
    árbol oficial; un código nacional no se acepta como código del producto;
    cada país guarda su línea oficial, versión, regla, impuestos y regulaciones."""
    ps = interno.get("/productos", params={"size": 50}).json()["items"]
    p = next(x for x in ps if x["estado"] not in ("aprobado", "corregido") and x["tipo"] == "calzado")
    det = interno.get(f"/productos/{p['id']}").json()
    # Un código nacional de 12 dígitos no es una línea SAC
    r = interno.post(f"/productos/{p['id']}/aprobar", {"version": det["version"], "codigo": "640419900090", "forzar": True})
    assert r.status_code == 422 and r.json()["codigo"] == "no_es_linea_sac", r.text
    partidas = {"GT": {"codigo": "6404199000", "estado": "ok", "dai": "15"},
                "PA": {"codigo": "640419970000", "estado": "ok", "manual": True, "sugerido": "640419910000", "motivo": "Revisado con la nota 4"}}
    r = interno.post(f"/productos/{p['id']}/aprobar", {"version": det["version"], "codigo": "6404199000", "partidas": partidas, "forzar": True})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["codigo"] == "6404.19" and d["sac_codigo"] == "6404.19.90.00"
    assert d["evidencia"]["version_arancel"] == "SAC-2025-V6" and "R-SYS-001" in d["evidencia"]["reglas_sistema"]
    gt, pa = d["partidas"]["GT"], d["partidas"]["PA"]
    assert gt["inciso_id"] and gt["evidencia"]["linea_oficial"] and any(i["tipo"] == "IVA" for i in gt["evidencia"]["impuestos"])
    assert gt["sugerido"] == "6404199000" and gt["aprobado_en"]
    assert pa["codigo"] == "640419970000" and pa["sugerido"] == "640419910000" and pa["motivo"] == "Revisado con la nota 4"


def test_configuracion_custom_cambia_la_clasificacion(interno):
    """Criterio de aceptación: atributo custom → ámbito → regla → clasificar →
    aparece la pregunta → se responde → cambian los candidatos y el HS6."""
    base = {"texto": "calzado con suela y parte superior de caucho", "dominio": "FOOTWEAR"}
    antes = interno.post("/clasificacion/sesion", base).json()
    assert antes["hs6"] and antes["hs6"] != "640192"
    # 1. Atributo custom con opciones y ámbito obligatorio en el dominio
    a = interno.post("/aranceles/atributos", {"codigo": "waterproof_level", "etiqueta": "Waterproof level", "tipo_dato": "select", "dominio": "FOOTWEAR"}).json()
    for cod in ("NONE", "WATER_RESISTANT", "WATERPROOF"):
        interno.post(f"/aranceles/atributos/{a['id']}/opciones", {"codigo": cod, "etiqueta": cod.title()})
    interno.post(f"/aranceles/atributos/{a['id']}/ambitos", {"tipo_ambito": "DOMAIN", "codigo_ambito": "FOOTWEAR", "modo": "REQUIRE", "prioridad": 990})
    # Dependencia: otro atributo solo se pregunta si es impermeable
    b = interno.post("/aranceles/atributos", {"codigo": "sealed_seams", "etiqueta": "Sealed seams", "tipo_dato": "boolean", "dominio": "FOOTWEAR"}).json()
    r = interno.post(f"/aranceles/atributos/{b['id']}/ambitos", {"tipo_ambito": "DOMAIN", "codigo_ambito": "FOOTWEAR", "modo": "SHOW",
                                                                 "condicion": [{"campo": "waterproof_level", "operador": "EQUAL", "valor": "WATERPROOF"}]})
    assert r.status_code == 200, r.text
    # 2. Reglas custom: impermeable → solo 6401.92; no impermeable → nunca 6401
    r1 = interno.post("/aranceles/reglas", {"tipo_ambito": "DOMAIN", "codigo_ambito": "FOOTWEAR", "tipo_regla": "HARD_CONSTRAINT", "prioridad": 950,
                                             "efecto": "Waterproof footwear goes to 64.01",
                                             "condiciones": [{"campo": "waterproof_level", "operador": "EQUAL", "valor": "WATERPROOF"}],
                                             "accion": {"tipo": "RESTRICT", "codigos": ["640192"]}})
    assert r1.status_code == 200, r1.text
    r1 = r1.json()
    assert r1["codigo"].startswith("R-USR-") and r1["accion"]["codigos"] == ["640192"]
    r2 = interno.post("/aranceles/reglas", {"tipo_ambito": "DOMAIN", "codigo_ambito": "FOOTWEAR", "tipo_regla": "HARD_CONSTRAINT", "prioridad": 940,
                                             "condiciones": [{"campo": "waterproof_level", "operador": "IN", "valor": ["NONE", "WATER_RESISTANT"]}],
                                             "accion": {"tipo": "EXCLUDE", "codigos": ["6401"]}}).json()
    # 3. Clasificar: la pregunta aparece primero (obligatoria y discriminante); la dependiente no
    s = interno.post("/clasificacion/sesion", base).json()
    pregs = {p["codigo"]: p for p in s["preguntas"]}
    assert s["preguntas"][0]["codigo"] == "waterproof_level" and pregs["waterproof_level"]["discrimina"] and pregs["waterproof_level"]["modo"] == "REQUIRE"
    assert "sealed_seams" not in pregs
    assert any(t["regla"] == r1["codigo"] and t["resultado"] is None for t in s["reglas"])
    # 4. Respondo impermeable: cambia el HS6, sube la confianza y aparece la dependiente
    s = interno.post("/clasificacion/sesion", {**base, "respuestas": {"waterproof_level": "WATERPROOF"}}).json()
    assert s["hs6"] == "640192" and s["confianza"] == "high" and [c["codigo"] for c in s["candidatos"]] == ["640192"]
    assert "sealed_seams" in {p["codigo"] for p in s["preguntas"]}
    assert s["paises"] and all(p["pais"] for p in s["paises"])
    # 5. Respondo no impermeable: 64.01 queda fuera
    s = interno.post("/clasificacion/sesion", {**base, "respuestas": {"waterproof_level": "NONE"}}).json()
    assert s["hs6"] and not any(c["codigo"].startswith("6401") for c in s["candidatos"])
    # 6. Apagar la regla devuelve el comportamiento anterior
    interno.patch(f"/aranceles/reglas/{r1['id']}", {"activo": False})
    s = interno.post("/clasificacion/sesion", {**base, "respuestas": {"waterproof_level": "WATERPROOF"}}).json()
    assert s["hs6"] == antes["hs6"]
    # 7. Las reglas del sistema gobiernan: sin R-SYS-009 el dominio pasa a filtrar capítulos
    sis = next(x for x in interno.get("/aranceles/reglas", params={"q": "R-SYS-009"}).json()["items"] if x["codigo"] == "R-SYS-009")
    caps_dom = {c["capitulo"] for d in interno.get("/aranceles/oficial/dominios").json() if d["codigo"] == "FOOTWEAR" for c in d["capitulos"]}
    con = interno.post("/clasificacion/sesion", {"texto": "bolsa de papel", "dominio": "FOOTWEAR"}).json()
    assert any(c["capitulo"] not in caps_dom for c in con["candidatos"])  # con R-SYS-009 el dominio solo ordena
    interno.patch(f"/aranceles/reglas/{sis['id']}", {"activo": False})
    s = interno.post("/clasificacion/sesion", {"texto": "bolsa de papel", "dominio": "FOOTWEAR"}).json()
    assert s["candidatos"] == [] or all(c["capitulo"] in caps_dom for c in s["candidatos"])
    interno.patch(f"/aranceles/reglas/{sis['id']}", {"activo": True})
    interno.patch(f"/aranceles/reglas/{r2['id']}", {"activo": False})
    # Validaciones de reglas propias
    assert interno.post("/aranceles/reglas", {"tipo_regla": "SOFT_SIGNAL", "accion": {"tipo": "EXCLUDE", "codigos": ["64"]}}).status_code == 422
    assert interno.post("/aranceles/reglas", {"tipo_regla": "QUESTION_GATE", "accion": {"tipo": "ASK"}}).status_code == 422


def test_dominio_nuevo_solo_con_configuracion(interno):
    """Un dominio nuevo (ELECTRONICS) con su categoría, atributo, ámbito, regla y
    capítulos aparece en la ficha y clasifica sin programar nada."""
    d = interno.post("/aranceles/oficial/dominios", {"codigo": "electronics", "nombre": "Electronics", "modo": "AUTO"})
    assert d.status_code == 200, d.text
    d = d.json()
    assert d["codigo"] == "ELECTRONICS"
    assert interno.post("/aranceles/oficial/dominios", {"codigo": "ELECTRONICS", "nombre": "x"}).status_code == 422
    interno.put(f"/aranceles/oficial/dominios/{d['id']}/capitulos/85", {"relevancia": "PRIMARY", "habilitado": True})
    caps = {c["capitulo"]: c for c in interno.get("/aranceles/oficial/capitulos").json()["items"]}
    interno.patch("/aranceles/oficial/capitulos", {"ids": [caps["85"]["id"]], "clasificacion": True, "candidato_auto": True})
    c = interno.post("/aranceles/categorias", {"nombre": "Batteries and power banks", "dominio": "ELECTRONICS", "grupo": "Electronics"}).json()
    assert c["codigo"] == "batteries_and_power_banks" and not c["ficha_motor"]
    assert any(x["codigo"] == c["codigo"] for x in interno.get("/clasificacion/contexto").json()["categorias"])
    a = interno.post("/aranceles/atributos", {"codigo": "battery_chemistry", "etiqueta": "Battery chemistry", "tipo_dato": "select", "dominio": "ELECTRONICS"}).json()
    for o in ("LITHIUM_ION", "LEAD_ACID", "NICKEL"):
        interno.post(f"/aranceles/atributos/{a['id']}/opciones", {"codigo": o, "etiqueta": o.replace("_", " ").title()})
    interno.post(f"/aranceles/atributos/{a['id']}/ambitos", {"tipo_ambito": "CATEGORY", "codigo_ambito": c["codigo"], "modo": "REQUIRE", "prioridad": 900})
    interno.post("/aranceles/reglas", {"tipo_ambito": "CATEGORY", "codigo_ambito": c["codigo"], "tipo_regla": "HARD_CONSTRAINT",
                                        "condiciones": [{"campo": "battery_chemistry", "operador": "EQUAL", "valor": "LITHIUM_ION"}],
                                        "accion": {"tipo": "RESTRICT", "codigos": ["850760"]}})
    base = {"texto": "batería recargable", "dominio": "ELECTRONICS", "categoria": c["codigo"]}
    s = interno.post("/clasificacion/sesion", base).json()
    assert s["preguntas"][0]["codigo"] == "battery_chemistry"
    s = interno.post("/clasificacion/sesion", {**base, "respuestas": {"battery_chemistry": "LITHIUM_ION"}}).json()
    assert s["hs6"] == "850760" and s["confianza"] == "high"
    # Archivar el dominio y la categoría: dejan de ofrecerse
    interno.patch(f"/aranceles/oficial/dominios/{d['id']}", {"activo": False})
    interno.patch(f"/aranceles/categorias/{c['id']}", {"activo": False})
    ctx = interno.get("/clasificacion/contexto").json()
    assert not any(x["codigo"] == "ELECTRONICS" for x in ctx["dominios_genericos"])
    assert not any(x["codigo"] == c["codigo"] for x in ctx["categorias"])
    assert any(x["ficha_motor"] and x["codigo"] == "calzado" for x in ctx["categorias"])
