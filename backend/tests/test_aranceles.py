"""Aranceles: países con sus dígitos, subpartidas SAC y códigos nacionales
editables, cargas desde Excel y exportación con filtros."""
import io

from openpyxl import Workbook, load_workbook


def _xlsx(filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for f in filas:
        ws.append(f)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _version_borrador(api, codigo, ambito):
    """Una versión nacional en borrador (las publicadas no cambian)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Versions"
    ws.append(["Version ID", "Dataset", "Version label", "Status", "Valid from", "Valid to", "Source ID", "Scope"])
    ws.append([codigo, f"{ambito} national tariff (draft)", "Draft", "Draft", "2026-01-01", None, "SRC-SIECA-ACI", ambito])
    b = io.BytesIO()
    wb.save(b)
    r = api.c.post("/api/aranceles/oficial/importar", headers=api.h, files={"archivo": ("v.xlsx", io.BytesIO(b.getvalue()), "application/octet-stream")})
    assert r.status_code == 200, r.text


def _subir(api, url, contenido, **params):
    return api.c.post(f"/api{url}", headers=api.h, params=params,
                      files={"archivo": ("datos.xlsx", io.BytesIO(contenido), "application/octet-stream")})


def test_paises_sac_y_codigos(interno, vans):
    ps = interno.get("/aranceles/paises").json()
    assert {p["iso"] for p in ps} >= {"GT", "SV", "HN", "NI", "CR", "PA"}
    # Solo hay líneas nacionales donde una fuente oficial las publicó: los países que aplican
    # el SAC regional a 10 dígitos (ACI de SIECA); NI, CR y PA esperan su arancel nacional
    por = {p["iso"]: p for p in ps}
    assert por["PA"]["digitos"] == 12 and all(por[i]["codigos"] > 0 for i in ("GT", "SV", "HN"))
    assert all(por[i]["codigos"] == 0 for i in ("NI", "PA"))  # CR puede tener las de su arancel oficial cargado en otra prueba
    # Un país nuevo con sus propios dígitos; el proveedor no puede
    assert vans.post("/aranceles/paises", {"iso": "DO", "nombre": "Dominican Republic", "digitos": 8}).status_code == 403
    r = interno.post("/aranceles/paises", {"iso": "DO", "nombre": "Dominican Republic", "digitos": 8, "impuesto": "ITBIS 18%"})
    assert r.status_code == 200, r.text
    ctx = interno.get("/clasificacion/contexto").json()
    assert any(d["iso"] == "DO" and d["digitos"] == 8 for d in ctx["destinos"])

    # SAC: filtros por capítulo y búsqueda; editar el texto de una subpartida
    sac = interno.get("/aranceles/sac", params={"capitulo": "64", "nivel": "6"}).json()
    assert sac["total"] > 5 and all(x["codigo"].startswith("64") for x in sac["items"])
    x = next(x for x in sac["items"] if x["codigo"] == "640419")
    assert x["nacionales"] >= 3
    # Lo propio es un override con motivo: el texto oficial no se toca
    assert interno.put(f"/aranceles/sac/{x['id']}", {"codigo": "640419", "descripcion": "Los demás (texto corregido)"}).status_code == 422
    assert interno.put(f"/aranceles/sac/{x['id']}", {"codigo": "640419", "descripcion": "Los demás (texto corregido)",
                                                     "motivo": "Descripción interna de compras"}).status_code == 200
    s = interno.post("/clasificacion/sesion", {"categoria": "calzado", "ficha": {"estilo_calzado": "tenis", "comp": {"corte": "100% canvas", "suela": "100% rubber"}},
                                               "paises": False}).json()
    assert s["clasificacion"]["hs6"]["descripcion"] == "Los demás (texto corregido)" and s["clasificacion"]["hs6"]["descripcion_oficial"]
    y = next(y for y in interno.get("/aranceles/sac", params={"q": "6404.19", "nivel": "6"}).json()["items"] if y["codigo"] == "640419")
    assert y["custom"] and y["descripcion"] == "Los demás (texto corregido)" and y["descripcion_oficial"] != y["descripcion"]
    # Un código que no está en el arancel oficial no se crea aquí
    assert interno.post("/aranceles/sac", {"codigo": "999999", "descripcion": "Inventado", "motivo": "x"}).status_code == 422
    # Quitar lo custom vuelve al oficial
    interno.delete_(f"/aranceles/sac/{x['id']}")
    y = next(y for y in interno.get("/aranceles/sac", params={"q": "6404.19", "nivel": "6"}).json()["items"] if y["codigo"] == "640419")
    assert not y["custom"] and y["descripcion"] == y["descripcion_oficial"]

    # Códigos nacionales: agregar con condiciones y validar los dígitos del país
    assert interno.post("/aranceles/codigos", {"pais": "DO", "codigo": "6404199"}).status_code == 422
    # Una línea nacional es dato oficial: sin fuente y versión no se crea
    sin = interno.post("/aranceles/codigos", {"pais": "DO", "codigo": "6404.19.00", "dai": "20%"})
    assert sin.status_code == 422 and sin.json()["codigo"] == "sin_procedencia"
    # Tampoco una que no cuelga del árbol oficial
    assert interno.post("/aranceles/codigos", {"pais": "DO", "codigo": "9999.99.00", "fuente": "SRC-SIECA-ACI", "version": "SAC-2025-V6"}).status_code == 422
    # Ni se agregan líneas a una versión publicada: van a una versión en borrador
    r = interno.post("/aranceles/codigos", {"pais": "DO", "codigo": "6404.19.00", "fuente": "SRC-SIECA-ACI", "version": "SAC-2025-V6"})
    assert r.status_code == 422 and r.json()["codigo"] == "version_publicada"
    _version_borrador(interno, "DO-DRAFT", "DO")
    r = interno.post("/aranceles/codigos", {"pais": "DO", "codigo": "6404.19.00", "dai": "20%", "fuente": "SRC-SIECA-ACI", "version": "DO-DRAFT",
                                            "cond": {"genero": "M", "edad": "adulto"}})
    assert r.status_code == 200, r.text
    lista = interno.get("/aranceles/codigos", params={"pais": "DO"}).json()
    assert lista["total"] == 1 and lista["items"][0]["cond_txt"].startswith("For men or for women: Men")
    # Varios países a la vez (filtro con varios valores)
    assert interno.get("/aranceles/codigos", params={"pais": "SV,PA", "capitulo": "64"}).json()["total"] > 0
    # Un país con códigos no se borra: se desactiva y deja de pedirse en las fichas
    do = next(p for p in interno.get("/aranceles/paises").json() if p["iso"] == "DO")
    assert interno.delete_(f"/aranceles/paises/{do['id']}").status_code == 409
    assert interno.put(f"/aranceles/paises/{do['id']}", {**{k: do[k] for k in ("iso", "nombre", "digitos", "mcca")},
                                                         "activo": False}).status_code == 200
    assert not any(d["iso"] == "DO" for d in interno.get("/clasificacion/contexto").json()["destinos"])


def test_cargar_y_exportar(interno):
    plantilla = interno.get("/aranceles/codigos/plantilla", params={"pais": "SV"})
    assert plantilla.status_code == 200
    wb = load_workbook(io.BytesIO(plantilla.content))
    enc = [c.value for c in wb["Data"][1]]
    assert enc[:2] == ["Country *", "Code *"] and "For men or for women" in enc and "Instructions" in wb.sheetnames
    contenido = _xlsx([["Country", "Code", "Description", "Duty (DAI %)", "For men or for women", "Footwear style"],
                       ["SV", "6402.99.10.00", "Para hombre", "15", "Men", "Sneaker or athletic shoe"],
                       ["SV", "6402991", "Mal", "", "", ""],
                       ["XX", "6402.99.10.00", "", "", "", ""]])
    # Una carga de líneas nacionales es la publicación oficial de un país: sin fuente ni versión no entra
    sin = _subir(interno, "/aranceles/codigos/importar", contenido)
    assert sin.status_code == 422 and sin.json()["codigo"] == "sin_procedencia"
    # Una versión publicada no cambia (ni su texto ni su DAI)
    r = _subir(interno, "/aranceles/codigos/importar", contenido, fuente="SRC-SIECA-ACI", version="SAC-2025-V6")
    assert r.status_code == 422 and r.json()["codigo"] == "version_publicada"
    _version_borrador(interno, "SV-DRAFT", "SV")
    r = _subir(interno, "/aranceles/codigos/importar", contenido, fuente="SRC-SIECA-ACI", version="SV-DRAFT")
    assert r.status_code == 200, r.text
    r = r.json()
    assert r["creados"] == 1 and r["actualizados"] == 0 and len(r["errores"]) == 2
    items = interno.get("/aranceles/codigos", params={"pais": "SV", "q": "6402991000"}).json()["items"]
    x = next(i for i in items if i["version"] == "SV-DRAFT")
    publicada = next(i for i in items if i["version"] == "SAC-2025-V6")
    assert x["cond"] == {"genero": "M", "estilo_calzado": "tenis"} and x["fuente"] == "oficial" and x["oficial"]
    # La misma línea en la misma versión se actualiza (no se duplica)
    assert _subir(interno, "/aranceles/codigos/importar", contenido, fuente="SRC-SIECA-ACI", version="SV-DRAFT").json()["actualizados"] == 1
    # La versión en borrador no es la vigente: El Salvador sigue clasificando con la regional
    s = interno.post("/clasificacion/sesion", {"categoria": "calzado", "ficha": {}, "hs6": "640299"}).json()
    assert all(p["version"] != "SV-DRAFT" for p in (s.get("clasificacion") or {}).get("paises") or [])
    for formato in ("xlsx", "pdf"):
        r = interno.get("/aranceles/codigos/exportar", params={"pais": "SV", "capitulo": "64", "formato": formato})
        assert r.status_code == 200 and len(r.content) > 1000
    r = _subir(interno, "/aranceles/sac/importar", _xlsx([["Code", "Description"], ["9999.99", "Prueba"], ["6403.51", "Botas de cuero (interno)"]])).json()
    assert r["actualizados"] == 1 and len(r["errores"]) == 1  # 9999.99 no es oficial: no se inventa
    assert interno.get("/aranceles/sac/exportar", params={"q": "6403", "formato": "pdf"}).status_code == 200
    # Una línea oficial de una versión publicada no se borra; la de un borrador sí
    assert interno.post("/aranceles/codigos/borrar", {"ids": [publicada["id"]]}).status_code == 422
    assert interno.post("/aranceles/codigos/borrar", {"ids": [x["id"]]}).status_code == 200


def test_notas_sac(interno):
    r = interno.get("/aranceles/notas", params={"capitulo": "64"}).json()
    nums = {(n["codigo"], n["numero"]) for n in r["items"]}
    # Las del capítulo 64 (materia de la parte superior y de la suela) y las reglas generales
    assert ("64", "4") in nums and ("RGI", "3") in nums and ("61", "9") not in nums
    assert any(n["codigo"] == "64" for n in interno.get("/clasificacion/contexto").json()["notas_sac"])
    n = interno.post("/aranceles/notas", {"ambito": "capitulo", "codigo": "64", "numero": "X", "texto": "Nota de prueba",
                                          "capitulos": ["64"]})
    assert n.status_code == 200, n.text
    nid = n.json()["id"]
    assert interno.put(f"/aranceles/notas/{nid}", {"ambito": "capitulo", "codigo": "64", "numero": "X",
                                                   "texto": "Nota editada", "capitulos": ["64"], "activo": False}).json()["activo"] is False
    assert interno.post("/aranceles/notas", {"ambito": "otro", "codigo": "64", "texto": "x"}).status_code == 422
    assert interno.delete_(f"/aranceles/notas/{nid}").status_code == 200


def test_notas_sac_excel(interno):
    assert len(interno.get("/aranceles/notas").json()["items"]) >= 100
    v = interno.get("/aranceles/notas/plantilla", params={"vista": 1}).json()
    assert v["hojas"][0]["columnas"][0]["nombre"] == "Kind"
    wb = Workbook()
    ws = wb.active
    ws.append(["Kind", "Section or chapter", "Number", "Text", "Applies to chapters"])
    ws.append(["Chapter note", "64", "4", "Texto oficial actualizado de la nota 4", "64"])
    ws.append(["Chapter note", "", "9", "Sin capítulo", ""])
    b = io.BytesIO()
    wb.save(b)
    r = _subir(interno, "/aranceles/notas/importar", b.getvalue()).json()
    assert r["actualizados"] == 1 and r["creados"] == 0 and len(r["errores"]) == 1, r
    n4 = [n for n in interno.get("/aranceles/notas", params={"capitulo": "64"}).json()["items"] if n["codigo"] == "64" and n["numero"] == "4"]
    # La nota oficial no se pisa: el texto del archivo queda como capa custom
    assert n4[0]["texto"].startswith("Texto oficial") and n4[0]["custom"] and n4[0]["oficial"] and n4[0]["texto_oficial"]
    x = interno.get("/aranceles/notas/exportar", params={"formato": "xlsx", "capitulo": "64"})
    assert x.status_code == 200 and x.content[:2] == b"PK"


def test_condiciones_por_pais_y_subpartida(interno):
    r = interno.get("/aranceles/condiciones", params={"pais": "GT", "codigo": "6404.19"}).json()
    assert "estilo_calzado" in r["aplican"] and "manga" not in r["aplican"] and "cifMax" in r["aplican"]
    assert r["subpartida"]["codigo"] == "6404.19" and all(h["codigo"].startswith("640419") for h in r["hermanos"])
    prendas = interno.get("/aranceles/condiciones", params={"pais": "SV", "codigo": "620342"}).json()
    assert "largo" in prendas["aplican"] and "estilo_calzado" not in prendas["aplican"]
    assert len(interno.get("/aranceles/condiciones", params={"pais": "GT", "codigo": "64"}).json()["aplican"]) > 10


def test_base_oficial_sieca(interno):
    """El SAC oficial (ACI SIECA VII Enmienda): subpartidas, notas y aperturas con DAI."""
    sac = interno.get("/aranceles/sac", params={"q": "6404", "size": 50}).json()["items"]
    assert any(x["codigo"] == "640419" and x["fuente"] in ("oficial", "manual") for x in sac)
    gt = interno.get("/aranceles/codigos", params={"pais": "GT", "q": "640419", "size": 20}).json()["items"]
    cods = {x["codigo"]: x for x in gt}
    assert "6404191000" in cods and cods["6404191000"]["dai"] in ("0", "0.0") and "Cubrecalzado" in cods["6404191000"]["descripcion"]
    notas = interno.get("/aranceles/notas", params={"capitulo": "61"}).json()["items"]
    assert any(n["ambito"] == "capitulo" and n["numero"] == "9" and "izquierda sobre derecha" in n["texto"] for n in notas)
    assert len(interno.get("/aranceles/notas").json()["items"]) > 400


def test_acuerdos_por_origen(interno, vans):
    """Los acuerdos comerciales dicen en qué destinos entra con preferencia un origen."""
    ctx = interno.get("/clasificacion/contexto").json()
    ac = ctx["acuerdos"]
    cubre = lambda o, d: [a["codigo"] for a in ac if o in a["origenes"] and d in a["destinos"]]  # noqa: E731
    assert cubre("CN", "CR") == ["CN-CR"] and cubre("CN", "NI") == ["CN-NI"]
    assert not cubre("CN", "GT") and not cubre("CN", "SV") and not cubre("VN", "PA")
    assert "CAFTA-DR" in cubre("US", "SV") and not cubre("US", "PA")
    assert "UE-CA" in cubre("DE", "PA")
    # Se mantienen en Master data → Trade agreements
    assert vans.post("/catalogos/acuerdos", {"codigo": "VN-XX", "nombre": "x", "origenes": "VN", "destinos": "GT"}).status_code == 403
    assert interno.post("/catalogos/acuerdos", {"codigo": "VN-XX", "nombre": "x", "origenes": "vn", "destinos": "gt, sv"}).status_code in (200, 201)
    assert interno.post("/catalogos/acuerdos", {"codigo": "BAD", "nombre": "x", "origenes": "VNM", "destinos": "GT"}).status_code == 422
    ac = interno.get("/clasificacion/contexto").json()["acuerdos"]
    assert any(a["codigo"] == "VN-XX" and a["destinos"] == ["GT", "SV"] for a in ac)


def test_soporte_de_clasificacion(interno, vans):
    """La ficha se clasifica con el motor y se revisa con su soporte: base legal
    de cada destino y notas legales y explicativas de la partida. No hay
    búsqueda libre en todo el SAC desde la ficha."""
    assert vans.get("/clasificacion/sac", params={"q": "8471.30"}).status_code in (404, 405)
    ctx = interno.get("/clasificacion/contexto").json()
    # La base legal solo existe si el paquete oficial la declara: nunca se arma con el nombre de la fuente
    assert not any(d["base_legal"] and " — " in d["base_legal"] for d in ctx["destinos"])
    # Notas explicativas por partida (resumen propio, no el texto oficial)
    notas = interno.get("/aranceles/notas", params={"capitulo": "64"}).json()
    ne = [n for n in notas["items"] if n["ambito"] == "explicativa"]
    assert any(n["codigo"] == "6404" for n in ne) and "explicativa" in notas["ambitos"]
