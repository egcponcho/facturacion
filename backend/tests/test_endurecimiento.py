"""Seguridad: permisos de lectura, datos ocultos en la configuración que solo
se consulta, archivos subidos, exportaciones, alcance, contraseñas, limpieza
de tablas técnicas y bitácora general."""
import io
import zipfile

from conftest import Api
from openpyxl import load_workbook


def _usuario_con_rol(client, admin, email, permisos):
    rol = admin.post("/roles", {"nombre": f"Rol {email}", "permisos": permisos}).json()
    u = admin.post("/usuarios", {"email": email, "nombre": "Prueba", "rol_id": rol["id"],
                                 "telefono": "+50370000123", "dos_pasos": True}).json()
    api = Api(client, email, u["password_temporal"])
    assert api.post("/auth/password", {"actual": u["password_temporal"], "nueva": "Segura#Clave2026"}).status_code == 200
    return api


def _factura_con_listas(interno) -> tuple[dict, dict]:
    """Una factura activa con al menos una lista de empaque."""
    for f in interno.get("/facturas", params={"size": 200}).json()["items"]:
        if f["estado"] != "CANCELADA":
            pls = [p for p in interno.get(f"/facturas/{f['id']}").json()["packing_lists"] if p["estado"] != "CANCELADO"]
            if pls:
                return f, pls[0]
    raise AssertionError("La demostración no tiene facturas con listas de empaque")


def test_leer_facturas_exige_su_permiso(client, admin, interno):
    factura, pl = _factura_con_listas(interno)
    transporte = _usuario_con_rol(client, admin, "solo.transporte@demo.com", ["transporte.gestionar"])
    assert transporte.get("/facturas").status_code == 403
    assert transporte.get(f"/facturas/{factura['id']}").status_code == 403
    assert transporte.get(f"/facturas/{factura['id']}/exportar").status_code == 403
    assert transporte.get(f"/packing-lists/{pl['id']}").status_code == 403
    assert "factura.ver" not in transporte.get("/auth/me").json()["permisos"]
    # La bodega ve la lista de empaque para recibirla, pero no las facturas
    bodega = _usuario_con_rol(client, admin, "bodega@demo.com", ["recepcion.registrar"])
    assert bodega.get(f"/packing-lists/{pl['id']}").status_code == 200
    assert bodega.get(f"/facturas/{factura['id']}").status_code == 403
    # Las plantillas de caja son de quien empaca
    assert transporte.get("/plantillas").status_code == 403


def test_inicio_sin_permiso_de_alertas(client, admin):
    """Un rol interno sin alertas.ver ve su página de inicio (antes respondía 403)."""
    comprador = _usuario_con_rol(client, admin, "comprador@demo.com", ["oc.ver", "factura.ver", "producto.ver"])
    r = comprador.get("/dashboard")
    assert r.status_code == 200, r.text
    assert r.json()["alertas"] == []


def test_configuracion_del_arancel_oculta_impuestos_a_quien_solo_consulta(interno, vans):
    inciso = interno.get("/aranceles/codigos", params={"size": 1}).json()["items"][0]
    params = {"pais": inciso["pais"], "codigo": inciso["codigo"]}
    completo = interno.get("/aranceles/requisitos", params=params).json()
    filtrado = vans.get("/aranceles/requisitos", params=params).json()
    assert "impuestos" in completo and "dai" in completo
    assert "impuestos" not in filtrado and "dai" not in filtrado


def test_archivos_se_validan_por_su_contenido(interno, vans):
    factura, _ = _factura_con_listas(interno)

    def subir(nombre, contenido, tipo="OTRO"):
        return interno.c.post(f"/api/facturas/{factura['id']}/archivos", headers=interno.h,
                              data={"tipo": tipo}, files={"archivo": (nombre, io.BytesIO(contenido), "application/pdf")})

    assert subir("programa.exe", b"MZ\x90\x00").status_code == 415
    assert subir("falso.pdf", b"<script>alert(1)</script>").status_code == 415
    assert subir("factura.pdf", b"%PDF-1.4 prueba", tipo="CUALQUIERA").status_code == 422
    assert subir("factura.pdf", b"%PDF-1.4 prueba").status_code == 200
    # Una foto que dice ser PNG pero no lo es
    producto = vans.get("/productos", params={"size": 1}).json()["items"][0]
    r = vans.c.post(f"/api/productos/{producto['id']}/fotos", headers=vans.h,
                    files={"archivo": ("foto.png", io.BytesIO(b"<svg onload=alert(1)>"), "image/png")})
    assert r.status_code == 422


def test_excel_con_bomba_de_descompresion(interno, monkeypatch):
    from app.core import archivos

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/worksheets/sheet1.xml", b"0" * 5_000_000)
    monkeypatch.setattr(archivos, "MAX_DESCOMPRIMIDO", 1_000_000)
    r = interno.c.post("/api/ordenes/importar/previa", headers=interno.h,
                       files={"archivo": ("ocs.xlsx", io.BytesIO(buf.getvalue()), "application/octet-stream")})
    assert r.status_code == 413, r.text


def test_exportar_no_ejecuta_formulas(interno, vans):
    from test_edicion import _factura_borrador

    fid = _factura_borrador(vans)
    try:
        detalle = interno.get(f"/facturas/{fid}").json()
        formula = '=HYPERLINK("http://ejemplo.invalid","clic")'
        r = interno.patch(f"/facturas/{fid}", {"version": detalle["version"], "condiciones": formula})
        assert r.status_code == 200, r.text
        wb = load_workbook(io.BytesIO(interno.get(f"/facturas/{fid}/exportar").content))
        celdas = [c for ws in wb.worksheets for fila in ws.iter_rows() for c in fila if c.value == formula]
        assert celdas and all(c.data_type == "s" for c in celdas)
    finally:
        vans.post(f"/facturas/{fid}/cancelar", {"motivo": "prueba de exportación"})


def test_alerta_fuera_del_alcance(admin, interno, client):
    from app.core.db import SessionLocal
    from app.modelos import Alerta, Proveedor

    with SessionLocal() as db:
        tnf = db.query(Proveedor).filter_by(codigo="TNF").one()
        a = Alerta(tipo="PRUEBA", mensaje="Alerta de TNF", proveedor_id=tnf.id)
        db.add(a)
        db.commit()
        alerta_id = a.id
    vans_id = next(p["id"] for p in admin.get("/proveedores").json() if p["codigo"] == "VANS")
    usuario = next(u for u in admin.get("/usuarios").json() if u["email"] == "interno@demo.com")
    admin.patch(f"/usuarios/{usuario['id']}", {"alcance": {"proveedores": [vans_id]}})
    try:
        assert interno.post(f"/alertas/{alerta_id}/resolver").status_code == 404
    finally:
        admin.patch(f"/usuarios/{usuario['id']}", {"alcance": usuario["alcance"]})
    assert interno.post(f"/alertas/{alerta_id}/resolver").status_code == 200


def test_hash_anterior_se_actualiza_al_entrar(client, admin):
    import base64
    import hashlib

    from app.core.db import SessionLocal
    from app.core.seguridad import ITERACIONES
    from app.modelos import Usuario

    u = admin.post("/usuarios", {"email": "antiguo@demo.com", "nombre": "Antiguo", "rol": "interno",
                                 "telefono": "+50370000124", "dos_pasos": True}).json()
    salt = b"0123456789abcdef"
    viejo = "pbkdf2${}${}".format(base64.b64encode(salt).decode(), base64.b64encode(
        hashlib.pbkdf2_hmac("sha256", b"Antigua#Clave2020", salt, 200_000)).decode())
    with SessionLocal() as db:
        x = db.get(Usuario, u["id"])
        x.password_hash, x.clave_temporal = viejo, False
        db.commit()
    Api(client, "antiguo@demo.com", "Antigua#Clave2020")
    with SessionLocal() as db:
        assert db.get(Usuario, u["id"]).password_hash.startswith(f"pbkdf2_sha256${ITERACIONES}$")
    assert Api(client, "antiguo@demo.com", "Antigua#Clave2020").get("/auth/me").status_code == 200


def test_limpieza_de_tablas_tecnicas():
    from datetime import timedelta

    from app.core.db import SessionLocal
    from app.instalacion.mantenimiento import purgar
    from app.modelos import Idempotencia, ahora

    with SessionLocal() as db:
        db.add(Idempotencia(clave="vieja:1", usuario_id=1, respuesta={}, creada_en=ahora() - timedelta(days=3)))
        db.add(Idempotencia(clave="nueva:1", usuario_id=1, respuesta={}, creada_en=ahora()))
        db.commit()
        assert purgar(db)["idempotencia"] >= 1
        assert db.get(Idempotencia, "vieja:1") is None and db.get(Idempotencia, "nueva:1") is not None


def test_bitacora_general(admin, interno):
    nuevo = admin.post("/usuarios", {"email": "auditado@demo.com", "nombre": "Auditado", "rol": "interno",
                                     "telefono": "+50370000125", "dos_pasos": True}).json()
    admin.patch(f"/usuarios/{nuevo['id']}", {"cargo": "Analyst", "password": "Nueva#Clave2026x"})
    r = admin.get("/auditoria", params={"entidad": "usuario", "entidad_id": nuevo["id"]}).json()
    acciones = [x["accion"] for x in r["items"]]
    assert acciones[-1] == "creado" and "actualizado" in acciones
    cambio = next(x for x in r["items"] if x["accion"] == "actualizado")
    assert cambio["detalle"]["cargo"]["despues"] == "Analyst" and cambio["detalle"]["clave"] == "cambiada"
    assert "Nueva#Clave2026x" not in str(r)
    assert r["items"][0]["usuario"] and r["items"][0]["entidad_txt"] == "User"
    assert any(e["valor"] == "usuario" for e in admin.get("/auditoria/opciones").json()["entidades"])
    assert interno.get("/auditoria").status_code == 403
