"""Datos visibles por rol: el servidor quita de las respuestas y de los
reportes los grupos de datos que el rol del usuario no ve."""
import io

from openpyxl import load_workbook


def _rol(admin, nombre):
    return next(r for r in admin.get("/roles").json()["roles"] if r["nombre"] == nombre)


def _claves(x, acc=None):
    acc = set() if acc is None else acc
    if isinstance(x, dict):
        for k, v in x.items():
            acc.add(k)
            _claves(v, acc)
    elif isinstance(x, list):
        for v in x:
            _claves(v, acc)
    return acc


def test_proveedor_no_ve_fechas_internas_ni_liberaciones(tnf, admin):
    yo = tnf.get("/auth/me").json()
    assert set(yo["datos_ocultos"]) == {"fechas_internas", "liberaciones", "impuestos"}
    assert admin.get("/auth/me").json()["datos_ocultos"] == []
    oc = tnf.get("/ordenes").json()["items"][0]
    assert not {"fecha_tienda", "tienda_estimada", "liberacion_comercial", "liberacion_logistica"} & set(oc)
    # Sigue sabiendo si la OC se puede facturar, y ve sus importes
    assert "liberada" in oc and "importe" in oc
    # El equipo interno ve todo
    oc = admin.get("/ordenes").json()["items"][0]
    assert {"fecha_tienda", "liberacion_comercial"} <= set(oc)
    # El tablero no le muestra el riesgo frente a la fecha en tienda
    assert "riesgo" not in {k["clave"] for k in tnf.get("/dashboard").json()["kpis"]}


def test_ocultar_precios_a_un_rol(tnf, admin):
    rol = _rol(admin, tnf.get("/auth/me").json()["rol_nombre"])
    try:
        r = admin.patch(f"/roles/{rol['id']}", {"datos_ocultos": rol["datos_ocultos"] + ["precios", "no_existe"]})
        assert r.status_code == 200, r.text
        assert set(r.json()["datos_ocultos"]) == set(rol["datos_ocultos"]) | {"precios"}
        fid = tnf.get("/facturas").json()["items"][0]["id"]
        claves = _claves(tnf.get(f"/facturas/{fid}").json()) | _claves(tnf.get("/ordenes").json())
        assert not {"precio_unitario", "importe", "precio_oc"} & claves
        kpis = tnf.get("/dashboard").json()["kpis"]
        assert all(k.get("formato") != "moneda" for k in kpis)
        # La configuración no se filtra: ahí se administran los datos
        assert "datos" in admin.get("/roles").json()
    finally:
        admin.patch(f"/roles/{rol['id']}", {"datos_ocultos": rol["datos_ocultos"]})


def test_el_administrador_siempre_ve_todo(admin):
    rol = _rol(admin, "Administrator")
    assert admin.patch(f"/roles/{rol['id']}", {"datos_ocultos": ["precios"]}).status_code == 422


def test_reportes_sin_columnas_ocultas(interno, admin):
    rol = _rol(admin, interno.get("/auth/me").json()["rol_nombre"])  # otras pruebas pueden renombrarlo
    try:
        admin.patch(f"/roles/{rol['id']}", {"datos_ocultos": ["fechas_internas", "codigos_internos"]})
        r = interno.get("/seguimiento/ordenes/exportar", params={"formato": "xlsx"})
        assert r.status_code == 200, r.text
        ws = load_workbook(io.BytesIO(r.content)).active
        textos = {str(c.value) for fila in ws.iter_rows() for c in fila if c.value}
        assert "PO" in textos and not {"In store", "Vs. store", "Co. · plant", "Late to store"} & textos
        r = interno.get("/seguimiento/ordenes/exportar", params={"formato": "pdf"})
        assert r.status_code == 200 and r.content[:4] == b"%PDF"
    finally:
        admin.patch(f"/roles/{rol['id']}", {"datos_ocultos": rol["datos_ocultos"]})
