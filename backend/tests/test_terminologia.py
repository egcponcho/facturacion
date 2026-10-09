"""Terminología propia, umbrales configurables y un solo límite de subida."""
import io

import pytest
from openpyxl import load_workbook

from app.core.config import settings
from app.core.errores import ErrorNegocio


def _textos(contenido: bytes) -> str:
    wb = load_workbook(io.BytesIO(contenido))
    return " ".join(str(c.value) for ws in wb for fila in ws.iter_rows() for c in fila if c.value)


def test_terminologia_propia_en_documentos(admin, interno):
    assert admin.put("/organizacion", {"textos": {"fr": {"Invoice": "Facture"}}}).status_code == 422
    r = admin.put("/organizacion", {"textos": {"es": {"Commercial invoice {0} · {1}": "Factura {0}"}}})
    assert r.status_code == 422 and "placeholders" in r.text
    fid = interno.get("/facturas").json()["items"][0]["id"]
    try:
        r = admin.put("/organizacion", {"textos": {"es": {"COMMERCIAL INVOICE": "FACTURA DE EXPORTACIÓN", "Vacío": ""},
                                                   "en": {"COMMERCIAL INVOICE": "EXPORT INVOICE"}}})
        assert r.status_code == 200 and r.json()["textos"]["es"] == {"COMMERCIAL INVOICE": "FACTURA DE EXPORTACIÓN"}
        assert admin.get("/auth/me").json()["organizacion"]["textos"]["en"] == {"COMMERCIAL INVOICE": "EXPORT INVOICE"}
        es = _textos(interno.get(f"/facturas/{fid}/exportar", params={"idioma": "es"}).content)
        en = _textos(interno.get(f"/facturas/{fid}/exportar", params={"idioma": "en"}).content)
        assert "FACTURA DE EXPORTACIÓN" in es and "FACTURA COMERCIAL" not in es
        assert "EXPORT INVOICE" in en and "COMMERCIAL INVOICE" not in en
    finally:
        admin.put("/organizacion", {"textos": {}})
    assert "COMMERCIAL INVOICE" in _textos(interno.get(f"/facturas/{fid}/exportar", params={"idioma": "en"}).content)


def test_dias_de_aviso_de_la_fecha_en_tienda(admin):
    reglas = {r["clave"]: r["valor"] for r in admin.get("/organizacion").json()["reglas"]}
    assert reglas["DIAS_AVISO_TIENDA"] == 30
    assert admin.put("/organizacion", {"reglas": {"DIAS_AVISO_TIENDA": -1}}).status_code == 422
    try:
        assert admin.put("/organizacion", {"reglas": {"DIAS_AVISO_TIENDA": 45}}).status_code == 200
        assert admin.get("/auth/me").json()["config"]["dias_aviso_tienda"] == 45
    finally:
        admin.put("/organizacion", {"reglas": {"DIAS_AVISO_TIENDA": 30}})


def test_un_solo_limite_de_subida(admin, monkeypatch):
    from app.core.archivos import exigir_tamano

    assert admin.get("/auth/me").json()["max_subida_mb"] == settings.MAX_SUBIDA_MB
    monkeypatch.setattr(settings, "MAX_SUBIDA_MB", 1)
    exigir_tamano(b"x" * 1024 * 1024)
    with pytest.raises(ErrorNegocio) as e:
        exigir_tamano(b"x" * (1024 * 1024 + 1))
    assert e.value.status == 413 and "maximum 1 MB" in e.value.mensaje
