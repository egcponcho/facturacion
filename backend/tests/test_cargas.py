"""Cargas masivas: artículos con su ficha técnica y cualquier catálogo de
datos maestros desde Excel; exportación con filtros."""
import io
from datetime import date, timedelta

from openpyxl import Workbook, load_workbook


def _xlsx(filas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for f in filas:
        ws.append(f)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _subir(api, url, contenido):
    return api.c.post(f"/api{url}", headers=api.h, files={"archivo": ("datos.xlsx", io.BytesIO(contenido), "application/octet-stream")})


def test_articulos_con_ficha(interno):
    pl = interno.get("/catalogos/articulos/plantilla")
    wb = load_workbook(io.BytesIO(pl.content))
    enc = [c.value for c in wb["Generics"][1]]
    assert enc[0] == "Generic code *" and "Category" in enc and "Upper" in enc
    assert "Supplier SKU" in [c.value for c in wb["Sizes"][1]]
    # También se acepta una hoja con una fila por artículo (formato anterior)
    enc = ["Item code", "Supplier SKU", "Style", "Color", "Size", "Brand", "Item group", "Supplier", "Unit",
           "Commercial name", "Category", "Gender", "Who it is for", "Country of origin", "Upper", "Sole", "Footwear style"]
    filas = [enc,
             ["30099980001", "VN0A5KRFBLK-8", "VN0A5KRF", "Black", "8", "VANS", "CALZ-CAS", "VANS", "PAR", "Sk8-Hi canvas sneaker",
              "Footwear: sneakers, boots, shoes, sandals", "Unisex", "Adult", "Vietnam", "100% canvas", "100% rubber", "Sneaker"],
             ["30099980002", "VN0A5KRFBLK-9", "VN0A5KRF", "Black", "9", "VANS", "CALZ-CAS", "VANS", "PAR", "", "", "", "", "", "", "", ""],
             ["99 *", "", "X", "Y", "1", "VANS", "CALZ-CAS", "VANS", "PAR", "", "", "", "", "", "", "", ""],
             ["30099980003", "", "VN0A5KRF", "Black", "10", "NOPE", "CALZ-CAS", "VANS", "PAR", "", "", "", "", "", "", "", ""]]
    r = _subir(interno, "/catalogos/articulos/importar", _xlsx(filas))
    assert r.status_code == 200, r.text
    r = r.json()
    assert r["creados"] == 2 and len(r["errores"]) == 2 and len(r["productos"]) == 1
    det = interno.get(f"/productos/{r['productos'][0]}").json()
    assert det["tipo"] == "calzado" and det["pais_origen"] == "VN" and det["nombre"] == "Sk8-Hi canvas sneaker"
    assert det["ficha"]["comp"] == {"corte": "100% canvas", "suela": "100% rubber"} and det["ficha"]["estiloCalz"] == "tenis"
    assert det["ficha"]["genero"] == "U" and {a["sku_proveedor"] for a in det["articulos"]} == {"VN0A5KRFBLK-8", "VN0A5KRFBLK-9"}
    # El navegador completa la ficha con el motor y guarda la clasificación
    resultado = {"sugerido": "640419", "confianza": "high", "completa": True, "descripcion_comercial": "Vans Sk8-Hi · Unisex sneaker",
                 "partidas": {"SV": {"codigo": "6404199000", "estado": "ok"}}}
    x = interno.post("/productos/clasificar", {"items": [{"id": det["id"], "tipo": "calzado", "resultado": resultado,
                                                          "ficha": {"altura": "tobillo", "genero": "M", "comp": {"forro": "100% textile"}}}]})
    assert x.status_code == 200, x.text
    det = interno.get(f"/productos/{det['id']}").json()
    # Completa lo que faltaba sin pisar lo cargado
    assert det["estado"] == "sugerida" and det["ficha"]["altura"] == "tobillo" and det["ficha"]["genero"] == "U"
    assert det["ficha"]["comp"]["forro"] == "100% textile" and det["descripcion_comercial"] == "Vans Sk8-Hi · Unisex sneaker"
    # El código de artículo es libre (letras y números), pero no acepta cualquier carácter
    assert any("Letters and numbers" in e["mensaje"] for e in r["errores"])


def test_catalogos_por_excel(interno):
    pl = interno.get("/catalogos/marcas/plantilla")
    assert pl.status_code == 200 and "Instructions" in load_workbook(io.BytesIO(pl.content)).sheetnames
    r = _subir(interno, "/catalogos/marcas/importar", _xlsx([["Code", "Name", "Active"], ["TIMB", "Timberland", "Yes"],
                                                             ["VANS", "Vans Off The Wall", "Yes"], ["", "", ""]]))
    assert r.status_code == 200, r.text
    assert r.json()["creados"] == 1 and r.json()["actualizados"] == 1
    r = _subir(interno, "/catalogos/centros/importar", _xlsx([["Code", "Company", "Name", "Country", "Type"],
                                                              ["0999", "8000", "Test DC", "SV", "Local warehouse"],
                                                              ["0998", "XXXX", "Bad", "SV", "Store"]]))
    assert r.json()["creados"] == 1 and "XXXX" in r.json()["errores"][0]["mensaje"], r.text
    nuevo = interno.get("/catalogos/centros", params={"q": "Test DC"}).json()["items"][0]
    assert interno.delete_(f"/catalogos/centros/{nuevo['id']}").status_code == 200
    for formato in ("xlsx", "pdf"):
        r = interno.get("/catalogos/marcas/exportar", params={"q": "Vans", "formato": formato})
        assert r.status_code == 200 and len(r.content) > 800
    r = interno.get("/catalogos/articulos/exportar", params={"marca_id": 1, "formato": "xlsx"})
    assert r.status_code == 200


def test_tablero_por_periodo(interno, tnf):
    from datetime import date, timedelta

    hoy = date.today()
    d = interno.get("/dashboard").json()["periodo"]
    assert d["desde"] == hoy.replace(day=1).isoformat() and d["facturado"]["grano"] in ("dia", "semana")
    assert {x["clave"] for x in d["resumen"]} == {"facturado", "pls", "llegadas", "clasificados"}
    anio = interno.get("/dashboard", params={"desde": (hoy - timedelta(days=364)).isoformat(), "hasta": hoy.isoformat()}).json()
    assert anio["periodo"]["facturado"]["grano"] == "mes" and len(anio["periodo"]["facturado"]["serie"]) >= 12
    total = sum(x["importe"] for x in anio["periodo"]["facturado"]["serie"])
    assert total > 0
    # Filtrar por marca reduce el total
    solo = interno.get("/dashboard", params={"desde": (hoy - timedelta(days=364)).isoformat(), "marcas": "VANS"}).json()
    assert sum(x["importe"] for x in solo["periodo"]["facturado"]["serie"]) < total
    semana = tnf.get("/dashboard", params={"desde": (hoy - timedelta(days=6)).isoformat()}).json()["periodo"]
    assert semana["facturado"]["grano"] == "dia" and len(semana["facturado"]["serie"]) == 7
    # Seguimiento con varios valores en un filtro
    r = interno.get("/seguimiento/ordenes", params={"proveedor": "The North Face,Vans"}).json()
    assert r["total"] >= interno.get("/seguimiento/ordenes", params={"proveedor": "Vans"}).json()["total"]


def test_genericos(interno):
    """El código de artículo: 8 dígitos de genérico (estilo-color) + 3 de talla.
    La clasificación es del genérico; sus tallas y prepacks la comparten."""
    m = {x["codigo"]: x["id"] for x in interno.get("/catalogos/marcas", params={"size": 100}).json()["items"]}
    g = {x["codigo"]: x["id"] for x in interno.get("/catalogos/grupos").json()["items"]}
    pv = {x["codigo"]: x["id"] for x in interno.get("/catalogos/proveedores").json()["items"]}
    base = {"estilo": "vn0a3wm3", "color": "Navy", "marca_id": m["VANS"], "grupo_id": g["CALZ-CAS"], "proveedor_id": pv["VANS"],
            "unidad": "PAR"}
    assert interno.post("/catalogos/genericos", {**base, "generico": "2009 997?"}).status_code == 422
    r = interno.post("/catalogos/genericos", {**base, "generico": "30099970", "nombre": "Era",
                                              "tallas": [{"talla": "8", "upc": "0196999000001", "sku_proveedor": "VN0A3WM3NVY-8"},
                                                         {"talla": "9.5"}, {"talla": "10", "sufijo": "101"}]})
    assert r.status_code == 200, r.text
    d = r.json()
    # Código usual de calzado: talla × 10 (8 → 080, 9.5 → 095); también se puede escribir
    assert [t["sku"] for t in d["tallas"]] == ["30099970080", "30099970095", "30099970101"] and d["siguiente"] == "001"
    assert d["estilo"] == "VN0A3WM3"
    # Agregar tallas: heredan los datos del genérico; una talla repetida no entra
    r = interno.post("/catalogos/genericos/30099970/tallas", {"tallas": [{"talla": "11"}]})
    assert r.status_code == 200 and r.json()["tallas"][-1]["sku"] == "30099970110"
    assert interno.post("/catalogos/genericos/30099970/tallas", {"tallas": [{"talla": "8"}]}).status_code == 422
    # Un artículo suelto con ese genérico debe ser del mismo estilo-color
    r = interno.post("/catalogos/articulos", {**base, "sku": "30099970020", "generico": "30099970", "talla": "12", "color": "Red", "tipo": "SOLIDO"})
    assert r.status_code == 422 and "Generic 30099970" in r.json()["detalle"][0]["mensaje"]
    # Todas las tallas son un solo producto (la clasificación es del genérico)
    prod = interno.get("/productos", params={"q": "30099970"}).json()["items"]
    assert len(prod) == 1 and prod[0]["codigo_generico"] == "30099970" and prod[0]["skus"] == 4
    # El prepack toma el genérico de sus sólidos, sea cual sea su código
    sol = interno.get("/catalogos/articulos", params={"q": "30099970080"}).json()["items"][0]
    r = interno.post("/catalogos/prepacks", {"sku": "PP-EE04", "codigo": "EE04", "estilo": "VN0A3WM3", "color": "Navy",
                                             "componentes": [{"articulo_id": sol["id"], "cantidad": 4}]})
    assert r.status_code == 200, r.text
    pp = interno.get("/catalogos/articulos", params={"q": "PP-EE04"}).json()["items"][0]
    assert pp["generico"] == "30099970"
    assert interno.get("/productos", params={"q": "30099970"}).json()["items"][0]["n_prepacks"] == 1


def test_carga_por_generico(interno):
    pl = load_workbook(io.BytesIO(interno.get("/catalogos/articulos/plantilla").content))
    assert pl.sheetnames[:2] == ["Generics", "Sizes"] and pl["Generics"]["A1"].value == "Generic code *"
    v = interno.get("/catalogos/articulos/plantilla", params={"vista": 1}).json()
    assert [h["nombre"] for h in v["hojas"]] == ["Generics", "Sizes"] and v["instrucciones"]
    cols = {c["nombre"]: c for c in v["hojas"][1]["columnas"]}
    assert cols["Generic code"]["req"] and "item code is empty" in cols["Size code"]["ayuda"] and v["hojas"][1]["filas"]
    assert "Commercial name" not in [c["nombre"] for c in v["hojas"][0]["columnas"]]
    assert interno.get("/aranceles/sac/plantilla", params={"vista": 1}).json()["hojas"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Generics"
    ws.append(["Generic code", "Style", "Color", "Brand", "Item group", "Supplier", "Unit", "Commercial name",
               "Country of origin", "Upper", "Sole"])
    ws.append(["30099960", "VN0A4U39", "True White", "VANS", "CALZ-CAS", "VANS", "PAR", "Old Skool Pro", "China",
               "100% suede", "100% rubber"])
    ws.append(["3009 996?", "X", "Y", "VANS", "CALZ-CAS", "VANS", "PAR", "", "", "", ""])
    t = wb.create_sheet("Sizes")
    t.append(["Generic code", "Size", "Size code", "UPC", "Supplier SKU"])
    t.append(["30099960", "7", "", "0196888000007", "VN0A4U39W-7"])
    t.append(["30099960", "8", "", "0196888000008", "VN0A4U39W-8"])
    t.append(["30099961", "9", "", "", ""])
    b = io.BytesIO()
    wb.save(b)
    r = _subir(interno, "/catalogos/articulos/importar", b.getvalue()).json()
    assert r["genericos_creados"] == 1 and r["creados"] == 2 and len(r["errores"]) == 2, r
    det = interno.get(f"/productos/{r['productos'][0]}").json()
    assert det["codigo_generico"] == "30099960" and det["ficha"]["comp"]["corte"] == "100% suede"
    assert sorted(a["sku"] for a in det["articulos"]) == ["30099960070", "30099960080"]
    # La misma carga otra vez actualiza (no duplica)
    b.seek(0)
    r = _subir(interno, "/catalogos/articulos/importar", b.getvalue()).json()
    assert r["genericos_actualizados"] == 1 and r["actualizados"] == 2 and r["creados"] == 0


def test_genericos_compacto_y_edicion(interno):
    r = interno.get("/catalogos/genericos", params={"q": "30095126"}).json()
    g = r["items"][0]
    assert r["total"] == 1 and g["generico"] == "30095126" and g["n_tallas"] >= 4 and g["n_prepacks"] >= 1
    assert g["rango_tallas"] and g["descripcion_comercial"].startswith("CALZADO VANS")
    det = interno.get("/catalogos/genericos/30095126").json()
    datos = {k: det[k] for k in ("estilo", "color", "marca_id", "grupo_id", "proveedor_id", "unidad")}
    # Con prepacks no cambia estilo ni color
    assert interno.put("/catalogos/genericos/30095126", json={**datos, "color": "Otro"}).status_code == 422
    grupos = interno.get("/catalogos/grupos/opciones").json()
    otro = next(x["id"] for x in grupos if x["id"] != datos["grupo_id"])
    r = interno.put("/catalogos/genericos/30095126", json={**datos, "grupo_id": otro})
    assert r.status_code == 200, r.text
    arts = interno.get("/catalogos/articulos", params={"q": "30095126", "size": 50}).json()["items"]
    assert arts and all(a["grupo_id"] == otro for a in arts)
    assert interno.put("/catalogos/genericos/30095126", json=datos).status_code == 200


def test_leadtimes_por_origen(interno):
    r = interno.get("/seguimiento/leadtimes", params={"size": 100}).json()
    assert [e["clave"] for e in r["etapas"]][:3] == ["comercial", "logistica", "despacho"]
    reg = {x["codigo"]: x for x in r["regiones"]}
    assert reg["ASIA"]["dias_liberacion"] == 21 and reg["CAM"]["dias_liberacion"] == 15
    ocs = {i["oc"]: i for i in r["items"]}
    # Asia: liberada 16 días antes de la XF (pide 21) -> tarde; 26 días -> a tiempo
    assert ocs["4400003702"]["lib_dias_antes_xf"] == 16 and ocs["4400003702"]["lib_a_tiempo"] is False
    assert ocs["4400003701"]["lib_a_tiempo"] is True
    # Centroamérica pide 15: 20 días antes está a tiempo
    assert ocs["4400003850"]["region"] == "CAM" and ocs["4400003850"]["lib_a_tiempo"] is True
    # Temprano/tarde contra la fecha límite en puerto (tienda - bodega - ingreso - reexportación)
    o = ocs["4400003702"]
    assert o["riesgo"] == "ATRASO" and o["holgura"] < 0
    assert o["limite_puerto"] == str(date.fromisoformat(o["fecha_tienda"]) - timedelta(days=10))
    hitos = {h["clave"]: h for h in o["hitos"]}
    assert hitos["arribo"]["estado"] == "tarde" and hitos["ingreso"]["fecha"] and hitos["tienda"]["estimada"]
    vn = next(x for x in r["origenes"] if x["origen"] == "VN")
    assert vn["etapas"]["transito"]["prom"] == 30 and vn["total_prom"] and vn["lib"]["meta"] == 21
    assert r["kpis"]["lib_total"] >= 10 and r["kpis"]["tarde"] >= 1
    # Filtros: región y solo las liberaciones tarde
    solo = interno.get("/seguimiento/leadtimes", params={"region": "CAM"}).json()
    assert {i["region"] for i in solo["items"]} == {"CAM"}
    tarde = interno.get("/seguimiento/leadtimes", params={"lib": "tarde", "size": 100}).json()["items"]
    assert tarde and all(i["lib_a_tiempo"] is False for i in tarde)
    # Las unidades de carga usan la misma fecha límite
    u = interno.get("/seguimiento/embarques").json()["items"]
    assert u
